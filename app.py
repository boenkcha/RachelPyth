from flask import Flask, jsonify, request, send_from_directory
import sqlite3
import os
import sys
import base64
import re
import io
import csv
from datetime import date, datetime, timedelta
import json
import uuid
from decimal import Decimal

def _base_dir():
    """Folder dasar aplikasi. Jalan normal (`python app.py`) → folder tempat
    app.py ini berada. Sudah di-compile PyInstaller → folder tempat file .exe
    berada (BUKAN folder temp ekstraksi PyInstaller, yg dihapus tiap keluar
    program) — supaya static/uploads/logo/backup konsisten nempel di sebelah
    .exe-nya, tidak hilang/salah tempat tiap restart server."""
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))

app = Flask(__name__, static_folder=os.path.join(_base_dir(), 'static'))
active_sessions = {}   # token → user dict

def get_current_user():
    token = request.headers.get('X-Token','') or request.headers.get('Authorization','').replace('Bearer ','')
    return active_sessions.get(token)
DB_PATH    = os.path.join(_base_dir(), 'rachel_clinic.db')
UPLOAD_DIR = os.path.join(_base_dir(), 'static', 'uploads', 'foto_pasien')
os.makedirs(UPLOAD_DIR, exist_ok=True)

try:
    from PIL import Image, ImageOps
    import io as _io
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

# ─── KONFIGURASI BACKEND DATABASE (SQLite default lama, MySQL baru) ──────────
# Ganti balik ke SQLite kapan saja: RC_DB_BACKEND=sqlite python3 app.py
DB_BACKEND = os.environ.get('RC_DB_BACKEND', 'mysql').strip().lower()  # 'sqlite' | 'mysql'
# Password MySQL TIDAK PUNYA default lagi (2026-10-26, sebelum repo ini dipublikasikan
# ke GitHub) — WAJIB diset lewat env var RC_MYSQL_PASSWORD saat backend=mysql, supaya
# tidak ada kredensial produksi tertanam di source code.
_MYSQL_PASSWORD = os.environ.get('RC_MYSQL_PASSWORD', '')
if DB_BACKEND == 'mysql' and not _MYSQL_PASSWORD:
    raise RuntimeError(
        "RC_MYSQL_PASSWORD wajib diset (env var) saat RC_DB_BACKEND=mysql — "
        "tidak ada password default lagi demi keamanan. Contoh:\n"
        "  RC_MYSQL_PASSWORD=xxxxx python3 app.py\n"
        "(atau set RC_DB_BACKEND=sqlite kalau mau pakai SQLite tanpa MySQL)"
    )
MYSQL_CONFIG = dict(
    host=os.environ.get('RC_MYSQL_HOST', 'localhost'),
    port=int(os.environ.get('RC_MYSQL_PORT', '3306')),
    user=os.environ.get('RC_MYSQL_USER', 'rachel_app'),
    password=_MYSQL_PASSWORD,
    database=os.environ.get('RC_MYSQL_DB', 'rachel_clinic'),
    charset='utf8mb4',
)

try:
    import pymysql
    import pymysql.cursors
except ImportError:
    pymysql = None

if DB_BACKEND == 'mysql' and pymysql is None:
    raise RuntimeError("RC_DB_BACKEND=mysql tapi module 'pymysql' belum terinstall. Jalankan: pip install pymysql")


class DictRow:
    """Row dual-akses (index & nama kolom) — meniru sqlite3.Row supaya kode lama
    yang pakai row['kolom'] ATAUPUN row[0] tetap jalan tanpa perubahan."""
    __slots__ = ('_keys', '_vals')

    def __init__(self, keys, vals):
        self._keys = keys
        self._vals = vals

    def __getitem__(self, key):
        if isinstance(key, (int, slice)):
            return self._vals[key]
        return self._vals[self._keys.index(key)]

    def keys(self):
        return self._keys

    def __iter__(self):
        return iter(self._vals)

    def __len__(self):
        return len(self._vals)

    def __contains__(self, key):
        return key in self._keys

    def __repr__(self):
        return repr(dict(zip(self._keys, self._vals)))


# ─── Dialect translator: SQL gaya SQLite → MySQL (aktif hanya saat DB_BACKEND=='mysql') ──
_DIALECT_PATTERNS = [
    (re.compile(r"PRAGMA\s+foreign_keys\s*=\s*OFF", re.I), "SET FOREIGN_KEY_CHECKS=0"),
    (re.compile(r"PRAGMA\s+foreign_keys\s*=\s*ON", re.I), "SET FOREIGN_KEY_CHECKS=1"),
    (re.compile(r"PRAGMA\s+journal_mode\s*=\s*\w+", re.I), "SELECT 1"),
    (re.compile(r"PRAGMA\s+synchronous\s*=\s*\w+", re.I), "SELECT 1"),
    (re.compile(r"INSERT\s+OR\s+IGNORE\s+INTO", re.I), "INSERT IGNORE INTO"),
    (re.compile(r"INSERT\s+OR\s+REPLACE\s+INTO", re.I), "REPLACE INTO"),
    (re.compile(r"datetime\(\s*'now'\s*\)", re.I), "NOW()"),
    (re.compile(r"date\(\s*'now'\s*\)", re.I), "CURDATE()"),
    # MAX(0,...)/MAX(0.0,...) = idiom "clamp ke non-negatif" (scalar, 2 argumen) —
    # SQLite izinkan MAX() scalar multi-arg, MySQL MAX() HANYA agregat 1-kolom.
    # Regex sempit: hanya cocok literal 0/0.0 sbg argumen pertama diikuti koma,
    # supaya SEMUA pemakaian MAX(kolom)/MIN(kolom) sbg agregat (dipakai luas di
    # laporan) tidak ikut ke-translate.
    (re.compile(r"\bMAX\(\s*(0(?:\.0)?)\s*,", re.I), r"GREATEST(\1,"),
]
_NAMED_PARAM_RE = re.compile(r":([A-Za-z_][A-Za-z0-9_]*)")


def _translate_sql(sql):
    """Terjemahkan satu query gaya SQLite ke MySQL. Urutan penting:
    1) escape literal '%' dulu (mis. LIKE 'retur%') supaya tidak tabrakan
       dengan gaya parameter %s milik PyMySQL,
    2) baru substitusi sintaks (PRAGMA/INSERT OR .../date func),
    3) baru ganti placeholder ':nama'/'?' → '%(nama)s'/'%s'."""
    if DB_BACKEND != 'mysql':
        return sql
    sql = sql.replace('%', '%%')
    for pat, repl in _DIALECT_PATTERNS:
        sql = pat.sub(repl, sql)
    if ':' in sql:
        sql = _NAMED_PARAM_RE.sub(r"%(\1)s", sql)
    sql = sql.replace('?', '%s')
    return sql


if pymysql:
    class RowCursor(pymysql.cursors.DictCursorMixin, pymysql.cursors.Cursor):
        """Cursor MySQL: menerjemahkan dialect SQL sebelum eksekusi, dan
        membungkus tiap baris hasil jadi DictRow (dual-akses index/nama).

        execute() mengembalikan `self` (bukan rowcount int spt PyMySQL asli) —
        meniru sqlite3.Cursor.execute() supaya idiom chaining yang dipakai
        luas di kode raw-cursor (mis. proses_retur, update_racikan):
            cur.execute(sql, params).fetchone()
        tetap jalan tanpa perlu diubah satu-satu.

        executemany() SENGAJA tidak pakai jalur bulk-insert bawaan PyMySQL
        (yang secara internal panggil self.execute() lagi dan expect int
        return — konflik dgn perubahan di atas); diganti loop execute per
        baris via super().execute() langsung. Lebih lambat dari bulk-insert
        asli tapi cukup krn dipakai hanya di seed data & migrasi CSV
        (bukan hot-path request)."""

        def execute(self, query, args=None):
            super().execute(_translate_sql(query), args)
            return self

        def executemany(self, query, args):
            translated = _translate_sql(query)
            total = 0
            for arg in args:
                total += super().execute(translated, arg)
            self.rowcount = total
            return total

        def _conv_row(self, row):
            if row is None:
                return None
            # MySQL/PyMySQL: SUM()/AVG() dll atas kolom INT mengembalikan
            # decimal.Decimal (bukan int/float spt SQLite selalu). Kalau ini
            # dicampur aritmetika dgn float biasa (mis. hasil kolom DOUBLE lain)
            # Python akan TypeError. Konversi ke float di satu titik ini supaya
            # SEMUA hasil query konsisten spt SQLite, tak perlu ditambal manual
            # di tiap tempat yg kebetulan melakukan aritmetika.
            row = tuple(float(v) if isinstance(v, Decimal) else v for v in row)
            return DictRow(self._fields, row)


    class MySQLConn:
        """Wrapper tipis di atas koneksi PyMySQL supaya kompatibel dengan pola
        sqlite3.Connection yang dipakai kode lama: con.execute(...), con.cursor(),
        con.commit(), con.close(), con.row_factory=... — pymysql.Connection asli
        tidak punya method .execute() langsung di connection."""

        def __init__(self, raw):
            self._raw = raw

        def execute(self, sql, params=()):
            cur = self._raw.cursor(RowCursor)
            cur.execute(sql, params)
            return cur

        def cursor(self):
            return self._raw.cursor(RowCursor)

        def commit(self):
            self._raw.commit()

        def rollback(self):
            self._raw.rollback()

        def close(self):
            self._raw.close()

        @property
        def row_factory(self):
            return None

        @row_factory.setter
        def row_factory(self, value):
            pass  # no-op; RowCursor sudah selalu bungkus DictRow


def _mysql_connect():
    # sql_mode dikosongkan: MySQL 8 default strict (mis. ONLY_FULL_GROUP_BY)
    # jauh lebih ketat dari SQLite (yang selalu longgar soal GROUP BY tanpa
    # agregat & koersi tipe) — banyak query laporan di app ini bergantung pada
    # perilaku longgar itu. Menyamakan sql_mode = zero perubahan perilaku.
    return pymysql.connect(cursorclass=RowCursor, autocommit=False,
                            init_command="SET sql_mode=''", **MYSQL_CONFIG)


# ─── DB HELPER ───────────────────────────────────────────────
def get_db():
    if DB_BACKEND == 'mysql':
        return MySQLConn(_mysql_connect())
    con = sqlite3.connect(DB_PATH, timeout=10, check_same_thread=False)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys = ON")
    con.execute("PRAGMA journal_mode = WAL")   # lebih baik untuk multi-koneksi bersamaan
    con.execute("PRAGMA synchronous = NORMAL")
    return con

def get_raw_db(timeout=30):
    """Untuk route yang butuh transaksi/koneksi manual (bypass query()/execute()),
    mis. proses_retur() dan route migrasi_*_eksekusi()."""
    if DB_BACKEND == 'mysql':
        return MySQLConn(_mysql_connect())
    return sqlite3.connect(DB_PATH, timeout=timeout)

def _bulk_executemany(con, cur, sql, args_list):
    """INSERT massal cepat utk migrasi CSV besar (ratusan ribu baris).
    MySQL: RowCursor.executemany() SENGAJA loop 1 query per baris (lihat docstring
    RowCursor di atas — override .execute() balikin `self`, bentrok dgn bulk-insert
    bawaan PyMySQL yg internal manggil self.execute() dan expect rowcount int).
    Itu aman utk seed data kecil, tapi utk CSV 700rb baris = 700rb round-trip network
    ke server, bisa berjam-jam. Solusinya: pakai cursor PyMySQL POLOS (bukan RowCursor)
    langsung dari sini — PyMySQL native executemany() otomatis MENGGABUNG banyak baris
    jadi 1 statement "INSERT ... VALUES (...),(...),(...)" per chunk (dibatasi
    max_allowed_packet), jadi cuma segelintir round-trip alih-alih satu per baris.
    SQLite: sqlite3.Cursor.executemany() bawaan sudah cukup cepat, dipakai apa adanya."""
    if not args_list:
        return 0
    if DB_BACKEND == 'mysql':
        # WAJIB pymysql.cursors.Cursor eksplisit — con._raw.cursor() TANPA argumen
        # ikut cursorclass default koneksi (RowCursor, di-set saat connect()), yg
        # justru override execute()/executemany() yg mau kita hindari di sini.
        raw_cur = con._raw.cursor(pymysql.cursors.Cursor)
        try:
            n = raw_cur.executemany(_translate_sql(sql), args_list)
        finally:
            raw_cur.close()
        return n or len(args_list)
    cur.executemany(sql, args_list)
    return len(args_list)


def query(sql, params=(), one=False):
    con = get_db()
    cur = con.execute(sql, params)
    rv = cur.fetchall()
    con.close()
    return (rv[0] if rv else None) if one else rv

def execute(sql, params=()):
    con = get_db()
    cur = con.execute(sql, params)
    con.commit()
    last_id = cur.lastrowid
    con.close()
    return last_id

def rows_to_list(rows):
    return [dict(r) for r in rows]

# ─── PENGATURAN (key-value) ──────────────────────────────────
def get_setting(key, default=None):
    # `key` di-backtick-quote krn reserved word di MySQL; backtick juga valid di SQLite.
    row = query("SELECT value FROM Pengaturan WHERE `key`=?", (key,), one=True)
    return row['value'] if row else default

def set_setting(key, value):
    execute("INSERT OR REPLACE INTO Pengaturan (`key`,value,updated_at) VALUES (?,?,datetime('now'))",
            (key, str(value)))

def get_point_rate():
    """Rp per 1 point — bisa diubah lewat Pengaturan. Default Rp50.000."""
    try:
        rate = float(get_setting('point_rate', 50000) or 50000)
        return rate if rate > 0 else 50000
    except (TypeError, ValueError):
        return 50000

def get_point_redeem_rate():
    """Rp per 1 point SAAT DITUKAR sbg pembayaran (2026-10-25) — kurs terpisah dari
    get_point_rate() (kurs dapat point), diubah sendiri lewat Pengaturan. Default
    Rp50.000 (sama dgn kurs dapat point, tapi independen — ganti salah satu tidak
    ikut mengubah yang lain)."""
    try:
        rate = float(get_setting('point_redeem_rate', 50000) or 50000)
        return rate if rate > 0 else 50000
    except (TypeError, ValueError):
        return 50000

# ─── AKTIVITAS LOG (audit trail — pengaman anti-manipulasi) ──
def catat_log(aksi, entitas='', ref='', keterangan='', data_lama=None, data_baru=None, actor=None):
    """Tulis 1 baris audit log. Dibungkus try/except — kegagalan logging TIDAK BOLEH
    pernah menggagalkan operasi bisnis yang sedang berjalan.
    `actor` dipakai utk kasus tanpa sesi token valid saat ini (mis. login/login gagal)."""
    try:
        user = actor if actor is not None else (get_current_user() or {})
        execute("""INSERT INTO Aktivitas_Log
            (waktu,username,nama_user,role,kd_lokasi,aksi,entitas,ref,keterangan,data_lama,data_baru,ip_address)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
            (str(datetime.now()), user.get('username',''), user.get('nama',''), user.get('role',''),
             user.get('kd_lokasi',''), aksi, entitas, str(ref), keterangan,
             json.dumps(data_lama, default=str) if data_lama is not None else None,
             json.dumps(data_baru, default=str) if data_baru is not None else None,
             request.remote_addr or ''))
    except Exception:
        pass

def _pasien_faktur(no_faktur):
    """Ambil kd_pasien+nm_pasien pemilik suatu faktur — dipakai catat_log() transaksi
    supaya admin bisa lihat pasien mana yang transaksinya diubah/diretur."""
    row = query("""SELECT p.kd_pasien, pa.nm_pasien FROM Penjualan p
                   LEFT JOIN Pasien pa ON p.kd_pasien=pa.kd_pasien
                   WHERE p.no_faktur=?""", (no_faktur,), one=True)
    if not row or not row['kd_pasien']:
        return ''
    return f"{row['nm_pasien'] or '—'} ({row['kd_pasien']})"

def _seed_barang_master(cur):
    """Isi master Barang dari Barang.csv sekali saja (jika tabel masih kosong)."""
    import csv
    cur.execute("SELECT COUNT(*) FROM Barang")
    if cur.fetchone()[0] > 0:
        return
    path = os.path.join(_base_dir(), 'Barang.csv')
    if not os.path.exists(path):
        return
    seen = {}   # kd_barang → row dict (dedup, prefer yang harga_jual terisi)
    with open(path, encoding='utf-8-sig', errors='replace') as f:
        for r in csv.DictReader(f):
            kd = (r.get('Kd_barang') or '').strip()
            if not kd:
                continue
            harga_raw = (r.get('Harga_jual') or '').strip()
            try:
                harga = float(harga_raw) if harga_raw else 0
            except ValueError:
                harga = 0
            try:
                isi = float((r.get('Isi') or '1').strip() or 1)
            except ValueError:
                isi = 1
            row = {
                'kd_barang':  kd,
                'nm_barang':  kd,                              # tampil = deskriptif
                'nm_generik': (r.get('Nm_barang') or '').strip(),
                'jenis':      (r.get('J_barang') or '').strip(),
                'merk':       (r.get('Merk') or '').strip(),
                'isi':        isi,
                'satuan':     (r.get('Satuan') or 'Pcs').strip() or 'Pcs',
                'harga_jual': harga,
            }
            prev = seen.get(kd)
            # Dedup: pertahankan baris dgn harga terisi
            if prev is None or (prev['harga_jual'] == 0 and harga > 0):
                seen[kd] = row
    if seen:
        rows_tuple = [
            (r['kd_barang'], r['nm_barang'], r['nm_generik'], r['jenis'],
             r['merk'], r['isi'], r['satuan'], r['harga_jual'])
            for r in seen.values()
        ]
        cur.executemany(
            """INSERT OR IGNORE INTO Barang
               (kd_barang,nm_barang,nm_generik,jenis,merk,isi,satuan,harga_jual,aktif)
               VALUES (?,?,?,?,?,?,?,?,1)""",
            rows_tuple)

# ─── INIT DATABASE ───────────────────────────────────────────
def init_db():
    """Entry point tunggal — dispatch ke skema SQLite (lama, tidak diubah) atau
    MySQL (baru) sesuai DB_BACKEND."""
    if DB_BACKEND == 'mysql':
        init_db_mysql()
    else:
        init_db_sqlite()

def init_db_sqlite():
    con = get_db()
    cur = con.cursor()

    cur.executescript("""
    CREATE TABLE IF NOT EXISTS Pasien (
        kd_pasien TEXT PRIMARY KEY,
        nm_pasien TEXT NOT NULL,
        alamat TEXT,
        telp TEXT,
        hp TEXT,
        tgl_lahir TEXT,
        jenis_kelamin TEXT,
        created_at TEXT DEFAULT (date('now'))
    );

    CREATE TABLE IF NOT EXISTS Member (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        kd_pasien TEXT NOT NULL,
        tgl_daftar TEXT,
        tgl_exp TEXT,
        FOREIGN KEY (kd_pasien) REFERENCES Pasien(kd_pasien)
    );

    CREATE TABLE IF NOT EXISTS DetailPoint (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        kd_pasien TEXT NOT NULL,
        keterangan TEXT,
        debit REAL DEFAULT 0,
        kredit REAL DEFAULT 0,
        tanggal TEXT DEFAULT (date('now'))
    );

    CREATE TABLE IF NOT EXISTS Karyawan (
        kd_karyawan TEXT PRIMARY KEY,
        nm_karyawan TEXT NOT NULL,
        nickname TEXT,
        pekerjaan TEXT,
        telp TEXT,
        aktif INTEGER DEFAULT 1
    );

    CREATE TABLE IF NOT EXISTS Stok_barang (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        kd_barang TEXT NOT NULL,
        nm_barang TEXT NOT NULL,
        harga_beli REAL DEFAULT 0,
        harga_beli_awal REAL DEFAULT 0,
        stok_awal INTEGER DEFAULT 0,
        stok_sisa INTEGER DEFAULT 0,
        stok_gudang INTEGER DEFAULT 0,
        stok_min INTEGER DEFAULT 0,
        harga REAL DEFAULT 0,
        disc TEXT DEFAULT '0',
        satuan TEXT,
        exp TEXT,
        DiscMem REAL DEFAULT 0,
        DiscLain REAL DEFAULT 0
    );

    CREATE TABLE IF NOT EXISTS produk_pelayanan (
        kd_pelayanan TEXT PRIMARY KEY,
        nm_pelayanan TEXT NOT NULL,
        harga REAL DEFAULT 0,
        komisi_dokter REAL DEFAULT 0,
        komisi_supervisor REAL DEFAULT 0,
        komisi_perawat REAL DEFAULT 0,
        aktif INTEGER DEFAULT 1
    );

    CREATE TABLE IF NOT EXISTS Suplier (
        kd_suplier TEXT PRIMARY KEY,
        nm_suplier TEXT NOT NULL,
        alamat TEXT,
        telp TEXT,
        kontak TEXT
    );

    CREATE TABLE IF NOT EXISTS Nomor (
        tanggal TEXT PRIMARY KEY,
        nota_jual INTEGER DEFAULT 0,
        nota_piutang INTEGER DEFAULT 0,
        nota_resep INTEGER DEFAULT 0
    );

    CREATE TABLE IF NOT EXISTS Antrian (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        kd_pasien TEXT,
        tanggal TEXT DEFAULT (date('now')),
        status TEXT DEFAULT 'wait',
        no_faktur TEXT,
        FOREIGN KEY (kd_pasien) REFERENCES Pasien(kd_pasien)
    );

    CREATE TABLE IF NOT EXISTS Penjualan (
        no_faktur TEXT PRIMARY KEY,
        tanggal TEXT,
        total_harga REAL DEFAULT 0,
        hpp REAL DEFAULT 0,
        dokter TEXT,
        kd_pasien TEXT,
        kasir TEXT,
        metode_bayar TEXT DEFAULT 'Tunai',
        status TEXT DEFAULT 'open'
    );

    CREATE TABLE IF NOT EXISTS R_Penjualan (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        no_faktur TEXT NOT NULL,
        tanggal TEXT,
        kd_barang TEXT,
        nm_barang TEXT,
        harga REAL DEFAULT 0,
        harga_asli REAL DEFAULT 0,
        qty INTEGER DEFAULT 1,
        disc TEXT DEFAULT '0',
        sub_total REAL DEFAULT 0,
        total_harbeli REAL DEFAULT 0,
        kd_karyawan TEXT,
        cashier TEXT,
        exp TEXT,
        jenis TEXT DEFAULT 'barang'
    );

    CREATE TABLE IF NOT EXISTS TKredit (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nota TEXT,
        konsumen TEXT,
        total REAL DEFAULT 0,
        bank TEXT,
        kasir TEXT,
        tanggal TEXT DEFAULT (date('now'))
    );

    CREATE TABLE IF NOT EXISTS TKunjunganBerikut (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        no_faktur TEXT,
        next_tanggal TEXT,
        kd_pasien TEXT
    );

    CREATE TABLE IF NOT EXISTS Point_karyawan (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        no_faktur TEXT,
        tanggal TEXT,
        kd_pasien TEXT,
        kd_pelayanan TEXT,
        nm_pelayanan TEXT,
        komisi REAL DEFAULT 0,
        kd_karyawan TEXT
    );

    CREATE TABLE IF NOT EXISTS Tindakan (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tanggal TEXT,
        kd_pelayanan TEXT,
        banyaknya INTEGER DEFAULT 0
    );

    CREATE TABLE IF NOT EXISTS Pembelian (
        no_faktur TEXT PRIMARY KEY,
        tanggal TEXT,
        total_harga REAL DEFAULT 0,
        sistem_bayar TEXT,
        lunas TEXT DEFAULT 'Belum',
        tempo TEXT,
        nm_suplier TEXT
    );

    CREATE TABLE IF NOT EXISTS R_Pembelian (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        no_faktur TEXT,
        tanggal TEXT,
        kd_suplier TEXT,
        nm_suplier TEXT,
        sistem_bayar TEXT,
        kd_barang TEXT,
        nm_barang TEXT,
        disc REAL DEFAULT 0,
        harga_beli_kotor REAL DEFAULT 0,
        harga_beli_bersih REAL DEFAULT 0,
        sub_total REAL DEFAULT 0,
        ppn REAL DEFAULT 0,
        tempo TEXT,
        qty INTEGER DEFAULT 0,
        harga_asli REAL DEFAULT 0,
        satuan TEXT,
        exp TEXT
    );

    CREATE TABLE IF NOT EXISTS RiwayatBarang (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tanggal TEXT,
        kd_barang TEXT,
        exp TEXT,
        keterangan TEXT,
        nilai INTEGER DEFAULT 0,
        saldo_e INTEGER DEFAULT 0,
        saldo_g INTEGER DEFAULT 0
    );

    CREATE TABLE IF NOT EXISTS Log (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tanggal TEXT DEFAULT (date('now')),
        keterangan TEXT,
        qty INTEGER DEFAULT 0,
        kasir TEXT,
        event TEXT
    );

    CREATE TABLE IF NOT EXISTS Aktivitas_Log (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        waktu TEXT DEFAULT (datetime('now')),
        username TEXT DEFAULT '',
        nama_user TEXT DEFAULT '',
        role TEXT DEFAULT '',
        kd_lokasi TEXT DEFAULT '',
        aksi TEXT DEFAULT '',
        entitas TEXT DEFAULT '',
        ref TEXT DEFAULT '',
        keterangan TEXT DEFAULT '',
        data_lama TEXT,
        data_baru TEXT,
        ip_address TEXT DEFAULT ''
    );

    CREATE TABLE IF NOT EXISTS ResepRacikan (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        id_rm INTEGER NOT NULL,
        no_faktur TEXT DEFAULT '',
        kd_pasien TEXT DEFAULT '',
        nama_produk TEXT NOT NULL,
        cara_pakai TEXT DEFAULT '',
        harga_total REAL DEFAULT 0,
        urutan INTEGER DEFAULT 1,
        FOREIGN KEY (id_rm) REFERENCES RekamMedik(id)
    );

    CREATE TABLE IF NOT EXISTS DetailResepRacikan (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        id_resep INTEGER NOT NULL,
        kd_barang TEXT NOT NULL,
        nm_bahan TEXT NOT NULL,
        qty REAL DEFAULT 0,
        satuan TEXT DEFAULT '',
        harga_satuan REAL DEFAULT 0,
        disc TEXT DEFAULT '0',
        sub_total REAL DEFAULT 0,
        exp TEXT DEFAULT '',
        FOREIGN KEY (id_resep) REFERENCES ResepRacikan(id)
    );

    CREATE TABLE IF NOT EXISTS RekamMedik (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        no_rm TEXT,
        kd_pasien TEXT NOT NULL,
        tanggal TEXT DEFAULT (date('now')),
        nm_dokter TEXT DEFAULT '',
        status TEXT DEFAULT 'open',
        S_keluhan TEXT DEFAULT '',
        O_jenis_kulit TEXT DEFAULT '',
        O_kondisi_kulit TEXT DEFAULT '',
        A_diagnosa TEXT DEFAULT '',
        P_tindakan TEXT DEFAULT '',
        P_catatan TEXT DEFAULT '',
        no_faktur TEXT DEFAULT '',
        FOREIGN KEY (kd_pasien) REFERENCES Pasien(kd_pasien)
    );

    CREATE TABLE IF NOT EXISTS Pengguna (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        password TEXT NOT NULL,
        nama TEXT,
        role TEXT DEFAULT 'kasir',
        aktif INTEGER DEFAULT 1
    );

    INSERT OR IGNORE INTO Pengguna (username, password, nama, role) 
    VALUES ('admin', 'admin123', 'Administrator', 'admin');
    """)

    # Tambah kolom racikan_ref ke R_Penjualan jika belum ada
    try:
        cur.execute("ALTER TABLE R_Penjualan ADD COLUMN racikan_ref INTEGER DEFAULT NULL")
    except Exception:
        pass
    # Tambah kolom aktif ke produk_pelayanan (master perawatan) jika belum ada
    try:
        cur.execute("ALTER TABLE produk_pelayanan ADD COLUMN aktif INTEGER DEFAULT 1")
    except Exception:
        pass
    # Penanda 'bukan obat' di master Barang (default 0 = obat → wajib resep)
    try:
        cur.execute("ALTER TABLE Barang ADD COLUMN non_obat INTEGER DEFAULT 0")
    except Exception:
        pass
    # Tautan resep digital ke faktur (1 transaksi = 1 resep)
    try:
        cur.execute("ALTER TABLE ResepDigital ADD COLUMN no_faktur TEXT DEFAULT ''")
    except Exception:
        pass
    # Tambah kolom disc_trx ke Penjualan jika belum ada
    try:
        cur.execute("ALTER TABLE Penjualan ADD COLUMN disc_trx REAL DEFAULT 0")
        cur.execute("ALTER TABLE Penjualan ADD COLUMN disc_trx_nilai REAL DEFAULT 0")
    except Exception:
        pass
    # Tambah kolom status + created_by ke ResepRacikan untuk fitur resep digital dokter
    try:
        cur.execute("ALTER TABLE ResepRacikan ADD COLUMN status TEXT DEFAULT 'confirmed'")
    except Exception:
        pass
    try:
        cur.execute("ALTER TABLE ResepRacikan ADD COLUMN created_by TEXT DEFAULT ''")
    except Exception:
        pass
    try:
        cur.execute("ALTER TABLE ResepRacikan ADD COLUMN catatan_kasir TEXT DEFAULT ''")
    except Exception:
        pass
    # Status pengerjaan racikan oleh apoteker — 'pending' (perlu diracik) | 'selesai'.
    # Beda dari kolom `status` di atas (draft/confirmed = status transaksi kasir).
    try:
        cur.execute("ALTER TABLE ResepRacikan ADD COLUMN status_racik TEXT DEFAULT 'pending'")
    except Exception:
        pass
    # Disc header (%) di file _P Terima Kiriman — kolom ke-3 yg sebelumnya diabaikan.
    # total_setelah_disc = total_rincian * (1 - disc_pct/100), dibandingkan ke total_header
    # (bukan total_rincian mentah) supaya diskon riil tidak dianggap "selisih"/mismatch palsu.
    try:
        cur.execute("ALTER TABLE Terima_Kiriman ADD COLUMN disc_pct REAL DEFAULT 0")
        cur.execute("ALTER TABLE Terima_Kiriman ADD COLUMN total_setelah_disc REAL DEFAULT 0")
    except Exception:
        pass
    # induk_id: penanda record "susulan" (proses ulang barang terlewat) — merujuk id
    # record ASLI-nya, supaya daftar Terima Kiriman bisa menampilkan susulan sbg
    # sub-baris di bawah induknya, bukan baris terpisah yg kelihatan spt "dobel"
    # (keluhan user 2026-10-01 — lihat terima_kiriman_proses_ulang()).
    try:
        cur.execute("ALTER TABLE Terima_Kiriman ADD COLUMN induk_id INTEGER DEFAULT NULL")
    except Exception:
        pass
    # Disc (%) per bahan racikan — bahan bisa saja punya diskon sendiri saat diracik.
    try:
        cur.execute("ALTER TABLE DetailResepRacikan ADD COLUMN disc TEXT DEFAULT '0'")
    except Exception:
        pass
    # Nomor faktur terakhir yg direserve per (kode,tanggal) — dipakai validasi
    # "sudah ada transaksi berjalan atau belum" sebelum increment counter baru
    # (lihat get_nomor_faktur) supaya buka-tutup POS sblm isi pasien/item tidak
    # membuang nomor faktur / bikin nomor lompat.
    try:
        cur.execute("ALTER TABLE Counter_Faktur ADD COLUMN last_faktur TEXT DEFAULT ''")
    except Exception:
        pass

    # ── Resep Digital (dokumen teks, tidak terkait stok) ──────────────────────
    cur.execute("""
        CREATE TABLE IF NOT EXISTS ResepDigital (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            no_resep    TEXT NOT NULL,
            kd_pasien   TEXT NOT NULL,
            id_rm       INTEGER,
            nm_dokter   TEXT DEFAULT '',
            tanggal     TEXT NOT NULL,
            created_by  TEXT DEFAULT '',
            catatan     TEXT DEFAULT '',
            no_faktur   TEXT DEFAULT ''
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS DetailResepDigital (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            id_resep    INTEGER NOT NULL,
            nama_obat   TEXT NOT NULL,
            dosis       TEXT DEFAULT '',
            satuan      TEXT DEFAULT '',
            cara_pakai  TEXT DEFAULT '',
            urutan      INTEGER DEFAULT 1
        )
    """)

    # ── Foto Pasien (before/after) ────────────────────────────
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Foto_Pasien (
            id         INTEGER PRIMARY KEY AUTOINCREMENT,
            kd_pasien  TEXT NOT NULL,
            id_rm      INTEGER,
            no_rm      TEXT DEFAULT '',
            tanggal    TEXT DEFAULT (date('now')),
            tipe       TEXT DEFAULT 'sebelum',
            area       TEXT DEFAULT '',
            keterangan TEXT DEFAULT '',
            filename   TEXT NOT NULL,
            created_at TEXT DEFAULT (datetime('now')),
            FOREIGN KEY (kd_pasien) REFERENCES Pasien(kd_pasien)
        )
    """)

    # ── Multi-Lokasi: Lokasi master ──────────────────────────────
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Lokasi (
            kd_lokasi   TEXT PRIMARY KEY,
            nm_lokasi   TEXT NOT NULL,
            tipe        TEXT DEFAULT 'cabang',
            alamat      TEXT DEFAULT '',
            aktif       INTEGER DEFAULT 1
        )
    """)
    cur.execute("INSERT OR IGNORE INTO Lokasi VALUES ('INDUK','Apotek Induk','induk','',1)")
    cur.execute("INSERT OR IGNORE INTO Lokasi VALUES ('CBG1','Klinik Rachel','cabang_dalam','',1)")
    cur.execute("INSERT OR IGNORE INTO Lokasi VALUES ('CBG2','KIBAR','cabang_dalam','',1)")
    # Rename nama lama CBG1 (Cabang Utama → Klinik Rachel) bila masih default
    cur.execute("UPDATE Lokasi SET nm_lokasi='Klinik Rachel' WHERE kd_lokasi='CBG1' AND nm_lokasi='Cabang Utama'")

    # Migrate: tambah kd_lokasi ke tabel yang ada
    for _sql in [
        "ALTER TABLE Stok_barang ADD COLUMN kd_lokasi TEXT DEFAULT 'CBG1'",
        "ALTER TABLE Pengguna    ADD COLUMN kd_lokasi TEXT DEFAULT 'CBG1'",
        "ALTER TABLE Penjualan   ADD COLUMN kd_lokasi TEXT DEFAULT 'CBG1'",
        "ALTER TABLE Pembelian   ADD COLUMN kd_lokasi TEXT DEFAULT 'INDUK'",
        "ALTER TABLE Pengguna    ADD COLUMN kd_karyawan TEXT DEFAULT ''",
        "ALTER TABLE Foto_Pasien ADD COLUMN uploaded_by TEXT DEFAULT ''",
        "ALTER TABLE Pasien      ADD COLUMN hp TEXT DEFAULT ''",
        "ALTER TABLE Pembelian   ADD COLUMN status TEXT DEFAULT 'approved'",
        "ALTER TABLE Pembelian   ADD COLUMN catatan_supervisor TEXT DEFAULT ''",
        "ALTER TABLE Pembelian   ADD COLUMN approved_by TEXT DEFAULT ''",
        "ALTER TABLE Pembelian   ADD COLUMN approved_at TEXT DEFAULT ''",
        "ALTER TABLE Pembelian   ADD COLUMN dibuat_oleh TEXT DEFAULT ''",
        "ALTER TABLE Perawatan_Bahan ADD COLUMN bahan_kunci INTEGER DEFAULT 0",
        "ALTER TABLE produk_pelayanan ADD COLUMN komisi_persen_kunci REAL DEFAULT 0",
        "ALTER TABLE Point_karyawan ADD COLUMN status TEXT DEFAULT 'released'",
    ]:
        try: cur.execute(_sql)
        except: pass
    # Jaring pengaman: faktur lama (sebelum kolom status ada) HARUS dianggap sudah
    # disetujui — stoknya sudah nyata dari dulu, jangan sampai ke-lock retroaktif.
    cur.execute("UPDATE Pembelian SET status='approved' WHERE status IS NULL OR status=''")
    cur.execute("UPDATE Point_karyawan SET status='released' WHERE status IS NULL OR status=''")

    # ── Serah-Terima Bahan Kunci (QR handshake) ───────────────────
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Serah_Terima_Kunci (
            id           INTEGER PRIMARY KEY AUTOINCREMENT,
            no_serah     TEXT NOT NULL,
            no_faktur    TEXT NOT NULL,
            kd_pelayanan TEXT NOT NULL,
            nm_pelayanan TEXT DEFAULT '',
            kd_barang    TEXT NOT NULL,
            nm_barang    TEXT DEFAULT '',
            qty          REAL NOT NULL DEFAULT 1,
            satuan       TEXT DEFAULT '',
            kd_pelaksana TEXT NOT NULL,
            nm_pelaksana TEXT DEFAULT '',
            kd_k_dipilih TEXT NOT NULL,
            kd_k_final   TEXT DEFAULT '',
            komisi_persen REAL DEFAULT 0,
            status       TEXT DEFAULT 'pending',
            token        TEXT DEFAULT '',
            token_exp    TEXT DEFAULT '',
            kd_pasien    TEXT DEFAULT '',
            nm_pasien    TEXT DEFAULT '',
            selesai_at   TEXT DEFAULT '',
            created_at   TEXT DEFAULT (datetime('now'))
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_serah_terima_status ON Serah_Terima_Kunci(status)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_serah_terima_faktur ON Serah_Terima_Kunci(no_faktur)")

    # Set lokasi default untuk data lama
    cur.execute("UPDATE Stok_barang SET kd_lokasi='CBG1' WHERE kd_lokasi IS NULL")
    cur.execute("UPDATE Penjualan   SET kd_lokasi='CBG1' WHERE kd_lokasi IS NULL")
    cur.execute("UPDATE Pembelian   SET kd_lokasi='INDUK' WHERE kd_lokasi IS NULL OR kd_lokasi='CBG1'")
    # Admin default ke INDUK, kasir ke CBG1
    cur.execute("UPDATE Pengguna SET kd_lokasi='INDUK' WHERE role='admin' AND (kd_lokasi IS NULL OR kd_lokasi='CBG1')")
    # Migrasi level Gudang Induk: role lama 'admin_induk' → 'gudang_induk'; nonaktifkan akun duplikat 'admin_induk'
    cur.execute("UPDATE Pengguna SET role='gudang_induk', nama='Gudang Induk' WHERE role='admin_induk'")
    cur.execute("UPDATE Pengguna SET aktif=0 WHERE username='admin_induk' AND role='admin'")

    # ── Stok Induk Multi-Pemilik (3 tempat, 1 gedung, 1 DB — beda kepemilikan) ──
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Induk_Pemilik (
            kd_pemilik  TEXT PRIMARY KEY,
            nm_pemilik  TEXT NOT NULL,
            keterangan  TEXT DEFAULT '',
            aktif       INTEGER DEFAULT 1
        )
    """)
    cur.execute("INSERT OR IGNORE INTO Induk_Pemilik VALUES ('P1','Induk Utama','Pemilik default (data lama)',1)")
    # Tag pemilik di stok induk & pembelian
    for _sql in [
        "ALTER TABLE Stok_barang ADD COLUMN pemilik TEXT DEFAULT ''",
        "ALTER TABLE Pembelian   ADD COLUMN pemilik TEXT DEFAULT ''",
    ]:
        try: cur.execute(_sql)
        except: pass
    # Data lama: semua stok & pembelian INDUK → pemilik utama 'P1'
    cur.execute("UPDATE Stok_barang SET pemilik='P1' WHERE kd_lokasi='INDUK' AND (pemilik IS NULL OR pemilik='')")
    cur.execute("UPDATE Pembelian   SET pemilik='P1' WHERE (pemilik IS NULL OR pemilik='')")

    # ── Konversi Satuan Stok Induk (box→pcs dll) + pindah antar pemilik ──
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Konversi_Stok (
            id              INTEGER PRIMARY KEY AUTOINCREMENT,
            no_konversi     TEXT,
            tgl             TEXT,
            sumber_kd       TEXT, sumber_nm TEXT, sumber_pemilik TEXT,
            qty_sumber      REAL, satuan_sumber TEXT, exp TEXT,
            target_kd       TEXT, target_nm TEXT, target_pemilik TEXT,
            qty_target      REAL, satuan_target TEXT,
            harga_beli_sumber REAL, harga_beli_target REAL,
            created_by      TEXT, created_at TEXT
        )
    """)

    # ── Formula Produksi (BOM) ────────────────────────────────────
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Formula_Produksi (
            id              INTEGER PRIMARY KEY AUTOINCREMENT,
            kd_produk_hasil TEXT NOT NULL,
            nm_produk_hasil TEXT DEFAULT '',
            nm_formula      TEXT NOT NULL,
            qty_hasil       REAL NOT NULL DEFAULT 1,
            satuan_hasil    TEXT DEFAULT 'pcs',
            keterangan      TEXT DEFAULT '',
            aktif           INTEGER DEFAULT 1,
            created_at      TEXT DEFAULT (datetime('now'))
        )
    """)
    try: cur.execute("ALTER TABLE Formula_Produksi ADD COLUMN nm_produk_hasil TEXT DEFAULT ''")
    except: pass
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Formula_Produksi_Detail (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            id_formula  INTEGER NOT NULL,
            kd_barang   TEXT NOT NULL,
            nm_barang   TEXT DEFAULT '',
            qty         REAL NOT NULL,
            satuan      TEXT DEFAULT '',
            FOREIGN KEY (id_formula) REFERENCES Formula_Produksi(id)
        )
    """)

    # ── Formula Perawatan (BOM bahan/alat per treatment) ──────────
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Perawatan_Bahan (
            id           INTEGER PRIMARY KEY AUTOINCREMENT,
            kd_pelayanan TEXT NOT NULL,
            kd_barang    TEXT NOT NULL,
            nm_barang    TEXT DEFAULT '',
            qty          REAL NOT NULL DEFAULT 1,
            satuan       TEXT DEFAULT '',
            tipe         TEXT DEFAULT 'bahan',     -- bahan | alat
            potong_stok  INTEGER DEFAULT 1,        -- 1=potong dari kit perawat, 0=checklist saja
            created_at   TEXT DEFAULT (datetime('now'))
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_r_penjualan_faktur ON R_Penjualan(no_faktur)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_perawatan_bahan_kd ON Perawatan_Bahan(kd_pelayanan)")
    # Ditemukan 2026-09-14 saat investigasi keluhan "kasir harus standby" (lambat) —
    # lihat catatan lebih lengkap di _MYSQL_INDEXES (dipakai kedua backend, sengaja
    # dicatat sekali saja di sana; SQLite jarang jadi masalah performa di skala ini
    # tapi tetap ditambah utk paritas/konsistensi antar backend).
    cur.execute("CREATE INDEX IF NOT EXISTS idx_stok_barang_lokasi ON Stok_barang(kd_lokasi, stok_sisa, harga)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_stok_barang_kode ON Stok_barang(kd_barang)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_penjualan_tanggal ON Penjualan(tanggal)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_penjualan_kdpasien ON Penjualan(kd_pasien)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_r_penjualan_kdbarang ON R_Penjualan(kd_barang)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_r_penjualan_racikanref ON R_Penjualan(racikan_ref)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_point_karyawan_kdkaryawan ON Point_karyawan(kd_karyawan)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_point_karyawan_tanggal ON Point_karyawan(tanggal)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_point_karyawan_faktur ON Point_karyawan(no_faktur)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_pasien_nama ON Pasien(nm_pasien)")

    # ── Kit per perawat (stok bahan/alat pribadi) ─────────────────
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Stok_Perawat (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            kd_karyawan TEXT NOT NULL,
            kd_barang   TEXT NOT NULL,
            nm_barang   TEXT DEFAULT '',
            exp         TEXT DEFAULT '',
            qty_sisa    REAL NOT NULL DEFAULT 0,
            satuan      TEXT DEFAULT '',
            updated_at  TEXT DEFAULT (datetime('now'))
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_stok_perawat ON Stok_Perawat(kd_karyawan, kd_barang, exp)")
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Kit_Log (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            tanggal     TEXT DEFAULT (datetime('now')),
            kd_karyawan TEXT NOT NULL,
            kd_barang   TEXT NOT NULL,
            nm_barang   TEXT DEFAULT '',
            exp         TEXT DEFAULT '',
            qty          REAL NOT NULL,        -- + masuk, - keluar
            saldo        REAL DEFAULT 0,       -- snapshot qty kit (barang+exp) setelah perubahan
            jenis        TEXT NOT NULL,        -- isi_ulang | pakai_treatment | retur_treatment | koreksi
            ref          TEXT DEFAULT '',      -- no_faktur / no_request / catatan
            kd_pelayanan TEXT DEFAULT '',      -- treatment sumber (untuk pakai/retur)
            kd_pasien    TEXT DEFAULT '',      -- pasien tujuan (untuk pakai_treatment)
            nm_pasien    TEXT DEFAULT '',
            kd_lokasi    TEXT DEFAULT '',
            operator     TEXT DEFAULT ''
        )
    """)
    try: cur.execute("ALTER TABLE Kit_Log ADD COLUMN kd_pelayanan TEXT DEFAULT ''")
    except: pass
    try: cur.execute("ALTER TABLE Kit_Log ADD COLUMN kd_pasien TEXT DEFAULT ''")
    except: pass
    try: cur.execute("ALTER TABLE Kit_Log ADD COLUMN nm_pasien TEXT DEFAULT ''")
    except: pass
    cur.execute("CREATE INDEX IF NOT EXISTS idx_kit_log ON Kit_Log(kd_karyawan, tanggal)")
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Kit_Request (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            no_request  TEXT UNIQUE NOT NULL,
            kd_karyawan TEXT NOT NULL,
            kd_lokasi   TEXT DEFAULT 'CBG1',   -- cabang sumber tarik stok
            tgl_request TEXT DEFAULT (date('now')),
            status      TEXT DEFAULT 'pending', -- pending | approved | rejected
            catatan     TEXT DEFAULT '',
            alasan      TEXT DEFAULT '',        -- alasan tolak
            created_at  TEXT DEFAULT (datetime('now')),
            approved_by TEXT DEFAULT '',
            approved_at TEXT DEFAULT ''
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Kit_Request_Detail (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            id_request  INTEGER NOT NULL,
            kd_barang   TEXT NOT NULL,
            nm_barang   TEXT DEFAULT '',
            qty         REAL NOT NULL,
            satuan      TEXT DEFAULT '',
            FOREIGN KEY (id_request) REFERENCES Kit_Request(id)
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_kit_request ON Kit_Request(status, kd_karyawan)")
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Pembelian_Koreksi (
            id           INTEGER PRIMARY KEY AUTOINCREMENT,
            no_faktur    TEXT NOT NULL,
            id_item      INTEGER,              -- NULL = hapus SELURUH faktur; terisi = hapus 1 item R_Pembelian
            jenis        TEXT NOT NULL,         -- hapus_item | hapus_faktur
            catatan      TEXT DEFAULT '',       -- alasan pengaju (wajib diisi, divalidasi di endpoint)
            status       TEXT DEFAULT 'pending',-- pending | approved | rejected
            requested_by TEXT DEFAULT '',
            requested_at TEXT DEFAULT (datetime('now')),
            approved_by  TEXT DEFAULT '',
            approved_at  TEXT DEFAULT '',
            alasan       TEXT DEFAULT ''        -- alasan admin kalau tolak
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Retur_Request (
            id             INTEGER PRIMARY KEY AUTOINCREMENT,
            no_faktur      TEXT NOT NULL,
            items_json     TEXT NOT NULL,        -- [{id, qty_retur}] baris R_Penjualan yg mau diretur
            kasir          TEXT DEFAULT '',
            status         TEXT DEFAULT 'pending', -- pending | disetujui | ditolak
            diajukan_oleh  TEXT DEFAULT '',
            diajukan_at    TEXT DEFAULT (datetime('now')),
            diproses_oleh  TEXT DEFAULT '',
            diproses_at    TEXT DEFAULT '',
            catatan_tolak  TEXT DEFAULT '',
            total_retur    REAL DEFAULT 0        -- terisi setelah disetujui & benar2 dieksekusi
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_retur_request_status ON Retur_Request(status)")

    # ── Produksi (batch run) ──────────────────────────────────────
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Produksi (
            id              INTEGER PRIMARY KEY AUTOINCREMENT,
            no_produksi     TEXT UNIQUE NOT NULL,
            id_formula      INTEGER NOT NULL,
            qty_batch       REAL NOT NULL DEFAULT 1,
            tgl_produksi    TEXT NOT NULL DEFAULT (date('now')),
            tgl_expired     TEXT NOT NULL DEFAULT '',
            status          TEXT DEFAULT 'draft',
            approved_by     TEXT DEFAULT '',
            approved_at     TEXT DEFAULT '',
            selesai_by      TEXT DEFAULT '',
            selesai_at      TEXT DEFAULT '',
            keterangan      TEXT DEFAULT '',
            created_by      TEXT DEFAULT '',
            created_at      TEXT DEFAULT (datetime('now')),
            FOREIGN KEY (id_formula) REFERENCES Formula_Produksi(id)
        )
    """)

    # ── Order Transfer (cabang → induk) ───────────────────────────
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Order_Transfer (
            id                      INTEGER PRIMARY KEY AUTOINCREMENT,
            no_order                TEXT UNIQUE NOT NULL,
            kd_lokasi_asal          TEXT NOT NULL,
            tgl_order               TEXT DEFAULT (date('now')),
            status                  TEXT DEFAULT 'draft',
            catatan_cabang          TEXT DEFAULT '',
            catatan_supervisor      TEXT DEFAULT '',
            approved_supervisor_by  TEXT DEFAULT '',
            approved_supervisor_at  TEXT DEFAULT '',
            fulfilled_by            TEXT DEFAULT '',
            fulfilled_at            TEXT DEFAULT '',
            is_from_import          INTEGER DEFAULT 0,
            created_by              TEXT DEFAULT '',
            created_at              TEXT DEFAULT (datetime('now'))
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Order_Transfer_Detail (
            id              INTEGER PRIMARY KEY AUTOINCREMENT,
            id_order        INTEGER NOT NULL,
            kd_barang       TEXT NOT NULL,
            nm_barang       TEXT DEFAULT '',
            qty_request     REAL NOT NULL,
            qty_approved    REAL DEFAULT NULL,
            qty_fulfilled   REAL DEFAULT 0,
            status_item     TEXT DEFAULT 'pending',
            catatan         TEXT DEFAULT '',
            FOREIGN KEY (id_order) REFERENCES Order_Transfer(id)
        )
    """)

    # ── Transfer Stok Log (audit trail pergerakan stok) ────────────
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Transfer_Stok_Log (
            id              INTEGER PRIMARY KEY AUTOINCREMENT,
            no_referensi    TEXT NOT NULL,
            tipe_referensi  TEXT NOT NULL,
            kd_barang       TEXT NOT NULL,
            nm_barang       TEXT DEFAULT '',
            kd_lokasi_dari  TEXT DEFAULT '',
            kd_lokasi_ke    TEXT NOT NULL,
            qty             REAL NOT NULL,
            exp             TEXT DEFAULT '',
            tgl             TEXT DEFAULT (datetime('now')),
            created_by      TEXT DEFAULT ''
        )
    """)

    # ── Master Barang (katalog produk — sumber kode konsisten) ─────
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Barang (
            kd_barang   TEXT PRIMARY KEY,
            nm_barang   TEXT NOT NULL,
            nm_generik  TEXT DEFAULT '',
            jenis       TEXT DEFAULT '',
            merk        TEXT DEFAULT '',
            isi         REAL DEFAULT 1,
            satuan      TEXT DEFAULT 'Pcs',
            harga_jual  REAL DEFAULT 0,
            aktif       INTEGER DEFAULT 1,
            non_obat    INTEGER DEFAULT 0,
            created_at  TEXT DEFAULT (datetime('now'))
        )
    """)
    _seed_barang_master(cur)

    # Seed akun dokter & staff
    cur.execute("INSERT OR IGNORE INTO Pengguna (username, password, nama, role) VALUES ('dokter', 'dokter123', 'Dokter', 'dokter')")
    # Test accounts
    for _u in [
        ('kasir1',      'kasir123',    'Kasir 1',       'kasir',       'CBG1', ''),
        ('supervisor1', 'super123',    'Supervisor',    'supervisor',  'CBG1', ''),
        ('induk',       'induk123',    'Gudang Induk',  'gudang_induk','INDUK',''),
        ('sari',        'sari123',     'Sari',          'perawat',     'CBG1', 'K003'),
        ('bunga',       'bunga123',    'Bunga',         'beautician',  'CBG1', 'K004'),
    ]:
        cur.execute("INSERT OR IGNORE INTO Pengguna (username,password,nama,role,kd_lokasi,kd_karyawan) VALUES (?,?,?,?,?,?)", _u)

    # Seed data contoh
    cur.execute("SELECT COUNT(*) FROM Pasien")
    if cur.fetchone()[0] == 0:
        sample_pasien = [
            ('P001', 'Siti Rahayu', 'Jl. Merdeka No.1', '08111111111', '1990-05-15', 'P'),
            ('P002', 'Dewi Kusuma', 'Jl. Sudirman No.2', '08222222222', '1985-08-20', 'P'),
            ('P003', 'Rina Marlina', 'Jl. Imam Bonjol No.3', '08333333333', '1995-03-10', 'P'),
        ]
        cur.executemany("INSERT OR IGNORE INTO Pasien (kd_pasien,nm_pasien,alamat,hp,tgl_lahir,jenis_kelamin,created_at) VALUES (?,?,?,?,?,?,date('now'))", sample_pasien)

        cur.execute("INSERT OR IGNORE INTO Member (kd_pasien, tgl_daftar, tgl_exp) VALUES ('P001','2024-01-01','2026-12-31')")

        sample_karyawan = [
            ('K001', 'dr. Natalia', 'Nat', 'Dokter', '081234567890'),
            ('K002', 'dr. Jeanette R', 'Jean', 'Dokter', '082345678901'),
            ('K003', 'Sari', 'Sari', 'Perawat', '083456789012'),
            ('K004', 'Bunga', 'Bunga', 'Beautition', '084567890123'),
        ]
        cur.executemany("INSERT OR IGNORE INTO Karyawan VALUES (?,?,?,?,?,1)", sample_karyawan)

        sample_barang = [
            ('B001', 'Krim Pemutih Wajah', 45000, 45000, 100, 80, 100, 5, 120000, '10', 'Pcs', '2026-12-31', 10, 10, 'INDUK'),
            ('B002', 'Serum Vitamin C', 75000, 75000, 50, 40, 50, 3, 180000, '0', 'Botol', '2026-06-30', 5, 5, 'INDUK'),
            ('B003', 'Sunscreen SPF50', 55000, 55000, 60, 55, 60, 5, 150000, '5', 'Pcs', '2026-09-30', 0, 0, 'INDUK'),
            ('B004', 'Masker Collagen', 30000, 30000, 200, 180, 200, 10, 85000, '0', 'Pcs', '2025-12-31', 0, 0, 'INDUK'),
        ]
        cur.executemany("INSERT OR IGNORE INTO Stok_barang (kd_barang,nm_barang,harga_beli,harga_beli_awal,stok_awal,stok_sisa,stok_gudang,stok_min,harga,disc,satuan,exp,DiscMem,DiscLain,kd_lokasi,pemilik) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,'P1')", sample_barang)

        sample_jasa = [
            ('J001', 'Facial Basic', 150000, 30000, 20000, 15000),
            ('J002', 'Laser Treatment', 500000, 100000, 75000, 50000),
            ('J003', 'Peeling Wajah', 250000, 50000, 40000, 30000),
            ('J004', 'Microdermabrasi', 350000, 70000, 55000, 40000),
            ('J005', 'Konsultasi Dokter', 100000, 50000, 0, 0),
        ]
        cur.executemany("INSERT OR IGNORE INTO produk_pelayanan (kd_pelayanan,nm_pelayanan,harga,komisi_dokter,komisi_supervisor,komisi_perawat) VALUES (?,?,?,?,?,?)", sample_jasa)

        cur.execute("INSERT OR IGNORE INTO Suplier VALUES ('SUP001','PT Pharma Utama','Jl. Industri No.10','021-5551234','Budi')")
        cur.execute("INSERT OR IGNORE INTO Suplier VALUES ('SUP002','CV Kosmetik Jaya','Jl. Raya No.5','021-5559876','Ani')")

    # ── Pengaturan sistem (key-value) ───────────────────────────
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Pengaturan (
            key        TEXT PRIMARY KEY,
            value      TEXT,
            updated_at TEXT DEFAULT (datetime('now'))
        )
    """)
    # Counter faktur per (kode terminal, tanggal) — anti-tabrakan multi-kasir
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Counter_Faktur (
            kode         TEXT NOT NULL,
            tanggal      TEXT NOT NULL,
            counter      INTEGER DEFAULT 0,
            last_faktur  TEXT DEFAULT '',
            PRIMARY KEY (kode, tanggal)
        )
    """)
    for _k, _v in [
        ('point_rate',   '50000'),
        ('faktur_prefix',''),
        ('klinik_nama',  'Rachel Clinic'),
        ('klinik_sub',   'Klinik Spesialis Kulit & Laser'),
        ('klinik_alamat','Jl. P. Diponegoro No. 92 Sumur Batu, Bandar Lampung'),
        ('klinik_telp',  'Telp. 0721-475600, 475601'),
        ('klinik_kota',  'Bandar Lampung'),
    ]:
        cur.execute("INSERT OR IGNORE INTO Pengaturan (key,value) VALUES (?,?)", (_k, _v))

    # ── Terima Kiriman Stok dari gedung lain (import file _P/_R per faktur) ────
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Terima_Kiriman (
            id                INTEGER PRIMARY KEY AUTOINCREMENT,
            no_faktur         TEXT NOT NULL UNIQUE,
            tgl_faktur        TEXT DEFAULT '',
            kd_lokasi_tujuan  TEXT DEFAULT 'CBG1',
            total_header      REAL DEFAULT 0,
            total_rincian     REAL DEFAULT 0,
            disc_pct          REAL DEFAULT 0,
            total_setelah_disc REAL DEFAULT 0,
            selisih           REAL DEFAULT 0,
            metode            TEXT DEFAULT '',
            sumber            TEXT DEFAULT '',
            status            TEXT DEFAULT 'pending_review',
            catatan           TEXT DEFAULT '',
            dibuat_oleh       TEXT DEFAULT '',
            created_at        TEXT DEFAULT (datetime('now')),
            diterima_at       TEXT DEFAULT '',
            induk_id          INTEGER DEFAULT NULL
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Terima_Kiriman_Detail (
            id                INTEGER PRIMARY KEY AUTOINCREMENT,
            id_kiriman        INTEGER NOT NULL,
            kd_barang         TEXT DEFAULT '',
            nm_file           TEXT DEFAULT '',
            nm_generik        TEXT DEFAULT '',
            harga             REAL DEFAULT 0,
            qty               REAL DEFAULT 0,
            exp               TEXT DEFAULT '',
            match_status      TEXT DEFAULT 'ok',
            expired           INTEGER DEFAULT 0,
            diterima          INTEGER DEFAULT 1,
            kd_barang_final   TEXT DEFAULT '',
            catatan_review    TEXT DEFAULT '',
            FOREIGN KEY (id_kiriman) REFERENCES Terima_Kiriman(id)
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_terima_kiriman_status ON Terima_Kiriman(status)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_terima_kiriman_detail ON Terima_Kiriman_Detail(id_kiriman)")

    # ── Pembayaran tersplit (multi-metode per faktur, mis. sebagian Tunai + sebagian
    # Transfer) — 1 baris per metode, SELALU diisi (bukan cuma kalau split) supaya semua
    # transaksi baru punya data terstruktur akurat. Faktur lama (sebelum fitur ini) tidak
    # punya baris di sini — pembaca WAJIB fallback ke Penjualan.metode_bayar (lihat
    # _bayar_lines_map()). 2026-10-25.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS Penjualan_Pembayaran (
            id         INTEGER PRIMARY KEY AUTOINCREMENT,
            no_faktur  TEXT NOT NULL,
            metode     TEXT NOT NULL DEFAULT 'Tunai',
            bank       TEXT DEFAULT '',
            nominal    REAL NOT NULL DEFAULT 0
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_penjualan_pembayaran_faktur ON Penjualan_Pembayaran(no_faktur)")

    con.commit()
    con.close()


# ─── INIT DATABASE (MySQL) ───────────────────────────────────
# Skema terpisah dari init_db_sqlite() (bukan hasil translasi otomatis) karena
# tipe kolom MySQL berbeda (VARCHAR wajib dibatasi utk PK/index, AUTOINCREMENT→
# AUTO_INCREMENT, dst). Setiap tabel sudah memuat kolom hasil "ALTER TABLE ADD
# COLUMN" historis di init_db_sqlite() (skema baru = bentuk akhir, bukan hasil
# evolusi bertahap) — urutan kolom disamakan dgn skema SQLite krn beberapa INSERT
# di _seed_common_data() pakai VALUES(...) tanpa daftar kolom eksplisit.
_MYSQL_TABLES = [
("Pasien", """
    CREATE TABLE IF NOT EXISTS Pasien (
        kd_pasien VARCHAR(64) PRIMARY KEY,
        nm_pasien VARCHAR(255) NOT NULL,
        alamat TEXT,
        telp VARCHAR(64),
        hp VARCHAR(64),
        tgl_lahir VARCHAR(32),
        jenis_kelamin VARCHAR(20),
        created_at VARCHAR(32) DEFAULT (CURDATE())
    )
"""),
("Member", """
    CREATE TABLE IF NOT EXISTS Member (
        id INT AUTO_INCREMENT PRIMARY KEY,
        kd_pasien VARCHAR(64) NOT NULL,
        tgl_daftar VARCHAR(32),
        tgl_exp VARCHAR(32),
        FOREIGN KEY (kd_pasien) REFERENCES Pasien(kd_pasien)
    )
"""),
("DetailPoint", """
    CREATE TABLE IF NOT EXISTS DetailPoint (
        id INT AUTO_INCREMENT PRIMARY KEY,
        kd_pasien VARCHAR(64) NOT NULL,
        keterangan VARCHAR(255),
        debit DOUBLE DEFAULT 0,
        kredit DOUBLE DEFAULT 0,
        tanggal VARCHAR(32) DEFAULT (CURDATE())
    )
"""),
("Karyawan", """
    CREATE TABLE IF NOT EXISTS Karyawan (
        kd_karyawan VARCHAR(64) PRIMARY KEY,
        nm_karyawan VARCHAR(255) NOT NULL,
        nickname VARCHAR(128),
        pekerjaan VARCHAR(64),
        telp VARCHAR(64),
        aktif INT DEFAULT 1
    )
"""),
("Stok_barang", """
    CREATE TABLE IF NOT EXISTS Stok_barang (
        id INT AUTO_INCREMENT PRIMARY KEY,
        kd_barang VARCHAR(64) COLLATE utf8mb4_bin NOT NULL,
        nm_barang VARCHAR(255) NOT NULL,
        harga_beli DOUBLE DEFAULT 0,
        harga_beli_awal DOUBLE DEFAULT 0,
        stok_awal INT DEFAULT 0,
        stok_sisa INT DEFAULT 0,
        stok_gudang INT DEFAULT 0,
        stok_min INT DEFAULT 0,
        harga DOUBLE DEFAULT 0,
        disc VARCHAR(16) DEFAULT '0',
        satuan VARCHAR(32),
        exp VARCHAR(32),
        DiscMem DOUBLE DEFAULT 0,
        DiscLain DOUBLE DEFAULT 0,
        kd_lokasi VARCHAR(16) DEFAULT 'CBG1',
        pemilik VARCHAR(16) DEFAULT ''
    )
"""),
("produk_pelayanan", """
    CREATE TABLE IF NOT EXISTS produk_pelayanan (
        kd_pelayanan VARCHAR(64) PRIMARY KEY,
        nm_pelayanan VARCHAR(255) NOT NULL,
        harga DOUBLE DEFAULT 0,
        komisi_dokter DOUBLE DEFAULT 0,
        komisi_supervisor DOUBLE DEFAULT 0,
        komisi_perawat DOUBLE DEFAULT 0,
        komisi_persen_kunci DOUBLE DEFAULT 0,
        aktif INT DEFAULT 1
    )
"""),
("Suplier", """
    CREATE TABLE IF NOT EXISTS Suplier (
        kd_suplier VARCHAR(64) PRIMARY KEY,
        nm_suplier VARCHAR(255) NOT NULL,
        alamat TEXT,
        telp VARCHAR(64),
        kontak VARCHAR(128)
    )
"""),
("Nomor", """
    CREATE TABLE IF NOT EXISTS Nomor (
        tanggal VARCHAR(32) PRIMARY KEY,
        nota_jual INT DEFAULT 0,
        nota_piutang INT DEFAULT 0,
        nota_resep INT DEFAULT 0
    )
"""),
("Antrian", """
    CREATE TABLE IF NOT EXISTS Antrian (
        id INT AUTO_INCREMENT PRIMARY KEY,
        kd_pasien VARCHAR(64),
        tanggal VARCHAR(32) DEFAULT (CURDATE()),
        status VARCHAR(32) DEFAULT 'wait',
        no_faktur VARCHAR(64),
        FOREIGN KEY (kd_pasien) REFERENCES Pasien(kd_pasien)
    )
"""),
("Penjualan", """
    CREATE TABLE IF NOT EXISTS Penjualan (
        no_faktur VARCHAR(64) PRIMARY KEY,
        tanggal VARCHAR(32),
        total_harga DOUBLE DEFAULT 0,
        hpp DOUBLE DEFAULT 0,
        dokter VARCHAR(255),
        kd_pasien VARCHAR(64),
        kasir VARCHAR(255),
        metode_bayar VARCHAR(32) DEFAULT 'Tunai',
        status VARCHAR(32) DEFAULT 'open',
        kd_lokasi VARCHAR(16) DEFAULT 'CBG1',
        disc_trx DOUBLE DEFAULT 0,
        disc_trx_nilai DOUBLE DEFAULT 0
    )
"""),
("R_Penjualan", """
    CREATE TABLE IF NOT EXISTS R_Penjualan (
        id INT AUTO_INCREMENT PRIMARY KEY,
        no_faktur VARCHAR(64) NOT NULL,
        tanggal VARCHAR(32),
        kd_barang VARCHAR(64) COLLATE utf8mb4_bin,
        nm_barang VARCHAR(255),
        harga DOUBLE DEFAULT 0,
        harga_asli DOUBLE DEFAULT 0,
        qty INT DEFAULT 1,
        disc VARCHAR(16) DEFAULT '0',
        sub_total DOUBLE DEFAULT 0,
        total_harbeli DOUBLE DEFAULT 0,
        kd_karyawan VARCHAR(64),
        cashier VARCHAR(255),
        exp VARCHAR(32),
        jenis VARCHAR(32) DEFAULT 'barang',
        racikan_ref INT DEFAULT NULL
    )
"""),
("TKredit", """
    CREATE TABLE IF NOT EXISTS TKredit (
        id INT AUTO_INCREMENT PRIMARY KEY,
        nota VARCHAR(64),
        konsumen VARCHAR(255),
        total DOUBLE DEFAULT 0,
        bank VARCHAR(64),
        kasir VARCHAR(255),
        tanggal VARCHAR(32) DEFAULT (CURDATE())
    )
"""),
("TKunjunganBerikut", """
    CREATE TABLE IF NOT EXISTS TKunjunganBerikut (
        id INT AUTO_INCREMENT PRIMARY KEY,
        no_faktur VARCHAR(64),
        next_tanggal VARCHAR(32),
        kd_pasien VARCHAR(64)
    )
"""),
("Point_karyawan", """
    CREATE TABLE IF NOT EXISTS Point_karyawan (
        id INT AUTO_INCREMENT PRIMARY KEY,
        no_faktur VARCHAR(64),
        tanggal VARCHAR(32),
        kd_pasien VARCHAR(64),
        kd_pelayanan VARCHAR(64),
        nm_pelayanan VARCHAR(255),
        komisi DOUBLE DEFAULT 0,
        kd_karyawan VARCHAR(64),
        status VARCHAR(16) DEFAULT 'released'
    )
"""),
("Tindakan", """
    CREATE TABLE IF NOT EXISTS Tindakan (
        id INT AUTO_INCREMENT PRIMARY KEY,
        tanggal VARCHAR(32),
        kd_pelayanan VARCHAR(64),
        banyaknya INT DEFAULT 0
    )
"""),
("Pembelian", """
    CREATE TABLE IF NOT EXISTS Pembelian (
        no_faktur VARCHAR(64) PRIMARY KEY,
        tanggal VARCHAR(32),
        total_harga DOUBLE DEFAULT 0,
        sistem_bayar VARCHAR(32),
        lunas VARCHAR(16) DEFAULT 'Belum',
        tempo VARCHAR(32),
        nm_suplier VARCHAR(255),
        kd_lokasi VARCHAR(16) DEFAULT 'INDUK',
        pemilik VARCHAR(16) DEFAULT '',
        status VARCHAR(16) DEFAULT 'approved',
        catatan_supervisor VARCHAR(500) DEFAULT '',
        approved_by VARCHAR(64) DEFAULT '',
        approved_at VARCHAR(32) DEFAULT '',
        dibuat_oleh VARCHAR(64) DEFAULT ''
    )
"""),
("R_Pembelian", """
    CREATE TABLE IF NOT EXISTS R_Pembelian (
        id INT AUTO_INCREMENT PRIMARY KEY,
        no_faktur VARCHAR(64),
        tanggal VARCHAR(32),
        kd_suplier VARCHAR(64),
        nm_suplier VARCHAR(255),
        sistem_bayar VARCHAR(32),
        kd_barang VARCHAR(64) COLLATE utf8mb4_bin,
        nm_barang VARCHAR(255),
        disc DOUBLE DEFAULT 0,
        harga_beli_kotor DOUBLE DEFAULT 0,
        harga_beli_bersih DOUBLE DEFAULT 0,
        sub_total DOUBLE DEFAULT 0,
        ppn DOUBLE DEFAULT 0,
        tempo VARCHAR(32),
        qty INT DEFAULT 0,
        harga_asli DOUBLE DEFAULT 0,
        satuan VARCHAR(32),
        exp VARCHAR(32)
    )
"""),
("RiwayatBarang", """
    CREATE TABLE IF NOT EXISTS RiwayatBarang (
        id INT AUTO_INCREMENT PRIMARY KEY,
        tanggal VARCHAR(32),
        kd_barang VARCHAR(64) COLLATE utf8mb4_bin,
        exp VARCHAR(32),
        keterangan VARCHAR(255),
        nilai INT DEFAULT 0,
        saldo_e INT DEFAULT 0,
        saldo_g INT DEFAULT 0
    )
"""),
("Log", """
    CREATE TABLE IF NOT EXISTS Log (
        id INT AUTO_INCREMENT PRIMARY KEY,
        tanggal VARCHAR(32) DEFAULT (CURDATE()),
        keterangan VARCHAR(255),
        qty INT DEFAULT 0,
        kasir VARCHAR(255),
        event VARCHAR(64)
    )
"""),
("Aktivitas_Log", """
    CREATE TABLE IF NOT EXISTS Aktivitas_Log (
        id INT AUTO_INCREMENT PRIMARY KEY,
        waktu VARCHAR(32) DEFAULT (CURRENT_TIMESTAMP),
        username VARCHAR(64) DEFAULT '',
        nama_user VARCHAR(255) DEFAULT '',
        role VARCHAR(32) DEFAULT '',
        kd_lokasi VARCHAR(16) DEFAULT '',
        aksi VARCHAR(64) DEFAULT '',
        entitas VARCHAR(64) DEFAULT '',
        ref VARCHAR(255) DEFAULT '',
        keterangan VARCHAR(500) DEFAULT '',
        data_lama TEXT,
        data_baru TEXT,
        ip_address VARCHAR(64) DEFAULT ''
    )
"""),
("RekamMedik", """
    CREATE TABLE IF NOT EXISTS RekamMedik (
        id INT AUTO_INCREMENT PRIMARY KEY,
        no_rm VARCHAR(64),
        kd_pasien VARCHAR(64) NOT NULL,
        tanggal VARCHAR(32) DEFAULT (CURDATE()),
        nm_dokter VARCHAR(255) DEFAULT '',
        status VARCHAR(32) DEFAULT 'open',
        S_keluhan TEXT,
        O_jenis_kulit TEXT,
        O_kondisi_kulit TEXT,
        A_diagnosa TEXT,
        P_tindakan TEXT,
        P_catatan TEXT,
        no_faktur VARCHAR(64) DEFAULT '',
        FOREIGN KEY (kd_pasien) REFERENCES Pasien(kd_pasien)
    )
"""),
("ResepRacikan", """
    CREATE TABLE IF NOT EXISTS ResepRacikan (
        id INT AUTO_INCREMENT PRIMARY KEY,
        id_rm INT NOT NULL,
        no_faktur VARCHAR(64) DEFAULT '',
        kd_pasien VARCHAR(64) DEFAULT '',
        nama_produk VARCHAR(255) NOT NULL,
        cara_pakai TEXT,
        harga_total DOUBLE DEFAULT 0,
        urutan INT DEFAULT 1,
        status VARCHAR(16) DEFAULT 'confirmed',
        created_by VARCHAR(64) DEFAULT '',
        catatan_kasir TEXT,
        status_racik VARCHAR(16) DEFAULT 'pending',
        FOREIGN KEY (id_rm) REFERENCES RekamMedik(id)
    )
"""),
("DetailResepRacikan", """
    CREATE TABLE IF NOT EXISTS DetailResepRacikan (
        id INT AUTO_INCREMENT PRIMARY KEY,
        id_resep INT NOT NULL,
        kd_barang VARCHAR(64) COLLATE utf8mb4_bin NOT NULL,
        nm_bahan VARCHAR(255) NOT NULL,
        qty DOUBLE DEFAULT 0,
        satuan VARCHAR(32) DEFAULT '',
        harga_satuan DOUBLE DEFAULT 0,
        disc VARCHAR(16) DEFAULT '0',
        sub_total DOUBLE DEFAULT 0,
        exp VARCHAR(32) DEFAULT '',
        FOREIGN KEY (id_resep) REFERENCES ResepRacikan(id)
    )
"""),
("Pengguna", """
    CREATE TABLE IF NOT EXISTS Pengguna (
        id INT AUTO_INCREMENT PRIMARY KEY,
        username VARCHAR(64) UNIQUE NOT NULL,
        password VARCHAR(255) NOT NULL,
        nama VARCHAR(255),
        role VARCHAR(32) DEFAULT 'kasir',
        aktif INT DEFAULT 1,
        kd_lokasi VARCHAR(16) DEFAULT 'CBG1',
        kd_karyawan VARCHAR(64) DEFAULT ''
    )
"""),
("ResepDigital", """
    CREATE TABLE IF NOT EXISTS ResepDigital (
        id INT AUTO_INCREMENT PRIMARY KEY,
        no_resep VARCHAR(64) NOT NULL,
        kd_pasien VARCHAR(64) NOT NULL,
        id_rm INT,
        nm_dokter VARCHAR(255) DEFAULT '',
        tanggal VARCHAR(32) NOT NULL,
        created_by VARCHAR(64) DEFAULT '',
        catatan TEXT,
        no_faktur VARCHAR(64) DEFAULT ''
    )
"""),
("DetailResepDigital", """
    CREATE TABLE IF NOT EXISTS DetailResepDigital (
        id INT AUTO_INCREMENT PRIMARY KEY,
        id_resep INT NOT NULL,
        nama_obat VARCHAR(255) NOT NULL,
        dosis VARCHAR(128) DEFAULT '',
        satuan VARCHAR(32) DEFAULT '',
        cara_pakai TEXT,
        urutan INT DEFAULT 1
    )
"""),
("Foto_Pasien", """
    CREATE TABLE IF NOT EXISTS Foto_Pasien (
        id INT AUTO_INCREMENT PRIMARY KEY,
        kd_pasien VARCHAR(64) NOT NULL,
        id_rm INT,
        no_rm VARCHAR(64) DEFAULT '',
        tanggal VARCHAR(32) DEFAULT (CURDATE()),
        tipe VARCHAR(32) DEFAULT 'sebelum',
        area VARCHAR(128) DEFAULT '',
        keterangan VARCHAR(255) DEFAULT '',
        filename VARCHAR(255) NOT NULL,
        created_at VARCHAR(32) DEFAULT (CURRENT_TIMESTAMP),
        uploaded_by VARCHAR(128) DEFAULT '',
        FOREIGN KEY (kd_pasien) REFERENCES Pasien(kd_pasien)
    )
"""),
("Lokasi", """
    CREATE TABLE IF NOT EXISTS Lokasi (
        kd_lokasi VARCHAR(16) PRIMARY KEY,
        nm_lokasi VARCHAR(255) NOT NULL,
        tipe VARCHAR(32) DEFAULT 'cabang',
        alamat VARCHAR(255) DEFAULT '',
        aktif INT DEFAULT 1
    )
"""),
("Induk_Pemilik", """
    CREATE TABLE IF NOT EXISTS Induk_Pemilik (
        kd_pemilik VARCHAR(16) PRIMARY KEY,
        nm_pemilik VARCHAR(255) NOT NULL,
        keterangan VARCHAR(255) DEFAULT '',
        aktif INT DEFAULT 1
    )
"""),
("Konversi_Stok", """
    CREATE TABLE IF NOT EXISTS Konversi_Stok (
        id INT AUTO_INCREMENT PRIMARY KEY,
        no_konversi VARCHAR(64),
        tgl VARCHAR(32),
        sumber_kd VARCHAR(64), sumber_nm VARCHAR(255), sumber_pemilik VARCHAR(16),
        qty_sumber DOUBLE, satuan_sumber VARCHAR(32), exp VARCHAR(32),
        target_kd VARCHAR(64), target_nm VARCHAR(255), target_pemilik VARCHAR(16),
        qty_target DOUBLE, satuan_target VARCHAR(32),
        harga_beli_sumber DOUBLE, harga_beli_target DOUBLE,
        created_by VARCHAR(64), created_at VARCHAR(32)
    )
"""),
("Formula_Produksi", """
    CREATE TABLE IF NOT EXISTS Formula_Produksi (
        id INT AUTO_INCREMENT PRIMARY KEY,
        kd_produk_hasil VARCHAR(64) NOT NULL,
        nm_produk_hasil VARCHAR(255) DEFAULT '',
        nm_formula VARCHAR(255) NOT NULL,
        qty_hasil DOUBLE NOT NULL DEFAULT 1,
        satuan_hasil VARCHAR(32) DEFAULT 'pcs',
        keterangan VARCHAR(255) DEFAULT '',
        aktif INT DEFAULT 1,
        created_at VARCHAR(32) DEFAULT (CURRENT_TIMESTAMP)
    )
"""),
("Formula_Produksi_Detail", """
    CREATE TABLE IF NOT EXISTS Formula_Produksi_Detail (
        id INT AUTO_INCREMENT PRIMARY KEY,
        id_formula INT NOT NULL,
        kd_barang VARCHAR(64) COLLATE utf8mb4_bin NOT NULL,
        nm_barang VARCHAR(255) DEFAULT '',
        qty DOUBLE NOT NULL,
        satuan VARCHAR(32) DEFAULT '',
        FOREIGN KEY (id_formula) REFERENCES Formula_Produksi(id)
    )
"""),
("Perawatan_Bahan", """
    CREATE TABLE IF NOT EXISTS Perawatan_Bahan (
        id INT AUTO_INCREMENT PRIMARY KEY,
        kd_pelayanan VARCHAR(64) NOT NULL,
        kd_barang VARCHAR(64) COLLATE utf8mb4_bin NOT NULL,
        nm_barang VARCHAR(255) DEFAULT '',
        qty DOUBLE NOT NULL DEFAULT 1,
        satuan VARCHAR(32) DEFAULT '',
        tipe VARCHAR(16) DEFAULT 'bahan',
        potong_stok INT DEFAULT 1,
        bahan_kunci INT DEFAULT 0,
        created_at VARCHAR(32) DEFAULT (CURRENT_TIMESTAMP)
    )
"""),
("Stok_Perawat", """
    CREATE TABLE IF NOT EXISTS Stok_Perawat (
        id INT AUTO_INCREMENT PRIMARY KEY,
        kd_karyawan VARCHAR(64) NOT NULL,
        kd_barang VARCHAR(64) COLLATE utf8mb4_bin NOT NULL,
        nm_barang VARCHAR(255) DEFAULT '',
        exp VARCHAR(32) DEFAULT '',
        qty_sisa DOUBLE NOT NULL DEFAULT 0,
        satuan VARCHAR(32) DEFAULT '',
        updated_at VARCHAR(32) DEFAULT (CURRENT_TIMESTAMP)
    )
"""),
("Kit_Log", """
    CREATE TABLE IF NOT EXISTS Kit_Log (
        id INT AUTO_INCREMENT PRIMARY KEY,
        tanggal VARCHAR(32) DEFAULT (CURRENT_TIMESTAMP),
        kd_karyawan VARCHAR(64) NOT NULL,
        kd_barang VARCHAR(64) COLLATE utf8mb4_bin NOT NULL,
        nm_barang VARCHAR(255) DEFAULT '',
        exp VARCHAR(32) DEFAULT '',
        qty DOUBLE NOT NULL,
        saldo DOUBLE DEFAULT 0,
        jenis VARCHAR(32) NOT NULL,
        ref VARCHAR(255) DEFAULT '',
        kd_pelayanan VARCHAR(64) DEFAULT '',
        kd_pasien VARCHAR(64) DEFAULT '',
        nm_pasien VARCHAR(255) DEFAULT '',
        kd_lokasi VARCHAR(16) DEFAULT '',
        operator VARCHAR(255) DEFAULT ''
    )
"""),
("Kit_Request", """
    CREATE TABLE IF NOT EXISTS Kit_Request (
        id INT AUTO_INCREMENT PRIMARY KEY,
        no_request VARCHAR(64) UNIQUE NOT NULL,
        kd_karyawan VARCHAR(64) NOT NULL,
        kd_lokasi VARCHAR(16) DEFAULT 'CBG1',
        tgl_request VARCHAR(32) DEFAULT (CURDATE()),
        status VARCHAR(32) DEFAULT 'pending',
        catatan VARCHAR(255) DEFAULT '',
        alasan VARCHAR(255) DEFAULT '',
        created_at VARCHAR(32) DEFAULT (CURRENT_TIMESTAMP),
        approved_by VARCHAR(64) DEFAULT '',
        approved_at VARCHAR(32) DEFAULT ''
    )
"""),
("Kit_Request_Detail", """
    CREATE TABLE IF NOT EXISTS Kit_Request_Detail (
        id INT AUTO_INCREMENT PRIMARY KEY,
        id_request INT NOT NULL,
        kd_barang VARCHAR(64) COLLATE utf8mb4_bin NOT NULL,
        nm_barang VARCHAR(255) DEFAULT '',
        qty DOUBLE NOT NULL,
        satuan VARCHAR(32) DEFAULT '',
        FOREIGN KEY (id_request) REFERENCES Kit_Request(id)
    )
"""),
("Serah_Terima_Kunci", """
    CREATE TABLE IF NOT EXISTS Serah_Terima_Kunci (
        id INT AUTO_INCREMENT PRIMARY KEY,
        no_serah VARCHAR(64) NOT NULL,
        no_faktur VARCHAR(64) NOT NULL,
        kd_pelayanan VARCHAR(64) NOT NULL,
        nm_pelayanan VARCHAR(255) DEFAULT '',
        kd_barang VARCHAR(64) COLLATE utf8mb4_bin NOT NULL,
        nm_barang VARCHAR(255) DEFAULT '',
        qty DOUBLE NOT NULL DEFAULT 1,
        satuan VARCHAR(32) DEFAULT '',
        kd_pelaksana VARCHAR(64) NOT NULL,
        nm_pelaksana VARCHAR(255) DEFAULT '',
        kd_k_dipilih VARCHAR(64) NOT NULL,
        kd_k_final VARCHAR(64) DEFAULT '',
        komisi_persen DOUBLE DEFAULT 0,
        status VARCHAR(16) DEFAULT 'pending',
        token VARCHAR(64) DEFAULT '',
        token_exp VARCHAR(32) DEFAULT '',
        kd_pasien VARCHAR(64) DEFAULT '',
        nm_pasien VARCHAR(255) DEFAULT '',
        selesai_at VARCHAR(32) DEFAULT '',
        created_at VARCHAR(32) DEFAULT (CURRENT_TIMESTAMP)
    )
"""),
("Pembelian_Koreksi", """
    CREATE TABLE IF NOT EXISTS Pembelian_Koreksi (
        id INT AUTO_INCREMENT PRIMARY KEY,
        no_faktur VARCHAR(64) NOT NULL,
        id_item INT,
        jenis VARCHAR(32) NOT NULL,
        catatan VARCHAR(255) DEFAULT '',
        status VARCHAR(32) DEFAULT 'pending',
        requested_by VARCHAR(64) DEFAULT '',
        requested_at VARCHAR(32) DEFAULT (CURRENT_TIMESTAMP),
        approved_by VARCHAR(64) DEFAULT '',
        approved_at VARCHAR(32) DEFAULT '',
        alasan VARCHAR(255) DEFAULT ''
    )
"""),
("Retur_Request", """
    CREATE TABLE IF NOT EXISTS Retur_Request (
        id INT AUTO_INCREMENT PRIMARY KEY,
        no_faktur VARCHAR(64) NOT NULL,
        items_json TEXT NOT NULL,
        kasir VARCHAR(64) DEFAULT '',
        status VARCHAR(32) DEFAULT 'pending',
        diajukan_oleh VARCHAR(64) DEFAULT '',
        diajukan_at VARCHAR(32) DEFAULT (CURRENT_TIMESTAMP),
        diproses_oleh VARCHAR(64) DEFAULT '',
        diproses_at VARCHAR(32) DEFAULT '',
        catatan_tolak VARCHAR(255) DEFAULT '',
        total_retur DOUBLE DEFAULT 0
    )
"""),
("Produksi", """
    CREATE TABLE IF NOT EXISTS Produksi (
        id INT AUTO_INCREMENT PRIMARY KEY,
        no_produksi VARCHAR(64) UNIQUE NOT NULL,
        id_formula INT NOT NULL,
        qty_batch DOUBLE NOT NULL DEFAULT 1,
        tgl_produksi VARCHAR(32) NOT NULL DEFAULT (CURDATE()),
        tgl_expired VARCHAR(32) NOT NULL DEFAULT '',
        status VARCHAR(32) DEFAULT 'draft',
        approved_by VARCHAR(64) DEFAULT '',
        approved_at VARCHAR(32) DEFAULT '',
        selesai_by VARCHAR(64) DEFAULT '',
        selesai_at VARCHAR(32) DEFAULT '',
        keterangan VARCHAR(255) DEFAULT '',
        created_by VARCHAR(64) DEFAULT '',
        created_at VARCHAR(32) DEFAULT (CURRENT_TIMESTAMP),
        FOREIGN KEY (id_formula) REFERENCES Formula_Produksi(id)
    )
"""),
("Order_Transfer", """
    CREATE TABLE IF NOT EXISTS Order_Transfer (
        id INT AUTO_INCREMENT PRIMARY KEY,
        no_order VARCHAR(64) UNIQUE NOT NULL,
        kd_lokasi_asal VARCHAR(16) NOT NULL,
        tgl_order VARCHAR(32) DEFAULT (CURDATE()),
        status VARCHAR(32) DEFAULT 'draft',
        catatan_cabang VARCHAR(255) DEFAULT '',
        catatan_supervisor VARCHAR(255) DEFAULT '',
        approved_supervisor_by VARCHAR(64) DEFAULT '',
        approved_supervisor_at VARCHAR(32) DEFAULT '',
        fulfilled_by VARCHAR(64) DEFAULT '',
        fulfilled_at VARCHAR(32) DEFAULT '',
        is_from_import INT DEFAULT 0,
        created_by VARCHAR(64) DEFAULT '',
        created_at VARCHAR(32) DEFAULT (CURRENT_TIMESTAMP)
    )
"""),
("Order_Transfer_Detail", """
    CREATE TABLE IF NOT EXISTS Order_Transfer_Detail (
        id INT AUTO_INCREMENT PRIMARY KEY,
        id_order INT NOT NULL,
        kd_barang VARCHAR(64) COLLATE utf8mb4_bin NOT NULL,
        nm_barang VARCHAR(255) DEFAULT '',
        qty_request DOUBLE NOT NULL,
        qty_approved DOUBLE DEFAULT NULL,
        qty_fulfilled DOUBLE DEFAULT 0,
        status_item VARCHAR(32) DEFAULT 'pending',
        catatan VARCHAR(255) DEFAULT '',
        FOREIGN KEY (id_order) REFERENCES Order_Transfer(id)
    )
"""),
("Transfer_Stok_Log", """
    CREATE TABLE IF NOT EXISTS Transfer_Stok_Log (
        id INT AUTO_INCREMENT PRIMARY KEY,
        no_referensi VARCHAR(64) NOT NULL,
        tipe_referensi VARCHAR(32) NOT NULL,
        kd_barang VARCHAR(64) COLLATE utf8mb4_bin NOT NULL,
        nm_barang VARCHAR(255) DEFAULT '',
        kd_lokasi_dari VARCHAR(16) DEFAULT '',
        kd_lokasi_ke VARCHAR(16) NOT NULL,
        qty DOUBLE NOT NULL,
        exp VARCHAR(32) DEFAULT '',
        tgl VARCHAR(32) DEFAULT (CURRENT_TIMESTAMP),
        created_by VARCHAR(64) DEFAULT ''
    )
"""),
("Barang", """
    CREATE TABLE IF NOT EXISTS Barang (
        kd_barang VARCHAR(64) COLLATE utf8mb4_bin PRIMARY KEY,
        nm_barang VARCHAR(255) NOT NULL,
        nm_generik VARCHAR(255) DEFAULT '',
        jenis VARCHAR(64) DEFAULT '',
        merk VARCHAR(128) DEFAULT '',
        isi DOUBLE DEFAULT 1,
        satuan VARCHAR(32) DEFAULT 'Pcs',
        harga_jual DOUBLE DEFAULT 0,
        aktif INT DEFAULT 1,
        non_obat INT DEFAULT 0,
        created_at VARCHAR(32) DEFAULT (CURRENT_TIMESTAMP)
    )
"""),
("Pengaturan", """
    CREATE TABLE IF NOT EXISTS Pengaturan (
        `key` VARCHAR(64) PRIMARY KEY,
        value TEXT,
        updated_at VARCHAR(32) DEFAULT (CURRENT_TIMESTAMP)
    )
"""),
("Counter_Faktur", """
    CREATE TABLE IF NOT EXISTS Counter_Faktur (
        kode VARCHAR(64) NOT NULL,
        tanggal VARCHAR(32) NOT NULL,
        counter INT DEFAULT 0,
        last_faktur VARCHAR(64) DEFAULT '',
        PRIMARY KEY (kode, tanggal)
    )
"""),
("Terima_Kiriman", """
    CREATE TABLE IF NOT EXISTS Terima_Kiriman (
        id INT AUTO_INCREMENT PRIMARY KEY,
        no_faktur VARCHAR(64) NOT NULL UNIQUE,
        tgl_faktur VARCHAR(32) DEFAULT '',
        kd_lokasi_tujuan VARCHAR(16) DEFAULT 'CBG1',
        total_header DOUBLE DEFAULT 0,
        total_rincian DOUBLE DEFAULT 0,
        disc_pct DOUBLE DEFAULT 0,
        total_setelah_disc DOUBLE DEFAULT 0,
        selisih DOUBLE DEFAULT 0,
        metode VARCHAR(64) DEFAULT '',
        sumber VARCHAR(128) DEFAULT '',
        status VARCHAR(24) DEFAULT 'pending_review',
        catatan TEXT,
        dibuat_oleh VARCHAR(64) DEFAULT '',
        created_at VARCHAR(32) DEFAULT (CURRENT_TIMESTAMP),
        diterima_at VARCHAR(32) DEFAULT '',
        induk_id INT DEFAULT NULL
    )
"""),
("Terima_Kiriman_Detail", """
    CREATE TABLE IF NOT EXISTS Terima_Kiriman_Detail (
        id INT AUTO_INCREMENT PRIMARY KEY,
        id_kiriman INT NOT NULL,
        kd_barang VARCHAR(64) COLLATE utf8mb4_bin DEFAULT '',
        nm_file VARCHAR(255) DEFAULT '',
        nm_generik VARCHAR(255) DEFAULT '',
        harga DOUBLE DEFAULT 0,
        qty DOUBLE DEFAULT 0,
        exp VARCHAR(32) DEFAULT '',
        match_status VARCHAR(24) DEFAULT 'ok',
        expired INT DEFAULT 0,
        diterima INT DEFAULT 1,
        kd_barang_final VARCHAR(64) COLLATE utf8mb4_bin DEFAULT '',
        catatan_review TEXT
    )
"""),
("Penjualan_Pembayaran", """
    CREATE TABLE IF NOT EXISTS Penjualan_Pembayaran (
        id INT AUTO_INCREMENT PRIMARY KEY,
        no_faktur VARCHAR(64) NOT NULL,
        metode VARCHAR(32) NOT NULL DEFAULT 'Tunai',
        bank VARCHAR(32) DEFAULT '',
        nominal DOUBLE NOT NULL DEFAULT 0
    )
"""),
]

_MYSQL_INDEXES = [
    "CREATE INDEX idx_perawatan_bahan_kd ON Perawatan_Bahan(kd_pelayanan)",
    "CREATE INDEX idx_stok_perawat ON Stok_Perawat(kd_karyawan, kd_barang, exp)",
    "CREATE INDEX idx_kit_log ON Kit_Log(kd_karyawan, tanggal)",
    "CREATE INDEX idx_kit_request ON Kit_Request(status, kd_karyawan)",
    "CREATE INDEX idx_retur_request_status ON Retur_Request(status)",
    # R_Penjualan (1,17 juta baris) TIDAK punya index di no_faktur sebelumnya —
    # setiap "WHERE no_faktur=?" (dipakai di HAMPIR SEMUA route POS/retur/laporan)
    # jadi full table scan. Ditemukan saat investigasi mismatch total_harga (2026-08-28).
    "CREATE INDEX idx_r_penjualan_faktur ON R_Penjualan(no_faktur)",
    "CREATE INDEX idx_serah_terima_status ON Serah_Terima_Kunci(status)",
    "CREATE INDEX idx_serah_terima_faktur ON Serah_Terima_Kunci(no_faktur)",
    "CREATE INDEX idx_terima_kiriman_status ON Terima_Kiriman(status)",
    "CREATE INDEX idx_terima_kiriman_detail ON Terima_Kiriman_Detail(id_kiriman)",
    # Ditemukan 2026-09-14 saat investigasi keluhan "kasir harus standby" (lambat) —
    # Stok_barang (20rb+ baris) TIDAK PUNYA index sama sekali selain PK id, padahal
    # jadi tempat query TERSIBUK di seluruh app (tiap buka/cari POS, racikan, stok
    # cabang). Query utama filter dulu by (kd_lokasi, stok_sisa, harga) SEBELUM LIKE
    # kode/nama — index gabungan ini bikin MySQL bisa persempit ke ratusan baris dulu
    # drpd scan 20rb baris + subquery SUM/COUNT/MIN(exp) berulang per baris.
    # Terbukti /api/barang-jual (load awal layar POS) 1.3 detik → dites turun jauh
    # setelah index ini + idx_stok_kdbarang ditambahkan.
    "CREATE INDEX idx_stok_barang_lokasi ON Stok_barang(kd_lokasi, stok_sisa, harga)",
    "CREATE INDEX idx_stok_barang_kode ON Stok_barang(kd_barang)",
    # Penjualan (338rb+ baris) cuma ber-PK no_faktur — semua Laporan Penjualan/
    # Dashboard/Monitoring yg filter rentang tanggal ("WHERE tanggal BETWEEN ?
    # AND ?") full-scan seluruh tabel tiap dipanggil.
    "CREATE INDEX idx_penjualan_tanggal ON Penjualan(tanggal)",
    "CREATE INDEX idx_penjualan_kdpasien ON Penjualan(kd_pasien)",
    # R_Penjualan (1,17 juta baris) — sudah ada index no_faktur, tapi kartu stok/
    # riwayat per barang & rekap komisi per racikan/karyawan masih full-scan tanpa ini.
    "CREATE INDEX idx_r_penjualan_kdbarang ON R_Penjualan(kd_barang)",
    "CREATE INDEX idx_r_penjualan_racikanref ON R_Penjualan(racikan_ref)",
    # Point_karyawan (243rb+ baris) cuma ber-PK id — Laporan Komisi per karyawan
    # (query paling sering dibuka perawat/beautician) full-scan tanpa ini.
    "CREATE INDEX idx_point_karyawan_kdkaryawan ON Point_karyawan(kd_karyawan)",
    "CREATE INDEX idx_point_karyawan_tanggal ON Point_karyawan(tanggal)",
    "CREATE INDEX idx_point_karyawan_faktur ON Point_karyawan(no_faktur)",
    # Pasien (42rb+ baris) — pencarian nama pasien (dipakai tiap transaksi POS &
    # Rekam Medik) full-scan tanpa ini.
    "CREATE INDEX idx_pasien_nama ON Pasien(nm_pasien)",
    "CREATE INDEX idx_penjualan_pembayaran_faktur ON Penjualan_Pembayaran(no_faktur)",
]


def _seed_common_data(cur):
    """Seed data awal (admin, lokasi, pengaturan, akun contoh) — SQL ditulis gaya
    SQLite (?, INSERT OR IGNORE, date('now')); _translate_sql() menerjemahkan
    otomatis saat DB_BACKEND=='mysql', jadi fungsi ini dipakai kedua backend."""
    cur.execute("INSERT OR IGNORE INTO Pengguna (username, password, nama, role) VALUES ('admin', 'admin123', 'Administrator', 'admin')")
    cur.execute("INSERT OR IGNORE INTO Lokasi VALUES ('INDUK','Apotek Induk','induk','',1)")
    cur.execute("INSERT OR IGNORE INTO Lokasi VALUES ('CBG1','Klinik Rachel','cabang_dalam','',1)")
    cur.execute("INSERT OR IGNORE INTO Lokasi VALUES ('CBG2','KIBAR','cabang_dalam','',1)")
    cur.execute("INSERT OR IGNORE INTO Induk_Pemilik VALUES ('P1','Induk Utama','Pemilik default (data lama)',1)")
    _seed_barang_master(cur)

    cur.execute("INSERT OR IGNORE INTO Pengguna (username, password, nama, role) VALUES ('dokter', 'dokter123', 'Dokter', 'dokter')")
    for _u in [
        ('kasir1',      'kasir123',    'Kasir 1',       'kasir',       'CBG1', ''),
        ('supervisor1', 'super123',    'Supervisor',    'supervisor',  'CBG1', ''),
        ('induk',       'induk123',    'Gudang Induk',  'gudang_induk','INDUK',''),
        ('sari',        'sari123',     'Sari',          'perawat',     'CBG1', 'K003'),
        ('bunga',       'bunga123',    'Bunga',         'beautician',  'CBG1', 'K004'),
    ]:
        cur.execute("INSERT OR IGNORE INTO Pengguna (username,password,nama,role,kd_lokasi,kd_karyawan) VALUES (?,?,?,?,?,?)", _u)

    cur.execute("SELECT COUNT(*) FROM Pasien")
    if cur.fetchone()[0] == 0:
        sample_pasien = [
            ('P001', 'Siti Rahayu', 'Jl. Merdeka No.1', '08111111111', '1990-05-15', 'P'),
            ('P002', 'Dewi Kusuma', 'Jl. Sudirman No.2', '08222222222', '1985-08-20', 'P'),
            ('P003', 'Rina Marlina', 'Jl. Imam Bonjol No.3', '08333333333', '1995-03-10', 'P'),
        ]
        cur.executemany("INSERT OR IGNORE INTO Pasien (kd_pasien,nm_pasien,alamat,hp,tgl_lahir,jenis_kelamin,created_at) VALUES (?,?,?,?,?,?,date('now'))", sample_pasien)

        cur.execute("INSERT OR IGNORE INTO Member (kd_pasien, tgl_daftar, tgl_exp) VALUES ('P001','2024-01-01','2026-12-31')")

        sample_karyawan = [
            ('K001', 'dr. Natalia', 'Nat', 'Dokter', '081234567890'),
            ('K002', 'dr. Jeanette R', 'Jean', 'Dokter', '082345678901'),
            ('K003', 'Sari', 'Sari', 'Perawat', '083456789012'),
            ('K004', 'Bunga', 'Bunga', 'Beautition', '084567890123'),
        ]
        cur.executemany("INSERT OR IGNORE INTO Karyawan VALUES (?,?,?,?,?,1)", sample_karyawan)

        sample_barang = [
            ('B001', 'Krim Pemutih Wajah', 45000, 45000, 100, 80, 100, 5, 120000, '10', 'Pcs', '2026-12-31', 10, 10, 'INDUK'),
            ('B002', 'Serum Vitamin C', 75000, 75000, 50, 40, 50, 3, 180000, '0', 'Botol', '2026-06-30', 5, 5, 'INDUK'),
            ('B003', 'Sunscreen SPF50', 55000, 55000, 60, 55, 60, 5, 150000, '5', 'Pcs', '2026-09-30', 0, 0, 'INDUK'),
            ('B004', 'Masker Collagen', 30000, 30000, 200, 180, 200, 10, 85000, '0', 'Pcs', '2025-12-31', 0, 0, 'INDUK'),
        ]
        cur.executemany("INSERT OR IGNORE INTO Stok_barang (kd_barang,nm_barang,harga_beli,harga_beli_awal,stok_awal,stok_sisa,stok_gudang,stok_min,harga,disc,satuan,exp,DiscMem,DiscLain,kd_lokasi,pemilik) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,'P1')", sample_barang)

        sample_jasa = [
            ('J001', 'Facial Basic', 150000, 30000, 20000, 15000),
            ('J002', 'Laser Treatment', 500000, 100000, 75000, 50000),
            ('J003', 'Peeling Wajah', 250000, 50000, 40000, 30000),
            ('J004', 'Microdermabrasi', 350000, 70000, 55000, 40000),
            ('J005', 'Konsultasi Dokter', 100000, 50000, 0, 0),
        ]
        cur.executemany("INSERT OR IGNORE INTO produk_pelayanan (kd_pelayanan,nm_pelayanan,harga,komisi_dokter,komisi_supervisor,komisi_perawat) VALUES (?,?,?,?,?,?)", sample_jasa)

        cur.execute("INSERT OR IGNORE INTO Suplier VALUES ('SUP001','PT Pharma Utama','Jl. Industri No.10','021-5551234','Budi')")
        cur.execute("INSERT OR IGNORE INTO Suplier VALUES ('SUP002','CV Kosmetik Jaya','Jl. Raya No.5','021-5559876','Ani')")

    for _k, _v in [
        ('point_rate',   '50000'),
        ('faktur_prefix',''),
        ('klinik_nama',  'Rachel Clinic'),
        ('klinik_sub',   'Klinik Spesialis Kulit & Laser'),
        ('klinik_alamat','Jl. P. Diponegoro No. 92 Sumur Batu, Bandar Lampung'),
        ('klinik_telp',  'Telp. 0721-475600, 475601'),
        ('klinik_kota',  'Bandar Lampung'),
    ]:
        cur.execute("INSERT OR IGNORE INTO Pengaturan (`key`,value) VALUES (?,?)", (_k, _v))


def init_db_mysql():
    raw = _mysql_connect()
    cur = raw.cursor()
    for _name, _ddl in _MYSQL_TABLES:
        cur.execute(_ddl)
    for _idx_sql in _MYSQL_INDEXES:
        try:
            cur.execute(_idx_sql)
        except Exception:
            pass   # index sudah ada dari run sebelumnya
    # Kolom baru utk tabel yg SUDAH ada dari sebelumnya (CREATE TABLE IF NOT EXISTS
    # tidak retroaktif) — try/except per kolom krn MySQL versi lama tak dukung
    # "ADD COLUMN IF NOT EXISTS" secara konsisten.
    for _alter_sql in [
        "ALTER TABLE Pembelian ADD COLUMN status VARCHAR(16) DEFAULT 'approved'",
        "ALTER TABLE Pembelian ADD COLUMN catatan_supervisor VARCHAR(500) DEFAULT ''",
        "ALTER TABLE Pembelian ADD COLUMN approved_by VARCHAR(64) DEFAULT ''",
        "ALTER TABLE Pembelian ADD COLUMN approved_at VARCHAR(32) DEFAULT ''",
        "ALTER TABLE Pembelian ADD COLUMN dibuat_oleh VARCHAR(64) DEFAULT ''",
        "ALTER TABLE Perawatan_Bahan ADD COLUMN bahan_kunci INT DEFAULT 0",
        "ALTER TABLE ResepRacikan ADD COLUMN status_racik VARCHAR(16) DEFAULT 'pending'",
        "ALTER TABLE produk_pelayanan ADD COLUMN komisi_persen_kunci DOUBLE DEFAULT 0",
        "ALTER TABLE Point_karyawan ADD COLUMN status VARCHAR(16) DEFAULT 'released'",
        "ALTER TABLE Terima_Kiriman ADD COLUMN disc_pct DOUBLE DEFAULT 0",
        "ALTER TABLE Terima_Kiriman ADD COLUMN total_setelah_disc DOUBLE DEFAULT 0",
        "ALTER TABLE Terima_Kiriman ADD COLUMN induk_id INT DEFAULT NULL",
        "ALTER TABLE DetailResepRacikan ADD COLUMN disc VARCHAR(16) DEFAULT '0'",
        "ALTER TABLE Counter_Faktur ADD COLUMN last_faktur VARCHAR(64) DEFAULT ''",
    ]:
        try:
            cur.execute(_alter_sql)
        except Exception:
            pass   # kolom sudah ada dari run sebelumnya
    # Jaring pengaman: faktur lama HARUS dianggap sudah disetujui (stoknya sudah nyata).
    cur.execute("UPDATE Pembelian SET status='approved' WHERE status IS NULL OR status=''")
    cur.execute("UPDATE Point_karyawan SET status='released' WHERE status IS NULL OR status=''")
    _seed_common_data(cur)
    raw.commit()
    raw.close()

# ─── SERVE FRONTEND ──────────────────────────────────────────
@app.after_request
def _no_cache_html(response):
    """Matikan cache browser utk halaman HTML (index.html/mobile.html/dokter.html/dst).
    App ini dikembangkan aktif — file statis sering diedit lalu langsung dites ulang
    lewat refresh, tapi browser sering nyimpen versi lama di cache (apalagi krn
    send_from_directory kirim Last-Modified/ETag yg bisa bikin browser pede pakai
    cache tanpa nanya ulang ke server) sehingga perubahan kelihatan "belum kepakai"
    padahal file di server sudah benar. Tidak berlaku ke response non-HTML (JSON API,
    dll) supaya tidak mengganggu performa endpoint lain."""
    ctype = response.content_type or ''
    if ctype.startswith('text/html'):
        response.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
        response.headers['Pragma'] = 'no-cache'
        response.headers['Expires'] = '0'
    return response

@app.route('/')
def index():
    return send_from_directory('static', 'index.html')

@app.route('/dokter')
def dokter_page():
    return send_from_directory('static', 'dokter.html')

@app.route('/migrasi')
def migrasi_page():
    return send_from_directory('static', 'migrasi.html')

# ─── AUTH ─────────────────────────────────────────────────────
@app.route('/api/login', methods=['POST'])
def login():
    d = request.json
    row = query("SELECT * FROM Pengguna WHERE username=? AND password=? AND aktif=1",
                (d.get('username'), d.get('password')), one=True)
    if row:
        token = str(uuid.uuid4())
        user  = dict(row)
        active_sessions[token] = user
        catat_log('login', 'Pengguna', user.get('username',''), f"Login berhasil: {user.get('nama','')}", actor=user)
        return jsonify({'ok': True, 'user': user, 'token': token})
    catat_log('login_gagal', 'Pengguna', d.get('username',''), f"Percobaan login gagal utk username '{d.get('username','')}'",
              actor={'username': d.get('username',''), 'nama': '', 'role': '', 'kd_lokasi': ''})
    return jsonify({'ok': False, 'msg': 'Username/password salah'}), 401

@app.route('/api/logout', methods=['POST'])
def logout():
    token = request.headers.get('X-Token','') or request.headers.get('Authorization','').replace('Bearer ','')
    user = active_sessions.get(token)
    if user:
        catat_log('logout', 'Pengguna', user.get('username',''), f"Logout: {user.get('nama','')}", actor=user)
    active_sessions.pop(token, None)
    return jsonify({'ok': True})

@app.route('/api/me')
def api_me():
    """Validasi sesi + kembalikan data user terkini (mis. untuk segarkan kd_karyawan
    di mobile setelah reload). 401 bila sesi tak valid (mis. server habis restart)."""
    user = get_current_user()
    if not user:
        return jsonify({'ok': False, 'msg': 'Sesi tidak valid'}), 401
    return jsonify({'ok': True, 'user': user})

# ─── USER MANAGEMENT ──────────────────────────────────────────
@app.route('/api/pengguna')
def list_pengguna():
    rows = query("SELECT id,username,nama,role,aktif,kd_lokasi,kd_karyawan FROM Pengguna ORDER BY nama")
    return jsonify(rows_to_list(rows))

@app.route('/api/pengguna', methods=['POST'])
def create_pengguna():
    d = request.json
    if not d.get('username') or not d.get('password'):
        return jsonify({'ok': False, 'msg': 'Username dan password wajib'}), 400
    try:
        uid = execute(
            "INSERT INTO Pengguna (username,password,nama,role,kd_lokasi,kd_karyawan) VALUES (?,?,?,?,?,?)",
            (d['username'], d['password'], d.get('nama',''), d.get('role','kasir'),
             d.get('kd_lokasi','CBG1'), d.get('kd_karyawan','')))
        catat_log('user_create', 'Pengguna', d['username'], f"Buat user baru '{d['username']}' role={d.get('role','kasir')}",
                   data_baru={'username': d['username'], 'nama': d.get('nama',''), 'role': d.get('role','kasir'),
                              'kd_lokasi': d.get('kd_lokasi','CBG1')})
        return jsonify({'ok': True, 'id': uid})
    except Exception as e:
        return jsonify({'ok': False, 'msg': str(e)}), 400

@app.route('/api/pengguna/<int:uid>', methods=['PUT'])
def update_pengguna(uid):
    d = request.json
    lama = query("SELECT username,nama,role,aktif,kd_lokasi,kd_karyawan FROM Pengguna WHERE id=?", (uid,), one=True)
    execute("UPDATE Pengguna SET nama=?,role=?,aktif=?,kd_lokasi=?,kd_karyawan=? WHERE id=?",
            (d.get('nama',''), d.get('role','kasir'), int(d.get('aktif',1)),
             d.get('kd_lokasi','CBG1'), d.get('kd_karyawan',''), uid))
    catat_log('user_update', 'Pengguna', lama['username'] if lama else uid,
              f"Update user '{lama['username'] if lama else uid}'",
              data_lama=dict(lama) if lama else None,
              data_baru={'nama': d.get('nama',''), 'role': d.get('role','kasir'), 'aktif': int(d.get('aktif',1)),
                         'kd_lokasi': d.get('kd_lokasi','CBG1'), 'kd_karyawan': d.get('kd_karyawan','')})
    if d.get('password'):
        execute("UPDATE Pengguna SET password=? WHERE id=?", (d['password'], uid))
        catat_log('user_ubah_password', 'Pengguna', lama['username'] if lama else uid,
                  f"Password user '{lama['username'] if lama else uid}' diubah (tanpa mencatat isi password)")
    return jsonify({'ok': True})

@app.route('/api/pengguna/<int:uid>', methods=['DELETE'])
def deactivate_pengguna(uid):
    lama = query("SELECT username,nama FROM Pengguna WHERE id=?", (uid,), one=True)
    execute("UPDATE Pengguna SET aktif=0 WHERE id=?", (uid,))
    catat_log('user_nonaktifkan', 'Pengguna', lama['username'] if lama else uid,
              f"Nonaktifkan user '{lama['username'] if lama else uid}'")
    return jsonify({'ok': True})

# ─── DASHBOARD ───────────────────────────────────────────────
@app.route('/api/dashboard')
def dashboard():
    today = str(date.today())
    kemarin = str(date.today() - timedelta(days=1))
    awal_bulan = str(date.today().replace(day=1))
    # metode_bayar='Internal' = transaksi gratis (internal/testing) — nilai barang/jasa
    # TETAP tercatat apa adanya di struk & Penjualan.total_harga (utk audit stok/HPP),
    # tapi SENGAJA dikecualikan dari semua perhitungan OMZET (bukan transaksi riil).
    omzet = query("SELECT COALESCE(SUM(total_harga),0) as v FROM Penjualan WHERE tanggal=? AND status='done' AND metode_bayar!='Internal'", (today,), one=True)
    trx   = query("SELECT COUNT(*) as v FROM Penjualan WHERE tanggal=?", (today,), one=True)
    pasien= query("SELECT COUNT(*) as v FROM Pasien", one=True)
    stok_min = query("SELECT COUNT(*) as v FROM Stok_barang WHERE stok_sisa <= stok_min AND stok_min > 0", one=True)
    antrian  = query("SELECT COUNT(*) as v FROM Antrian WHERE tanggal=? AND status='wait'", (today,), one=True)
    sejak_7hari  = str(date.today() - timedelta(days=6))
    sejak_30hari = str(date.today() - timedelta(days=30))
    sejak_29hari = str(date.today() - timedelta(days=29))
    omzet7 = query("""
        SELECT tanggal, COALESCE(SUM(total_harga),0) as total
        FROM Penjualan WHERE tanggal >= ? AND status='done' AND metode_bayar!='Internal'
        GROUP BY tanggal ORDER BY tanggal
    """, (sejak_7hari,))
    top_jasa = query("""
        SELECT nm_barang, SUM(qty) as total
        FROM R_Penjualan WHERE tanggal >= ? AND jenis='jasa'
        GROUP BY nm_barang ORDER BY total DESC LIMIT 5
    """, (sejak_30hari,))

    # ── Ringkasan Hari Ini lengkap ──
    trx_selesai = query("SELECT COUNT(*) as v FROM Penjualan WHERE tanggal=? AND status='done'", (today,), one=True)
    trx_open    = query("SELECT COUNT(*) as v FROM Penjualan WHERE tanggal=? AND status='open'", (today,), one=True)
    retur_hari  = query("""SELECT COUNT(*) as cnt, COALESCE(SUM(sub_total),0) as nilai
                           FROM R_Penjualan WHERE tanggal=? AND jenis LIKE 'retur%'""", (today,), one=True)
    # Split-aware tunai/non-tunai (2026-10-25): faktur split (sebagian Tunai + sebagian
    # Transfer) dipecah per baris Penjualan_Pembayaran, bukan ditempel ke 1 metode saja.
    bayar_rows = query("""SELECT no_faktur, metode_bayar, total_harga FROM Penjualan
                          WHERE tanggal=? AND status='done' AND metode_bayar!='Internal'""", (today,))
    bayar_lines_map = _bayar_lines_map([r['no_faktur'] for r in bayar_rows])
    bs_tunai = bs_non_tunai = 0.0
    for r in bayar_rows:
        for ln in _bayar_lines_for_row(r['no_faktur'], r['metode_bayar'], r['total_harga'], bayar_lines_map):
            if (ln['metode'] or 'Tunai').strip() == 'Tunai':
                bs_tunai += ln['nominal']
            else:
                bs_non_tunai += ln['nominal']
    bayar_split = {'tunai': bs_tunai, 'non_tunai': bs_non_tunai}
    omzet_kmr = query("SELECT COALESCE(SUM(total_harga),0) as v FROM Penjualan WHERE tanggal=? AND status='done' AND metode_bayar!='Internal'", (kemarin,), one=True)

    # ── Grafik & tren tambahan ──
    omzet_bln = query("SELECT COALESCE(SUM(total_harga),0) as v FROM Penjualan WHERE tanggal >= ? AND status='done' AND metode_bayar!='Internal'", (awal_bulan,), one=True)
    omzet30 = query("""
        SELECT tanggal, COALESCE(SUM(total_harga),0) as total
        FROM Penjualan WHERE tanggal >= ? AND status='done' AND metode_bayar!='Internal'
        GROUP BY tanggal ORDER BY tanggal
    """, (sejak_29hari,))
    top_barang = query("""
        SELECT nm_barang, SUM(qty) as total
        FROM R_Penjualan WHERE tanggal >= ? AND jenis='barang'
        GROUP BY nm_barang ORDER BY total DESC LIMIT 5
    """, (sejak_30hari,))
    kasir_hari = query("""
        SELECT COALESCE(NULLIF(kasir,''),'—') as kasir, COUNT(*) as cnt,
               COALESCE(SUM(total_harga),0) as total
        FROM Penjualan WHERE tanggal=? AND status='done' AND metode_bayar!='Internal'
        GROUP BY kasir ORDER BY total DESC LIMIT 6
    """, (today,))

    return jsonify({
        'omzet_hari_ini': omzet['v'],
        'transaksi_hari_ini': trx['v'],
        'total_pasien': pasien['v'],
        'stok_warning': stok_min['v'],
        'antrian': antrian['v'],
        'omzet_7hari': rows_to_list(omzet7),
        'top_jasa': rows_to_list(top_jasa),
        'ringkasan_hari_ini': {
            'trx_selesai': trx_selesai['v'],
            'trx_open': trx_open['v'],
            'retur_cnt': retur_hari['cnt'],
            'retur_nilai': retur_hari['nilai'],
            'bayar_tunai': bayar_split['tunai'],
            'bayar_non_tunai': bayar_split['non_tunai'],
            'omzet_kemarin': omzet_kmr['v'],
        },
        'omzet_bulan_ini': omzet_bln['v'],
        'omzet_30hari': rows_to_list(omzet30),
        'top_barang': rows_to_list(top_barang),
        'kasir_hari_ini': rows_to_list(kasir_hari),
    })

# ─── MONITORING SUPERVISOR ────────────────────────────────────
@app.route('/api/monitor/realtime')
def monitor_realtime():
    today = str(date.today())
    now   = datetime.now().strftime('%H:%M:%S')

    active_rows = query("""
        SELECT p.no_faktur, p.kasir, p.total_harga,
               pa.nm_pasien,
               (SELECT COUNT(*) FROM R_Penjualan WHERE no_faktur=p.no_faktur AND jenis!='bahan_racikan') as jml_item
        FROM Penjualan p
        LEFT JOIN Pasien pa ON p.kd_pasien = pa.kd_pasien
        WHERE p.status='open' AND p.tanggal=?
        ORDER BY p.no_faktur DESC
    """, (today,))
    active_list = []
    for row in rows_to_list(active_rows):
        items = query("""
            SELECT nm_barang, jenis, qty, sub_total
            FROM R_Penjualan WHERE no_faktur=? AND jenis!='bahan_racikan'
            ORDER BY id
        """, (row['no_faktur'],))
        row['items'] = rows_to_list(items)
        active_list.append(row)

    antrian_rows = query("SELECT status, COUNT(*) as cnt FROM Antrian WHERE tanggal=? GROUP BY status", (today,))
    antrian = {r['status']: r['cnt'] for r in antrian_rows}

    # metode_bayar='Internal' = transaksi gratis — dikecualikan dari omzet (lihat
    # catatan sama di dashboard()/laporan_penjualan()); cnt/pasien_cnt tetap hitung
    # semua (aktivitas riil, cuma bukan pendapatan).
    sales = query("""
        SELECT COALESCE(SUM(CASE WHEN metode_bayar!='Internal' THEN total_harga ELSE 0 END),0) as total,
               COUNT(*) as cnt,
               COUNT(DISTINCT kd_pasien) as pasien_cnt
        FROM Penjualan WHERE tanggal=? AND status='done'
    """, (today,), one=True)
    sales = dict(sales)
    # Split-aware tunai/non-tunai (sama pola dgn dashboard()) — pecah per baris
    # Penjualan_Pembayaran, bukan ditempel ke 1 metode faktur saja.
    sales_bayar_rows = query("""SELECT no_faktur, metode_bayar, total_harga FROM Penjualan
                               WHERE tanggal=? AND status='done' AND metode_bayar!='Internal'""", (today,))
    sales_lines_map = _bayar_lines_map([r['no_faktur'] for r in sales_bayar_rows])
    sales['tunai'] = sales['non_tunai'] = 0.0
    for r in sales_bayar_rows:
        for ln in _bayar_lines_for_row(r['no_faktur'], r['metode_bayar'], r['total_harga'], sales_lines_map):
            if (ln['metode'] or 'Tunai').strip() == 'Tunai':
                sales['tunai'] += ln['nominal']
            else:
                sales['non_tunai'] += ln['nominal']
    # Point dipakai hari ini (2026-10-25) — utk monitoring/audit real-time, lihat
    # _point_dipakai_map(). Sengaja dihitung terpisah dari sales_bayar_rows di atas
    # (yg sudah per-BARIS pembayaran) krn DetailPoint-nya per-FAKTUR, bukan per-baris.
    sales_pm_fakturs = [r['no_faktur'] for r in sales_bayar_rows
                        if (r['metode_bayar'] or '').strip() == 'Point Member']
    sales['point_dipakai'] = sum(_point_dipakai_map(sales_pm_fakturs).values())

    alerts = query("""
        SELECT kd_barang, nm_barang, stok_sisa, stok_min, satuan
        FROM Stok_barang
        WHERE stok_min > 0 AND stok_sisa <= stok_min
        ORDER BY CAST(stok_sisa AS REAL)/stok_min ASC LIMIT 10
    """)

    recent_rows = query("""
        SELECT p.no_faktur, p.kasir, p.total_harga, p.metode_bayar,
               pa.nm_pasien,
               (SELECT COALESCE(SUM(-sub_total),0) FROM R_Penjualan
                WHERE no_faktur=p.no_faktur AND jenis LIKE 'retur%') as nilai_retur
        FROM Penjualan p
        LEFT JOIN Pasien pa ON p.kd_pasien=pa.kd_pasien
        WHERE p.tanggal=? AND p.status='done'
        ORDER BY p.no_faktur DESC LIMIT 12
    """, (today,))
    recent_rows_list = rows_to_list(recent_rows)
    recent_pm_fakturs = [r['no_faktur'] for r in recent_rows_list
                         if (r['metode_bayar'] or '').strip() == 'Point Member']
    recent_point_map = _point_dipakai_map(recent_pm_fakturs)
    recent_list = []
    for row in recent_rows_list:
        # Preview item SENGAJA cuma yg benar2 terjual (bukan baris retur) — kalau tidak,
        # item yg diretur muncul dobel (nama sama persis dgn baris asalnya) dan total
        # jadi Rp0/berkurang tanpa keterangan knp (bingung dikira transaksi aneh/error).
        preview = query("""
            SELECT nm_barang, qty FROM R_Penjualan
            WHERE no_faktur=? AND jenis!='bahan_racikan' AND jenis NOT LIKE 'retur%'
            ORDER BY id LIMIT 4
        """, (row['no_faktur'],))
        row['items_preview'] = rows_to_list(preview)
        row['ada_retur']  = (row.get('nilai_retur') or 0) > 0
        row['total_awal'] = (row.get('total_harga') or 0) + (row.get('nilai_retur') or 0)
        row['point_dipakai'] = recent_point_map.get(row['no_faktur'], 0) \
            if (row['metode_bayar'] or '').strip() == 'Point Member' else 0
        recent_list.append(row)

    top_today = query("""
        SELECT nm_barang, jenis, SUM(qty) as total_qty, SUM(sub_total) as total_nilai
        FROM R_Penjualan
        WHERE tanggal=? AND jenis NOT IN ('bahan_racikan') AND sub_total > 0
        GROUP BY kd_barang, nm_barang
        ORDER BY total_nilai DESC LIMIT 5
    """, (today,))

    # Order Transfer — pending & recent activity
    ot_status_rows = query("""
        SELECT status, COUNT(*) as cnt FROM Order_Transfer GROUP BY status
    """)
    ot_status = {r['status']: r['cnt'] for r in ot_status_rows}

    ot_pending = query("""
        SELECT o.id, o.no_order, o.kd_lokasi_asal, l.nm_lokasi, o.status,
               o.tgl_order, o.created_by, o.catatan_cabang,
               (SELECT COUNT(*) FROM Order_Transfer_Detail WHERE id_order=o.id) as jml_item,
               (SELECT COALESCE(SUM(qty_request),0) FROM Order_Transfer_Detail WHERE id_order=o.id) as total_qty_req
        FROM Order_Transfer o
        LEFT JOIN Lokasi l ON o.kd_lokasi_asal=l.kd_lokasi
        WHERE o.status IN ('draft','submitted','approved')
        ORDER BY CASE o.status WHEN 'submitted' THEN 1 WHEN 'approved' THEN 2 ELSE 3 END,
                 o.created_at DESC
        LIMIT 20
    """)

    tgl_7hari = str(date.today() - timedelta(days=7))
    ot_recent = query("""
        SELECT o.id, o.no_order, o.kd_lokasi_asal, l.nm_lokasi, o.status,
               o.tgl_order, o.fulfilled_by, o.fulfilled_at,
               o.approved_supervisor_by, o.catatan_supervisor,
               (SELECT COUNT(*) FROM Order_Transfer_Detail WHERE id_order=o.id) as jml_item,
               (SELECT COALESCE(SUM(qty_fulfilled),0) FROM Order_Transfer_Detail WHERE id_order=o.id) as total_qty_sent
        FROM Order_Transfer o
        LEFT JOIN Lokasi l ON o.kd_lokasi_asal=l.kd_lokasi
        WHERE o.status IN ('fulfilled','partial','batal')
          AND COALESCE(o.fulfilled_at, o.created_at) >= ?
        ORDER BY COALESCE(o.fulfilled_at, o.created_at) DESC
        LIMIT 10
    """, (tgl_7hari,))

    # ── Aktivitas Retur hari ini (apa yang dikembalikan) ──
    retur_sum = query("""
        SELECT COUNT(DISTINCT no_faktur) AS faktur_cnt,
               COUNT(*) AS item_cnt,
               COALESCE(SUM(-sub_total),0) AS total_nilai,
               COALESCE(SUM(qty),0) AS total_qty
        FROM R_Penjualan
        WHERE tanggal=? AND jenis LIKE 'retur%'
    """, (today,), one=True)

    retur_lines = query("""
        SELECT rp.id, rp.no_faktur, rp.nm_barang, rp.qty, rp.jenis,
               -rp.sub_total AS nilai, COALESCE(p.kasir,'—') AS kasir,
               COALESCE(pa.nm_pasien,'—') AS nm_pasien
        FROM R_Penjualan rp
        LEFT JOIN Penjualan p  ON rp.no_faktur = p.no_faktur
        LEFT JOIN Pasien    pa ON p.kd_pasien  = pa.kd_pasien
        WHERE rp.tanggal=? AND rp.jenis LIKE 'retur%'
        ORDER BY rp.id DESC
    """, (today,))
    retur_map, retur_order = {}, []
    for r in rows_to_list(retur_lines):
        f = r['no_faktur']
        if f not in retur_map:
            retur_map[f] = {'no_faktur': f, 'kasir': r['kasir'], 'nm_pasien': r['nm_pasien'],
                            'items': [], 'total_nilai': 0, 'total_qty': 0}
            retur_order.append(f)
        g = retur_map[f]
        g['items'].append({'nm_barang': r['nm_barang'], 'qty': r['qty'],
                           'jenis': r['jenis'], 'nilai': r['nilai']})
        g['total_nilai'] += r['nilai']
        g['total_qty']   += r['qty']
    retur_recent = [retur_map[f] for f in retur_order][:15]

    # ── Konversi / pemindahan stok induk (pengawasan keluar-masuk barang) ──
    konversi_sum = query("""
        SELECT COUNT(*) AS cnt,
               COALESCE(SUM(qty_sumber),0) AS qty_keluar,
               COALESCE(SUM(qty_target),0) AS qty_masuk
        FROM Konversi_Stok WHERE tgl=?
    """, (today,), one=True)

    konversi_recent = query("""
        SELECT k.no_konversi, k.tgl, k.created_at, k.created_by, k.exp,
               k.sumber_kd, k.sumber_nm, k.qty_sumber, k.satuan_sumber,
               k.target_kd, k.target_nm, k.qty_target, k.satuan_target,
               COALESCE(ps.nm_pemilik, k.sumber_pemilik) AS sumber_nm_pemilik,
               COALESCE(pt.nm_pemilik, k.target_pemilik) AS target_nm_pemilik
        FROM Konversi_Stok k
        LEFT JOIN Induk_Pemilik ps ON k.sumber_pemilik = ps.kd_pemilik
        LEFT JOIN Induk_Pemilik pt ON k.target_pemilik = pt.kd_pemilik
        WHERE k.tgl >= ?
        ORDER BY k.id DESC LIMIT 15
    """, (tgl_7hari,))

    # ── Retur Kit hari ini (bahan kembali ke kit perawat) ──
    kit_retur_sum = query("""
        SELECT COUNT(*) AS cnt, COALESCE(SUM(qty),0) AS total_qty,
               COUNT(DISTINCT ref) AS faktur_cnt
        FROM Kit_Log WHERE jenis='retur_treatment' AND date(tanggal)=?
    """, (today,), one=True)
    kit_retur_recent = query("""
        SELECT kl.tanggal, kl.kd_karyawan,
               COALESCE(k.nickname,k.nm_karyawan,kl.kd_karyawan) AS nama,
               kl.nm_barang, kl.kd_barang, kl.exp, kl.qty,
               kl.ref AS no_faktur, kl.nm_pasien, kl.operator
        FROM Kit_Log kl LEFT JOIN Karyawan k ON kl.kd_karyawan=k.kd_karyawan
        WHERE kl.jenis='retur_treatment' AND date(kl.tanggal)=?
        ORDER BY kl.id DESC LIMIT 15
    """, (today,))

    return jsonify({
        'active': active_list,
        'antrian': antrian,
        'sales': dict(sales) if sales else {},
        'alerts': rows_to_list(alerts),
        'recent': recent_list,
        'top_today': rows_to_list(top_today),
        'ot_status': ot_status,
        'ot_pending': rows_to_list(ot_pending),
        'ot_recent': rows_to_list(ot_recent),
        'retur': dict(retur_sum) if retur_sum else {},
        'retur_recent': retur_recent,
        'konversi': dict(konversi_sum) if konversi_sum else {},
        'konversi_recent': rows_to_list(konversi_recent),
        'kit_retur': dict(kit_retur_sum) if kit_retur_sum else {},
        'kit_retur_recent': rows_to_list(kit_retur_recent),
        'timestamp': now,
        'today': today,
    })

def _split_metode_bank(metode_bayar_str):
    """'Metode - Bank' → (metode, bank); bank='' kalau tak ada ' - '. Dipakai bareng
    dgn _bayar_lines_map() sbg fallback utk faktur yg tak punya baris di
    Penjualan_Pembayaran (transaksi lama, pra-fitur split pembayaran 2026-10-25)."""
    mb = (metode_bayar_str or 'Tunai').strip() or 'Tunai'
    if ' - ' in mb:
        metode, bank = mb.split(' - ', 1)
        return metode, bank
    return mb, ''

def _bayar_lines_map(no_fakturs):
    """Bulk-fetch {no_faktur: [{'metode','bank','nominal'}, ...]} dari
    Penjualan_Pembayaran utk list no_faktur (1 query, BUKAN N+1 per faktur).
    Faktur yg tak ada di hasil (dict tak py key-nya) = blm py baris split tercatat
    (transaksi lama) — pemanggil WAJIB fallback ke _split_metode_bank(metode_bayar)
    + nilai total faktur itu sbg 1 baris sintetis."""
    no_fakturs = [f for f in (no_fakturs or []) if f]
    if not no_fakturs:
        return {}
    ph = ','.join(['?'] * len(no_fakturs))
    rows = rows_to_list(query(
        f"SELECT no_faktur, metode, bank, nominal FROM Penjualan_Pembayaran WHERE no_faktur IN ({ph})",
        no_fakturs))
    m = {}
    for r in rows:
        m.setdefault(r['no_faktur'], []).append(r)
    return m

def _bayar_lines_for_row(no_faktur, metode_bayar, nilai, lines_map):
    """Baris pembayaran utk 1 faktur — dari lines_map (Penjualan_Pembayaran) kalau ada
    (sadar-split), else 1 baris sintetis dari metode_bayar+nilai (faktur lama)."""
    lines = lines_map.get(no_faktur)
    if lines:
        return lines
    metode, bank = _split_metode_bank(metode_bayar)
    return [{'metode': metode, 'bank': bank, 'nominal': nilai}]

def _point_dipakai_map(no_fakturs_point_member):
    """Bulk-fetch {no_faktur: jumlah_point} dari DetailPoint utk faktur yg metode
    bayarnya 'Point Member' (2026-10-25) — dipakai laporan utk audit/monitoring: bisa
    disilang-cek apakah point yg benar² terpotong dari saldo pasien (DetailPoint.kredit)
    konsisten dgn nilai Rupiah 'Point Member' di breakdown (bukan sekadar lihat Rupiah).
    PENTING: caller WAJIB hanya mengirim faktur yg metode_bayar='Point Member' — faktur
    LAIN yg diretur jg bisa punya baris DetailPoint.kredit dgn keterangan=no_faktur yg
    sama (itu klaim-balik point krn retur, bukan pembayaran-pakai-point), jadi tidak
    boleh dicampur di sini (lihat _eksekusi_retur() utk konteks klaim-balik itu)."""
    no_fakturs = [f for f in (no_fakturs_point_member or []) if f]
    if not no_fakturs:
        return {}
    ph = ','.join(['?'] * len(no_fakturs))
    rows = rows_to_list(query(
        f"SELECT keterangan, SUM(kredit) as pts FROM DetailPoint WHERE keterangan IN ({ph}) GROUP BY keterangan",
        no_fakturs))
    return {r['keterangan']: (r['pts'] or 0) for r in rows}

def _monitor_grup_mentah(tgl1, tgl2, kolom):
    """Ambil baris mentah utk pengelompokan per kasir/dokter (dipakai monitor_per_kasir(),
    monitor_per_dokter(), & sheet 'Per Dokter' laporan komprehensif) — SENGAJA ambil
    total_harga_stored + SUM R_Penjualan sekaligus (bukan GROUP BY langsung di SQL spt
    sebelumnya), krn sebagian faktur migrasi py Penjualan.total_harga=0 padahal item
    riil ada di R_Penjualan (kasus nyata: 55 faktur migrasi akhiran 'C' per audit
    2026-10-03) — kalau di-SUM mentah di SQL, kontribusi faktur2 itu ke omzet kasir/
    dokter jadi 0, understate total scr signifikan. Pola rekonsiliasi SAMA PERSIS
    laporan_penjualan(). `no_faktur` disertakan supaya _monitor_grup_hitung() bisa
    lookup _bayar_lines_map() utk breakdown tunai/non_tunai sadar-split (2026-10-25)."""
    kol_sql = 'kasir' if kolom == 'kasir' else "COALESCE(NULLIF(dokter,''),'')"
    raw = query(f"""
        SELECT {kol_sql} as grup, p.kd_pasien, p.metode_bayar, p.no_faktur,
               p.total_harga AS total_harga_stored,
               (SELECT SUM(rp.sub_total) FROM R_Penjualan rp
                WHERE rp.no_faktur=p.no_faktur AND rp.jenis!='bahan_racikan') AS total_item
        FROM Penjualan p
        WHERE p.tanggal >= ? AND p.tanggal <= ? AND p.status='done'
    """, (tgl1, tgl2))
    return rows_to_list(raw)

def _monitor_grup_hitung(raw, kolom, exclude_internal=False):
    """Kelompokkan baris mentah dari _monitor_grup_mentah() per kasir/dokter, dgn
    total_harga sudah direkonsiliasi (prefer SUM R_Penjualan drpd kolom tersimpan).
    exclude_internal=True utk laporan omzet resmi (konsisten dgn laporan_penjualan()/
    Ringkasan); False (default) utk tool monitoring operasional — 'Internal' tetap
    diikutkan spt keputusan desain sebelumnya. tunai/non_tunai dihitung per BARIS
    pembayaran (sadar pembayaran tersplit, 2026-10-25) — bukan per faktur utuh lagi,
    supaya faktur split (mis. separuh Tunai separuh Transfer) terbagi benar ke 2
    bucket, tidak nempel 100% ke 1 metode spt sebelumnya."""
    # Internal TIDAK bisa digabung split (keputusan desain fitur split, 2026-10-25) —
    # jadi cek exclude_internal tetap di level FAKTUR (metode_bayar), persis pola lama:
    # faktur Internal selalu py tepat 1 baris pembayaran (metode='Internal'), tak pernah
    # campur dgn metode lain, jadi skip di sini == skip semua baris turunannya jg.
    lines_map = _bayar_lines_map([r['no_faktur'] for r in raw])
    groups = {}
    for r in raw:
        if exclude_internal and (r['metode_bayar'] or '').strip() == 'Internal':
            continue
        nilai = r['total_item'] if r['total_item'] is not None else r['total_harga_stored']
        key = (r['grup'] or '').strip() or '—'
        g = groups.setdefault(key, {kolom: key, 'total_trx': 0, 'total_nilai': 0.0,
                                     '_pasien': set(), 'tunai': 0.0, 'non_tunai': 0.0})
        g['total_trx'] += 1
        g['_pasien'].add(r['kd_pasien'])
        for ln in _bayar_lines_for_row(r['no_faktur'], r['metode_bayar'], nilai, lines_map):
            g['total_nilai'] += ln['nominal']
            if (ln['metode'] or 'Tunai').strip() == 'Tunai':
                g['tunai'] += ln['nominal']
            else:
                g['non_tunai'] += ln['nominal']
    out = []
    for g in groups.values():
        g['total_pasien'] = len(g.pop('_pasien'))
        out.append(g)
    out.sort(key=lambda x: -x['total_nilai'])
    return out

@app.route('/api/monitor/per-kasir')
def monitor_per_kasir():
    dari   = request.args.get('dari', str(date.today()))
    sampai = request.args.get('sampai', str(date.today()))
    raw = _monitor_grup_mentah(dari, sampai, 'kasir')
    rows = _monitor_grup_hitung(raw, 'kasir')
    total = sum(r['total_nilai'] for r in rows)
    return jsonify({'rows': rows, 'total': total})

@app.route('/api/monitor/per-dokter')
def monitor_per_dokter():
    """Kinerja per dokter (Penjualan.dokter — nama dokter yg dipilih saat bayar,
    lihat validasi wajib di bayar_penjualan()), utk panel Monitoring > Analisis
    Periode. Pola SAMA PERSIS dgn monitor_per_kasir() (lihat _monitor_grup_mentah()/
    _monitor_grup_hitung()), cuma beda kolom pengelompokan — termasuk 'Internal'
    tetap ikut ke total_nilai spt di per-kasir, krn tool ini utk monitoring
    operasional, bukan laporan omzet resmi."""
    dari   = request.args.get('dari', str(date.today()))
    sampai = request.args.get('sampai', str(date.today()))
    raw = _monitor_grup_mentah(dari, sampai, 'dokter')
    rows = _monitor_grup_hitung(raw, 'dokter')
    total = sum(r['total_nilai'] for r in rows)
    return jsonify({'rows': rows, 'total': total})

@app.route('/api/monitor/per-hari')
def monitor_per_hari():
    dari   = request.args.get('dari', str(date.today()))
    sampai = request.args.get('sampai', str(date.today()))
    rows = query("""
        SELECT tanggal,
               COUNT(*) as total_trx,
               COALESCE(SUM(total_harga),0) as total_nilai,
               COUNT(DISTINCT kd_pasien) as total_pasien
        FROM Penjualan
        WHERE tanggal >= ? AND tanggal <= ? AND status='done'
        GROUP BY tanggal ORDER BY tanggal ASC
    """, (dari, sampai))
    return jsonify(rows_to_list(rows))

@app.route('/api/monitor/aktivitas')
def monitor_aktivitas():
    dari   = request.args.get('dari', str(date.today()))
    sampai = request.args.get('sampai', str(date.today()))

    jual = rows_to_list(query("""
        SELECT 'penjualan' as tipe, p.no_faktur as ref, p.tanggal,
               COALESCE(p.kasir,'—') as aktor,
               COALESCE(pa.nm_pasien,'—') as detail,
               p.total_harga as nilai, p.metode_bayar as keterangan,
               (SELECT COUNT(*) FROM R_Penjualan WHERE no_faktur=p.no_faktur
                AND jenis!='bahan_racikan') as jml_item
        FROM Penjualan p
        LEFT JOIN Pasien pa ON p.kd_pasien=pa.kd_pasien
        WHERE p.tanggal >= ? AND p.tanggal <= ? AND p.status='done'
        ORDER BY p.no_faktur DESC LIMIT 200
    """, (dari, sampai)))
    # Point dipakai (2026-10-25) — sama pola dgn monitor_realtime(), lihat
    # _point_dipakai_map(). Ditempel ke keterangan supaya langsung kelihatan di
    # feed aktivitas gabungan ini tanpa perlu buka Laporan Point terpisah.
    jual_pm_fakturs = [r['ref'] for r in jual if (r['keterangan'] or '').strip() == 'Point Member']
    jual_point_map = _point_dipakai_map(jual_pm_fakturs)
    for r in jual:
        if (r['keterangan'] or '').strip() == 'Point Member':
            pts = jual_point_map.get(r['ref'], 0)
            if pts:
                r['keterangan'] = f"Point Member — {int(pts)} point dipotong"

    beli = query("""
        SELECT 'pembelian' as tipe, no_faktur as ref, tanggal,
               COALESCE(nm_suplier,'—') as aktor,
               sistem_bayar as detail, total_harga as nilai,
               lunas as keterangan,
               (SELECT COUNT(*) FROM R_Pembelian WHERE no_faktur=pb.no_faktur) as jml_item
        FROM Pembelian pb
        WHERE tanggal >= ? AND tanggal <= ?
        ORDER BY tanggal DESC LIMIT 100
    """, (dari, sampai))

    retur = query("""
        SELECT 'retur' as tipe, rp.no_faktur as ref, rp.tanggal,
               COALESCE(p.kasir,'—') as aktor,
               COALESCE(pa.nm_pasien,'—') as detail,
               SUM(-rp.sub_total) as nilai,
               'pengembalian' as keterangan,
               COUNT(*) as jml_item
        FROM R_Penjualan rp
        LEFT JOIN Penjualan p  ON rp.no_faktur = p.no_faktur
        LEFT JOIN Pasien    pa ON p.kd_pasien  = pa.kd_pasien
        WHERE rp.tanggal >= ? AND rp.tanggal <= ? AND rp.jenis LIKE 'retur%'
        GROUP BY rp.no_faktur
        ORDER BY rp.tanggal DESC LIMIT 100
    """, (dari, sampai))

    hapus_raw = query("""
        SELECT 'hapus' as tipe, ref, waktu as tanggal,
               COALESCE(nama_user,'—') as aktor,
               data_lama,
               'item dihapus' as keterangan,
               0 as jml_item
        FROM Aktivitas_Log
        WHERE aksi='hapus_item_penjualan' AND waktu >= ? AND waktu < ?
        ORDER BY waktu DESC LIMIT 100
    """, (dari + ' 00:00:00', sampai + ' 23:59:59'))
    hapus = rows_to_list(hapus_raw)
    for h in hapus:
        try:
            dl = json.loads(h.pop('data_lama') or '{}')
        except Exception:
            dl = {}
            h.pop('data_lama', None)
        h['nilai'] = -float(dl.get('sub_total') or 0)
        qty = dl.get('qty')
        h['detail'] = (dl.get('nm_barang') or '?') \
            + (f" (qty {float(qty):g})" if qty else '') \
            + (f" — {dl['pasien']}" if dl.get('pasien') else '')
        if h.get('tanggal'):
            h['tanggal'] = h['tanggal'][:16]

    koreksi_raw = query("""
        SELECT 'koreksi' as tipe, ref, waktu as tanggal,
               COALESCE(nama_user,'—') as aktor,
               keterangan as detail,
               0 as nilai,
               '' as keterangan,
               0 as jml_item
        FROM Aktivitas_Log
        WHERE aksi='koreksi_kit' AND waktu >= ? AND waktu < ?
        ORDER BY waktu DESC LIMIT 100
    """, (dari + ' 00:00:00', sampai + ' 23:59:59'))
    koreksi = rows_to_list(koreksi_raw)
    for k in koreksi:
        if k.get('tanggal'):
            k['tanggal'] = k['tanggal'][:16]

    manual_raw = query("""
        SELECT 'manual' as tipe, ref, waktu as tanggal,
               COALESCE(nama_user,'—') as aktor,
               data_baru, keterangan,
               0 as jml_item
        FROM Aktivitas_Log
        WHERE aksi='input_stok_manual' AND waktu >= ? AND waktu < ?
        ORDER BY waktu DESC LIMIT 100
    """, (dari + ' 00:00:00', sampai + ' 23:59:59'))
    manual = rows_to_list(manual_raw)
    for m in manual:
        try:
            db = json.loads(m.pop('data_baru') or '{}')
        except Exception:
            db = {}
            m.pop('data_baru', None)
        qty = db.get('qty')
        m['nilai'] = float(db.get('harga_beli') or 0) * float(qty or 0)
        m['detail'] = (db.get('nm_barang') or m.get('ref') or '?') \
            + (f" (qty {float(qty):g} {db.get('satuan','')})" if qty else '') \
            + (f" @ {db['kd_lokasi']}" if db.get('kd_lokasi') else '')
        # Keterangan tampilkan alasan asli yg diketik admin (bukan teks generik) —
        # supaya langsung kelihatan "kenapa" tanpa perlu buka detail terpisah.
        m['keterangan'] = db.get('catatan') or m.get('keterangan') or ''
        if m.get('tanggal'):
            m['tanggal'] = m['tanggal'][:16]

    result = jual + rows_to_list(beli) + rows_to_list(retur) + hapus + koreksi + manual
    result.sort(key=lambda x: (x.get('tanggal',''), x.get('ref','')), reverse=True)
    return jsonify(result[:300])

@app.route('/api/monitor/top-produk')
def monitor_top_produk():
    dari   = request.args.get('dari', str(date.today()))
    sampai = request.args.get('sampai', str(date.today()))
    rows = query("""
        SELECT kd_barang, nm_barang, jenis,
               SUM(qty) as total_qty,
               SUM(sub_total) as total_nilai,
               COALESCE(SUM(total_harbeli),0) as total_hpp,
               SUM(sub_total) - COALESCE(SUM(total_harbeli),0) as gross_margin,
               COUNT(DISTINCT no_faktur) as total_trx
        FROM R_Penjualan
        WHERE tanggal >= ? AND tanggal <= ?
          AND jenis NOT IN ('bahan_racikan')
          AND sub_total > 0
        GROUP BY kd_barang, nm_barang, jenis
        ORDER BY total_nilai DESC LIMIT 20
    """, (dari, sampai))
    total_all = sum(r['total_nilai'] for r in rows)
    result = rows_to_list(rows)
    for r in result:
        r['kontribusi'] = round(r['total_nilai'] / total_all * 100, 1) if total_all else 0
        r['margin_pct']  = round(r['gross_margin'] / r['total_nilai'] * 100, 1) if r['total_nilai'] else 0
        r['avg_per_trx'] = round(r['total_nilai'] / r['total_trx'], 0) if r['total_trx'] else 0
    return jsonify({'rows': result, 'total': total_all})

@app.route('/api/monitor/kategori')
def monitor_kategori():
    dari   = request.args.get('dari', str(date.today()))
    sampai = request.args.get('sampai', str(date.today()))
    rows = query("""
        SELECT CASE WHEN jenis='barang'  THEN 'Produk Fisik'
                    WHEN jenis='jasa'    THEN 'Jasa / Treatment'
                    WHEN jenis='racikan' THEN 'Racikan'
                    ELSE jenis END as kategori,
               SUM(qty) as total_qty,
               SUM(sub_total) as total_nilai,
               COUNT(DISTINCT no_faktur) as total_trx
        FROM R_Penjualan
        WHERE tanggal >= ? AND tanggal <= ?
          AND jenis NOT IN ('bahan_racikan')
          AND sub_total > 0
        GROUP BY kategori
        ORDER BY total_nilai DESC
    """, (dari, sampai))
    total = sum(r['total_nilai'] for r in rows)
    result = rows_to_list(rows)
    for r in result:
        r['pct'] = round(r['total_nilai'] / total * 100, 1) if total else 0
    return jsonify({'rows': result, 'total': total})

# ─── FOTO PASIEN ─────────────────────────────────────────────
def _save_image(image_data_b64, kd_pasien, tipe):
    """Decode base64 image, optionally resize with Pillow, save to disk. Return filename."""
    match = re.match(r'data:image/(\w+);base64,', image_data_b64)
    raw = re.sub(r'^data:image/\w+;base64,', '', image_data_b64)
    img_bytes = base64.b64decode(raw)

    if HAS_PIL:
        img = Image.open(_io.BytesIO(img_bytes))
        img = ImageOps.exif_transpose(img)          # fix phone rotation
        if img.width > 1600 or img.height > 1600:
            img.thumbnail((1600, 1600), Image.LANCZOS)
        if img.mode in ('RGBA', 'P', 'LA'):
            img = img.convert('RGB')
        buf = _io.BytesIO()
        img.save(buf, format='JPEG', quality=85, optimize=True)
        img_bytes = buf.getvalue()
        ext = 'jpg'
    else:
        ext = (match.group(1) if match else 'jpg').replace('jpeg', 'jpg')

    folder = os.path.join(UPLOAD_DIR, kd_pasien)
    os.makedirs(folder, exist_ok=True)
    ts = datetime.now().strftime('%Y%m%d_%H%M%S_%f')[:19]
    filename = f"{tipe}_{ts}.{ext}"
    with open(os.path.join(folder, filename), 'wb') as f:
        f.write(img_bytes)
    return f"{kd_pasien}/{filename}"

@app.route('/api/foto-pasien', methods=['POST'])
def upload_foto():
    d = request.json
    kd   = d.get('kd_pasien', '')
    tipe = d.get('tipe', 'sebelum')
    if not kd or not d.get('image_data'):
        return jsonify({'ok': False, 'msg': 'kd_pasien dan image_data wajib'}), 400
    user = get_current_user()
    uploaded_by = (user.get('nama') or user.get('username') or '') if user else d.get('uploaded_by', '')
    filename = _save_image(d['image_data'], kd, tipe)
    fid = execute("""
        INSERT INTO Foto_Pasien (kd_pasien, id_rm, no_rm, tanggal, tipe, area, keterangan, filename, uploaded_by)
        VALUES (?,?,?,?,?,?,?,?,?)
    """, (kd, d.get('id_rm') or None, d.get('no_rm',''),
          d.get('tanggal', str(date.today())),
          tipe, d.get('area',''), d.get('keterangan',''), filename, uploaded_by))
    return jsonify({'ok': True, 'id': fid, 'url': f'/static/uploads/foto_pasien/{filename}'})

@app.route('/api/foto-pasien/<path:kd_pasien>')
def list_foto(kd_pasien):
    rows = query("""
        SELECT f.*, rm.no_rm as rm_no_rm, rm.tanggal as rm_tanggal
        FROM Foto_Pasien f
        LEFT JOIN RekamMedik rm ON f.id_rm = rm.id
        WHERE f.kd_pasien = ?
        ORDER BY f.tanggal DESC, f.created_at DESC
    """, (kd_pasien,))
    user = get_current_user()
    role = user.get('role','') if user else ''
    result = rows_to_list(rows)
    for r in result:
        r['url'] = f"/static/uploads/foto_pasien/{r['filename']}"
        # can_delete: admin/supervisor always; others only within 24h of upload
        created_str = r.get('created_at','')
        if role in ('admin','supervisor'):
            r['can_delete'] = True
        elif created_str:
            try:
                age_h = (datetime.now() - datetime.strptime(created_str[:19],'%Y-%m-%d %H:%M:%S')).total_seconds() / 3600
                r['can_delete'] = age_h <= 24
            except:
                r['can_delete'] = True
        else:
            r['can_delete'] = True
    return jsonify(result)

@app.route('/api/foto-pasien/<int:fid>', methods=['PUT'])
def update_foto(fid):
    d = request.json
    execute("UPDATE Foto_Pasien SET tipe=?, area=?, keterangan=?, tanggal=? WHERE id=?",
            (d.get('tipe'), d.get('area',''), d.get('keterangan',''), d.get('tanggal'), fid))
    return jsonify({'ok': True})

@app.route('/api/foto-pasien/<int:fid>', methods=['DELETE'])
def delete_foto(fid):
    user = get_current_user()
    role = user.get('role','') if user else ''
    row = query("SELECT filename, created_at FROM Foto_Pasien WHERE id=?", (fid,), one=True)
    if row:
        created_str = dict(row).get('created_at','')
        if role not in ('admin','supervisor') and created_str:
            try:
                age_h = (datetime.now() - datetime.strptime(created_str[:19],'%Y-%m-%d %H:%M:%S')).total_seconds() / 3600
                if age_h > 24:
                    return jsonify({'ok': False, 'msg': 'Foto lebih dari 1 hari hanya bisa dihapus oleh admin/supervisor'}), 403
            except:
                pass
        path = os.path.join(UPLOAD_DIR, dict(row)['filename'])
        if os.path.exists(path):
            os.remove(path)
        execute("DELETE FROM Foto_Pasien WHERE id=?", (fid,))
    return jsonify({'ok': True})

@app.route('/mobile')
def mobile_page():
    return send_from_directory('static', 'mobile.html')

@app.route('/mobile/foto/<path:kd_pasien>')
def mobile_foto_deep(kd_pasien):
    return send_from_directory('static', 'mobile.html')

# ─── PASIEN ───────────────────────────────────────────────────
@app.route('/api/pasien')
def get_pasien():
    q = request.args.get('q', '')
    rows = query("SELECT p.*, CASE WHEN m.kd_pasien IS NOT NULL THEN 1 ELSE 0 END as is_member FROM Pasien p LEFT JOIN Member m ON p.kd_pasien=m.kd_pasien AND m.tgl_exp >= date('now') WHERE p.kd_pasien LIKE ? OR p.nm_pasien LIKE ? ORDER BY p.nm_pasien LIMIT 50",
                 (f'%{q}%', f'%{q}%'))
    return jsonify(rows_to_list(rows))

@app.route('/api/pasien/<path:kd>')
def get_pasien_one(kd):
    row = query("SELECT p.*, CASE WHEN m.kd_pasien IS NOT NULL THEN 1 ELSE 0 END as is_member FROM Pasien p LEFT JOIN Member m ON p.kd_pasien=m.kd_pasien AND m.tgl_exp >= date('now') WHERE p.kd_pasien=?", (kd,), one=True)
    if not row: return jsonify({'error': 'not found'}), 404
    point = query("SELECT COALESCE(SUM(debit),0)-COALESCE(SUM(kredit),0) as pts FROM DetailPoint WHERE kd_pasien=?", (kd,), one=True)
    d = dict(row)
    d['point'] = point['pts'] if point else 0
    return jsonify(d)

@app.route('/api/pasien', methods=['POST'])
def add_pasien():
    d = request.json
    try:
        execute("INSERT INTO Pasien (kd_pasien,nm_pasien,alamat,telp,tgl_lahir,jenis_kelamin) VALUES (?,?,?,?,?,?)",
                (d['kd_pasien'], d['nm_pasien'], d.get('alamat',''), d.get('telp',''), d.get('tgl_lahir',''), d.get('jenis_kelamin','P')))
        return jsonify({'ok': True})
    except Exception as e:
        return jsonify({'ok': False, 'msg': str(e)}), 400

@app.route('/api/pasien/<path:kd>', methods=['PUT'])
def update_pasien(kd):
    d = request.json
    execute("UPDATE Pasien SET nm_pasien=?,alamat=?,telp=?,tgl_lahir=?,jenis_kelamin=? WHERE kd_pasien=?",
            (d['nm_pasien'], d.get('alamat',''), d.get('telp',''), d.get('tgl_lahir',''), d.get('jenis_kelamin','P'), kd))
    return jsonify({'ok': True})

# ─── STOK BARANG ─────────────────────────────────────────────
@app.route('/api/stok')
def get_stok():
    """Pencarian kombinasi kode & nama (spt POS/racikan) — lihat _kata_cocok_kolom().
    Tiap kata yg diketik wajib cocok ke SALAH SATU (kode ATAU nama), antar kata
    digabung AND — jadi "cream pso" tetap ketemu "Cream Uriage PSO" walau gabungan
    katanya bukan substring utuh (dulu cuma LIKE '%q%' mentah, gagal utk query
    multi-kata — keluhan user 2026-10-01)."""
    filter_type = request.args.get('filter', 'all')
    q = request.args.get('q', '').strip()
    kd_lokasi = request.args.get('kd_lokasi', '')
    lok_clause = " AND kd_lokasi=?" if kd_lokasi else ""
    lok_params = (kd_lokasi,) if kd_lokasi else ()
    q_where, q_params = _kata_cocok_kolom('kd_barang', 'nm_barang', q.split())
    if filter_type == 'aktif':
        sql = f"SELECT * FROM Stok_barang WHERE harga != 0 AND stok_sisa > 0 AND {q_where}{lok_clause} ORDER BY stok_sisa"
    elif filter_type == 'warning':
        sql = f"SELECT * FROM Stok_barang WHERE (harga = 0 OR harga_beli >= harga) AND {q_where}{lok_clause} ORDER BY kd_barang"
    else:
        sql = f"SELECT * FROM Stok_barang WHERE {q_where}{lok_clause} ORDER BY nm_barang"
    rows = query(sql, list(q_params) + list(lok_params))
    return jsonify(rows_to_list(rows))

@app.route('/api/stok/counts')
def get_stok_counts():
    """Jumlah baris per filter (ringan, untuk badge tab — hindari load semua baris)."""
    kd_lokasi = request.args.get('kd_lokasi', '')
    lok = " AND kd_lokasi=?" if kd_lokasi else ""
    p   = (kd_lokasi,) if kd_lokasi else ()
    semua   = query(f"SELECT COUNT(*) c FROM Stok_barang WHERE 1=1{lok}", p, one=True)['c']
    ada     = query(f"SELECT COUNT(*) c FROM Stok_barang WHERE harga!=0 AND stok_sisa>0{lok}", p, one=True)['c']
    warning = query(f"SELECT COUNT(*) c FROM Stok_barang WHERE (harga=0 OR harga_beli>=harga){lok}", p, one=True)['c']
    return jsonify({'semua': semua, 'ada_stok': ada, 'warning': warning})

@app.route('/api/stok/<int:id>', methods=['PUT'])
def update_stok(id):
    d = request.json
    lama = query("SELECT kd_barang,nm_barang,stok_min,harga,disc,DiscMem,DiscLain FROM Stok_barang WHERE id=?", (id,), one=True)
    execute("UPDATE Stok_barang SET stok_min=?,harga=?,disc=?,DiscMem=?,DiscLain=? WHERE id=?",
            (d.get('stok_min',0), d.get('harga',0), d.get('disc','0'), d.get('DiscMem',0), d.get('DiscLain',0), id))
    catat_log('ubah_harga_stok', 'Stok_barang', lama['kd_barang'] if lama else id,
              f"Ubah harga/disc '{lama['nm_barang'] if lama else id}'",
              data_lama=dict(lama) if lama else None,
              data_baru={'stok_min': d.get('stok_min',0), 'harga': d.get('harga',0), 'disc': d.get('disc','0'),
                         'DiscMem': d.get('DiscMem',0), 'DiscLain': d.get('DiscLain',0)})
    return jsonify({'ok': True})

@app.route('/api/exp-bermasalah')
def exp_bermasalah_list():
    """Versi UI dari _cek_format_exp_asing() (yg sebelumnya cuma print ke log saat
    server start) — list semua baris Stok_barang yg exp-nya TIDAK cocok format mana
    pun di _EXP_FORMATS_DIKENAL, supaya admin bisa lihat & perbaiki manual lewat
    halaman, bukan nunggu developer baca log server. Sekalian kasih 'saran_tanggal'
    (hasil coba-tebak via _ki_tgl(), normalizer toleran banyak format) — TETAP wajib
    dikonfirmasi manual oleh admin, bukan auto-fix, krn menyentuh data stok asli."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor'):
        return jsonify({'ok': False, 'msg': 'Tidak berhak'}), 403
    rows = query("""SELECT id, kd_barang, nm_barang, exp, stok_sisa, stok_awal, kd_lokasi, harga
                     FROM Stok_barang WHERE exp IS NOT NULL AND exp != '' ORDER BY stok_sisa DESC""")
    result = []
    for r in rows_to_list(rows):
        e = (r['exp'] or '').strip()
        dikenal = False
        for fmt in _EXP_FORMATS_DIKENAL:
            try:
                datetime.strptime(e, fmt)
                dikenal = True
                break
            except ValueError:
                continue
        if dikenal:
            continue
        tebakan = _ki_tgl(e)   # '' kalau _ki_tgl jg tak bisa nebak sama sekali
        saran_iso = ''
        if tebakan:
            try:
                saran_iso = datetime.strptime(tebakan, '%m/%d/%Y').strftime('%Y-%m-%d')
            except ValueError:
                pass
        r['saran_tanggal'] = saran_iso
        result.append(r)
    return jsonify(result)

@app.route('/api/exp-bermasalah/<int:sid>/perbaiki', methods=['POST'])
def exp_bermasalah_perbaiki(sid):
    """Admin konfirmasi tanggal exp yg benar utk 1 baris Stok_barang bermasalah.
    Input selalu dari <input type=\"date\"> di frontend (jadi selalu ISO) — dinormalkan
    via _ki_tgl() ke format konvensi app ini (M/D/YYYY) sebelum disimpan, pola sama
    persis dgn input_stok_manual() (lihat memory bug-fefo-exp-string-sort.md)."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor'):
        return jsonify({'ok': False, 'msg': 'Tidak berhak'}), 403
    d = request.json or {}
    exp_baru_raw = (d.get('exp_baru') or '').strip()
    if not exp_baru_raw:
        return jsonify({'ok': False, 'msg': 'Tanggal exp yang benar wajib diisi'}), 400
    exp_baru = _ki_tgl(exp_baru_raw)
    valid = any(_cocok_format(exp_baru, fmt) for fmt in _EXP_FORMATS_DIKENAL)
    if not valid:
        return jsonify({'ok': False, 'msg': f'Tanggal tidak valid/tidak bisa dikenali: "{exp_baru_raw}"'}), 400
    lama = query("SELECT kd_barang, exp FROM Stok_barang WHERE id=?", (sid,), one=True)
    if not lama:
        return jsonify({'ok': False, 'msg': 'Baris stok tidak ditemukan'}), 404
    execute("UPDATE Stok_barang SET exp=? WHERE id=?", (exp_baru, sid))
    catat_log('perbaiki_exp_asing', 'Stok_barang', lama['kd_barang'],
              f"Perbaiki format tanggal exp '{lama['kd_barang']}': '{lama['exp']}' → '{exp_baru}'",
              data_lama={'exp': lama['exp']}, data_baru={'exp': exp_baru})
    return jsonify({'ok': True, 'exp_baru': exp_baru})

@app.route('/api/stok-manual', methods=['POST'])
def input_stok_manual():
    """'Suntik' stok manual langsung ke suatu lokasi (biasanya cabang) — dipakai
    utk koreksi kalau ada barang yg SEHARUSNYA ada tapi hilang/tidak tercatat
    (mis. gap migrasi data). BUKAN jalur normal (Pembelian→approve, Order
    Transfer, Terima Kiriman) — jadi WAJIB ada alasan, dan SELALU dicatat
    lengkap ke RiwayatBarang + Transfer_Stok_Log + Aktivitas_Log, supaya tidak
    jadi "data siluman" tanpa jejak spt kasus lama (lihat memory
    koreksi-data-satuan-disc-legacy — akar masalahnya persis krn migrasi dulu
    TIDAK mencatat jejak audit). Admin/supervisor saja."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor yang bisa input stok manual'}), 403

    d = request.json or {}
    kd_barang = (d.get('kd_barang') or '').strip()
    nm_barang = (d.get('nm_barang') or '').strip()
    kd_lokasi = (d.get('kd_lokasi') or '').strip() or (user.get('kd_lokasi') or 'CBG1')
    catatan   = (d.get('catatan') or '').strip()
    # Input exp di form ini pakai <input type="date">, yg SELALU keluar format ISO
    # 'YYYY-MM-DD' — beda dari konvensi 'M/D/YYYY' yg dipakai di SELURUH Stok_barang.exp
    # lainnya. Tanpa dinormalkan, FEFO (_exp_key) gagal baca & batch ini keliru dianggap
    # "tak pernah expired" (ketahuan 2026-10-02 dari peringatan _cek_format_exp_asing()
    # saat user betulan pakai fitur ini). _ki_tgl() sudah menangani '%Y-%m-%d' sbg salah
    # satu formatnya, reuse langsung drpd bikin normalizer baru.
    exp       = _ki_tgl((d.get('exp') or '').strip())
    try:    qty = float(d.get('qty') or 0)
    except (TypeError, ValueError): qty = 0
    try:    harga_beli = float(d.get('harga_beli') or 0)
    except (TypeError, ValueError): harga_beli = 0
    try:    harga = float(d.get('harga') or 0)
    except (TypeError, ValueError): harga = 0
    try:    stok_min = float(d.get('stok_min') or 0)
    except (TypeError, ValueError): stok_min = 0
    satuan = _migr_satuan_aman(d.get('satuan', ''), 'Pcs')
    disc   = _migr_disc_aman(d.get('disc', '0'))

    if not kd_barang or not nm_barang:
        return jsonify({'ok': False, 'msg': 'Kode & nama barang wajib diisi'}), 400
    if qty <= 0:
        return jsonify({'ok': False, 'msg': 'Qty harus lebih dari 0'}), 400
    if not catatan:
        return jsonify({'ok': False, 'msg': 'Alasan/catatan wajib diisi — ini jalur koreksi manual, bukan transaksi normal'}), 400

    # Daftarkan ke master Barang kalau belum ada — konsisten dgn perilaku Pembelian
    # (_pembelian_posting_stok), supaya barang ini ikut kebaca di fitur lain yg
    # merujuk master (mis. pencarian POS, Terima Kiriman correlate).
    if not query("SELECT 1 FROM Barang WHERE kd_barang=?", (kd_barang,), one=True):
        execute("""INSERT OR IGNORE INTO Barang (kd_barang,nm_barang,nm_generik,satuan,harga_jual,aktif)
                   VALUES (?,?,?,?,?,1)""", (kd_barang, nm_barang, '', satuan, harga))

    existing = query("SELECT id FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                      (kd_barang, exp, kd_lokasi), one=True)
    if existing:
        execute("""UPDATE Stok_barang SET stok_awal=stok_awal+?, stok_sisa=stok_sisa+?,
                   stok_gudang=stok_gudang+?, harga_beli=?, harga=?, disc=?, satuan=?, stok_min=?
                   WHERE id=?""",
                (qty, qty, qty, harga_beli, harga, disc, satuan, stok_min, existing['id']))
    else:
        execute("""INSERT INTO Stok_barang (kd_barang,nm_barang,harga_beli,harga_beli_awal,
                   stok_awal,stok_sisa,stok_gudang,stok_min,harga,disc,satuan,exp,DiscMem,DiscLain,kd_lokasi,pemilik)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,0,0,?,'')""",
                (kd_barang, nm_barang, harga_beli, harga_beli, qty, qty, qty, stok_min,
                 harga, disc, satuan, exp, kd_lokasi))

    today  = str(date.today())
    # Pakai mikrodetik (bukan cuma detik) biar no_referensi tidak pernah dobel
    # kalau 2 input manual kejadian di detik yg sama.
    no_ref = f"MANUAL{datetime.now().strftime('%y%m%d%H%M%S%f')}"
    execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
            (today, kd_barang, exp, f"Input stok manual ({no_ref}) — {catatan}", qty, qty, qty))
    execute("""INSERT INTO Transfer_Stok_Log (no_referensi,tipe_referensi,kd_barang,nm_barang,
               kd_lokasi_dari,kd_lokasi_ke,qty,exp,created_by)
               VALUES (?,?,?,?,?,?,?,?,?)""",
            (no_ref, 'input_manual', kd_barang, nm_barang, 'MANUAL', kd_lokasi, qty, exp, user.get('nama', '')))
    catat_log('input_stok_manual', 'Stok_barang', kd_barang,
              f"Input stok manual {qty:g} {satuan} '{nm_barang}' @ {kd_lokasi} — alasan: {catatan}",
              data_baru={'kd_barang': kd_barang, 'nm_barang': nm_barang, 'qty': qty,
                         'harga_beli': harga_beli, 'harga': harga, 'disc': disc, 'satuan': satuan,
                         'exp': exp, 'kd_lokasi': kd_lokasi, 'catatan': catatan, 'no_referensi': no_ref})
    return jsonify({'ok': True, 'no_referensi': no_ref})

@app.route('/api/stok-manual/riwayat')
def riwayat_stok_manual():
    """Riwayat input stok manual (halaman Input Stok Manual) — admin/supervisor.
    Filter rentang tanggal (?dari=&sampai=) — pola sama persis dgn get_log_aktivitas()
    (waktu >= awal hari s/d waktu < akhir hari), defaultnya 30 hari terakhir."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor'):
        return jsonify({'ok': False, 'msg': 'Tidak berhak'}), 403
    tgl1 = request.args.get('dari',   str(date.today() - timedelta(days=30)))
    tgl2 = request.args.get('sampai', str(date.today()))
    rows = query("""SELECT waktu, nama_user, role, ref AS kd_barang, keterangan, data_baru
                     FROM Aktivitas_Log WHERE aksi='input_stok_manual'
                       AND waktu >= ? AND waktu < ?
                     ORDER BY id DESC LIMIT 300""", (tgl1 + ' 00:00:00', tgl2 + ' 23:59:59'))
    result = rows_to_list(rows)
    for r in result:
        try:
            r['detail'] = json.loads(r['data_baru']) if r.get('data_baru') else None
        except Exception:
            r['detail'] = None
        del r['data_baru']
    return jsonify(result)

def _exp_key(exp_str):
    """Kunci urut FEFO yg BENAR (tanggal asli, bukan string sort). Stok_barang.exp disimpan
    'M/D/YYYY' tanpa leading zero (lihat _ki_tgl) — 'ORDER BY exp ASC' / 'MIN(exp)' di SQL
    itu SORT STRING MENTAH, salah tiap kali digit bulan/tanggal beda panjang (mis.
    '1/1/2027' vs '12/1/2026' — string treat '1/1/2027' lbh kecil krn '/' < '2', padahal
    Jan 2027 lbh BELAKANGAN drpd Des 2026 — barang lbh baru bisa kepilih duluan drpd yg
    beneran lbh dulu expired). Exp kosong/tak valid dianggap paling akhir (date.max) —
    supaya batch tanpa exp tidak "menyerobot" duluan di urutan FEFO.

    PENTING: sebagian batch (hasil migrasi CSV / jalur input beda) tersimpan dgn jam
    ikut ('M/D/YYYY 0:00:00'), bukan cuma 'M/D/YYYY' polos. Ditemukan 2026-09-27 saat
    audit kd_barang 'Cream White SR gr' — 167 dari 202 batch aktif (stok_sisa>0)
    ternyata pakai format berjam, dan SEBELUM perbaikan ini semuanya gagal parse →
    jatuh ke date.max (dianggap "tak pernah expired") → FEFO salah pilih batch yg
    justru lebih baru drpd batch yg beneran lbh dulu kedaluwarsa (mis. batch exp
    4/30/2027 kalah prioritas drpd batch exp 3/30/2028 krn yg pertama gagal parse).
    Makanya WAJIB coba kedua format, bukan cuma '%m/%d/%Y'."""
    e = (exp_str or '').strip()
    for fmt in _EXP_FORMATS_DIKENAL:
        try:
            return datetime.strptime(e, fmt).date()
        except ValueError:
            continue
    return date.max

# Daftar format tanggal exp yg SUDAH terbukti ada di data (lihat _exp_key). Dipakai jg
# oleh _cek_format_exp_asing() — satu sumber kebenaran, jangan duplikasi list-nya.
_EXP_FORMATS_DIKENAL = ('%m/%d/%Y %H:%M:%S', '%m/%d/%Y')

def _cek_format_exp_asing():
    """Pengaman utk migrasi data BERIKUTNYA: cek semua Stok_barang.exp yg masih aktif
    (stok_sisa>0) terhadap _EXP_FORMATS_DIKENAL. Kalau ada format LAIN yg belum pernah
    dikenal sistem (mis. migrasi dari sumber baru yg formatnya beda lagi), _exp_key()
    akan DIAM-DIAM menjatuhkannya ke date.max ("dianggap tak pernah expired") — FEFO
    salah urut tanpa ketahuan, persis kasus 2026-09-27 (167/202 batch aktif ternyata
    pakai format berjam yg blm dikenal saat itu, baru ketahuan krn user tanya soal 1
    barang spesifik). Dipanggil sekali saat server start (lihat __main__) — kalau nemu
    format asing, cetak peringatan jelas ke log supaya ketahuan LANGSUNG, bukan nunggu
    ditemukan gak sengaja lagi. Aman dipanggil kapan saja (read-only, tak pernah raise)."""
    try:
        rows = query("""SELECT DISTINCT kd_barang, exp FROM Stok_barang
                        WHERE stok_sisa > 0 AND exp IS NOT NULL AND exp != ''""")
    except Exception:
        return   # tabel/koneksi belum siap — jangan sampai bikin startup gagal gara2 ini
    asing = []
    for r in rows:
        e = (r['exp'] or '').strip()
        if not any(_cocok_format(e, fmt) for fmt in _EXP_FORMATS_DIKENAL):
            asing.append((r['kd_barang'], e))
    if not asing:
        return
    print(f"\n⚠️  PERINGATAN: {len(asing)} kombinasi kd_barang+exp aktif pakai format TANGGAL")
    print(f"   ASING yg belum dikenal _exp_key() — FEFO utk batch ini bisa SALAH URUT (batch")
    print(f"   dianggap 'tak pernah expired' padahal exp-nya valid, cuma formatnya beda).")
    print(f"   Kemungkinan besar dari migrasi data baru. Contoh:")
    for kd, e in asing[:10]:
        print(f"     - {kd!r}: exp={e!r}")
    if len(asing) > 10:
        print(f"     ... dan {len(asing) - 10} kombinasi lainnya")
    print(f"   → Tambahkan format ini ke _EXP_FORMATS_DIKENAL (app.py, dekat _exp_key).\n")

def _cocok_format(s, fmt):
    try:
        datetime.strptime(s, fmt)
        return True
    except ValueError:
        return False

def _kata_cocok_kolom(kolom_kode, kolom_nama, words):
    """Klausa SQL: query dipecah per-kata, tiap kata wajib cocok ke SALAH SATU
    (kode ATAU nama) — antar kata digabung AND. Jadi mis. kasir ketik "TR001 facial"
    (kode + sepotong nama sekaligus) tetap ketemu, meski gabungan kata itu bukan
    substring utuh di kode maupun nama manapun sendiri-sendiri. Dipakai di pencarian
    barang/jasa POS (get_barang_jual) & pencarian bahan racikan (get_bahan_racikan).

    WAJIB pakai LOWER() di kedua sisi — kolom kd_barang di MySQL sengaja dikasih
    collation utf8mb4_bin (case-SENSITIVE, biar exact-match/join presisi), jadi LIKE
    polos ikut jadi case-sensitive juga: "cream" (huruf kecil semua) TIDAK match
    "Cream M" krn 'C' besar ≠ 'c' kecil scr biner. Ditemukan nyata saat cari
    "cream m" tidak ketemu produk "Cream M" yg justru py stok. LOWER() bikin
    pencarian selalu case-insensitive spt yg diharapkan user, apa pun collation
    kolomnya (aman jg di SQLite yg dari awal memang case-insensitive)."""
    if not words:
        return '1=1', []
    klausa = ' AND '.join([f'(LOWER({kolom_kode}) LIKE ? OR LOWER({kolom_nama}) LIKE ?)'] * len(words))
    params = []
    for w in words:
        wl = w.lower()
        params += [f'%{wl}%', f'%{wl}%']
    return klausa, params

def _prioritas_kode(kolom_kode, q):
    """ORDER BY CASE — item yg KODE-nya PERSIS sama dgn query (case-insensitive)
    naik ke urutan paling atas, lalu yg kodenya DIAWALI query, baru sisanya
    alfabetis. Tanpa ini, kasir ketik kode barang yg tepat (mis. "Cream M") bisa
    "tenggelam" di bawah LIMIT kalau ada banyak barang lain yg SAMA-SAMA
    mengandung kata yg sama (mis. semua barang "Cream ..."). Dipakai bareng
    _kata_cocok_kolom() supaya kode yg diketik pasti muncul duluan."""
    ql = (q or '').strip().lower()
    return (f"CASE WHEN LOWER({kolom_kode})=? THEN 0 WHEN LOWER({kolom_kode}) LIKE ? THEN 1 ELSE 2 END",
            [ql, f'{ql}%'])

@app.route('/api/barang-jual')
def get_barang_jual():
    """Combined barang + jasa for POS — filtered by kd_lokasi (cabang stok).
    Pencarian mengkombinasi kode & nama — lihat _kata_cocok_kolom().

    PENTING: batch FEFO (yg exp/harga-nya dipakai sbg baris utama) dipilih di PYTHON
    pakai _exp_key(), BUKAN 'exp = (SELECT MIN(exp) ...)' di SQL — exp disimpan teks
    'M/D/YYYY' (kadang + jam ikut, 'M/D/YYYY 0:00:00'), jadi MIN() SQL itu SORT STRING
    MENTAH & salah tiap kali panjang digit/format beda. Ditemukan nyata 2026-09-27 di
    kd_barang 'Cream White SR gr': batch exp 4/30/2027 (lbh dulu expired, seharusnya
    dijual duluan) malah "tersembunyi" di balik batch exp 3/30/2028 krn '3' < '4' scr
    string — sama persis pola yg sudah diperbaiki di get_bahan_racikan()/get_stok_batches()
    (lihat _exp_key() utk detail), skrg dipakai di sini juga krn ini titik PALING sering
    disentuh (pencarian barang utama tiap transaksi POS)."""
    q         = request.args.get('q', '').strip()
    kd_lokasi = request.args.get('kd_lokasi', 'CBG1')
    words = q.split()

    barang_where, barang_params = _kata_cocok_kolom('s.kd_barang', 's.nm_barang', words)
    barang_prio, barang_prio_params = _prioritas_kode('s.kd_barang', q)

    kd_rows = query(f"""
        SELECT DISTINCT s.kd_barang FROM Stok_barang s
        WHERE s.harga > 0 AND s.stok_sisa > 0 AND s.kd_lokasi=? AND {barang_where}
        ORDER BY {barang_prio}, s.nm_barang LIMIT 30
    """, (kd_lokasi, *barang_params, *barang_prio_params))
    kds = [r['kd_barang'] for r in kd_rows]

    barang = []
    if kds:
        ph = ','.join('?' * len(kds))
        all_batches = rows_to_list(query(f"""
            SELECT s.*,
                   COALESCE((SELECT b.non_obat FROM Barang b WHERE b.kd_barang=s.kd_barang),0) as non_obat
            FROM Stok_barang s
            WHERE s.kd_barang IN ({ph}) AND s.stok_sisa > 0 AND s.harga > 0 AND s.kd_lokasi=?
        """, (*kds, kd_lokasi)))
        per_kd = {}
        for b in all_batches:
            per_kd.setdefault(b['kd_barang'], []).append(b)
        for lst in per_kd.values():
            lst.sort(key=lambda b: _exp_key(b['exp']))   # FEFO yg benar

        for kd in kds:   # pertahankan urutan prioritas pencarian dari kd_rows
            batches = per_kd.get(kd)
            if not batches:
                continue
            main = dict(batches[0])   # batch FEFO paling awal → baris utama (harga/exp acuan)
            barang.append({
                'kd': main['kd_barang'], 'nama': main['nm_barang'],
                'harga': main['harga'], 'disc': main['disc'], 'jenis': 'barang',
                'stok': main['stok_sisa'], 'exp': main['exp'], 'satuan': main['satuan'],
                'DiscMem': main['DiscMem'], 'non_obat': main['non_obat'],
                'total_stok': sum(b['stok_sisa'] for b in batches),
                'jml_batch': len(batches),
            })

    jasa_where, jasa_params = _kata_cocok_kolom('kd_pelayanan', 'nm_pelayanan', words)
    jasa_prio, jasa_prio_params = _prioritas_kode('kd_pelayanan', q)
    jasa = query(f"""SELECT kd_pelayanan as kd, nm_pelayanan as nama, harga, '0' as disc,
        'jasa' as jenis, 999 as stok, 1 as non_obat,
        komisi_dokter, komisi_supervisor, komisi_perawat
        FROM produk_pelayanan WHERE harga > 0 AND COALESCE(aktif,1)=1
          AND {jasa_where}
        ORDER BY {jasa_prio}, nm_pelayanan LIMIT 20""", jasa_params + jasa_prio_params)
    return jsonify(barang + rows_to_list(jasa))

# ─── KARYAWAN ────────────────────────────────────────────────
@app.route('/api/karyawan')
def get_karyawan():
    q = request.args.get('q', '')
    if q:
        rows = query("""SELECT * FROM Karyawan WHERE aktif=1
            AND (kd_karyawan LIKE ? OR nm_karyawan LIKE ? OR nickname LIKE ?)
            ORDER BY pekerjaan, nm_karyawan LIMIT 10""",
            (f'%{q}%', f'%{q}%', f'%{q}%'))
    else:
        rows = query("SELECT * FROM Karyawan WHERE aktif=1 ORDER BY pekerjaan, nm_karyawan")
    return jsonify(rows_to_list(rows))

@app.route('/api/master/karyawan')
def master_karyawan_list():
    q         = request.args.get('q', '')
    aktif_f   = request.args.get('aktif', '1')
    pekerjaan = request.args.get('pekerjaan', '')
    where, params = [], []
    if aktif_f != 'all':
        where.append('aktif=?'); params.append(int(aktif_f))
    if q:
        where.append('(kd_karyawan LIKE ? OR nm_karyawan LIKE ? OR nickname LIKE ? OR telp LIKE ?)')
        params.extend([f'%{q}%']*4)
    if pekerjaan:
        where.append('pekerjaan=?'); params.append(pekerjaan)
    sql = 'SELECT * FROM Karyawan'
    if where:
        sql += ' WHERE ' + ' AND '.join(where)
    sql += ' ORDER BY pekerjaan, nm_karyawan'
    return jsonify(rows_to_list(query(sql, params)))


@app.route('/api/master/karyawan', methods=['POST'])
def master_karyawan_create():
    d  = request.json or {}
    kd = str(d.get('kd_karyawan', '')).strip().upper()
    nm = str(d.get('nm_karyawan', '')).strip()
    if not kd or not nm:
        return jsonify({'ok': False, 'msg': 'Kode dan nama karyawan wajib diisi'}), 400
    # Cek dobel dgn SELECT dulu (bukan tangkap exception DB) — portabel di kedua backend;
    # `except sqlite3.IntegrityError` sebelumnya TIDAK PERNAH kena di MySQL (exception-nya
    # pymysql.err.IntegrityError), jadi kode dobel malah jatuh ke error 500 mentah.
    if query("SELECT 1 FROM Karyawan WHERE kd_karyawan=?", (kd,), one=True):
        return jsonify({'ok': False, 'msg': f'Kode karyawan "{kd}" sudah dipakai'}), 409
    execute("INSERT INTO Karyawan (kd_karyawan,nm_karyawan,nickname,pekerjaan,telp,aktif) VALUES (?,?,?,?,?,1)",
            (kd, nm, str(d.get('nickname','')).strip(),
             str(d.get('pekerjaan','')).strip(), str(d.get('telp','')).strip()))
    return jsonify({'ok': True})


@app.route('/api/master/karyawan/<kd>', methods=['PUT'])
def master_karyawan_update(kd):
    d  = request.json or {}
    nm = str(d.get('nm_karyawan', '')).strip()
    if not nm:
        return jsonify({'ok': False, 'msg': 'Nama karyawan wajib diisi'}), 400
    execute("UPDATE Karyawan SET nm_karyawan=?,nickname=?,pekerjaan=?,telp=?,aktif=? WHERE kd_karyawan=?",
            (nm, str(d.get('nickname','')).strip(),
             str(d.get('pekerjaan','')).strip(), str(d.get('telp','')).strip(),
             int(d.get('aktif', 1)), kd))
    return jsonify({'ok': True})


@app.route('/api/master/karyawan/<kd>', methods=['DELETE'])
def master_karyawan_toggle(kd):
    row = query("SELECT aktif FROM Karyawan WHERE kd_karyawan=?", (kd,), one=True)
    if not row:
        return jsonify({'ok': False, 'msg': 'Karyawan tidak ditemukan'}), 404
    new_aktif = 0 if row['aktif'] else 1
    execute("UPDATE Karyawan SET aktif=? WHERE kd_karyawan=?", (new_aktif, kd))
    return jsonify({'ok': True, 'aktif': new_aktif})

@app.route('/api/master/karyawan/<kd>/hapus', methods=['DELETE'])
def master_karyawan_hapus(kd):
    """Hapus PERMANEN (bukan sekadar nonaktifkan) — admin-only, hanya boleh kalau
    karyawan ini tidak punya jejak riwayat apa pun (belum pernah transaksi/komisi/
    pegang kit), supaya tidak merusak laporan historis yg mereferensikan nama/kode
    karyawan ini."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role != 'admin':
        return jsonify({'ok': False, 'msg': 'Hanya admin yang bisa menghapus permanen'}), 403
    row = query("SELECT nm_karyawan FROM Karyawan WHERE kd_karyawan=?", (kd,), one=True)
    if not row:
        return jsonify({'ok': False, 'msg': 'Karyawan tidak ditemukan'}), 404
    checks = [
        ("transaksi penjualan", "SELECT 1 FROM R_Penjualan WHERE kd_karyawan=? LIMIT 1", (kd,)),
        ("komisi", "SELECT 1 FROM Point_karyawan WHERE kd_karyawan=? LIMIT 1", (kd,)),
        ("stok kit perawat", "SELECT 1 FROM Stok_Perawat WHERE kd_karyawan=? LIMIT 1", (kd,)),
        ("log kit perawat", "SELECT 1 FROM Kit_Log WHERE kd_karyawan=? LIMIT 1", (kd,)),
        ("permintaan kit", "SELECT 1 FROM Kit_Request WHERE kd_karyawan=? LIMIT 1", (kd,)),
        ("serah-terima bahan kunci", "SELECT 1 FROM Serah_Terima_Kunci WHERE kd_pelaksana=? OR kd_k_dipilih=? OR kd_k_final=? LIMIT 1", (kd, kd, kd)),
        ("akun pengguna tertaut", "SELECT 1 FROM Pengguna WHERE kd_karyawan=? LIMIT 1", (kd,)),
    ]
    found = [label for label, sql, params in checks if query(sql, params, one=True)]
    if found:
        return jsonify({'ok': False, 'msg': f"Tidak bisa dihapus permanen — karyawan ini masih punya riwayat: {', '.join(found)}. Gunakan Nonaktifkan saja."}), 400
    execute("DELETE FROM Karyawan WHERE kd_karyawan=?", (kd,))
    catat_log('karyawan_hapus', 'Karyawan', kd, f"Hapus permanen karyawan '{row['nm_karyawan']}' ({kd})")
    return jsonify({'ok': True})


# ─── MASTER PERAWATAN (produk_pelayanan / jasa treatment) ────
def _num(v):
    """Konversi nilai harga/komisi ke angka, terima string/None."""
    try:
        return float(str(v).replace('.', '').replace(',', '.')) if str(v).strip() else 0
    except (TypeError, ValueError):
        return 0

@app.route('/api/master/perawatan')
def master_perawatan_list():
    q       = request.args.get('q', '')
    aktif_f = request.args.get('aktif', '1')
    where, params = [], []
    if aktif_f != 'all':
        where.append('COALESCE(aktif,1)=?'); params.append(int(aktif_f))
    if q:
        where.append('(kd_pelayanan LIKE ? OR nm_pelayanan LIKE ?)')
        params.extend([f'%{q}%']*2)
    sql = 'SELECT kd_pelayanan,nm_pelayanan,harga,komisi_dokter,komisi_supervisor,komisi_perawat,COALESCE(komisi_persen_kunci,0) AS komisi_persen_kunci,COALESCE(aktif,1) AS aktif FROM produk_pelayanan'
    if where:
        sql += ' WHERE ' + ' AND '.join(where)
    sql += ' ORDER BY nm_pelayanan'
    return jsonify(rows_to_list(query(sql, params)))


@app.route('/api/master/perawatan', methods=['POST'])
def master_perawatan_create():
    d  = request.json or {}
    kd = str(d.get('kd_pelayanan', '')).strip().upper()
    nm = str(d.get('nm_pelayanan', '')).strip()
    if not kd or not nm:
        return jsonify({'ok': False, 'msg': 'Kode dan nama perawatan wajib diisi'}), 400
    # Cek dobel dgn SELECT dulu — lihat catatan di master_karyawan_create() soal
    # `except sqlite3.IntegrityError` yg tak pernah kena di backend MySQL.
    if query("SELECT 1 FROM produk_pelayanan WHERE kd_pelayanan=?", (kd,), one=True):
        return jsonify({'ok': False, 'msg': f'Kode perawatan "{kd}" sudah dipakai'}), 409
    execute("""INSERT INTO produk_pelayanan
               (kd_pelayanan,nm_pelayanan,harga,komisi_dokter,komisi_supervisor,komisi_perawat,komisi_persen_kunci,aktif)
               VALUES (?,?,?,?,?,?,?,1)""",
            (kd, nm, _num(d.get('harga')), _num(d.get('komisi_dokter')),
             _num(d.get('komisi_supervisor')), _num(d.get('komisi_perawat')), _num(d.get('komisi_persen_kunci'))))
    return jsonify({'ok': True})


@app.route('/api/master/perawatan/<kd>', methods=['PUT'])
def master_perawatan_update(kd):
    d  = request.json or {}
    nm = str(d.get('nm_pelayanan', '')).strip()
    if not nm:
        return jsonify({'ok': False, 'msg': 'Nama perawatan wajib diisi'}), 400
    execute("""UPDATE produk_pelayanan
               SET nm_pelayanan=?,harga=?,komisi_dokter=?,komisi_supervisor=?,komisi_perawat=?,komisi_persen_kunci=?,aktif=?
               WHERE kd_pelayanan=?""",
            (nm, _num(d.get('harga')), _num(d.get('komisi_dokter')),
             _num(d.get('komisi_supervisor')), _num(d.get('komisi_perawat')), _num(d.get('komisi_persen_kunci')),
             int(d.get('aktif', 1)), kd))
    return jsonify({'ok': True})


@app.route('/api/master/perawatan/<kd>', methods=['DELETE'])
def master_perawatan_toggle(kd):
    row = query("SELECT COALESCE(aktif,1) AS aktif FROM produk_pelayanan WHERE kd_pelayanan=?", (kd,), one=True)
    if not row:
        return jsonify({'ok': False, 'msg': 'Perawatan tidak ditemukan'}), 404
    new_aktif = 0 if row['aktif'] else 1
    execute("UPDATE produk_pelayanan SET aktif=? WHERE kd_pelayanan=?", (new_aktif, kd))
    return jsonify({'ok': True, 'aktif': new_aktif})


# ─── FORMULA PERAWATAN (bahan/alat per treatment) ────────────
@app.route('/api/master/perawatan/<kd>/formula')
def perawatan_formula_get(kd):
    rows = query("""SELECT id,kd_pelayanan,kd_barang,nm_barang,qty,satuan,tipe,
                           COALESCE(potong_stok,1) AS potong_stok,
                           COALESCE(bahan_kunci,0) AS bahan_kunci
                    FROM Perawatan_Bahan WHERE kd_pelayanan=? ORDER BY tipe DESC, id""",
                 (kd,))
    return jsonify(rows_to_list(rows))


@app.route('/api/master/perawatan/<kd>/formula', methods=['PUT'])
def perawatan_formula_save(kd):
    """Replace-all: kirim daftar baris bahan/alat lengkap untuk satu perawatan."""
    d = request.json or {}
    rows = d.get('rows', [])
    execute("DELETE FROM Perawatan_Bahan WHERE kd_pelayanan=?", (kd,))
    for r in rows:
        kb = str(r.get('kd_barang', '')).strip()
        if not kb:
            continue
        tipe = r.get('tipe', 'bahan')
        if tipe not in ('bahan', 'alat'):
            tipe = 'bahan'
        try:
            qty = float(r.get('qty', 1)) or 1
        except (TypeError, ValueError):
            qty = 1
        # Alat default tak potong stok; bahan default potong. Override bila dikirim.
        potong = r.get('potong_stok')
        potong = int(potong) if potong is not None else (0 if tipe == 'alat' else 1)
        # Bahan kunci cuma masuk akal utk tipe='bahan' (alat tak pernah "diserahterimakan").
        kunci = 1 if (tipe == 'bahan' and r.get('bahan_kunci')) else 0
        execute("""INSERT INTO Perawatan_Bahan
                   (kd_pelayanan,kd_barang,nm_barang,qty,satuan,tipe,potong_stok,bahan_kunci)
                   VALUES (?,?,?,?,?,?,?,?)""",
                (kd, kb, r.get('nm_barang', ''), qty, r.get('satuan', ''), tipe, potong, kunci))
    return jsonify({'ok': True, 'count': len([r for r in rows if str(r.get('kd_barang','')).strip()])})


# ─── KIT PERAWAT (stok bahan/alat pribadi) ───────────────────
def _kit_apply(kd_karyawan, kd_barang, nm_barang, exp, delta, satuan='',
               jenis='koreksi', ref='', kd_lokasi='', operator='', kd_pelayanan='',
               kd_pasien='', nm_pasien=''):
    """Ubah stok kit perawat (upsert per karyawan+barang+exp) + tulis Kit_Log.
       delta>0 masuk, delta<0 keluar. Saldo boleh minus (sesuai aturan kit kurang)."""
    exp = exp or ''
    row = query("""SELECT id, qty_sisa FROM Stok_Perawat
                   WHERE kd_karyawan=? AND kd_barang=? AND COALESCE(exp,'')=? LIMIT 1""",
                (kd_karyawan, kd_barang, exp), one=True)
    if row:
        new_saldo = (row['qty_sisa'] or 0) + delta
        execute("""UPDATE Stok_Perawat
                   SET qty_sisa=?, nm_barang=COALESCE(NULLIF(?,''),nm_barang),
                       satuan=COALESCE(NULLIF(?,''),satuan), updated_at=datetime('now')
                   WHERE id=?""",
                (new_saldo, nm_barang, satuan, row['id']))
    else:
        new_saldo = delta
        execute("""INSERT INTO Stok_Perawat (kd_karyawan,kd_barang,nm_barang,exp,qty_sisa,satuan)
                   VALUES (?,?,?,?,?,?)""",
                (kd_karyawan, kd_barang, nm_barang, exp, new_saldo, satuan))
    execute("""INSERT INTO Kit_Log (kd_karyawan,kd_barang,nm_barang,exp,qty,saldo,jenis,ref,kd_pelayanan,kd_pasien,nm_pasien,kd_lokasi,operator)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (kd_karyawan, kd_barang, nm_barang, exp, delta, new_saldo, jenis, ref, kd_pelayanan, kd_pasien, nm_pasien, kd_lokasi, operator))
    return new_saldo


def _kit_take_fefo(kd_karyawan, kd_barang, nm_barang, qty_needed, satuan='',
                   jenis='pakai_treatment', ref='', kd_pelayanan='', operator='',
                   kd_pasien='', nm_pasien=''):
    """Potong qty dari kit perawat FEFO (exp terdekat dulu). Boleh minus:
       bila kit tak cukup, sisa kekurangan dibebankan ke batch no-exp (jadi negatif → 'tandai kurang').
       Return (kurang) jumlah yang tak tertutup stok kit positif."""
    batches = query("""SELECT exp, qty_sisa FROM Stok_Perawat
                       WHERE kd_karyawan=? AND kd_barang=? AND qty_sisa>0
                       ORDER BY CASE WHEN COALESCE(exp,'')='' THEN 1 ELSE 0 END, exp""",
                    (kd_karyawan, kd_barang))
    remaining = qty_needed
    for b in batches:
        if remaining <= 0:
            break
        take = min(b['qty_sisa'], remaining)
        if take <= 0:
            continue
        _kit_apply(kd_karyawan, kd_barang, nm_barang, b['exp'] or '', -take, satuan=satuan,
                   jenis=jenis, ref=ref, kd_pelayanan=kd_pelayanan, operator=operator,
                   kd_pasien=kd_pasien, nm_pasien=nm_pasien)
        remaining -= take
    if remaining > 0:
        # Kit kurang → bebankan ke bucket no-exp (saldo jadi minus = tanda perlu isi ulang)
        _kit_apply(kd_karyawan, kd_barang, nm_barang, '', -remaining, satuan=satuan,
                   jenis=jenis, ref=ref, kd_pelayanan=kd_pelayanan, operator=operator,
                   kd_pasien=kd_pasien, nm_pasien=nm_pasien)
    return remaining


def _kit_settle(kd_karyawan, kd_barang, operator=''):
    """Lunasi 'utang' kit: saldo minus di bucket exp='' ditutup pakai stok batch ber-exp (FEFO).
       Net tetap (konsolidasi) supaya tak ada minus menggantung saat sudah ada isi ulang."""
    neg = query("""SELECT qty_sisa FROM Stok_Perawat
                   WHERE kd_karyawan=? AND kd_barang=? AND COALESCE(exp,'')='' AND qty_sisa<0""",
                (kd_karyawan, kd_barang), one=True)
    if not neg:
        return
    debt = -(neg['qty_sisa'] or 0)
    if debt <= 0:
        return
    pos = query("""SELECT exp, qty_sisa FROM Stok_Perawat
                   WHERE kd_karyawan=? AND kd_barang=? AND qty_sisa>0 AND COALESCE(exp,'')!=''
                   ORDER BY exp""",
                (kd_karyawan, kd_barang))
    for b in pos:
        if debt <= 0:
            break
        take = min(b['qty_sisa'], debt)
        if take <= 0:
            continue
        _kit_apply(kd_karyawan, kd_barang, '', b['exp'] or '', -take,
                   jenis='koreksi', ref='Pelunasan minus kit', operator=operator)
        _kit_apply(kd_karyawan, kd_barang, '', '', take,
                   jenis='koreksi', ref='Pelunasan minus kit', operator=operator)
        debt -= take
    # Hapus bucket exp='' bila sudah lunas (0)
    execute("""DELETE FROM Stok_Perawat
               WHERE kd_karyawan=? AND kd_barang=? AND COALESCE(exp,'')='' AND qty_sisa=0""",
            (kd_karyawan, kd_barang))


# ─── SERAH-TERIMA BAHAN KUNCI (QR handshake) ──────────────────
# Bahan treatment yg scarce/terkontrol ("kunci") tidak langsung dipotong dari kit
# pelaksana spt bahan biasa (lihat kit_pakai_batch) — harus diserahterimakan lewat
# QR sekali-pakai dari Perawat pemegang bahan (K) ke Pelaksana (A) dulu. Komisi
# treatment terkait ditahan (Point_karyawan.status='pending') sampai SEMUA bahan
# kunci treatment itu selesai diserahterimakan.

def _is_perawat(kd_karyawan):
    row = query("SELECT pekerjaan FROM Karyawan WHERE kd_karyawan=?", (kd_karyawan,), one=True)
    return bool(row) and row['pekerjaan'] == 'Perawat'

def _is_key_material(kd_barang):
    row = query("SELECT 1 FROM Perawatan_Bahan WHERE kd_barang=? AND bahan_kunci=1 LIMIT 1", (kd_barang,), one=True)
    return bool(row)

def _perawat_punya_bahan(kd_karyawan, kd_barang):
    """True kalau Perawat ini punya stok kd_barang > 0 di kit-nya — dipakai utk validasi
    K yg dipilih kasir, supaya tidak bisa asal pilih perawat yg sebenarnya tak pegang
    bahan kuncinya sama sekali."""
    row = query("SELECT 1 FROM Stok_Perawat WHERE kd_karyawan=? AND kd_barang=? AND qty_sisa>0 LIMIT 1",
                (kd_karyawan, kd_barang), one=True)
    return bool(row)

def _serah_no():
    ddmmyy = datetime.now().strftime('%d%m%y')
    pfx = f'ST{ddmmyy}'
    last = query("SELECT no_serah FROM Serah_Terima_Kunci WHERE no_serah LIKE ? ORDER BY no_serah DESC LIMIT 1",
                (f'{pfx}%',), one=True)
    seq = int(last['no_serah'][-3:]) + 1 if last else 1
    return f'{pfx}{seq:03d}'

def _release_komisi_kunci(no_faktur, kd_pelayanan, kd_k, persen, nm_pelayanan='', kd_pasien=''):
    """Cairkan komisi treatment ini: tiap baris Point_karyawan pending utk (no_faktur,
    kd_pelayanan) dipecah — persen% jadi baris baru milik K (released), sisanya tetap
    di pemilik asli (released)."""
    rows = query("SELECT * FROM Point_karyawan WHERE no_faktur=? AND kd_pelayanan=? AND status='pending'",
                 (no_faktur, kd_pelayanan))
    for r in rows:
        total = r['komisi'] or 0
        k_share = round(total * persen / 100, 2)
        sisa = total - k_share
        execute("UPDATE Point_karyawan SET komisi=?, status='released' WHERE id=?", (sisa, r['id']))
        if k_share > 0:
            execute("""INSERT INTO Point_karyawan (no_faktur,tanggal,kd_pasien,kd_pelayanan,nm_pelayanan,komisi,kd_karyawan,status)
                       VALUES (?,?,?,?,?,?,?,'released')""",
                    (no_faktur, str(date.today()), kd_pasien, kd_pelayanan, nm_pelayanan, k_share, kd_k))

def _serah_execute(rec, op=''):
    """Eksekusi 1 baris serah-terima: potong bahan dari kit K (FEFO, boleh minus — reuse
    _kit_take_fefo yg sudah ada), tandai selesai, lalu cek apakah SEMUA bahan kunci
    treatment ini sudah selesai → kalau ya, cairkan komisinya."""
    _kit_take_fefo(rec['kd_k_final'], rec['kd_barang'], rec['nm_barang'], rec['qty'],
                   satuan=rec['satuan'] or '', jenis='pakai_treatment', ref=rec['no_faktur'],
                   kd_pelayanan=rec['kd_pelayanan'], operator=op,
                   kd_pasien=rec['kd_pasien'] or '', nm_pasien=rec['nm_pasien'] or '')
    execute("UPDATE Serah_Terima_Kunci SET status='selesai', selesai_at=? WHERE id=?",
            (str(datetime.now()), rec['id']))
    remaining = query("""SELECT COUNT(*) as c FROM Serah_Terima_Kunci
                         WHERE no_faktur=? AND kd_pelayanan=? AND status IN ('pending','terlewat')""",
                      (rec['no_faktur'], rec['kd_pelayanan']), one=True)
    if not remaining or remaining['c'] == 0:
        _release_komisi_kunci(rec['no_faktur'], rec['kd_pelayanan'], rec['kd_k_final'],
                              rec['komisi_persen'] or 0, rec['nm_pelayanan'] or '', rec['kd_pasien'] or '')

@app.route('/api/serah-terima/perawat-punya')
def serah_perawat_punya():
    """Daftar Perawat yang punya stok bahan kunci ini di kit-nya (utk kasir pilih K)."""
    kd_barang = request.args.get('kd_barang', '')
    rows = query("""
        SELECT sp.kd_karyawan, COALESCE(k.nickname, k.nm_karyawan, sp.kd_karyawan) as nama,
               sp.qty_sisa, sp.satuan
        FROM Stok_Perawat sp
        JOIN Karyawan k ON sp.kd_karyawan=k.kd_karyawan
        WHERE sp.kd_barang=? AND k.pekerjaan='Perawat' AND k.aktif=1 AND sp.qty_sisa>0
        ORDER BY sp.qty_sisa DESC
    """, (kd_barang,))
    return jsonify(rows_to_list(rows))

@app.route('/api/serah-terima/create', methods=['POST'])
def serah_terima_create():
    """kd_k_dipilih KOSONG = kasir tidak pilih Perawat K sama sekali (dibolehkan, sesuai
    keputusan: bahan kunci tidak wajib/tidak blokir bayar). Tetap dicatat sbg status
    'terlewat' (bukan 'pending') — bahan TIDAK terpotong dari kit siapa pun, dan komisi
    treatment ini (sudah dibekukan 'pending' oleh frontend saat bayar) TETAP dibekukan
    sampai kasir menyelesaikannya lewat /pilih-perawat (lihat _serah_execute — 'terlewat'
    ikut dihitung sbg penghalang pencairan)."""
    d = request.json or {}
    kd_k = (d.get('kd_k_dipilih') or '').strip()
    kd_barang = d.get('kd_barang', '')
    if kd_k:
        if not _is_perawat(kd_k):
            return jsonify({'ok': False, 'msg': 'Perawat K tidak valid — harus role Perawat aktif'}), 400
        if not _perawat_punya_bahan(kd_k, kd_barang):
            return jsonify({'ok': False, 'msg': 'Perawat ini tidak punya stok bahan kunci tsb di kit-nya'}), 400
        status = 'pending'
    else:
        status = 'terlewat'
    no_serah = _serah_no()
    rid = execute("""INSERT INTO Serah_Terima_Kunci
        (no_serah,no_faktur,kd_pelayanan,nm_pelayanan,kd_barang,nm_barang,qty,satuan,
         kd_pelaksana,nm_pelaksana,kd_k_dipilih,komisi_persen,kd_pasien,nm_pasien,status)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (no_serah, d.get('no_faktur',''), d.get('kd_pelayanan',''), d.get('nm_pelayanan',''),
         d.get('kd_barang',''), d.get('nm_barang',''), float(d.get('qty',1) or 1), d.get('satuan',''),
         d.get('kd_pelaksana',''), d.get('nm_pelaksana',''), kd_k, float(d.get('komisi_persen',0) or 0),
         d.get('kd_pasien',''), d.get('nm_pasien',''), status))
    return jsonify({'ok': True, 'id': rid, 'no_serah': no_serah, 'status': status})

@app.route('/api/serah-terima')
def serah_terima_list():
    status = request.args.get('status', '')
    no_faktur = request.args.get('no_faktur', '')
    where, params = [], []
    if status: where.append("stk.status=?"); params.append(status)
    if no_faktur: where.append("stk.no_faktur=?"); params.append(no_faktur)
    w = ('WHERE ' + ' AND '.join(where)) if where else ''
    rows = query(f"""
        SELECT stk.*, COALESCE(k.nickname, k.nm_karyawan) as nm_k
        FROM Serah_Terima_Kunci stk
        LEFT JOIN Karyawan k ON k.kd_karyawan = CASE WHEN stk.kd_k_final!='' THEN stk.kd_k_final ELSE stk.kd_k_dipilih END
        {w}
        ORDER BY stk.created_at DESC LIMIT 300
    """, params)
    return jsonify(rows_to_list(rows))

@app.route('/api/serah-terima/pending')
def serah_terima_pending():
    kd_karyawan = request.args.get('kd_karyawan', '')
    peran = request.args.get('peran', 'kunci')
    if peran == 'terlewat':
        # Peringatan utk Pelaksana: bahan kunci yg dipakainya tidak lewat serah-terima
        # resmi (K tak dipilih kasir) — murni informasi, tidak ada aksi.
        rows = query("""SELECT * FROM Serah_Terima_Kunci WHERE kd_pelaksana=? AND status='terlewat'
                        ORDER BY created_at DESC LIMIT 30""", (kd_karyawan,))
    elif peran == 'kunci':
        # Perawat K perlu tahu berapa komisi yg BAKAL didapat begitu serah-terima ini
        # selesai — dihitung dari komisi pelaksana yg masih pending × komisi_persen.
        rows = query("""
            SELECT stk.*,
                   COALESCE((SELECT SUM(pk.komisi) FROM Point_karyawan pk
                             WHERE pk.no_faktur=stk.no_faktur AND pk.kd_pelayanan=stk.kd_pelayanan
                               AND pk.kd_karyawan=stk.kd_pelaksana AND pk.status='pending'), 0) as komisi_pelaksana_pending
            FROM Serah_Terima_Kunci stk
            WHERE stk.kd_k_dipilih=? AND stk.status='pending'
            ORDER BY stk.created_at ASC
        """, (kd_karyawan,))
        result = rows_to_list(rows)
        for r in result:
            r['komisi_k_estimasi'] = round(r['komisi_pelaksana_pending'] * (r['komisi_persen'] or 0) / 100, 2)
        return jsonify(result)
    else:
        rows = query("SELECT * FROM Serah_Terima_Kunci WHERE kd_pelaksana=? AND status='pending' ORDER BY created_at ASC",
                    (kd_karyawan,))
    return jsonify(rows_to_list(rows))

@app.route('/api/serah-terima/<int:sid>/pilih-perawat', methods=['POST'])
def serah_terima_pilih_perawat(sid):
    """Kasir memilih (atau mengganti) Perawat K. Juga dipakai utk menyelesaikan record
    'terlewat' (K tak dipilih saat bayar) — begitu K dipilih di sini, status balik jadi
    'pending' dan alur QR/scan normal berjalan, komisi yg sempat dibekukan ikut cair
    setelah serah-terima ini selesai."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor', 'kasir'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor/kasir'}), 403
    d = request.json or {}
    kd_k = (d.get('kd_k_dipilih') or '').strip()
    if not _is_perawat(kd_k):
        return jsonify({'ok': False, 'msg': 'Perawat K tidak valid — harus role Perawat aktif'}), 400
    rec = query("SELECT status, kd_barang FROM Serah_Terima_Kunci WHERE id=?", (sid,), one=True)
    if not rec or rec['status'] not in ('pending', 'terlewat'):
        return jsonify({'ok': False, 'msg': 'Serah-terima tidak ditemukan / sudah diproses'}), 400
    if not _perawat_punya_bahan(kd_k, rec['kd_barang']):
        return jsonify({'ok': False, 'msg': 'Perawat ini tidak punya stok bahan kunci tsb di kit-nya'}), 400
    # Reset token/QR yg mungkin sudah digenerate K sebelumnya — kalau K diganti, QR lama
    # (kalau ada) jadi tidak valid lagi, K yg baru wajib generate ulang dari mobile-nya.
    execute("""UPDATE Serah_Terima_Kunci SET kd_k_dipilih=?, status='pending',
               token='', token_exp='', kd_k_final='' WHERE id=?""", (kd_k, sid))
    return jsonify({'ok': True})

@app.route('/api/serah-terima/qr', methods=['POST'])
def serah_terima_qr():
    """Perawat K generate 1 token QR utk SEMUA baris pending miliknya pd faktur+pelaksana yg sama."""
    d = request.json or {}
    no_faktur = d.get('no_faktur', '')
    kd_k = (d.get('kd_karyawan') or '').strip()
    kd_pelaksana = d.get('kd_pelaksana', '')
    rows = query("""SELECT id FROM Serah_Terima_Kunci
                    WHERE no_faktur=? AND kd_k_dipilih=? AND kd_pelaksana=? AND status='pending'""",
                (no_faktur, kd_k, kd_pelaksana))
    if not rows:
        return jsonify({'ok': False, 'msg': 'Tidak ada serah-terima pending utk kombinasi ini'}), 400
    token = str(uuid.uuid4())
    exp = str(datetime.now() + timedelta(minutes=5))
    ids = [r['id'] for r in rows]
    for rid in ids:
        execute("UPDATE Serah_Terima_Kunci SET token=?, token_exp=?, kd_k_final=? WHERE id=?",
                (token, exp, kd_k, rid))
    items = rows_to_list(query(f"SELECT nm_barang, qty, satuan FROM Serah_Terima_Kunci WHERE id IN ({','.join('?'*len(ids))})", ids))
    return jsonify({'ok': True, 'token': token, 'exp': exp, 'count': len(ids), 'items': items})

@app.route('/api/serah-terima/scan', methods=['POST'])
def serah_terima_scan():
    """Pelaksana A konfirmasi terima — WAJIB login sbg kd_pelaksana yg tertera di token."""
    user = get_current_user()
    d = request.json or {}
    token = (d.get('token') or '').strip()
    if not token:
        return jsonify({'ok': False, 'msg': 'Kode/token wajib diisi'}), 400
    rows = query("SELECT * FROM Serah_Terima_Kunci WHERE token=? AND status='pending'", (token,))
    if not rows:
        return jsonify({'ok': False, 'msg': 'Token tidak valid atau sudah dipakai'}), 400
    now = str(datetime.now())
    if rows[0]['token_exp'] and rows[0]['token_exp'] < now:
        return jsonify({'ok': False, 'msg': 'Token sudah kedaluwarsa (>5 menit) — minta K generate ulang'}), 400
    kd_pelaksana = rows[0]['kd_pelaksana']
    if not user or user.get('kd_karyawan') != kd_pelaksana:
        return jsonify({'ok': False, 'msg': f"Harus login sebagai pelaksana yang dituju ({rows[0]['nm_pelaksana']})"}), 403
    op = user.get('nama', '')
    for rec in rows:
        _serah_execute(dict(rec), op=op)
    return jsonify({'ok': True, 'count': len(rows)})

@app.route('/api/kit/perawat')
def kit_perawat_list():
    """Ringkasan kit semua perawat yang punya isi (untuk admin/supervisor)."""
    rows = query("""
        SELECT sp.kd_karyawan,
               COALESCE(k.nickname, k.nm_karyawan, sp.kd_karyawan) AS nama,
               k.nm_karyawan, k.pekerjaan,
               COUNT(*) AS n_item,
               SUM(sp.qty_sisa) AS total_qty,
               SUM(CASE WHEN sp.qty_sisa < 0 THEN 1 ELSE 0 END) AS n_minus
        FROM Stok_Perawat sp
        LEFT JOIN Karyawan k ON sp.kd_karyawan = k.kd_karyawan
        GROUP BY sp.kd_karyawan
        ORDER BY nama
    """)
    return jsonify(rows_to_list(rows))


@app.route('/api/kit/<kd_karyawan>')
def kit_perawat_detail(kd_karyawan):
    kar = query("SELECT kd_karyawan, nm_karyawan, nickname, pekerjaan FROM Karyawan WHERE kd_karyawan=?",
                (kd_karyawan,), one=True)
    items = query("""SELECT kd_barang, nm_barang, exp, qty_sisa, satuan
                     FROM Stok_Perawat WHERE kd_karyawan=?
                     ORDER BY (qty_sisa < 0) DESC, nm_barang, exp""", (kd_karyawan,))
    result = rows_to_list(items)
    # Tandai per barang berapa yg sudah "dipesan" lewat serah-terima kunci pending (K=user
    # ini) — supaya kelihatan stok yg BENAR-BENAR bebas dipakai, bukan cuma qty_sisa mentah.
    reserved_rows = query("""SELECT kd_barang, COALESCE(SUM(qty),0) as r
                             FROM Serah_Terima_Kunci WHERE kd_k_dipilih=? AND status='pending'
                             GROUP BY kd_barang""", (kd_karyawan,))
    reserved_map = {r['kd_barang']: r['r'] for r in reserved_rows}
    # Detail siapa yg menunggu (per barang) — supaya "X dipesan" bisa nyebut namanya, bukan
    # cuma angka mentah.
    reserved_detail_rows = query("""SELECT kd_barang, nm_pelaksana, COALESCE(SUM(qty),0) as r
                                    FROM Serah_Terima_Kunci WHERE kd_k_dipilih=? AND status='pending'
                                    GROUP BY kd_barang, nm_pelaksana""", (kd_karyawan,))
    reserved_detail_map = {}
    for r in reserved_detail_rows:
        qv = r['r']
        qtxt = str(int(qv)) if float(qv) == int(qv) else str(qv)
        reserved_detail_map.setdefault(r['kd_barang'], []).append(
            f"{r['nm_pelaksana'] or '—'} ({qtxt})")
    for it in result:
        it['reserved_serah'] = reserved_map.get(it['kd_barang'], 0)
        it['reserved_pelaksana'] = ', '.join(reserved_detail_map.get(it['kd_barang'], []))
    pending_count = query("SELECT COUNT(*) as c FROM Serah_Terima_Kunci WHERE kd_k_dipilih=? AND status='pending'",
                          (kd_karyawan,), one=True)
    return jsonify({
        'karyawan': dict(kar) if kar else {'kd_karyawan': kd_karyawan},
        'items': result,
        'pending_serah_count': pending_count['c'] if pending_count else 0,
    })


@app.route('/api/kit/<kd_karyawan>/log')
def kit_perawat_log(kd_karyawan):
    limit = int(request.args.get('limit', 50))
    jenis = request.args.get('jenis', '')
    kd_barang = request.args.get('kd_barang', '')
    where, params = ["kd_karyawan=?"], [kd_karyawan]
    if jenis:
        where.append("jenis=?"); params.append(jenis)
    if kd_barang:
        where.append("kd_barang=?"); params.append(kd_barang)
    params.append(limit)
    # via_serah_terima=1 → bahan ini keluar dari kit krn dipakai BARENG perawat lain lewat
    # QR handshake (bahan kunci: pelaksana tanpa kit ini wajib lewat pemegangnya, bukan
    # "pinjam" krn barangnya langsung habis dipakai, bukan kembali) — bukan dipakai sendiri.
    rows = query(f"""SELECT kl.tanggal,kl.kd_barang,kl.nm_barang,kl.exp,kl.qty,kl.saldo,kl.jenis,kl.ref,
                            kl.kd_pelayanan,COALESCE(pp.nm_pelayanan,kl.kd_pelayanan) as nm_pelayanan,
                            kl.kd_pasien,kl.nm_pasien,kl.operator,
                            EXISTS(SELECT 1 FROM Serah_Terima_Kunci stk
                                   WHERE stk.no_faktur=kl.ref AND stk.kd_pelayanan=kl.kd_pelayanan
                                     AND stk.kd_k_final=kl.kd_karyawan) as via_serah_terima,
                            (SELECT stk.nm_pelaksana FROM Serah_Terima_Kunci stk
                             WHERE stk.no_faktur=kl.ref AND stk.kd_pelayanan=kl.kd_pelayanan
                               AND stk.kd_k_final=kl.kd_karyawan LIMIT 1) as st_nm_pelaksana
                     FROM Kit_Log kl
                     LEFT JOIN produk_pelayanan pp ON kl.kd_pelayanan=pp.kd_pelayanan
                     WHERE {' AND '.join(f'kl.{w}' for w in where)} ORDER BY kl.id DESC LIMIT ?""",
                 params)
    return jsonify(rows_to_list(rows))


@app.route('/api/kit/koreksi', methods=['POST'])
def kit_koreksi():
    """Admin/supervisor/kasir: koreksi/seed manual stok kit perawat.
       mode 'add' → tambah qty; mode 'set' → jadikan qty target (hitung selisih)."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor', 'kasir'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor/kasir'}), 403
    d = request.json or {}
    kd_kar = str(d.get('kd_karyawan', '')).strip()
    kd_brg = str(d.get('kd_barang', '')).strip()
    if not kd_kar or not kd_brg:
        return jsonify({'ok': False, 'msg': 'Perawat & barang wajib diisi'}), 400
    try:
        qty = float(d.get('qty', 0))
    except (TypeError, ValueError):
        qty = 0
    exp  = d.get('exp', '') or ''
    mode = d.get('mode', 'add')
    if mode == 'set':
        cur = query("""SELECT qty_sisa FROM Stok_Perawat
                       WHERE kd_karyawan=? AND kd_barang=? AND COALESCE(exp,'')=? LIMIT 1""",
                    (kd_kar, kd_brg, exp), one=True)
        delta = qty - (cur['qty_sisa'] if cur else 0)
    else:
        delta = qty
    if delta == 0:
        return jsonify({'ok': True, 'msg': 'Tidak ada perubahan'})
    op = user.get('nama') if user else ''
    saldo = _kit_apply(kd_kar, kd_brg, d.get('nm_barang', ''), exp, delta,
                       satuan=d.get('satuan', ''), jenis='koreksi',
                       ref=d.get('catatan', 'koreksi manual'), operator=op)
    if delta > 0:
        _kit_settle(kd_kar, kd_brg, op)   # lunasi minus menggantung bila ada
        saldo = query("""SELECT COALESCE(SUM(qty_sisa),0) AS s FROM Stok_Perawat
                         WHERE kd_karyawan=? AND kd_barang=?""", (kd_kar, kd_brg), one=True)['s']
    catat_log('koreksi_kit', 'Stok_Perawat', f"{kd_kar}/{kd_brg}",
              f"Koreksi kit perawat {kd_kar} - '{d.get('nm_barang','')}': delta {delta:+g}, saldo akhir {saldo}",
              data_baru={'kd_karyawan': kd_kar, 'kd_barang': kd_brg, 'delta': delta, 'saldo': saldo,
                         'catatan': d.get('catatan','')})
    return jsonify({'ok': True, 'saldo': saldo})


# ─── ISI ULANG KIT (request perawat → approval admin) ────────
def _cabang_take_fefo(kd_barang, kd_lokasi, qty_needed, keterangan):
    """Tarik qty dari stok cabang secara FEFO (exp paling awal dulu).
       Kurangi Stok_barang.stok_sisa + tulis RiwayatBarang per batch (−).
       Return (taken_list, shortage). taken_list=[{'exp','qty'}]; shortage>0 bila stok kurang.
       TIDAK membuat stok cabang minus (hanya ambil yang ada)."""
    today = str(date.today())
    batches = query("""SELECT id, exp, stok_sisa FROM Stok_barang
                       WHERE kd_barang=? AND kd_lokasi=? AND stok_sisa>0
                       ORDER BY CASE WHEN COALESCE(exp,'')='' THEN 1 ELSE 0 END, exp, id""",
                    (kd_barang, kd_lokasi))
    taken, remaining = [], qty_needed
    for b in batches:
        if remaining <= 0:
            break
        take = min(b['stok_sisa'], remaining)
        if take <= 0:
            continue
        execute("UPDATE Stok_barang SET stok_sisa=stok_sisa-? WHERE id=?", (take, b['id']))
        sn = query("SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE id=?", (b['id'],), one=True)
        execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                (today, kd_barang, b['exp'] or '', keterangan, -take,
                 sn['stok_sisa'] if sn else 0, sn['stok_gudang'] if sn else 0))
        taken.append({'exp': b['exp'] or '', 'qty': take})
        remaining -= take
    return taken, remaining


@app.route('/api/kit/request', methods=['POST'])
def kit_request_create():
    """Perawat ajukan isi ulang kit (dari mobile). kd_karyawan/kd_lokasi diambil dari sesi bila tak dikirim."""
    user = get_current_user()
    d = request.json or {}
    kd_kar = (d.get('kd_karyawan') or (user.get('kd_karyawan') if user else '') or '').strip()
    kd_lok = (d.get('kd_lokasi') or (user.get('kd_lokasi') if user else '') or 'CBG1').strip()
    items = [it for it in d.get('items', [])
             if str(it.get('kd_barang', '')).strip() and float(it.get('qty', 0) or 0) > 0]
    if not kd_kar:
        return jsonify({'ok': False, 'msg': 'Akun tidak tertaut ke karyawan'}), 400
    if not items:
        return jsonify({'ok': False, 'msg': 'Minimal 1 barang'}), 400
    today = date.today()
    cnt = query("SELECT COUNT(*) AS c FROM Kit_Request WHERE tgl_request=?", (str(today),), one=True)
    no_req = f"KR{today.strftime('%d%m%y')}{(cnt['c']+1):03d}"
    rid = execute("""INSERT INTO Kit_Request (no_request,kd_karyawan,kd_lokasi,catatan,status)
                     VALUES (?,?,?,?,'pending')""",
                  (no_req, kd_kar, kd_lok, d.get('catatan', '')))
    for it in items:
        execute("""INSERT INTO Kit_Request_Detail (id_request,kd_barang,nm_barang,qty,satuan)
                   VALUES (?,?,?,?,?)""",
                (rid, str(it['kd_barang']).strip(), it.get('nm_barang', ''),
                 float(it.get('qty', 0) or 0), it.get('satuan', '')))
    catat_log('kit_request_ajukan', 'Kit_Request', no_req,
              f"Ajukan isi ulang kit {no_req} — {len(items)} barang")
    return jsonify({'ok': True, 'id': rid, 'no_request': no_req})


@app.route('/api/kit/request')
def kit_request_list():
    status = request.args.get('status', '')
    kd_kar = request.args.get('kd_karyawan', '')
    where, params = [], []
    if status:
        where.append('r.status=?'); params.append(status)
    if kd_kar:
        where.append('r.kd_karyawan=?'); params.append(kd_kar)
    sql = """SELECT r.*, COALESCE(k.nickname,k.nm_karyawan,r.kd_karyawan) AS nama,
                    (SELECT COUNT(*) FROM Kit_Request_Detail d WHERE d.id_request=r.id) AS n_item
             FROM Kit_Request r LEFT JOIN Karyawan k ON r.kd_karyawan=k.kd_karyawan"""
    if where:
        sql += ' WHERE ' + ' AND '.join(where)
    sql += ' ORDER BY r.id DESC LIMIT 200'
    return jsonify(rows_to_list(query(sql, params)))


@app.route('/api/kit/request/<int:rid>')
def kit_request_detail(rid):
    r = query("""SELECT r.*, COALESCE(k.nickname,k.nm_karyawan,r.kd_karyawan) AS nama
                 FROM Kit_Request r LEFT JOIN Karyawan k ON r.kd_karyawan=k.kd_karyawan
                 WHERE r.id=?""", (rid,), one=True)
    if not r:
        return jsonify(None)
    res = dict(r)
    res['items'] = []
    for it in query("SELECT * FROM Kit_Request_Detail WHERE id_request=?", (rid,)):
        di = dict(it)
        av = query("SELECT COALESCE(SUM(stok_sisa),0) AS s FROM Stok_barang WHERE kd_barang=? AND kd_lokasi=?",
                   (it['kd_barang'], r['kd_lokasi']), one=True)
        di['tersedia'] = av['s'] if av else 0
        res['items'].append(di)
    return jsonify(res)


@app.route('/api/kit/request/<int:rid>/approve', methods=['POST'])
def kit_request_approve(rid):
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor', 'kasir'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor/kasir'}), 403
    r = query("SELECT * FROM Kit_Request WHERE id=?", (rid,), one=True)
    if not r:
        return jsonify({'ok': False, 'msg': 'Permintaan tidak ditemukan'}), 404
    if r['status'] != 'pending':
        return jsonify({'ok': False, 'msg': f"Permintaan sudah {r['status']}"}), 400
    items = query("SELECT * FROM Kit_Request_Detail WHERE id_request=?", (rid,))
    kd_lok = r['kd_lokasi']
    # Validasi ketersediaan DULU agar stok cabang tak jadi minus & tak ada perubahan parsial
    short = []
    for it in items:
        av = query("SELECT COALESCE(SUM(stok_sisa),0) AS s FROM Stok_barang WHERE kd_barang=? AND kd_lokasi=?",
                   (it['kd_barang'], kd_lok), one=True)
        avail = av['s'] if av else 0
        if it['qty'] > avail:
            short.append({'kd_barang': it['kd_barang'], 'nm_barang': it['nm_barang'],
                          'minta': it['qty'], 'tersedia': avail})
    if short:
        return jsonify({'ok': False, 'msg': 'Stok cabang tidak cukup untuk sebagian barang',
                        'shortages': short}), 400
    op = user.get('nama') if user else ''
    for it in items:
        taken, _ = _cabang_take_fefo(it['kd_barang'], kd_lok, it['qty'],
                                     f"Isi kit {r['no_request']} → {r['kd_karyawan']}")
        for t in taken:
            _kit_apply(r['kd_karyawan'], it['kd_barang'], it['nm_barang'], t['exp'], t['qty'],
                       satuan=it['satuan'], jenis='isi_ulang', ref=r['no_request'],
                       kd_lokasi=kd_lok, operator=op)
        _kit_settle(r['kd_karyawan'], it['kd_barang'], op)   # lunasi minus menggantung bila ada
    execute("UPDATE Kit_Request SET status='approved', approved_by=?, approved_at=datetime('now') WHERE id=?",
            (op, rid))
    catat_log('kit_request_approve', 'Kit_Request', r['no_request'],
              f"Setujui isi ulang kit {r['no_request']} untuk {r['kd_karyawan']}")
    return jsonify({'ok': True})


@app.route('/api/kit/request/<int:rid>/reject', methods=['POST'])
def kit_request_reject(rid):
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor', 'kasir'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor/kasir'}), 403
    d = request.json or {}
    r = query("SELECT status, no_request, kd_karyawan FROM Kit_Request WHERE id=?", (rid,), one=True)
    if not r:
        return jsonify({'ok': False, 'msg': 'Tidak ditemukan'}), 404
    if r['status'] != 'pending':
        return jsonify({'ok': False, 'msg': f"Permintaan sudah {r['status']}"}), 400
    execute("UPDATE Kit_Request SET status='rejected', alasan=?, approved_by=?, approved_at=datetime('now') WHERE id=?",
            (d.get('alasan', ''), (user.get('nama') if user else ''), rid))
    catat_log('kit_request_reject', 'Kit_Request', r['no_request'],
              f"Tolak isi ulang kit {r['no_request']} untuk {r['kd_karyawan']}" + (f" — {d.get('alasan')}" if d.get('alasan') else ''))
    return jsonify({'ok': True})


@app.route('/api/kit/request/<int:rid>', methods=['DELETE'])
def kit_request_delete(rid):
    """Batalkan permintaan yang masih pending (perawat sendiri / admin)."""
    r = query("SELECT status, no_request, kd_karyawan FROM Kit_Request WHERE id=?", (rid,), one=True)
    if not r:
        return jsonify({'ok': False, 'msg': 'Tidak ditemukan'}), 404
    if r['status'] != 'pending':
        return jsonify({'ok': False, 'msg': 'Hanya permintaan pending bisa dibatalkan'}), 400
    execute("DELETE FROM Kit_Request_Detail WHERE id_request=?", (rid,))
    execute("DELETE FROM Kit_Request WHERE id=?", (rid,))
    catat_log('kit_request_batal', 'Kit_Request', r['no_request'],
              f"Batalkan permintaan isi ulang kit {r['no_request']} untuk {r['kd_karyawan']}")
    return jsonify({'ok': True})


# ─── POTONG KIT SAAT TREATMENT TERJUAL (POS) ─────────────────
@app.route('/api/perawatan/with-formula')
def perawatan_with_formula():
    """Daftar kd_pelayanan yang punya bahan potong_stok — POS pakai untuk tahu treatment mana yang potong kit."""
    rows = query("""SELECT DISTINCT kd_pelayanan FROM Perawatan_Bahan
                    WHERE tipe='bahan' AND COALESCE(potong_stok,1)=1""")
    return jsonify([r['kd_pelayanan'] for r in rows])

@app.route('/api/perawatan/with-kunci')
def perawatan_with_kunci():
    """Daftar kd_pelayanan yang punya bahan KUNCI — POS pakai utk tampilkan kontrol pilih K."""
    rows = query("SELECT DISTINCT kd_pelayanan FROM Perawatan_Bahan WHERE tipe='bahan' AND bahan_kunci=1")
    return jsonify([r['kd_pelayanan'] for r in rows])

@app.route('/api/perawatan/<kd_pelayanan>/bahan-kunci')
def perawatan_bahan_kunci(kd_pelayanan):
    """Rincian bahan kunci utk 1 treatment (kasir pilih K per baris) + persen komisi kunci-nya."""
    bahan = query("""SELECT kd_barang, nm_barang, qty, satuan FROM Perawatan_Bahan
                     WHERE kd_pelayanan=? AND tipe='bahan' AND bahan_kunci=1""", (kd_pelayanan,))
    pel = query("SELECT komisi_persen_kunci FROM produk_pelayanan WHERE kd_pelayanan=?", (kd_pelayanan,), one=True)
    return jsonify({'bahan': rows_to_list(bahan), 'komisi_persen_kunci': (pel['komisi_persen_kunci'] if pel else 0) or 0})


@app.route('/api/kit/pakai/batch', methods=['POST'])
def kit_pakai_batch():
    """Potong bahan kit untuk treatment yang terjual (dipanggil saat bayar).
       body: {no_faktur, items:[{kd_pelayanan, kd_karyawan(kit), qty}]}.
       Boleh minus → kekurangan dikembalikan di 'kurang' untuk ditandai."""
    user = get_current_user()
    op = user.get('nama') if user else ''
    d = request.json or {}
    no_faktur = d.get('no_faktur', '')
    kd_pasien = d.get('kd_pasien', '')
    nm_pasien = d.get('nm_pasien', '')
    kurang_all = []
    for it in d.get('items', []):
        kd_pel = str(it.get('kd_pelayanan', '')).strip()
        kd_kar = str(it.get('kd_karyawan', '')).strip()
        try:
            qty_tr = float(it.get('qty', 1) or 1)
        except (TypeError, ValueError):
            qty_tr = 1
        if not kd_pel or not kd_kar:
            continue
        # Bahan kunci TIDAK ikut potong langsung di sini — jalurnya lewat serah-terima QR
        # (lihat serah_terima_create/_serah_execute), baru terpotong setelah A konfirmasi.
        bahan = query("""SELECT kd_barang, nm_barang, qty, satuan FROM Perawatan_Bahan
                         WHERE kd_pelayanan=? AND tipe='bahan' AND COALESCE(potong_stok,1)=1
                           AND COALESCE(bahan_kunci,0)=0""",
                      (kd_pel,))
        for b in bahan:
            need = (b['qty'] or 0) * qty_tr
            if need <= 0:
                continue
            kurang = _kit_take_fefo(kd_kar, b['kd_barang'], b['nm_barang'], need,
                                    satuan=b['satuan'], jenis='pakai_treatment',
                                    ref=no_faktur, kd_pelayanan=kd_pel, operator=op,
                                    kd_pasien=kd_pasien, nm_pasien=nm_pasien)
            if kurang > 0:
                kurang_all.append({'kd_karyawan': kd_kar, 'kd_barang': b['kd_barang'],
                                   'nm_barang': b['nm_barang'], 'kurang': kurang})
    n_treatment = len([1 for it in d.get('items', [])
                       if str(it.get('kd_pelayanan','')).strip() and str(it.get('kd_karyawan','')).strip()])
    if n_treatment:
        catat_log('kit_pakai_treatment', 'Kit_Log', no_faktur,
                  f"Potong bahan kit perawat utk {n_treatment} treatment di transaksi {no_faktur}"
                  + (f" — kurang: {', '.join(sorted({k['nm_barang'] for k in kurang_all}))}" if kurang_all else ''),
                  data_baru={'no_faktur': no_faktur, 'kd_pasien': kd_pasien,
                             'items': d.get('items', []), 'kurang': kurang_all})
    return jsonify({'ok': True, 'kurang': kurang_all})


def _kit_retur(no_faktur, kd_pelayanan, kd_karyawan='', rasio=1, operator=''):
    """Kembalikan bahan kit yang dipakai treatment (baca Kit_Log pakai_treatment, kembalikan ×rasio).
       Return jumlah baris yang dikembalikan."""
    where = "ref=? AND kd_pelayanan=? AND jenis='pakai_treatment'"
    params = [no_faktur, kd_pelayanan]
    if kd_karyawan:
        where += " AND kd_karyawan=?"; params.append(kd_karyawan)
    rows = query(f"SELECT kd_karyawan,kd_barang,nm_barang,exp,qty,kd_pasien,nm_pasien FROM Kit_Log WHERE {where}", params)
    n = 0
    for r in rows:
        kembali = -(r['qty'] or 0) * rasio   # qty pakai negatif → kembalikan positif
        if kembali == 0:
            continue
        _kit_apply(r['kd_karyawan'], r['kd_barang'], r['nm_barang'], r['exp'] or '', kembali,
                   jenis='retur_treatment', ref=no_faktur, kd_pelayanan=kd_pelayanan, operator=operator,
                   kd_pasien=dict(r).get('kd_pasien',''), nm_pasien=dict(r).get('nm_pasien',''))
        n += 1
    return n


@app.route('/api/kit/retur', methods=['POST'])
def kit_retur():
    """Kembalikan bahan kit saat treatment diretur.
       body: {no_faktur, kd_pelayanan, kd_karyawan, rasio(0..1 default 1)}."""
    user = get_current_user()
    op = user.get('nama') if user else ''
    d = request.json or {}
    no_faktur = str(d.get('no_faktur', '')).strip()
    kd_pel = str(d.get('kd_pelayanan', '')).strip()
    kd_kar = str(d.get('kd_karyawan', '')).strip()
    try:
        rasio = float(d.get('rasio', 1) or 1)
    except (TypeError, ValueError):
        rasio = 1
    if not no_faktur or not kd_pel:
        return jsonify({'ok': False, 'msg': 'no_faktur & kd_pelayanan wajib'}), 400
    return jsonify({'ok': True, 'restored': _kit_retur(no_faktur, kd_pel, kd_kar, rasio, op)})


# ─── LAPORAN KIT (Tahap 5) ───────────────────────────────────
@app.route('/api/kit/minus')
def kit_minus():
    """Semua item kit yang saldo-nya minus (perlu isi ulang) — lintas perawat."""
    rows = query("""SELECT sp.kd_karyawan, COALESCE(k.nickname,k.nm_karyawan,sp.kd_karyawan) AS nama,
                           sp.kd_barang, sp.nm_barang, sp.exp, sp.qty_sisa, sp.satuan
                    FROM Stok_Perawat sp LEFT JOIN Karyawan k ON sp.kd_karyawan=k.kd_karyawan
                    WHERE sp.qty_sisa < 0
                    ORDER BY nama, sp.nm_barang""")
    return jsonify(rows_to_list(rows))


@app.route('/api/kit/laporan/pemakaian')
def kit_lap_pemakaian():
    """Pemakaian bahan kit per perawat dalam periode (NET = pakai − retur)."""
    dari   = request.args.get('dari', '')
    sampai = request.args.get('sampai', '')
    where, params = ["jenis IN ('pakai_treatment','retur_treatment')"], []
    if dari:
        where.append("date(tanggal)>=?"); params.append(dari)
    if sampai:
        where.append("date(tanggal)<=?"); params.append(sampai)
    rows = query(f"""
        SELECT kl.kd_karyawan, COALESCE(k.nickname,k.nm_karyawan,kl.kd_karyawan) AS nama,
               kl.kd_barang, kl.nm_barang,
               SUM(-kl.qty) AS terpakai,
               COUNT(DISTINCT CASE WHEN kl.jenis='pakai_treatment' THEN kl.ref END) AS n_transaksi
        FROM Kit_Log kl LEFT JOIN Karyawan k ON kl.kd_karyawan=k.kd_karyawan
        WHERE {' AND '.join(where)}
        GROUP BY kl.kd_karyawan, kl.kd_barang
        HAVING SUM(-kl.qty) <> 0
        ORDER BY nama, terpakai DESC
    """, params)
    return jsonify(rows_to_list(rows))


@app.route('/api/kit/pakai-history')
def kit_pakai_history():
    """Riwayat pemakaian bahan kit (pakai_treatment) per transaksi, lintas semua perawat,
    dalam periode tertentu — versi rinci (baris per baris) dari /laporan/pemakaian yg
    cuma total agregat. Dipakai di Laporan Kit Perawat."""
    dari   = request.args.get('dari', '')
    sampai = request.args.get('sampai', '')
    limit  = int(request.args.get('limit', 300))
    where, params = ["kl.jenis='pakai_treatment'"], []
    if dari:
        where.append("date(kl.tanggal)>=?"); params.append(dari)
    if sampai:
        where.append("date(kl.tanggal)<=?"); params.append(sampai)
    params.append(limit)
    rows = query(f"""
        SELECT kl.tanggal, kl.kd_karyawan,
               COALESCE(k.nickname,k.nm_karyawan,kl.kd_karyawan) AS nama,
               kl.kd_barang, kl.nm_barang, kl.exp, kl.qty,
               kl.ref AS no_faktur, kl.kd_pelayanan,
               COALESCE(pp.nm_pelayanan, kl.kd_pelayanan) as nm_pelayanan,
               kl.kd_pasien, kl.nm_pasien, kl.operator,
               EXISTS(SELECT 1 FROM Serah_Terima_Kunci stk
                      WHERE stk.no_faktur=kl.ref AND stk.kd_pelayanan=kl.kd_pelayanan
                        AND stk.kd_k_final=kl.kd_karyawan) as via_serah_terima,
               (SELECT stk.nm_pelaksana FROM Serah_Terima_Kunci stk
                WHERE stk.no_faktur=kl.ref AND stk.kd_pelayanan=kl.kd_pelayanan
                  AND stk.kd_k_final=kl.kd_karyawan LIMIT 1) as st_nm_pelaksana
        FROM Kit_Log kl
        LEFT JOIN Karyawan k ON kl.kd_karyawan=k.kd_karyawan
        LEFT JOIN produk_pelayanan pp ON kl.kd_pelayanan=pp.kd_pelayanan
        WHERE {' AND '.join(where)}
        ORDER BY kl.id DESC LIMIT ?
    """, params)
    return jsonify(rows_to_list(rows))


@app.route('/api/kit/retur-history')
def kit_retur_history():
    """Riwayat retur kit (bahan dikembalikan ke kit perawat) untuk monitoring."""
    dari   = request.args.get('dari', '')
    sampai = request.args.get('sampai', '')
    limit  = int(request.args.get('limit', 150))
    where, params = ["kl.jenis='retur_treatment'"], []
    if dari:
        where.append("date(kl.tanggal)>=?"); params.append(dari)
    if sampai:
        where.append("date(kl.tanggal)<=?"); params.append(sampai)
    params.append(limit)
    rows = query(f"""
        SELECT kl.tanggal, kl.kd_karyawan,
               COALESCE(k.nickname,k.nm_karyawan,kl.kd_karyawan) AS nama,
               kl.kd_barang, kl.nm_barang, kl.exp, kl.qty,
               kl.ref AS no_faktur, kl.kd_pelayanan, kl.nm_pasien, kl.kd_pasien, kl.operator
        FROM Kit_Log kl LEFT JOIN Karyawan k ON kl.kd_karyawan=k.kd_karyawan
        WHERE {' AND '.join(where)}
        ORDER BY kl.id DESC LIMIT ?
    """, params)
    return jsonify(rows_to_list(rows))


# ─── ANTRIAN ─────────────────────────────────────────────────
@app.route('/api/antrian')
def get_antrian():
    today = str(date.today())
    tanggal = request.args.get('tanggal', today)
    rows = query("SELECT a.*, p.nm_pasien FROM Antrian a JOIN Pasien p ON a.kd_pasien=p.kd_pasien WHERE a.tanggal=? ORDER BY a.id", (tanggal,))
    return jsonify(rows_to_list(rows))

@app.route('/api/antrian', methods=['POST'])
def add_antrian():
    d = request.json
    execute("INSERT INTO Antrian (kd_pasien, tanggal, status) VALUES (?,date('now'),'wait')", (d['kd_pasien'],))
    return jsonify({'ok': True})

@app.route('/api/antrian/<int:id>/panggil', methods=['POST'])
def panggil_antrian(id):
    execute("UPDATE Antrian SET status='dipanggil' WHERE id=?", (id,))
    return jsonify({'ok': True})

# ─── NOMOR FAKTUR ────────────────────────────────────────────
def get_faktur_prefix(kd_lokasi=''):
    """Prefix faktur per cabang (Pengaturan key faktur_prefix_<kd_lokasi>); fallback ke global faktur_prefix."""
    if kd_lokasi:
        p = (get_setting(f'faktur_prefix_{kd_lokasi}', '') or '').strip()
        if p:
            return p
    return (get_setting('faktur_prefix', '') or '').strip()

def _kd_lokasi_boleh_jual(kd_lokasi):
    """Penjualan ke pasien HANYA boleh dari Cabang lewat kasir — Apotek Induk murni
    gudang/distribusi, tidak boleh transaksi langsung ke pasien (walau role admin).
    True kalau kd_lokasi ini valid utk berjualan."""
    if not kd_lokasi:
        return False
    row = query("SELECT tipe FROM Lokasi WHERE kd_lokasi=?", (kd_lokasi,), one=True)
    if not row:
        return kd_lokasi != 'INDUK'
    return row['tipe'] != 'induk'

def _kode_terminal(terminal, kd_lokasi=''):
    """Gabungan prefix cabang (per kd_lokasi) + kode terminal kasir → key counter unik."""
    prefix = get_faktur_prefix(kd_lokasi)
    return f"{prefix}{(terminal or '1').strip()}", prefix

def _reserve_or_reuse_faktur(kode, tanggal, builder):
    """Atomik: validasi dulu apakah nomor TERAKHIR yg direserve utk (kode,tanggal)
    ini sudah benar-benar dipakai (ada baris Penjualan-nya). Kalau BELUM (mis.
    kasir buka POS lalu tutup lagi sebelum pasien/item diisi) → pakai ULANG nomor
    yang sama, JANGAN naikkan counter — supaya buka-tutup POS tidak membuang
    nomor faktur / bikin nomor lompat. Counter cuma naik kalau nomor terakhir
    memang sudah terpakai (atau belum pernah ada nomor sama sekali).
    `builder(nomor_int)` membentuk string faktur lengkap dari nomor urut.
    Locking sama seperti versi lama (anti tabrakan kasir paralel)."""
    if DB_BACKEND == 'mysql':
        return _reserve_or_reuse_faktur_mysql(kode, tanggal, builder)
    con = get_db()
    con.isolation_level = None   # autocommit → kita kontrol transaksi manual
    try:
        con.execute("BEGIN IMMEDIATE")
        con.execute("INSERT OR IGNORE INTO Counter_Faktur (kode,tanggal,counter,last_faktur) VALUES (?,?,0,'')", (kode, tanggal))
        row = con.execute("SELECT counter, last_faktur FROM Counter_Faktur WHERE kode=? AND tanggal=?", (kode, tanggal)).fetchone()
        counter, last_faktur = row[0], (row[1] or '')
        reused = False
        if last_faktur:
            dipakai = con.execute("SELECT 1 FROM Penjualan WHERE no_faktur=?", (last_faktur,)).fetchone()
            if not dipakai:
                faktur, nomor, reused = last_faktur, counter, True
        if not reused:
            con.execute("UPDATE Counter_Faktur SET counter=counter+1 WHERE kode=? AND tanggal=?", (kode, tanggal))
            nomor  = con.execute("SELECT counter FROM Counter_Faktur WHERE kode=? AND tanggal=?", (kode, tanggal)).fetchone()[0]
            faktur = builder(nomor)
            con.execute("UPDATE Counter_Faktur SET last_faktur=? WHERE kode=? AND tanggal=?", (faktur, kode, tanggal))
        con.execute("COMMIT")
        return faktur, nomor, reused
    except Exception:
        try: con.execute("ROLLBACK")
        except Exception: pass
        raise
    finally:
        con.close()

def _reserve_or_reuse_faktur_mysql(kode, tanggal, builder):
    """Versi MySQL: SELECT ... FOR UPDATE (row lock InnoDB) menggantikan
    BEGIN IMMEDIATE SQLite — jaminan anti-tabrakan setara untuk kasir paralel."""
    raw = _mysql_connect()
    try:
        cur = raw.cursor()
        # Placeholder ditulis gaya '?' (sama spt sisa app) krn RowCursor.execute()
        # otomatis terjemahkan lewat _translate_sql() — literal '%s' di sini akan
        # ikut ter-escape jadi '%%s' dan korup.
        cur.execute("INSERT IGNORE INTO Counter_Faktur (kode,tanggal,counter,last_faktur) VALUES (?,?,0,'')", (kode, tanggal))
        raw.commit()   # pastikan baris ada & ter-commit sebelum di-lock
        cur.execute("SELECT counter, last_faktur FROM Counter_Faktur WHERE kode=? AND tanggal=? FOR UPDATE", (kode, tanggal))
        counter, last_faktur = cur.fetchone()
        last_faktur = last_faktur or ''
        reused = False
        if last_faktur:
            cur.execute("SELECT 1 FROM Penjualan WHERE no_faktur=?", (last_faktur,))
            if not cur.fetchone():
                faktur, nomor, reused = last_faktur, counter, True
        if not reused:
            cur.execute("UPDATE Counter_Faktur SET counter=counter+1 WHERE kode=? AND tanggal=?", (kode, tanggal))
            cur.execute("SELECT counter FROM Counter_Faktur WHERE kode=? AND tanggal=?", (kode, tanggal))
            nomor  = cur.fetchone()[0]
            faktur = builder(nomor)
            cur.execute("UPDATE Counter_Faktur SET last_faktur=? WHERE kode=? AND tanggal=?", (faktur, kode, tanggal))
        raw.commit()
        return faktur, nomor, reused
    except Exception:
        raw.rollback()
        raise
    finally:
        raw.close()

@app.route('/api/nomor-faktur')
def get_nomor_faktur():
    today    = str(date.today())
    ddmmyy   = datetime.now().strftime('%d%m%y')
    terminal = request.args.get('terminal') or request.args.get('kasir') or '1'
    kd_lokasi = request.args.get('kd_lokasi', '')
    if not _kd_lokasi_boleh_jual(kd_lokasi):
        return jsonify({'ok': False, 'msg': 'Penjualan ke pasien hanya boleh dari Cabang — Apotek Induk tidak bisa transaksi langsung ke pasien.'}), 403
    kode, prefix = _kode_terminal(terminal, kd_lokasi)
    # Format: {prefix}{terminal}{ddmmyy}{NNN} — kode kasir/terminal digabung ke awalan (sebelum
    # tanggal), nomor urut 3 digit. prefix & terminal tetap jaga keunikan antar cabang/kasir.
    builder = lambda n: f"{prefix}{(terminal or '1').strip()}{ddmmyy}{n:03d}"
    # ATOMIK + tervalidasi: nomor terakhir yg belum kepakai (belum ada transaksi Penjualan-nya)
    # dipakai ulang, bukan bikin nomor baru tiap POS dibuka-tutup.
    faktur, nomor, reused = _reserve_or_reuse_faktur(kode, today, builder)
    return jsonify({'faktur': faktur, 'nomor': nomor, 'reused': reused})


@app.route('/api/server-info')
def server_info():
    """IP LAN server + URL akses mobile — untuk QR akses dari HP (bukan localhost)."""
    import socket
    ip = '127.0.0.1'
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(('8.8.8.8', 80))   # tak benar-benar mengirim; hanya pilih interface LAN aktif
        ip = s.getsockname()[0]
        s.close()
    except Exception:
        try:
            ip = socket.gethostbyname(socket.gethostname())
        except Exception:
            ip = '127.0.0.1'
    host = request.host or ''
    port = host.split(':', 1)[1] if ':' in host else '5000'
    base = f"http://{ip}:{port}"
    return jsonify({'ip': ip, 'port': port, 'base': base, 'mobile_url': base + '/mobile'})

# ─── PENJUALAN (POS) ─────────────────────────────────────────
@app.route('/api/penjualan/<faktur>')
def get_penjualan(faktur):
    header = query("SELECT * FROM Penjualan WHERE no_faktur=?", (faktur,), one=True)
    items  = query("SELECT * FROM R_Penjualan WHERE no_faktur=? ORDER BY id", (faktur,))
    if not header:
        return jsonify({'header': None, 'items': []})
    return jsonify({'header': dict(header), 'items': rows_to_list(items)})

@app.route('/api/penjualan/<faktur>/cetak-ulang')
def penjualan_cetak_ulang(faktur):
    """Rekonstruksi data faktur lengkap utk CETAK ULANG (mis. struk fisik hilang) —
    dipakai tombol 🖨 di Laporan Penjualan. Bentuknya dibuat PERSIS sama dgn fakturData
    yg dipakai cetakFaktur()/cetakFakturAuto() saat pembayaran asli (lihat konfirmBayar()
    di frontend), supaya hasil cetak ulang identik dgn struk yg pertama kali keluar —
    baris retur SENGAJA dikecualikan krn itu bukan bagian dari faktur asli saat dicetak
    pertama kali (retur terjadi belakangan, punya jejak audit sendiri)."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor', 'kasir'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor/kasir'}), 403

    header = query("SELECT * FROM Penjualan WHERE no_faktur=?", (faktur,), one=True)
    if not header:
        return jsonify({'ok': False, 'msg': 'Faktur tidak ditemukan'}), 404

    pasien = query("SELECT * FROM Pasien WHERE kd_pasien=?", (header['kd_pasien'],), one=True) \
             if header['kd_pasien'] else None

    items_raw = rows_to_list(query("""
        SELECT * FROM R_Penjualan WHERE no_faktur=?
          AND jenis NOT LIKE 'retur%' AND jenis != 'bahan_racikan'
        ORDER BY id
    """, (faktur,)))

    # Nama beautician per jasa (Point_karyawan + Karyawan.nickname) — bisa lebih dari
    # 1 petugas per jasa yg sama, gabung koma spt tampilan struk aslinya.
    pk_rows = query("""SELECT pk.kd_pelayanan, k.nickname FROM Point_karyawan pk
                        LEFT JOIN Karyawan k ON pk.kd_karyawan=k.kd_karyawan
                        WHERE pk.no_faktur=?""", (faktur,))
    beautician_by_kd = {}
    for r in pk_rows:
        if r['nickname']:
            beautician_by_kd.setdefault(r['kd_pelayanan'], []).append(r['nickname'])

    items = [{
        'nama': it['nm_barang'], 'jenis': it['jenis'],
        'harga': it['harga'], 'disc': float(it['disc'] or 0), 'qty': it['qty'],
        'beautician': ', '.join(beautician_by_kd.get(it['kd_barang'], [])),
    } for it in items_raw]

    next_visit = query("SELECT next_tanggal FROM TKunjunganBerikut WHERE no_faktur=?", (faktur,), one=True)

    # Rincian pembayaran tersplit (kosong [] utk faktur lama pra-fitur — frontend
    # fallback tampilan 1-baris "Metode: ..." spt sebelumnya, lihat cetakFaktur*()).
    pembayaran = rows_to_list(query(
        "SELECT metode, bank, nominal FROM Penjualan_Pembayaran WHERE no_faktur=? ORDER BY id",
        (faktur,)))

    return jsonify({
        'ok': True,
        'faktur':     header['no_faktur'],
        'tanggal':    header['tanggal'],
        'dokter':     header['dokter'] or '',
        'kasir':      header['kasir'] or '',
        'metode':     header['metode_bayar'] or 'Tunai',
        'discTrx':    header['disc_trx'] or 0,
        'nextVisit':  next_visit['next_tanggal'] if next_visit else '',
        'pasien':     dict(pasien) if pasien else {},
        'items':      items,
        'pembayaran': pembayaran,
    })

@app.route('/api/penjualan/item', methods=['POST'])
def add_item_penjualan():
    d = request.json
    faktur = d['no_faktur']
    if not _kd_lokasi_boleh_jual(d.get('kd_lokasi', 'CBG1')):
        return jsonify({'ok': False, 'msg': 'Penjualan ke pasien hanya boleh dari Cabang — Apotek Induk tidak bisa transaksi langsung ke pasien.'}), 403
    today = str(date.today())
    harga = float(d.get('harga', 0))
    disc_pct = float(d.get('disc', 0))
    qty = int(d.get('qty', 1))
    harga_beli = float(d.get('harga_beli', 0))
    sub_total = harga * (1 - disc_pct/100) * qty

    # Upsert header
    existing = query("SELECT * FROM Penjualan WHERE no_faktur=?", (faktur,), one=True)
    if not existing:
        execute("INSERT INTO Penjualan (no_faktur,tanggal,total_harga,hpp,dokter,kd_pasien,kasir,status) VALUES (?,?,0,0,?,?,?,'open')",
                (faktur, today, d.get('dokter',''), d.get('kd_pasien',''), d.get('kasir','')))

    # Add item
    last_id = execute("""INSERT INTO R_Penjualan
               (no_faktur,tanggal,kd_barang,nm_barang,harga,harga_asli,qty,disc,sub_total,total_harbeli,kd_karyawan,cashier,exp,jenis)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (faktur, today, d['kd_barang'], d['nm_barang'], harga, harga, qty,
             str(disc_pct), sub_total, harga_beli*qty, d.get('kd_karyawan',''),
             d.get('kasir',''), d.get('exp',''), d.get('jenis','barang')))

    # Update total header (kecuali bahan_racikan)
    execute("""UPDATE Penjualan SET total_harga=(SELECT COALESCE(SUM(sub_total),0) FROM R_Penjualan
               WHERE no_faktur=? AND jenis!='bahan_racikan') WHERE no_faktur=?""",
            (faktur, faktur))

    # Kurangi stok cabang + catat RiwayatBarang
    kd_lokasi = d.get('kd_lokasi', 'CBG1')
    # Simpan kd_lokasi ke header Penjualan jika belum ada
    execute("UPDATE Penjualan SET kd_lokasi=? WHERE no_faktur=? AND (kd_lokasi IS NULL OR kd_lokasi='')",
            (kd_lokasi, faktur))
    if d.get('jenis') == 'barang':
        exp = d.get('exp', '')
        if exp:
            execute("UPDATE Stok_barang SET stok_sisa=stok_sisa-? WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                    (qty, d['kd_barang'], exp, kd_lokasi))
            stok_now = query("SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                             (d['kd_barang'], exp, kd_lokasi), one=True)
        else:
            execute("UPDATE Stok_barang SET stok_sisa=stok_sisa-? WHERE kd_barang=? AND kd_lokasi=?",
                    (qty, d['kd_barang'], kd_lokasi))
            stok_now = query("SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND kd_lokasi=? ORDER BY id LIMIT 1",
                             (d['kd_barang'], kd_lokasi), one=True)
        execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                (today, d['kd_barang'], exp, f"Jual {faktur}",
                 -qty,
                 stok_now['stok_sisa']   if stok_now else 0,
                 stok_now['stok_gudang'] if stok_now else 0))

    catat_log('tambah_item_penjualan', 'R_Penjualan', faktur,
              f"Tambah item ke transaksi {faktur}: {d.get('nm_barang','')} × {qty} @ {harga:g}"
              + (f" (disc {disc_pct:g}%)" if disc_pct else ''),
              data_baru={'kd_barang': d.get('kd_barang',''), 'nm_barang': d.get('nm_barang',''),
                         'qty': qty, 'harga': harga, 'disc': disc_pct, 'sub_total': sub_total,
                         'jenis': d.get('jenis','barang')})
    return jsonify({'ok': True, 'id': last_id})

@app.route('/api/penjualan/item-disc', methods=['PUT'])
def update_item_disc():
    """Update diskon per item dalam satu transaksi (kd_barang + no_faktur)."""
    d = request.json
    no_faktur = d['no_faktur']
    kd_barang  = d['kd_barang']
    disc_pct   = max(0.0, min(100.0, float(d.get('disc', 0))))

    # 'racikan' ikut diizinkan — disc di sini utk disc TRANSAKSI level baris racikan
    # (beda dari disc per-bahan saat racikan disusun); qty racikan selalu 1 jadi
    # formula sub_total di bawah tetap berlaku sama.
    rows = query("SELECT * FROM R_Penjualan WHERE no_faktur=? AND kd_barang=? AND jenis IN ('barang','jasa','racikan')",
                 (no_faktur, kd_barang))
    for row in rows:
        new_sub = row['harga'] * (1 - disc_pct / 100) * row['qty']
        execute("UPDATE R_Penjualan SET disc=?,sub_total=? WHERE id=?",
                (str(disc_pct), new_sub, row['id']))
        pasien_info = _pasien_faktur(no_faktur)
        catat_log('ubah_diskon_item', 'R_Penjualan', no_faktur,
                  f"Ubah diskon '{kd_barang}' di faktur {no_faktur}: {row['disc']}% → {disc_pct}%"
                  + (f" — Pasien: {pasien_info}" if pasien_info else ''),
                  data_lama={'disc': row['disc'], 'sub_total': row['sub_total']},
                  data_baru={'disc': disc_pct, 'sub_total': new_sub})

    execute("""UPDATE Penjualan SET
        total_harga=(SELECT COALESCE(SUM(sub_total),0) FROM R_Penjualan
                     WHERE no_faktur=? AND jenis!='bahan_racikan')
        WHERE no_faktur=?""", (no_faktur, no_faktur))

    return jsonify({'ok': True})

def _bersihkan_penjualan_kosong(faktur):
    """Kalau kasir menghapus item/racikan lewat tombol ✕ sampai keranjang kosong SAMA
    SEKALI tanpa pernah membayar (status masih 'open'), header Penjualan-nya jangan
    dibiarkan jadi bangkai kosong yg nongkrong selamanya di Laporan Penjualan berlabel
    'Proses' (padahal transaksinya sudah dibatalkan total, bukan sedang berjalan).
    Dipanggil setelah tiap hapus item/racikan — aman no-op kalau faktur masih ada isi
    ATAU sudah 'done' (transaksi riil, jangan pernah dihapus otomatis)."""
    p = query("SELECT status FROM Penjualan WHERE no_faktur=?", (faktur,), one=True)
    if not p or p['status'] != 'open':
        return
    sisa = query("SELECT COUNT(*) c FROM R_Penjualan WHERE no_faktur=?", (faktur,), one=True)
    if sisa and sisa['c'] == 0:
        execute("DELETE FROM Penjualan WHERE no_faktur=?", (faktur,))
        catat_log('penjualan_batal_kosong', 'Penjualan', faktur,
                  f"Faktur {faktur} dibatalkan otomatis — semua item dihapus sebelum dibayar")

@app.route('/api/penjualan/item/<int:id>', methods=['DELETE'])
def delete_item_penjualan(id):
    item = query("SELECT * FROM R_Penjualan WHERE id=?", (id,), one=True)
    if item:
        penjualan = query("SELECT kd_lokasi FROM Penjualan WHERE no_faktur=?", (item['no_faktur'],), one=True)
        kd_lokasi = (penjualan['kd_lokasi'] if penjualan else None) or 'CBG1'
        today = str(date.today())
        if item['jenis'] == 'barang':
            exp = item['exp'] or ''
            if exp:
                execute("UPDATE Stok_barang SET stok_sisa=stok_sisa+? WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                        (item['qty'], item['kd_barang'], exp, kd_lokasi))
                sn = query("SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                           (item['kd_barang'], exp, kd_lokasi), one=True)
            else:
                execute("UPDATE Stok_barang SET stok_sisa=stok_sisa+? WHERE kd_barang=? AND kd_lokasi=?",
                        (item['qty'], item['kd_barang'], kd_lokasi))
                sn = query("SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND kd_lokasi=? ORDER BY id LIMIT 1",
                           (item['kd_barang'], kd_lokasi), one=True)
            # Catat pembatalan item di Kartu Stok (nilai positif = stok kembali)
            execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                    (today, item['kd_barang'], exp, f"Batal item {item['no_faktur']}",
                     item['qty'], sn['stok_sisa'] if sn else 0, sn['stok_gudang'] if sn else 0))
        faktur = item['no_faktur']
        execute("DELETE FROM R_Penjualan WHERE id=?", (id,))
        execute("""UPDATE Penjualan SET total_harga=(SELECT COALESCE(SUM(sub_total),0)
                   FROM R_Penjualan WHERE no_faktur=? AND jenis!='bahan_racikan') WHERE no_faktur=?""",
                (faktur, faktur))
        pasien_info = _pasien_faktur(faktur)
        catat_log('hapus_item_penjualan', 'R_Penjualan', faktur,
                  f"Hapus item '{item['nm_barang']}' (qty {item['qty']}, Rp{item['sub_total']:,.0f}) dari faktur {faktur}"
                  + (f" — Pasien: {pasien_info}" if pasien_info else ''),
                  data_lama={'kd_barang': item['kd_barang'], 'nm_barang': item['nm_barang'],
                             'qty': item['qty'], 'harga': item['harga'], 'sub_total': item['sub_total'],
                             'jenis': item['jenis'], 'pasien': pasien_info})
        _bersihkan_penjualan_kosong(faktur)
    return jsonify({'ok': True})


# ─── TRANSAKSI 'OPEN' TERLEWAT (kasir mulai transaksi, ada isi, tapi tidak pernah ──
# dibayar berhari-hari — mis. lupa/pindah device sebelum sempat klik Bayar). Beda
# dari _bersihkan_penjualan_kosong() di atas yg cuma nangani kasus KOSONG SAMA SEKALI
# (semua item dihapus lewat ✕) — di sini transaksinya MASIH ADA ISI, cuma ditinggal.
# Halaman admin "Transaksi Belum Selesai" pakai 3 endpoint ini utk cari & selesaikan.
@app.route('/api/penjualan/open-list')
def penjualan_open_list():
    rows = rows_to_list(query("""
        SELECT p.no_faktur, p.tanggal, p.kd_pasien, ps.nm_pasien, p.dokter, p.kasir,
               p.total_harga, p.kd_lokasi,
               (SELECT COUNT(*) FROM R_Penjualan rp WHERE rp.no_faktur=p.no_faktur) as jml_item
        FROM Penjualan p LEFT JOIN Pasien ps ON p.kd_pasien=ps.kd_pasien
        WHERE p.status='open'
        ORDER BY p.tanggal ASC, p.no_faktur ASC
    """))
    today = date.today()
    for r in rows:
        try:
            r['umur_hari'] = (today - datetime.strptime(r['tanggal'], '%Y-%m-%d').date()).days
        except (ValueError, TypeError):
            r['umur_hari'] = None
    return jsonify(rows)

@app.route('/api/penjualan/open/<no_faktur>')
def penjualan_open_detail(no_faktur):
    """Detail lengkap 1 faktur open — dipakai halaman admin utk tombol 'Lanjutkan'
    (memuat ulang keranjang ke kasir supaya bisa diselesaikan/dibayar)."""
    header = query("SELECT * FROM Penjualan WHERE no_faktur=?", (no_faktur,), one=True)
    if not header or header['status'] != 'open':
        return jsonify({'ok': False, 'msg': 'Faktur open tidak ditemukan'}), 404
    pasien = query("SELECT * FROM Pasien WHERE kd_pasien=?", (header['kd_pasien'],), one=True) \
             if header['kd_pasien'] else None
    items = rows_to_list(query("""
        SELECT rp.*, COALESCE(b.non_obat,0) as non_obat
        FROM R_Penjualan rp LEFT JOIN Barang b ON rp.kd_barang=b.kd_barang
        WHERE rp.no_faktur=? AND rp.jenis != 'bahan_racikan'
        ORDER BY rp.id
    """, (no_faktur,)))
    for it in items:
        it['kd'] = it['kd_barang']; it['nama'] = it['nm_barang']; it['dbId'] = it['id']
        if it['jenis'] == 'racikan':
            it['id_resep'] = it['racikan_ref']
            rr = query("SELECT cara_pakai FROM ResepRacikan WHERE id=?", (it['racikan_ref'],), one=True)
            it['cara_pakai'] = rr['cara_pakai'] if rr else ''
    return jsonify({'ok': True, 'header': dict(header),
                    'pasien': dict(pasien) if pasien else None, 'items': items})

@app.route('/api/penjualan/open/<no_faktur>/batalkan', methods=['POST'])
def penjualan_open_batalkan(no_faktur):
    """Batalkan faktur 'open' yg terlanjur ada isinya tapi ditinggal — kembalikan SEMUA
    stok (barang + bahan racikan), spt kalau dihapus satu-satu lewat ✕ di kasir, lalu
    hapus seluruh baris transaksinya. Faktur 'done' TIDAK BOLEH lewat sini — itu harus
    lewat menu Retur (jejak auditnya beda: retur menyisakan faktur asli + baris retur)."""
    header = query("SELECT * FROM Penjualan WHERE no_faktur=?", (no_faktur,), one=True)
    if not header:
        return jsonify({'ok': False, 'msg': 'Faktur tidak ditemukan'}), 404
    if header['status'] != 'open':
        return jsonify({'ok': False, 'msg': 'Hanya transaksi berstatus open (belum dibayar) yang bisa dibatalkan di sini — transaksi yang sudah dibayar harus lewat menu Retur'}), 400

    kd_lokasi = header['kd_lokasi'] or 'CBG1'
    today = str(date.today())
    con = get_db()
    jml_item = 0
    try:
        cur = con.cursor()
        items = cur.execute("SELECT * FROM R_Penjualan WHERE no_faktur=?", (no_faktur,)).fetchall()
        jml_item = len(items)
        id_reseps_done = set()
        for row in items:
            it = dict(row)
            if it['jenis'] == 'barang':
                exp = it['exp'] or ''
                if exp:
                    cur.execute("UPDATE Stok_barang SET stok_sisa=stok_sisa+? WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                                (it['qty'], it['kd_barang'], exp, kd_lokasi))
                    sn = cur.execute("SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                                     (it['kd_barang'], exp, kd_lokasi)).fetchone()
                else:
                    cur.execute("UPDATE Stok_barang SET stok_sisa=stok_sisa+? WHERE kd_barang=? AND kd_lokasi=?",
                                (it['qty'], it['kd_barang'], kd_lokasi))
                    sn = cur.execute("SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND kd_lokasi=? ORDER BY id LIMIT 1",
                                     (it['kd_barang'], kd_lokasi)).fetchone()
                sn = dict(sn) if sn else None
                cur.execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                            (today, it['kd_barang'], exp, f"Batal faktur open {no_faktur} (terlewat)",
                             it['qty'], sn['stok_sisa'] if sn else 0, sn['stok_gudang'] if sn else 0))
            elif it['jenis'] == 'racikan' and it['racikan_ref'] and it['racikan_ref'] not in id_reseps_done:
                id_resep = it['racikan_ref']
                id_reseps_done.add(id_resep)
                bahans = cur.execute("SELECT kd_barang, qty, exp FROM DetailResepRacikan WHERE id_resep=?", (id_resep,)).fetchall()
                for brow in bahans:
                    b = dict(brow)
                    cur.execute("UPDATE Stok_barang SET stok_sisa=stok_sisa+? WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                                (b['qty'], b['kd_barang'], b['exp'], kd_lokasi))
                    sr = cur.execute("SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                                     (b['kd_barang'], b['exp'], kd_lokasi)).fetchone()
                    if sr:
                        sr = dict(sr)
                        cur.execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                                    (today, b['kd_barang'], b['exp'], f"Batal faktur open {no_faktur} (terlewat)",
                                     b['qty'], sr['stok_sisa'], sr['stok_gudang'] or 0))
                cur.execute("DELETE FROM DetailResepRacikan WHERE id_resep=?", (id_resep,))
                cur.execute("DELETE FROM ResepRacikan WHERE id=?", (id_resep,))
            # jenis 'jasa' tidak menyentuh stok; 'bahan_racikan' sudah ditangani lewat blok racikan

        cur.execute("DELETE FROM R_Penjualan WHERE no_faktur=?", (no_faktur,))
        cur.execute("DELETE FROM Penjualan WHERE no_faktur=?", (no_faktur,))
        con.commit()
    finally:
        con.close()

    catat_log('penjualan_batal_open', 'Penjualan', no_faktur,
              f"Batalkan faktur open {no_faktur} ({jml_item} baris item, tertunda sejak {header['tanggal']}) — stok dikembalikan")
    return jsonify({'ok': True, 'jml_item': jml_item})


@app.route('/api/penjualan/cari-retur')
def cari_penjualan_retur():
    """Cari transaksi done untuk proses retur — berdasarkan no_faktur atau nama/kode pasien."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor', 'kasir'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor/kasir'}), 403
    q = request.args.get('q', '').strip()
    if q:
        rows = query("""
            SELECT p.no_faktur, p.tanggal, p.total_harga, p.status,
                   p.kd_pasien, pa.nm_pasien, p.kasir, p.metode_bayar, p.disc_trx
            FROM Penjualan p LEFT JOIN Pasien pa ON p.kd_pasien=pa.kd_pasien
            WHERE p.status='done'
              AND (p.no_faktur LIKE ? OR pa.nm_pasien LIKE ? OR p.kd_pasien LIKE ?)
            ORDER BY p.tanggal DESC, p.no_faktur DESC LIMIT 30
        """, (f'%{q}%', f'%{q}%', f'%{q}%'))
    else:
        rows = query("""
            SELECT p.no_faktur, p.tanggal, p.total_harga, p.status,
                   p.kd_pasien, pa.nm_pasien, p.kasir, p.metode_bayar, p.disc_trx
            FROM Penjualan p LEFT JOIN Pasien pa ON p.kd_pasien=pa.kd_pasien
            WHERE p.status='done' AND p.tanggal=date('now')
            ORDER BY p.no_faktur DESC LIMIT 30
        """)
    return jsonify(rows_to_list(rows))


@app.route('/api/penjualan/retur-detail/<faktur>')
def retur_detail(faktur):
    """Ambil detail item sebuah faktur lengkap dengan qty yang sudah diretur."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor', 'kasir'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor/kasir'}), 403
    header = query("""
        SELECT p.*, pa.nm_pasien FROM Penjualan p
        LEFT JOIN Pasien pa ON p.kd_pasien=pa.kd_pasien
        WHERE p.no_faktur=?
    """, (faktur,), one=True)
    if not header:
        return jsonify({'ok': False, 'msg': 'Faktur tidak ditemukan'}), 404
    items = query("""
        SELECT * FROM R_Penjualan
        WHERE no_faktur=? AND jenis NOT LIKE 'retur%' AND jenis != 'bahan_racikan'
        ORDER BY id
    """, (faktur,))
    con = get_db()
    result = []
    for item in items:
        d = dict(item)
        already = con.execute("""
            SELECT COALESCE(SUM(qty), 0) FROM R_Penjualan
            WHERE no_faktur=? AND kd_karyawan=? AND jenis LIKE 'retur%'
        """, (faktur, f"RETUR:{d['id']}")).fetchone()[0]
        d['qty_retur_sudah'] = float(already)
        d['qty_available']   = max(0.0, float(d['qty']) - float(already))
        result.append(d)
    con.close()
    return jsonify({'ok': True, 'header': dict(header), 'items': result})


@app.route('/api/penjualan/retur', methods=['POST'])
def proses_retur():
    """Proses retur item dari transaksi yang sudah done. Kalau fakturnya sudah
    lewat 1 hari DAN yg mengajukan cuma kasir (bukan admin/supervisor), retur
    TIDAK langsung dieksekusi — disimpan sbg pengajuan 'pending' di Retur_Request,
    baru benar2 diproses setelah disetujui admin/supervisor lewat Approval Queue.
    Admin/supervisor yg retur sendiri (berapa pun umur fakturnya) langsung
    dieksekusi tanpa perlu approval tambahan — mereka sendiri sudah level approval."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor', 'kasir'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor/kasir'}), 403
    d         = request.json or {}
    no_faktur = d.get('no_faktur', '').strip()
    items_req = d.get('items', [])   # [{id, qty_retur}]
    kasir     = d.get('kasir', '')
    if not no_faktur or not items_req:
        return jsonify({'ok': False, 'msg': 'Data tidak lengkap'}), 400

    header0 = query("SELECT tanggal, status FROM Penjualan WHERE no_faktur=?", (no_faktur,), one=True)
    if not header0:
        return jsonify({'ok': False, 'msg': f'Faktur {no_faktur} tidak ditemukan'}), 404
    if header0['status'] != 'done':
        return jsonify({'ok': False, 'msg': 'Transaksi harus berstatus done untuk diretur'}), 400
    try:
        umur_hari = (date.today() - datetime.strptime(header0['tanggal'], '%Y-%m-%d').date()).days
    except (ValueError, TypeError):
        umur_hari = 0

    if role == 'kasir' and umur_hari > 1:
        rid = execute("""INSERT INTO Retur_Request (no_faktur, items_json, kasir, diajukan_oleh)
                          VALUES (?,?,?,?)""",
                       (no_faktur, json.dumps(items_req), kasir, user.get('nama', '') if user else ''))
        catat_log('retur_diajukan', 'Retur_Request', no_faktur,
                  f"Ajukan retur faktur {no_faktur} (umur {umur_hari} hari) — menunggu approval admin/supervisor")
        return jsonify({'ok': True, 'perlu_approval': True, 'id': rid, 'umur_hari': umur_hari,
                        'msg': f'Faktur ini sudah berumur {umur_hari} hari — retur diajukan, menunggu persetujuan admin/supervisor'})

    result, status = _eksekusi_retur(no_faktur, items_req, kasir)
    return jsonify(result), status

def _eksekusi_retur(no_faktur, items_req, kasir):
    """Inti proses retur (dipakai baik dari proses_retur() langsung MAUPUN dari
    persetujuan Retur_Request) — return (dict, http_status), BUKAN Response."""
    con = get_raw_db(timeout=30)
    con.row_factory = sqlite3.Row
    cur = con.cursor()
    try:
        header = cur.execute("SELECT * FROM Penjualan WHERE no_faktur=?", (no_faktur,)).fetchone()
        if not header:
            return {'ok': False, 'msg': f'Faktur {no_faktur} tidak ditemukan'}, 404
        if header['status'] != 'done':
            return {'ok': False, 'msg': 'Transaksi harus berstatus done untuk diretur'}, 400

        kd_lokasi   = header['kd_lokasi'] or 'CBG1'
        today       = str(date.today())
        total_retur = 0.0
        kit_returns = []   # treatment jasa yang diretur → restore kit setelah commit
        retur_log_items = []   # ringkasan utk Aktivitas_Log

        for req in items_req:
            orig_id   = req.get('id')
            qty_retur = float(req.get('qty_retur', 0))
            if qty_retur <= 0 or not orig_id:
                continue

            orig = cur.execute(
                "SELECT * FROM R_Penjualan WHERE id=? AND no_faktur=?", (orig_id, no_faktur)
            ).fetchone()
            if not orig or str(orig['jenis']).startswith('retur'):
                continue

            # Qty sudah diretur sebelumnya dari baris ini
            already = cur.execute("""
                SELECT COALESCE(SUM(qty), 0) FROM R_Penjualan
                WHERE no_faktur=? AND kd_karyawan=? AND jenis LIKE 'retur%'
            """, (no_faktur, f'RETUR:{orig_id}')).fetchone()[0]

            qty_retur = min(qty_retur, float(orig['qty']) - float(already))
            if qty_retur <= 0:
                continue

            # Nilai retur per-item (harga setelah disc item, sebelum disc nota)
            try:
                disc_pct = float(orig['disc'] or 0)
            except Exception:
                disc_pct = 0.0
            harga_eff  = float(orig['harga']) * (1 - disc_pct / 100)
            sub_retur  = round(harga_eff * qty_retur, 2)
            total_retur += sub_retur

            jenis_orig  = orig['jenis'] or 'barang'
            retur_jenis = f'retur_{jenis_orig}'

            # Sisipkan baris retur (sub_total negatif)
            cur.execute("""INSERT INTO R_Penjualan
                (no_faktur, tanggal, kd_barang, nm_barang, harga, harga_asli,
                 qty, disc, sub_total, total_harbeli, kd_karyawan, cashier, exp, jenis)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (no_faktur, today,
                 orig['kd_barang'], orig['nm_barang'],
                 orig['harga'], orig['harga_asli'],
                 qty_retur, orig['disc'] or '0',
                 -sub_retur, 0,
                 f'RETUR:{orig_id}', kasir,
                 orig['exp'] or '', retur_jenis))

            retur_log_items.append({'kd_barang': orig['kd_barang'], 'nm_barang': orig['nm_barang'],
                                     'qty_retur': qty_retur, 'sub_retur': sub_retur, 'jenis': jenis_orig})

            # Treatment (jasa): kumpulkan untuk restore bahan kit perawat setelah commit
            if jenis_orig == 'jasa':
                qo = float(orig['qty']) or 1.0
                kit_returns.append({'kd_pelayanan': orig['kd_barang'],
                                    'rasio': (qty_retur / qo) if qo else 1.0})
                # Komisi perawat/dokter/supervisor utk treatment ini TIDAK dihapus — cuma
                # ditandai 'diretur' supaya tak lagi terhitung di Laporan Komisi (klinik
                # tak bayar komisi atas omzet yg sudah di-refund), tapi tetap tampil di
                # riwayat dgn label DIRETUR utk transparansi (bukan hilang diam-diam).
                cur.execute("UPDATE Point_karyawan SET status='diretur' WHERE no_faktur=? AND kd_pelayanan=?",
                            (no_faktur, orig['kd_barang']))

            # Pulihkan stok untuk produk fisik + catat di Kartu Stok (RiwayatBarang)
            if jenis_orig == 'barang':
                exp = orig['exp'] or ''
                if exp:
                    cur.execute("""UPDATE Stok_barang SET stok_sisa=stok_sisa+?
                        WHERE kd_barang=? AND exp=? AND kd_lokasi=?""",
                        (qty_retur, orig['kd_barang'], exp, kd_lokasi))
                    stok_now = cur.execute(
                        "SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                        (orig['kd_barang'], exp, kd_lokasi)).fetchone()
                else:
                    cur.execute("""UPDATE Stok_barang SET stok_sisa=stok_sisa+?
                        WHERE kd_barang=? AND kd_lokasi=?""",
                        (qty_retur, orig['kd_barang'], kd_lokasi))
                    stok_now = cur.execute(
                        "SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND kd_lokasi=? ORDER BY id LIMIT 1",
                        (orig['kd_barang'], kd_lokasi)).fetchone()

                # Catat pengembalian stok di Kartu Stok cabang (nilai positif = stok masuk)
                qty_log = int(qty_retur) if float(qty_retur).is_integer() else qty_retur
                cur.execute("""INSERT INTO RiwayatBarang
                    (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)""",
                    (today, orig['kd_barang'], exp, f"Retur {no_faktur}",
                     qty_log,
                     stok_now['stok_sisa']   if stok_now else 0,
                     stok_now['stok_gudang'] if stok_now else 0))

            # Racikan: pulihkan stok SEMUA bahan (dari DetailResepRacikan) + catat di Kartu Stok
            elif jenis_orig == 'racikan' and orig['racikan_ref']:
                qty_orig = float(orig['qty']) or 1.0
                ratio    = qty_retur / qty_orig            # dukung retur sebagian (umumnya 1)
                bahan_rows = cur.execute(
                    "SELECT kd_barang, nm_bahan, qty, exp FROM DetailResepRacikan WHERE id_resep=?",
                    (orig['racikan_ref'],)).fetchall()
                for b in bahan_rows:
                    qty_b = round(float(b['qty']) * ratio, 4)
                    if qty_b <= 0:
                        continue
                    exp_b = b['exp'] or ''
                    if exp_b:
                        cur.execute("""UPDATE Stok_barang SET stok_sisa=stok_sisa+?
                            WHERE kd_barang=? AND exp=? AND kd_lokasi=?""",
                            (qty_b, b['kd_barang'], exp_b, kd_lokasi))
                        stok_b = cur.execute(
                            "SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                            (b['kd_barang'], exp_b, kd_lokasi)).fetchone()
                    else:
                        cur.execute("""UPDATE Stok_barang SET stok_sisa=stok_sisa+?
                            WHERE kd_barang=? AND kd_lokasi=?""",
                            (qty_b, b['kd_barang'], kd_lokasi))
                        stok_b = cur.execute(
                            "SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND kd_lokasi=? ORDER BY id LIMIT 1",
                            (b['kd_barang'], kd_lokasi)).fetchone()
                    qty_b_log = int(qty_b) if float(qty_b).is_integer() else qty_b
                    cur.execute("""INSERT INTO RiwayatBarang
                        (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)""",
                        (today, b['kd_barang'], exp_b,
                         f"Retur racikan {orig['nm_barang']} [{no_faktur}]",
                         qty_b_log,
                         stok_b['stok_sisa']   if stok_b else 0,
                         stok_b['stok_gudang'] if stok_b else 0))

                # Bila seluruh qty racikan sudah diretur → tandai resep 'returned' supaya
                # tidak lagi muncul di Rekam Medik & Laporan Stok Penjualan (rancu/menempel).
                # Sekalian batalkan di Antrian Racikan (status_racik='batal') — apoteker
                # tidak perlu lagi menyiapkan racikan yang sudah diretur pasien.
                if (float(already) + qty_retur) >= qty_orig:
                    cur.execute("UPDATE ResepRacikan SET status='returned', status_racik='batal' WHERE id=?",
                                (orig['racikan_ref'],))

        if total_retur > 0:
            cur.execute("""UPDATE Penjualan
                SET total_harga = MAX(0.0, total_harga - ?)
                WHERE no_faktur=?""", (total_retur, no_faktur))

        # ── Sesuaikan point pasien karena nominal transaksi berkurang ──
        point_dikurangi = 0
        kd_pasien_h = header['kd_pasien']
        if total_retur > 0 and kd_pasien_h:
            rate = get_point_rate()
            net_row = cur.execute(
                """SELECT COALESCE(SUM(debit),0)-COALESCE(SUM(kredit),0) AS net
                   FROM DetailPoint WHERE kd_pasien=? AND keterangan=?""",
                (kd_pasien_h, no_faktur)).fetchone()
            net_now = int(net_row['net']) if net_row and net_row['net'] else 0
            if net_now > 0 and rate > 0:
                total_after = max(0.0, float(header['total_harga']) - total_retur)
                deserved    = int(total_after / rate)        # point yang seharusnya setelah retur
                selisih     = net_now - deserved
                if selisih > 0:
                    cur.execute(
                        "INSERT INTO DetailPoint (kd_pasien,keterangan,kredit,tanggal) VALUES (?,?,?,?)",
                        (kd_pasien_h, no_faktur, selisih, today))
                    point_dikurangi = selisih

        con.commit()

        pasien_info = _pasien_faktur(no_faktur)
        catat_log('retur_penjualan', 'Penjualan', no_faktur,
                  f"Retur faktur {no_faktur}: {len(retur_log_items)} item, total Rp{total_retur:,.0f}"
                  + (f" — Pasien: {pasien_info}" if pasien_info else ''),
                  data_baru={'items': retur_log_items, 'total_retur': total_retur, 'pasien': pasien_info})

        # Ambil info pasien untuk mode "tukar"
        updated_p = cur.execute(
            "SELECT kd_pasien FROM Penjualan WHERE no_faktur=?", (no_faktur,)
        ).fetchone()
        pasien = None
        if updated_p and updated_p['kd_pasien']:
            p = cur.execute(
                "SELECT kd_pasien, nm_pasien, telp FROM Pasien WHERE kd_pasien=?",
                (updated_p['kd_pasien'],)
            ).fetchone()
            if p:
                pasien = dict(p)
        con.close()

        # Restore bahan kit perawat untuk treatment yang diretur (post-commit, koneksi terpisah)
        for kr in kit_returns:
            try:
                _kit_retur(no_faktur, kr['kd_pelayanan'], '', kr['rasio'], operator=kasir)
            except Exception:
                pass
            # Batalkan serah-terima kunci yg BELUM sempat dieksekusi (blm ada di Kit_Log,
            # jadi tak ikut ke-restore _kit_retur di atas) — cegah nyangkut pending selamanya.
            try:
                execute("""UPDATE Serah_Terima_Kunci SET status='batal'
                           WHERE no_faktur=? AND kd_pelayanan=? AND status='pending'""",
                        (no_faktur, kr['kd_pelayanan']))
            except Exception:
                pass

        return {'ok': True, 'total_retur': total_retur, 'no_faktur': no_faktur,
                'pasien': pasien, 'point_dikurangi': point_dikurangi}, 200

    except Exception as e:
        try: con.close()
        except Exception: pass
        return {'ok': False, 'msg': str(e)}, 500


@app.route('/api/retur-request/<int:rid>')
def retur_request_detail(rid):
    """Rincian 1 pengajuan retur — dipakai modal 'Review' sebelum supervisor/admin
    Setujui/Tolak, supaya isi barang/qty yg diretur kelihatan dulu, bukan asal klik."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor'}), 403
    req = query("""
        SELECT rr.*, p.tanggal as tgl_faktur, p.dokter, ps.nm_pasien
        FROM Retur_Request rr
        LEFT JOIN Penjualan p ON rr.no_faktur=p.no_faktur
        LEFT JOIN Pasien ps ON p.kd_pasien=ps.kd_pasien
        WHERE rr.id=?
    """, (rid,), one=True)
    if not req:
        return jsonify({'ok': False, 'msg': 'Pengajuan tidak ditemukan'}), 404
    req = dict(req)
    items_req = json.loads(req.pop('items_json'))

    items = []
    for it in items_req:
        orig = query("SELECT kd_barang, nm_barang, jenis, qty, harga, disc FROM R_Penjualan WHERE id=?",
                     (it.get('id'),), one=True)
        if not orig:
            continue
        try:
            disc_pct = float(orig['disc'] or 0)
        except (TypeError, ValueError):
            disc_pct = 0.0
        qty_retur = float(it.get('qty_retur', 0))
        sub_retur = round(float(orig['harga']) * (1 - disc_pct / 100) * qty_retur, 2)
        items.append({
            'kd_barang': orig['kd_barang'], 'nm_barang': orig['nm_barang'], 'jenis': orig['jenis'],
            'qty_asli': orig['qty'], 'qty_retur': qty_retur, 'harga': orig['harga'],
            'disc': disc_pct, 'sub_retur': sub_retur,
        })
    try:
        req['umur_hari'] = (date.today() - datetime.strptime(req['tgl_faktur'], '%Y-%m-%d').date()).days
    except (ValueError, TypeError):
        req['umur_hari'] = None
    req['items'] = items
    req['total_retur_estimasi'] = round(sum(i['sub_retur'] for i in items), 2)
    return jsonify({'ok': True, **req})

@app.route('/api/retur-request/<int:rid>/setujui', methods=['POST'])
def retur_request_setujui(rid):
    """Supervisor/admin menyetujui pengajuan retur faktur lama — baru DI SINI stok/
    poin/komisi benar2 disentuh (lewat _eksekusi_retur, sama persis dgn retur langsung)."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor yang bisa menyetujui'}), 403
    req = query("SELECT * FROM Retur_Request WHERE id=?", (rid,), one=True)
    if not req:
        return jsonify({'ok': False, 'msg': 'Pengajuan tidak ditemukan'}), 404
    if req['status'] != 'pending':
        return jsonify({'ok': False, 'msg': f"Pengajuan ini sudah {req['status']}"}), 400

    items_req = json.loads(req['items_json'])
    result, status = _eksekusi_retur(req['no_faktur'], items_req, req['kasir'])
    if status != 200 or not result.get('ok'):
        return jsonify(result), status

    execute("""UPDATE Retur_Request SET status='disetujui', diproses_oleh=?, diproses_at=?,
               total_retur=? WHERE id=?""",
            (user.get('nama', ''), str(datetime.now()), result.get('total_retur', 0), rid))
    catat_log('retur_disetujui', 'Retur_Request', req['no_faktur'],
              f"Setujui pengajuan retur faktur {req['no_faktur']} — total Rp{result.get('total_retur',0):,.0f}")
    return jsonify(result)

@app.route('/api/retur-request/<int:rid>/tolak', methods=['POST'])
def retur_request_tolak(rid):
    """Supervisor/admin menolak pengajuan retur — TIDAK ada stok/poin/komisi yg
    tersentuh sama sekali (memang belum pernah dieksekusi sejak awal diajukan)."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor yang bisa menolak'}), 403
    req = query("SELECT * FROM Retur_Request WHERE id=?", (rid,), one=True)
    if not req:
        return jsonify({'ok': False, 'msg': 'Pengajuan tidak ditemukan'}), 404
    if req['status'] != 'pending':
        return jsonify({'ok': False, 'msg': f"Pengajuan ini sudah {req['status']}"}), 400

    catatan = (request.json or {}).get('catatan', '').strip()
    execute("""UPDATE Retur_Request SET status='ditolak', diproses_oleh=?, diproses_at=?,
               catatan_tolak=? WHERE id=?""",
            (user.get('nama', ''), str(datetime.now()), catatan, rid))
    catat_log('retur_ditolak', 'Retur_Request', req['no_faktur'],
              f"Tolak pengajuan retur faktur {req['no_faktur']}" + (f" — Alasan: {catatan}" if catatan else ''))
    return jsonify({'ok': True})


@app.route('/api/barang-stok-batches/<kd_barang>')
def get_stok_batches(kd_barang):
    """Semua batch stok satu produk (FEFO order) — untuk validasi multi-batch di POS
    (autoFefoSplit cart utama & auto-split bahan racikan). Urutan FEFO dihitung di
    PYTHON via _exp_key() — BUKAN 'ORDER BY exp' SQL, yg salah krn exp string 'M/D/YYYY'
    tanpa leading zero (lihat docstring _exp_key)."""
    kd_lokasi = request.args.get('kd_lokasi', 'CBG1')
    rows = rows_to_list(query("""
        SELECT id, kd_barang, nm_barang, exp, stok_sisa, harga, satuan
        FROM Stok_barang
        WHERE kd_barang=? AND stok_sisa > 0 AND harga > 0 AND kd_lokasi=?
    """, (kd_barang, kd_lokasi)))
    rows.sort(key=lambda r: _exp_key(r['exp']))
    total = sum(r['stok_sisa'] for r in rows)
    return jsonify({'batches': rows, 'total_stok': total})

@app.route('/api/penjualan/item-qty', methods=['PUT'])
def update_item_qty():
    """Update qty satu baris R_Penjualan, sinkron stok Stok_barang."""
    d = request.json
    item_id = d['id']
    new_qty = int(d['new_qty'])
    today   = str(date.today())

    item = query("SELECT * FROM R_Penjualan WHERE id=?", (item_id,), one=True)
    if not item:
        return jsonify({'ok': False, 'msg': 'Item tidak ditemukan'}), 404

    old_qty = item['qty']
    delta   = new_qty - old_qty
    kd      = item['kd_barang']
    exp     = item['exp'] or ''

    penjualan = query("SELECT kd_lokasi FROM Penjualan WHERE no_faktur=?", (item['no_faktur'],), one=True)
    kd_lokasi = (penjualan['kd_lokasi'] if penjualan else None) or 'CBG1'

    if new_qty <= 0:
        if item['jenis'] == 'barang':
            if exp:
                execute("UPDATE Stok_barang SET stok_sisa=stok_sisa+? WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                        (old_qty, kd, exp, kd_lokasi))
                sn = query("SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                           (kd, exp, kd_lokasi), one=True)
            else:
                execute("UPDATE Stok_barang SET stok_sisa=stok_sisa+? WHERE kd_barang=? AND kd_lokasi=?",
                        (old_qty, kd, kd_lokasi))
                sn = query("SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND kd_lokasi=? ORDER BY id LIMIT 1",
                           (kd, kd_lokasi), one=True)
            # Catat pembatalan item di Kartu Stok (nilai positif = stok kembali)
            execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                    (today, kd, exp, f"Batal item {item['no_faktur']}",
                     old_qty, sn['stok_sisa'] if sn else 0, sn['stok_gudang'] if sn else 0))
        execute("DELETE FROM R_Penjualan WHERE id=?", (item_id,))
    else:
        if delta > 0 and item['jenis'] == 'barang':
            sr = query("SELECT stok_sisa FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?", (kd, exp, kd_lokasi), one=True) \
                 if exp else query("SELECT stok_sisa FROM Stok_barang WHERE kd_barang=? AND kd_lokasi=? ORDER BY exp ASC LIMIT 1", (kd, kd_lokasi), one=True)
            avail = sr['stok_sisa'] if sr else 0
            if avail < delta:
                return jsonify({'ok': False, 'msg': f'Stok hanya tersisa {avail} unit untuk batch ini',
                                'avail': avail, 'max_qty': old_qty + avail}), 400

        disc_pct  = float(item['disc'] or 0)
        new_sub   = item['harga'] * (1 - disc_pct/100) * new_qty
        hbeli_sat = (float(item['total_harbeli'] or 0) / old_qty) if old_qty else 0
        execute("UPDATE R_Penjualan SET qty=?, sub_total=?, total_harbeli=? WHERE id=?",
                (new_qty, new_sub, hbeli_sat * new_qty, item_id))

        if item['jenis'] == 'barang' and delta != 0:
            if exp:
                if delta > 0:
                    execute("UPDATE Stok_barang SET stok_sisa=stok_sisa-? WHERE kd_barang=? AND exp=? AND kd_lokasi=?", (delta, kd, exp, kd_lokasi))
                else:
                    execute("UPDATE Stok_barang SET stok_sisa=stok_sisa+? WHERE kd_barang=? AND exp=? AND kd_lokasi=?", (-delta, kd, exp, kd_lokasi))
                sn = query("SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?", (kd, exp, kd_lokasi), one=True)
            else:
                if delta > 0:
                    execute("UPDATE Stok_barang SET stok_sisa=stok_sisa-? WHERE kd_barang=? AND kd_lokasi=?", (delta, kd, kd_lokasi))
                else:
                    execute("UPDATE Stok_barang SET stok_sisa=stok_sisa+? WHERE kd_barang=? AND kd_lokasi=?", (-delta, kd, kd_lokasi))
                sn = query("SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND kd_lokasi=? ORDER BY id LIMIT 1", (kd, kd_lokasi), one=True)
            if sn:
                execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                        (today, kd, exp, f"Ubah qty jual {item['no_faktur']} ({old_qty}→{new_qty})", -delta, sn['stok_sisa'], sn['stok_gudang']))

    execute("""UPDATE Penjualan SET total_harga=(SELECT COALESCE(SUM(sub_total),0) FROM R_Penjualan
               WHERE no_faktur=? AND jenis!='bahan_racikan') WHERE no_faktur=?""",
            (item['no_faktur'], item['no_faktur']))

    pasien_info = _pasien_faktur(item['no_faktur'])
    catat_log('ubah_qty_item', 'R_Penjualan', item['no_faktur'],
              f"Ubah qty '{item['nm_barang']}' di faktur {item['no_faktur']}: {old_qty} → {new_qty}"
              + (f" — Pasien: {pasien_info}" if pasien_info else ''),
              data_lama={'qty': old_qty}, data_baru={'qty': new_qty})

    stok_sisa = None
    if item['jenis'] == 'barang' and new_qty > 0 and exp:
        sr = query("SELECT stok_sisa FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?", (kd, exp, kd_lokasi), one=True)
        if sr: stok_sisa = sr['stok_sisa']

    return jsonify({'ok': True, 'new_qty': new_qty, 'stok_sisa': stok_sisa})

@app.route('/api/penjualan/bayar', methods=['POST'])
def bayar_penjualan():
    d = request.json
    faktur = d['no_faktur']
    today = str(date.today())
    metode = d.get('metode_bayar', 'Tunai')
    kd_pasien = d.get('kd_pasien', '')

    # Pengaman anti-dobel (defense-in-depth di luar guard klik frontend): kalau faktur ini
    # SUDAH 'done', jangan proses ulang — endpoint ini py efek samping tidak idempoten
    # (poin member, TKredit, dll) shg dipanggil 2x (klik dobel yg lolos, retry jaringan,
    # atau 2 tab) bisa bikin poin/kredit tercatat dobel. Aman dipanggil berkali² setelah ini.
    sudah = query("SELECT status FROM Penjualan WHERE no_faktur=?", (faktur,), one=True)
    if sudah and sudah['status'] == 'done':
        return jsonify({'ok': True, 'sudah_dibayar': True})

    dokter = (d.get('dokter') or '').strip()
    if not dokter:
        return jsonify({'ok': False, 'msg': 'Dokter wajib dipilih sebelum transaksi bisa dibayar'}), 400

    disc_trx = max(0.0, min(100.0, float(d.get('disc_trx', 0))))
    penjualan_now = query("SELECT total_harga FROM Penjualan WHERE no_faktur=?", (faktur,), one=True)
    # Pengaman: kalau header belum ada (mis. transaksi hanya racikan), buat dulu & hitung total
    # dari R_Penjualan — supaya bayar tidak "sukses" tapi transaksi hilang dari laporan.
    if not penjualan_now:
        execute("""INSERT INTO Penjualan (no_faktur,tanggal,total_harga,hpp,kd_pasien,kasir,status)
                   VALUES (?,?,0,0,?,?,'open')""",
                (faktur, today, kd_pasien, d.get('kasir','')))
        execute("""UPDATE Penjualan SET total_harga=(SELECT COALESCE(SUM(sub_total),0)
                   FROM R_Penjualan WHERE no_faktur=? AND jenis!='bahan_racikan') WHERE no_faktur=?""",
                (faktur, faktur))
        penjualan_now = query("SELECT total_harga FROM Penjualan WHERE no_faktur=?", (faktur,), one=True)
    total_before   = penjualan_now['total_harga'] if penjualan_now else 0
    disc_trx_nilai = total_before * disc_trx / 100
    total_akhir    = total_before - disc_trx_nilai

    # ── Pembayaran tersplit (multi-metode per faktur, mis. sebagian Tunai + sebagian
    # Transfer) — 2026-10-25. `pembayaran` (opsional dari frontend baru): list
    # [{metode,bank,nominal}, ...]. Kalau dikirim, WAJIB jumlahnya persis = total_akhir
    # (toleransi ±Rp1 pembulatan float JS) — ini uang, tidak boleh auto-disesuaikan diam².
    # Kalau TIDAK dikirim (compat penuh dgn request lama), fallback ke `metode` tunggal
    # spt sebelumnya, lalu disintesis jadi 1 baris supaya Penjualan_Pembayaran TETAP
    # terisi utk SEMUA transaksi baru (bukan cuma yg split).
    pembayaran = d.get('pembayaran') or []
    # Pembayaran pakai saldo point pasien (2026-10-25) — TIDAK bisa digabung split
    # (keputusan desain: harus menutup PENUH 1 faktur, spt 'Internal'), jadi divalidasi
    # & di-redeem di sini sebelum ada perubahan data apapun. `pts_perlu` dipakai lagi
    # di bawah (setelah faktur 'done') utk benar2 memotong DetailPoint.
    pts_perlu = 0
    if pembayaran:
        if any((p.get('metode') or '').strip() == 'Point Member' for p in pembayaran):
            return jsonify({'ok': False, 'msg':
                'Pembayaran pakai Point Member tidak bisa digabung tersplit — matikan toggle split & pilih Point Member sendiri utk bayar penuh pakai point'
            }), 400
        total_bayar_sum = sum(float(p.get('nominal') or 0) for p in pembayaran)
        if abs(total_bayar_sum - total_akhir) > 1:
            return jsonify({'ok': False, 'msg':
                f"Total pembayaran (Rp{total_bayar_sum:,.0f}) tidak sama dengan total tagihan (Rp{total_akhir:,.0f})".replace(',', '.')
            }), 400
        labels = []
        for p in pembayaran:
            mt = (p.get('metode') or 'Tunai').strip() or 'Tunai'
            bk = (p.get('bank') or '').strip()
            labels.append(f"{mt} - {bk}" if bk else mt)
        metode = ' + '.join(labels) if len(labels) > 1 else labels[0]
    else:
        mt0, bk0 = _split_metode_bank(metode)
        if mt0 == 'Point Member':
            if not kd_pasien:
                return jsonify({'ok': False, 'msg': 'Pasien wajib dipilih untuk membayar pakai point'}), 400
            saldo_row = query("""SELECT COALESCE(SUM(debit),0)-COALESCE(SUM(kredit),0) as s
                                 FROM DetailPoint WHERE kd_pasien=?""", (kd_pasien,), one=True)
            saldo_pt = saldo_row['s'] if saldo_row else 0
            redeem_rate = get_point_redeem_rate()
            # Jumlah point WAJIB diisi kasir sendiri (bukan dihitung diam² oleh sistem) —
            # minimal cukup ceil(total_akhir/redeem_rate) supaya nilainya menutup PENUH
            # tagihan (pembulatan konversi point→Rupiah jarang pas persis). Kasir boleh
            # isi lebih dari minimal, tidak boleh kurang.
            pts_minimum = 0
            if total_akhir > 0:
                pts_minimum = int(total_akhir // redeem_rate) + (1 if total_akhir % redeem_rate > 0.0001 else 0)
            try:
                pts_perlu = int(d.get('point_jumlah') or 0)
            except (TypeError, ValueError):
                pts_perlu = 0
            if pts_perlu < pts_minimum:
                return jsonify({'ok': False, 'msg':
                    f"Jumlah point ({pts_perlu}) kurang dari minimal {pts_minimum} point utk menutup tagihan Rp{total_akhir:,.0f}".replace(',', '.')
                }), 400
            if pts_perlu > saldo_pt:
                return jsonify({'ok': False, 'msg':
                    f"Saldo point tidak cukup — diisi {pts_perlu} point, saldo pasien hanya {int(saldo_pt)} point"
                }), 400
        pembayaran = [{'metode': mt0, 'bank': bk0, 'nominal': total_akhir}]

    kd_lokasi_trx = (d.get('kd_lokasi', '') or '').strip()
    execute("""UPDATE Penjualan SET status='done', metode_bayar=?, dokter=?, kd_pasien=?,
               disc_trx=?, disc_trx_nilai=?, total_harga=?,
               kd_lokasi=CASE WHEN ?<>'' THEN ? ELSE kd_lokasi END WHERE no_faktur=?""",
            (metode, dokter, kd_pasien,
             disc_trx, disc_trx_nilai, total_akhir,
             kd_lokasi_trx, kd_lokasi_trx, faktur))

    for p in pembayaran:
        execute("INSERT INTO Penjualan_Pembayaran (no_faktur,metode,bank,nominal) VALUES (?,?,?,?)",
                (faktur, (p.get('metode') or 'Tunai').strip() or 'Tunai',
                 (p.get('bank') or '').strip(), float(p.get('nominal') or 0)))

    # Counter faktur TIDAK dinaikkan di sini — nomor sudah dikunci atomik saat Transaksi Baru
    # (lihat _reserve_counter / get_nomor_faktur) → anti tabrakan kasir paralel.

    # Kredit / bank — 'Internal' (gratis) & 'Tunai'/'Point Member' bukan kanal pembayaran
    # riil, jadi tidak perlu dicatat sbg kredit/bank. 1 baris TKredit PER metode non-tunai
    # (bukan 1 baris total faktur lagi) — supaya faktur split (mis. separuh Tunai separuh
    # Transfer) catat nominal Transfer-nya saja, bukan seluruh total faktur.
    for p in pembayaran:
        mt = (p.get('metode') or 'Tunai').strip() or 'Tunai'
        bk = (p.get('bank') or '').strip()
        if mt not in ('Tunai', 'Point Member', 'Internal') and float(p.get('nominal') or 0) > 0:
            execute("INSERT INTO TKredit (nota,konsumen,total,bank,kasir,tanggal) VALUES (?,?,?,?,?,?)",
                    (faktur, d.get('nm_pasien',''), float(p.get('nominal') or 0),
                     f"{mt} - {bk}" if bk else mt, d.get('kasir',''), today))

    # Potong saldo point (redeem) — kalau faktur ini dibayar pakai Point Member, validasi
    # saldo sudah lolos di atas, di sini tinggal eksekusi baris kredit-nya.
    paid_via_point = any((p.get('metode') or '').strip() == 'Point Member' for p in pembayaran)
    if paid_via_point and pts_perlu > 0:
        execute("INSERT INTO DetailPoint (kd_pasien,keterangan,kredit,tanggal) VALUES (?,?,?,?)",
                (kd_pasien, faktur, pts_perlu, today))

    # Point member — earn point baru dari total belanja. Dikecualikan kalau faktur ini
    # SENDIRI dibayar pakai Point Member (keputusan desain: tidak boleh "muter" — point
    # yg dipakai bayar tidak ikut menghasilkan point baru lagi).
    if kd_pasien and not paid_via_point:
        penjualan = query("SELECT * FROM Penjualan WHERE no_faktur=?", (faktur,), one=True)
        if penjualan and penjualan['total_harga'] > 0:
            point_rate = get_point_rate()   # Rp per 1 point — dari Pengaturan (default 50.000)
            pts = int(penjualan['total_harga'] / point_rate)
            if pts > 0:
                execute("INSERT INTO DetailPoint (kd_pasien,keterangan,debit,tanggal) VALUES (?,?,?,?)",
                        (kd_pasien, faktur, pts, today))

    # Kunjungan berikutnya
    if d.get('next_visit'):
        existing = query("SELECT * FROM TKunjunganBerikut WHERE kd_pasien=?", (kd_pasien,), one=True)
        if existing:
            execute("UPDATE TKunjunganBerikut SET no_faktur=?,next_tanggal=? WHERE kd_pasien=?",
                    (faktur, d['next_visit'], kd_pasien))
        else:
            execute("INSERT INTO TKunjunganBerikut (no_faktur,next_tanggal,kd_pasien) VALUES (?,?,?)",
                    (faktur, d['next_visit'], kd_pasien))

    # Catat jumlah point yg dipotong di keterangan log — supaya audit "ada yang bermain"
    # bisa langsung kelihatan dari Log Aktivitas tanpa perlu buka Laporan Point segala.
    point_ket = f" — {pts_perlu} point dipotong dari saldo pasien {kd_pasien}" if paid_via_point else ""
    catat_log('penjualan_bayar', 'Penjualan', faktur,
              f"Transaksi {faktur} dibayar via {metode} — total Rp{total_akhir:,.0f}".replace(',', '.') + point_ket)
    return jsonify({'ok': True})

# ─── PEMBELIAN ───────────────────────────────────────────────
@app.route('/api/pembelian')
def get_pembelian_list():
    dari    = request.args.get('dari', '')
    sampai  = request.args.get('sampai', '')
    q       = request.args.get('q', '').strip()
    status  = request.args.get('status', '')   # lunas | belum

    # Default: hanya 2 hari terakhir (biar ringan). Saat ada pencarian (q), cari ke semua riwayat.
    if not dari and not q:
        dari = str(date.today() - timedelta(days=2))

    filters, params = [], []
    if dari:
        filters.append("tanggal >= ?"); params.append(dari)
    if sampai:
        filters.append("tanggal <= ?"); params.append(sampai)
    if q:
        filters.append("(no_faktur LIKE ? OR nm_suplier LIKE ?)")
        params += [f'%{q}%', f'%{q}%']
    if status == 'lunas':
        filters.append("lunas='Sudah'")
    elif status == 'belum':
        filters.append("lunas='Belum'")

    where = ("WHERE " + " AND ".join(filters)) if filters else ""
    rows = query(f"""
        SELECT pb.*,
               (SELECT COUNT(*) FROM R_Pembelian WHERE no_faktur=pb.no_faktur) AS jumlah_item
        FROM Pembelian pb {where}
        ORDER BY tanggal DESC, no_faktur DESC LIMIT 200
    """, params)
    result = rows_to_list(rows)
    total       = sum(r['total_harga'] for r in result)
    total_belum = sum(r['total_harga'] for r in result if r['lunas'] == 'Belum')
    return jsonify({'rows': result, 'total': total, 'total_belum': total_belum})


@app.route('/api/laporan/pembelian')
def laporan_pembelian():
    dari    = request.args.get('dari', '')
    sampai  = request.args.get('sampai', '')
    q       = request.args.get('q', '').strip()
    status  = request.args.get('status', '')
    today   = str(date.today())
    week_later  = (date.today() + timedelta(days=7)).isoformat()
    month_later = (date.today() + timedelta(days=30)).isoformat()

    filters, params = [], []
    if dari:   filters.append("pb.tanggal >= ?"); params.append(dari)
    if sampai: filters.append("pb.tanggal <= ?"); params.append(sampai)
    if q:
        filters.append("(pb.no_faktur LIKE ? OR pb.nm_suplier LIKE ?)")
        params += [f'%{q}%', f'%{q}%']
    if status == 'lunas':
        filters.append("pb.lunas='Sudah'")
    elif status == 'belum':
        filters.append("pb.lunas='Belum'")
    elif status == 'overdue':
        filters.append(f"pb.lunas='Belum' AND pb.tempo!='' AND pb.tempo < '{today}'")

    where = ("WHERE " + " AND ".join(filters)) if filters else ""
    rows = query(f"""
        SELECT pb.*,
               (SELECT COUNT(*) FROM R_Pembelian WHERE no_faktur=pb.no_faktur) AS jumlah_item
        FROM Pembelian pb {where}
        ORDER BY
            CASE WHEN pb.lunas='Belum' AND pb.tempo!='' AND pb.tempo < '{today}' THEN 0
                 WHEN pb.lunas='Belum' AND pb.tempo!='' AND pb.tempo <= '{week_later}' THEN 1
                 WHEN pb.lunas='Belum' AND pb.tempo!='' THEN 2
                 WHEN pb.lunas='Belum' THEN 3
                 ELSE 4 END,
            pb.tempo ASC, pb.tanggal DESC
        LIMIT 500
    """, params)

    result = rows_to_list(rows)
    for r in result:
        t = r.get('tempo') or ''
        if r['lunas'] == 'Sudah':
            r['tempo_status'] = 'lunas'
        elif not t:
            r['tempo_status'] = 'notempo'
        elif t < today:
            r['tempo_status'] = 'overdue'
        elif t <= week_later:
            r['tempo_status'] = 'urgent'
        elif t <= month_later:
            r['tempo_status'] = 'upcoming'
        else:
            r['tempo_status'] = 'ok'

    total        = sum(r['total_harga'] for r in result)
    total_belum  = sum(r['total_harga'] for r in result if r['lunas'] == 'Belum')
    cnt_overdue  = sum(1 for r in result if r.get('tempo_status') == 'overdue')
    cnt_urgent   = sum(1 for r in result if r.get('tempo_status') == 'urgent')
    val_overdue  = sum(r['total_harga'] for r in result if r.get('tempo_status') == 'overdue')
    val_urgent   = sum(r['total_harga'] for r in result if r.get('tempo_status') == 'urgent')

    return jsonify({
        'rows': result, 'total': total, 'total_belum': total_belum,
        'cnt_overdue': cnt_overdue, 'cnt_urgent': cnt_urgent,
        'val_overdue': val_overdue, 'val_urgent': val_urgent,
        'total_faktur': len(result), 'today': today,
    })


@app.route('/api/pembelian/next-faktur')
def pembelian_next_faktur():
    today = str(date.today())
    ymd   = today.replace('-', '')[2:]          # YYMMDD
    last  = query("SELECT no_faktur FROM Pembelian WHERE no_faktur LIKE ? ORDER BY no_faktur DESC LIMIT 1",
                  (f'PB{ymd}%',), one=True)
    seq   = (int(last['no_faktur'][-3:]) + 1) if last else 1
    return jsonify({'faktur': f'PB{ymd}{seq:03d}', 'tanggal': today})


@app.route('/api/pembelian/<faktur>')
def get_pembelian(faktur):
    header = query("SELECT * FROM Pembelian WHERE no_faktur=?", (faktur,), one=True)
    items  = query("SELECT * FROM R_Pembelian WHERE no_faktur=? ORDER BY id", (faktur,))
    return jsonify({'header': dict(header) if header else None, 'items': rows_to_list(items)})


def _pembelian_item_sudah_bergerak(item, pemilik):
    """True kalau stok batch (kd_barang+exp+pemilik) dari item pembelian ini SUDAH
    berkurang di bawah qty yang dibeli — artinya sebagian/semua sudah terdistribusi
    ke cabang (order transfer), dipakai produksi, atau dikonversi. Kalau begitu,
    item ini TIDAK BOLEH diedit/dihapus lagi krn akan bikin data stok tidak akurat
    (qty pembelian dihapus tapi barangnya sudah kadung keluar/terpakai)."""
    stok = query("""SELECT COALESCE(SUM(stok_sisa),0) as s FROM Stok_barang
                    WHERE kd_barang=? AND exp=? AND kd_lokasi='INDUK' AND pemilik=?""",
                 (item['kd_barang'], item['exp'], pemilik), one=True)
    sisa = stok['s'] if stok else 0
    return sisa < item['qty']

def _pembelian_edit_guard(faktur):
    """Gerbang edit/hapus Pembelian, mempertimbangkan status approval + umur faktur.
    Admin/supervisor selalu boleh (termasuk utk cleanup). Selain itu: draft & rejected
    bebas diedit pembuatnya (rejected = supervisor tolak dgn alasan, koreksi jadi
    tugas user level INDUK yg mengajukan — bukan supervisor yg mengubah langsung),
    submitted diblokir total (sedang direview), approved pakai aturan lama (>1 hari
    → admin/supervisor saja). Return (response, status) bila TIDAK boleh, atau None
    bila boleh."""
    user = get_current_user()
    role = (user or {}).get('role', '')
    if role in ('admin', 'supervisor'):
        return None
    head = query("SELECT tanggal, status FROM Pembelian WHERE no_faktur=?", (faktur,), one=True)
    if not head:
        return None   # faktur baru / belum ada → bukan pengubahan data lama
    status = head['status'] or 'approved'
    if status in ('draft', 'rejected'):
        return None   # blm ada dampak stok — bebas diubah pembuatnya (termasuk revisi setelah ditolak)
    if status == 'submitted':
        return jsonify({'ok': False,
            'msg': 'Faktur ini sedang menunggu approval supervisor/admin — tidak bisa diubah sampai ada keputusan.'}), 403
    # status == 'approved' → aturan waktu lama
    try:
        tgl = datetime.strptime(str(head['tanggal'])[:10], '%Y-%m-%d').date()
    except Exception:
        return None
    if (date.today() - tgl).days > 1:
        return jsonify({'ok': False,
            'msg': 'Pembelian lebih dari 1 hari hanya bisa diubah/dihapus oleh Administrator atau Supervisor'}), 403
    return None

@app.route('/api/pembelian/<faktur>', methods=['PUT'])
def update_pembelian_header(faktur):
    blok = _pembelian_edit_guard(faktur)
    if blok: return blok
    d = request.json
    allowed = ['nm_suplier', 'kd_suplier', 'sistem_bayar', 'tempo', 'lunas']
    sets, params = [], []
    for f in allowed:
        if f in d:
            sets.append(f"{f}=?"); params.append(d[f])
    if not sets:
        return jsonify({'ok': False, 'msg': 'Tidak ada field yang diupdate'}), 400
    params.append(faktur)
    execute(f"UPDATE Pembelian SET {','.join(sets)} WHERE no_faktur=?", params)
    return jsonify({'ok': True})


def _hitung_harga_pembelian(harga_asli, disc, ppn, qty):
    """Satu sumber kebenaran rumus harga beli — dipakai saat input awal (draft)
    MAUPUN saat supervisor mengoreksi nilai sebelum approve, supaya keduanya
    tidak pernah menghitung dgn cara berbeda."""
    har_disc  = harga_asli * (100 - disc) / 100
    har_ppn   = har_disc   * (100 + ppn)  / 100
    sub_total = har_disc   * qty
    return har_disc, har_ppn, sub_total

def _pembelian_posting_stok(item, pemilik, by):
    """Posting stok SESUNGGUHNYA ke Stok Induk untuk 1 baris R_Pembelian — dipanggil
    HANYA sekali, tepat saat faktur disetujui (bukan saat item diinput). Sebelum
    fitur approval ini, blok ini jalan langsung di add_item_pembelian; sekarang
    ditunda sampai supervisor/admin klik Setujui."""
    kd_barang = item['kd_barang']; nm_barang = item['nm_barang'] or ''
    exp = item['exp'] or ''; qty = item['qty']; satuan = item['satuan'] or 'Pcs'
    harga_beli = item['harga_beli_bersih'] or 0
    today = str(date.today())

    mb = query("SELECT nm_barang, satuan, harga_jual FROM Barang WHERE kd_barang=?", (kd_barang,), one=True)
    if not mb:
        execute("""INSERT OR IGNORE INTO Barang (kd_barang,nm_barang,nm_generik,satuan,harga_jual,aktif)
                   VALUES (?,?,?,?,0,1)""", (kd_barang, nm_barang, '', satuan))
        harga_jual = 0
    else:
        harga_jual = mb['harga_jual'] or 0

    existing_stok = query("SELECT id FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi='INDUK' AND pemilik=?",
                          (kd_barang, exp, pemilik), one=True)
    if existing_stok:
        execute("UPDATE Stok_barang SET stok_awal=stok_awal+?,stok_sisa=stok_sisa+?,stok_gudang=stok_gudang+?,harga_beli=? WHERE id=?",
                (qty, qty, qty, harga_beli, existing_stok['id']))
    else:
        execute("""INSERT INTO Stok_barang (kd_barang,nm_barang,harga_beli,harga_beli_awal,stok_awal,stok_sisa,stok_gudang,stok_min,harga,disc,satuan,exp,DiscMem,DiscLain,kd_lokasi,pemilik)
                   VALUES (?,?,?,?,?,?,?,0,?,'0',?,?,0,0,'INDUK',?)""",
                (kd_barang, nm_barang, harga_beli, harga_beli, qty, qty, qty, harga_jual, satuan, exp, pemilik))

    execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
            (today, kd_barang, exp, f"Beli {item['no_faktur']}", qty, qty, qty))
    execute("INSERT INTO Transfer_Stok_Log (no_referensi,tipe_referensi,kd_barang,nm_barang,kd_lokasi_dari,kd_lokasi_ke,qty,exp,created_by) VALUES (?,?,?,?,?,?,?,?,?)",
            (item['no_faktur'], 'pembelian', kd_barang, nm_barang, '', 'INDUK', qty, exp, by))

@app.route('/api/pembelian/item', methods=['POST'])
def add_item_pembelian():
    d         = request.json
    faktur    = d['no_faktur']
    blok = _pembelian_edit_guard(faktur)   # faktur lama (>1 hari)/submitted/rejected → sesuai status
    if blok: return blok
    today     = str(date.today())
    harga_asli = float(d.get('harga_asli', 0))
    disc      = float(d.get('disc', 0))
    ppn       = float(d.get('ppn', 0))
    qty       = int(d.get('qty', 1))
    har_disc, har_ppn, sub_total = _hitung_harga_pembelian(harga_asli, disc, ppn, qty)

    kd_suplier   = d.get('kd_suplier', '')
    nm_suplier   = d.get('nm_suplier', '')
    sistem_bayar = d.get('sistem_bayar', 'Kredit')
    tempo        = d.get('tempo', today)
    exp          = d.get('exp', today)
    pemilik      = (d.get('pemilik') or 'P1').strip()   # stok induk ini ditujukan ke pemilik mana

    existing_h = query("SELECT no_faktur, status FROM Pembelian WHERE no_faktur=?", (faktur,), one=True)
    if not existing_h:
        lunas = 'Sudah' if sistem_bayar == 'Cash' else 'Belum'
        pembuat = (get_current_user() or {}).get('nama', '')
        execute("""INSERT INTO Pembelian (no_faktur,tanggal,total_harga,sistem_bayar,lunas,tempo,nm_suplier,pemilik,status,dibuat_oleh)
                   VALUES (?,?,0,?,?,?,?,?,'draft',?)""",
                (faktur, today, sistem_bayar, lunas, tempo, nm_suplier, pemilik, pembuat))
    else:
        if (existing_h['status'] or 'draft') not in ('draft', 'rejected'):
            return jsonify({'ok': False,
                'msg': f"Faktur ini berstatus '{existing_h['status']}' — tidak bisa menambah item lagi. Buat faktur baru."}), 400
        # Faktur sama → pastikan pemilik header konsisten (set bila masih kosong)
        execute("UPDATE Pembelian SET pemilik=? WHERE no_faktur=? AND (pemilik IS NULL OR pemilik='')",
                (pemilik, faktur))

    existing = query("SELECT id FROM R_Pembelian WHERE no_faktur=? AND kd_barang=? AND exp=?",
                     (faktur, d['kd_barang'], exp), one=True)
    if existing:
        execute("UPDATE R_Pembelian SET qty=qty+?,sub_total=sub_total+? WHERE id=?",
                (qty, sub_total, existing['id']))
    else:
        execute("""INSERT INTO R_Pembelian
                   (no_faktur,tanggal,kd_suplier,nm_suplier,sistem_bayar,kd_barang,nm_barang,
                    disc,harga_beli_kotor,harga_beli_bersih,sub_total,ppn,tempo,qty,harga_asli,satuan,exp)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (faktur, today, kd_suplier, nm_suplier, sistem_bayar,
                 d['kd_barang'], d['nm_barang'], disc, har_disc, har_ppn,
                 sub_total, ppn, tempo, qty, harga_asli, d.get('satuan', 'Pcs'), exp))

    execute("UPDATE Pembelian SET total_harga=(SELECT COALESCE(SUM(sub_total),0) FROM R_Pembelian WHERE no_faktur=?) WHERE no_faktur=?",
            (faktur, faktur))

    # CATATAN: stok TIDAK diposting di sini lagi — faktur masih draft, blm direview.
    # Stok baru masuk ke Stok Induk saat faktur ini disetujui (lihat approve_pembelian()).

    item = query("SELECT id FROM R_Pembelian WHERE no_faktur=? AND kd_barang=? AND exp=?",
                 (faktur, d['kd_barang'], exp), one=True)
    return jsonify({'ok': True, 'id': item['id'] if item else None,
                    'harga_beli': har_disc, 'sub_total': sub_total})


@app.route('/api/pembelian/<faktur>/submit', methods=['POST'])
def submit_pembelian(faktur):
    """Ajukan faktur draft (atau rejected yg sudah direvisi user INDUK) ke supervisor/
    admin utk direview — dari titik ini staf biasa tidak bisa lagi mengubahnya
    (lihat _pembelian_edit_guard)."""
    head = query("SELECT status FROM Pembelian WHERE no_faktur=?", (faktur,), one=True)
    if not head:
        return jsonify({'ok': False, 'msg': 'Faktur tidak ditemukan'}), 404
    if (head['status'] or 'draft') not in ('draft', 'rejected'):
        return jsonify({'ok': False, 'msg': 'Hanya faktur berstatus draft/ditolak yang bisa diajukan'}), 400
    jml = query("SELECT COUNT(*) as c FROM R_Pembelian WHERE no_faktur=?", (faktur,), one=True)
    if not jml or not jml['c']:
        return jsonify({'ok': False, 'msg': 'Tambahkan minimal 1 barang sebelum mengajukan'}), 400
    execute("UPDATE Pembelian SET status='submitted' WHERE no_faktur=?", (faktur,))
    catat_log('pembelian_ajukan', 'Pembelian', faktur, f"Ajukan faktur pembelian {faktur} ke supervisor/admin")
    return jsonify({'ok': True})

@app.route('/api/pembelian/<faktur>/approve', methods=['POST'])
def approve_pembelian(faktur):
    """Supervisor/admin HANYA setuju/tolak — tidak mengubah nilai yg diajukan.
    Kalau ada yg keliru, tolak dgn alasan; koreksi jadi tugas user level INDUK yg
    mengajukan (edit faktur yg statusnya balik ke 'rejected', lalu ajukan ulang).
    Stok Induk baru diposting TEPAT DI SINI — inilah gerbang approval yang sesungguhnya."""
    user = get_current_user()
    role = (user or {}).get('role', '')
    if role not in ('admin', 'supervisor'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor yang boleh approve'}), 403
    d = request.json or {}
    head = query("SELECT * FROM Pembelian WHERE no_faktur=?", (faktur,), one=True)
    if not head or (head['status'] or '') != 'submitted':
        return jsonify({'ok': False, 'msg': 'Hanya faktur berstatus submitted yang bisa di-approve'}), 400
    pemilik = head['pemilik'] if head['pemilik'] else 'P1'

    # Posting stok — sekali per item, inilah efek nyata yang ditunda sejak draft.
    # Kartu Stok mencatat "Oleh" = maker pertama (yg input faktur), BUKAN yg approve —
    # approver-nya sendiri tetap tersimpan terpisah di Pembelian.approved_by.
    items = query("SELECT * FROM R_Pembelian WHERE no_faktur=?", (faktur,))
    pembuat = head['dibuat_oleh'] or ''
    approver = d.get('approved_by') or (user or {}).get('nama', '')
    for item in items:
        _pembelian_posting_stok(item, pemilik, pembuat)

    execute("""UPDATE Pembelian SET status='approved',catatan_supervisor=?,approved_by=?,approved_at=?
               WHERE no_faktur=?""",
            (d.get('catatan_supervisor', ''), approver, str(datetime.now()), faktur))
    catat_log('pembelian_approve', 'Pembelian', faktur, f"Approve faktur pembelian {faktur}",
              data_baru={'jml_item': len(items), 'total_harga': head['total_harga']})
    return jsonify({'ok': True})

@app.route('/api/pembelian/<faktur>/reject', methods=['POST'])
def reject_pembelian(faktur):
    user = get_current_user()
    role = (user or {}).get('role', '')
    if role not in ('admin', 'supervisor'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor yang boleh menolak'}), 403
    d = request.json or {}
    head = query("SELECT status FROM Pembelian WHERE no_faktur=?", (faktur,), one=True)
    if not head or (head['status'] or '') != 'submitted':
        return jsonify({'ok': False, 'msg': 'Hanya faktur berstatus submitted yang bisa ditolak'}), 400
    catatan = d.get('catatan', '')
    execute("""UPDATE Pembelian SET status='rejected',catatan_supervisor=?,approved_by=?,approved_at=?
               WHERE no_faktur=?""",
            (catatan, d.get('by', ''), str(datetime.now()), faktur))
    catat_log('pembelian_reject', 'Pembelian', faktur, f"Tolak faktur pembelian {faktur}: {catatan}")
    return jsonify({'ok': True})


@app.route('/api/pembelian/item/<int:id>', methods=['DELETE'])
def delete_item_pembelian(id):
    item = query("SELECT * FROM R_Pembelian WHERE id=?", (id,), one=True)
    if not item:
        return jsonify({'ok': False, 'msg': 'Item tidak ditemukan'}), 404
    blok = _pembelian_edit_guard(item['no_faktur'])
    if blok: return blok
    head = query("SELECT pemilik, status FROM Pembelian WHERE no_faktur=?", (item['no_faktur'],), one=True)
    pemilik = (head['pemilik'] if head and head['pemilik'] else 'P1')
    if (head['status'] if head else 'approved') != 'approved':
        # Belum pernah diposting ke Stok_barang (masih draft/submitted) — hapus
        # langsung tanpa reversal stok & tanpa cek sudah-bergerak (tidak relevan).
        execute("DELETE FROM R_Pembelian WHERE id=?", (id,))
        execute("UPDATE Pembelian SET total_harga=(SELECT COALESCE(SUM(sub_total),0) FROM R_Pembelian WHERE no_faktur=?) WHERE no_faktur=?",
                (item['no_faktur'], item['no_faktur']))
        return jsonify({'ok': True})
    if _pembelian_item_sudah_bergerak(item, pemilik):
        return jsonify({'ok': False, 'locked_needs_approval': True,
            'msg': f"Barang '{item['nm_barang']}' dari pembelian ini sudah terdistribusi/terpakai "
                   f"(stok induk tersisa lebih sedikit dari qty pembelian) — tidak bisa dihapus langsung. "
                   f"Ajukan permintaan koreksi ke admin/supervisor."}), 409
    # Kembalikan stok induk milik pemilik terkait
    execute("UPDATE Stok_barang SET stok_awal=MAX(0,stok_awal-?),stok_sisa=MAX(0,stok_sisa-?),stok_gudang=MAX(0,stok_gudang-?) WHERE kd_barang=? AND exp=? AND kd_lokasi='INDUK' AND pemilik=?",
            (item['qty'], item['qty'], item['qty'], item['kd_barang'], item['exp'], pemilik))
    execute("DELETE FROM R_Pembelian WHERE id=?", (id,))
    execute("UPDATE Pembelian SET total_harga=(SELECT COALESCE(SUM(sub_total),0) FROM R_Pembelian WHERE no_faktur=?) WHERE no_faktur=?",
            (item['no_faktur'], item['no_faktur']))
    execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
            (str(date.today()), item['kd_barang'], item['exp'],
             f"Hapus item beli {item['no_faktur']}", -item['qty'], 0, 0))
    # Catat juga "keluar" di Transfer_Stok_Log — kalau tidak, Kartu Stok Induk (yg
    # murni dihitung dari log ini) akan terus menyimpan jejak "masuk" hantu yg tak
    # pernah di-offset, bikin Saldo Akhir Periode menyimpang dari Stok Aktual DB.
    execute("INSERT INTO Transfer_Stok_Log (no_referensi,tipe_referensi,kd_barang,nm_barang,kd_lokasi_dari,kd_lokasi_ke,qty,exp,created_by) VALUES (?,?,?,?,?,?,?,?,?)",
            (item['no_faktur'], 'pembelian_hapus', item['kd_barang'], item['nm_barang'], 'INDUK', '', item['qty'], item['exp'],
             (get_current_user() or {}).get('nama', '')))
    return jsonify({'ok': True})


@app.route('/api/pembelian/<faktur>', methods=['DELETE'])
def delete_pembelian(faktur):
    blok = _pembelian_edit_guard(faktur)
    if blok: return blok
    items = query("SELECT * FROM R_Pembelian WHERE no_faktur=?", (faktur,))
    today = str(date.today())
    head = query("SELECT pemilik, status FROM Pembelian WHERE no_faktur=?", (faktur,), one=True)
    pemilik = (head['pemilik'] if head and head['pemilik'] else 'P1')
    if (head['status'] if head else 'approved') != 'approved':
        # Belum pernah diposting ke Stok_barang — hapus langsung tanpa reversal stok.
        execute("DELETE FROM R_Pembelian WHERE no_faktur=?", (faktur,))
        execute("DELETE FROM Pembelian WHERE no_faktur=?", (faktur,))
        return jsonify({'ok': True})
    # Cek SEMUA item dulu (all-or-nothing) — kalau satu saja sudah terdistribusi/
    # terpakai, seluruh faktur ini tidak boleh dihapus.
    bermasalah = [item['nm_barang'] for item in items if _pembelian_item_sudah_bergerak(item, pemilik)]
    if bermasalah:
        return jsonify({'ok': False, 'locked_needs_approval': True,
            'msg': f"Faktur ini tidak bisa dihapus langsung — barang berikut sudah terdistribusi/terpakai: "
                   f"{', '.join(bermasalah)}. Ajukan permintaan koreksi ke admin/supervisor."}), 409
    by = (get_current_user() or {}).get('nama', '')
    for item in items:
        execute("UPDATE Stok_barang SET stok_awal=MAX(0,stok_awal-?),stok_sisa=MAX(0,stok_sisa-?),stok_gudang=MAX(0,stok_gudang-?) WHERE kd_barang=? AND exp=? AND kd_lokasi='INDUK' AND pemilik=?",
                (item['qty'], item['qty'], item['qty'], item['kd_barang'], item['exp'], pemilik))
        execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                (today, item['kd_barang'], item['exp'], f"Batal beli {faktur}", -item['qty'], 0, 0))
        # Catat "keluar" di Transfer_Stok_Log jg — lihat penjelasan di delete_item_pembelian.
        execute("INSERT INTO Transfer_Stok_Log (no_referensi,tipe_referensi,kd_barang,nm_barang,kd_lokasi_dari,kd_lokasi_ke,qty,exp,created_by) VALUES (?,?,?,?,?,?,?,?,?)",
                (faktur, 'pembelian_hapus', item['kd_barang'], item['nm_barang'], 'INDUK', '', item['qty'], item['exp'], by))
    execute("DELETE FROM R_Pembelian WHERE no_faktur=?", (faktur,))
    execute("DELETE FROM Pembelian WHERE no_faktur=?", (faktur,))
    return jsonify({'ok': True})


# ─── KOREKSI PEMBELIAN (barang sudah terdistribusi) — butuh approval ─
@app.route('/api/pembelian/koreksi', methods=['POST'])
def ajukan_koreksi_pembelian():
    """Kasir/staff ajukan permintaan hapus item/faktur pembelian yang sudah
    terkunci (stok sudah bergerak) — wajib isi alasan, menunggu approval
    admin/supervisor lewat Approval Queue."""
    d = request.json or {}
    no_faktur = (d.get('no_faktur') or '').strip()
    id_item   = d.get('id_item')  # None/kosong = ajukan hapus SELURUH faktur
    jenis     = 'hapus_item' if id_item else 'hapus_faktur'
    catatan   = (d.get('catatan') or '').strip()
    if not no_faktur:
        return jsonify({'ok': False, 'msg': 'No faktur wajib diisi'}), 400
    if not catatan:
        return jsonify({'ok': False, 'msg': 'Alasan/catatan wajib diisi'}), 400
    user = get_current_user()
    rid = execute("""INSERT INTO Pembelian_Koreksi (no_faktur,id_item,jenis,catatan,requested_by)
                     VALUES (?,?,?,?,?)""",
                  (no_faktur, id_item, jenis, catatan, (user or {}).get('nama', '')))
    catat_log('koreksi_pembelian_ajukan', 'Pembelian_Koreksi', no_faktur,
              f"Ajukan koreksi ({jenis}) faktur {no_faktur}: {catatan}")
    return jsonify({'ok': True, 'id': rid})


@app.route('/api/pembelian/koreksi/<int:kid>/approve', methods=['POST'])
def approve_koreksi_pembelian(kid):
    user = get_current_user()
    role = (user or {}).get('role', '')
    if role not in ('admin', 'supervisor'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor yang boleh approve'}), 403
    req = query("SELECT * FROM Pembelian_Koreksi WHERE id=?", (kid,), one=True)
    if not req or req['status'] != 'pending':
        return jsonify({'ok': False, 'msg': 'Request tidak ditemukan atau sudah diproses'}), 400

    head = query("SELECT pemilik FROM Pembelian WHERE no_faktur=?", (req['no_faktur'],), one=True)
    if not head:
        return jsonify({'ok': False, 'msg': 'Faktur pembelian sudah tidak ada'}), 400
    pemilik = head['pemilik'] if head['pemilik'] else 'P1'

    if req['jenis'] == 'hapus_item':
        item = query("SELECT * FROM R_Pembelian WHERE id=?", (req['id_item'],), one=True)
        if not item:
            return jsonify({'ok': False, 'msg': 'Item pembelian sudah tidak ada (mungkin sudah dihapus sebelumnya)'}), 400
        items_to_check = [item]
    else:
        items_to_check = query("SELECT * FROM R_Pembelian WHERE no_faktur=?", (req['no_faktur'],))
        if not items_to_check:
            return jsonify({'ok': False, 'msg': 'Faktur sudah tidak ada itemnya'}), 400

    # WAJIB dicek ULANG saat approve (bukan pakai kondisi saat pengajuan) — dan
    # digabung per batch (kd_barang+exp) supaya kalau 1 faktur punya >1 baris
    # yg kebetulan berbagi batch, totalnya tetap tidak boleh bikin stok minus.
    qty_by_batch  = {}
    nama_by_batch = {}
    for it in items_to_check:
        key = (it['kd_barang'], it['exp'])
        qty_by_batch[key]  = qty_by_batch.get(key, 0) + it['qty']
        nama_by_batch[key] = it['nm_barang']
    for (kd_barang, exp), total_qty in qty_by_batch.items():
        stok = query("""SELECT COALESCE(SUM(stok_sisa),0) as s FROM Stok_barang
                        WHERE kd_barang=? AND exp=? AND kd_lokasi='INDUK' AND pemilik=?""",
                     (kd_barang, exp, pemilik), one=True)
        sisa = stok['s'] if stok else 0
        if sisa - total_qty < 0:
            return jsonify({'ok': False,
                'msg': f"Tidak bisa disetujui — stok '{nama_by_batch[(kd_barang, exp)]}' akan menjadi minus "
                       f"(sisa saat ini {sisa}, qty yang mau dihapus {total_qty}). Approval dibatalkan, "
                       f"stok tidak diubah sama sekali."}), 409

    # Aman (sudah dipastikan tidak akan minus) — eksekusi penghapusan sesungguhnya.
    today = str(date.today())
    data_lama = rows_to_list(items_to_check)
    for it in items_to_check:
        execute("""UPDATE Stok_barang SET stok_awal=stok_awal-?,stok_sisa=stok_sisa-?,stok_gudang=stok_gudang-?
                   WHERE kd_barang=? AND exp=? AND kd_lokasi='INDUK' AND pemilik=?""",
                (it['qty'], it['qty'], it['qty'], it['kd_barang'], it['exp'], pemilik))
        execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                (today, it['kd_barang'], it['exp'],
                 f"Koreksi disetujui: hapus beli {req['no_faktur']}", -it['qty'], 0, 0))

    if req['jenis'] == 'hapus_item':
        execute("DELETE FROM R_Pembelian WHERE id=?", (req['id_item'],))
        execute("UPDATE Pembelian SET total_harga=(SELECT COALESCE(SUM(sub_total),0) FROM R_Pembelian WHERE no_faktur=?) WHERE no_faktur=?",
                (req['no_faktur'], req['no_faktur']))
    else:
        execute("DELETE FROM R_Pembelian WHERE no_faktur=?", (req['no_faktur'],))
        execute("DELETE FROM Pembelian WHERE no_faktur=?", (req['no_faktur'],))

    execute("UPDATE Pembelian_Koreksi SET status='approved',approved_by=?,approved_at=? WHERE id=?",
            ((user or {}).get('nama', ''), str(datetime.now()), kid))

    catat_log('koreksi_pembelian_approve', 'Pembelian_Koreksi', req['no_faktur'],
              f"Approve koreksi ({req['jenis']}) faktur {req['no_faktur']} — alasan pengaju: {req['catatan']}",
              data_lama=data_lama)
    return jsonify({'ok': True})


@app.route('/api/pembelian/koreksi/<int:kid>/reject', methods=['POST'])
def reject_koreksi_pembelian(kid):
    user = get_current_user()
    role = (user or {}).get('role', '')
    if role not in ('admin', 'supervisor'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor yang boleh menolak'}), 403
    d = request.json or {}
    req = query("SELECT no_faktur, status FROM Pembelian_Koreksi WHERE id=?", (kid,), one=True)
    if not req or req['status'] != 'pending':
        return jsonify({'ok': False, 'msg': 'Request tidak ditemukan atau sudah diproses'}), 400
    execute("UPDATE Pembelian_Koreksi SET status='rejected',alasan=?,approved_by=?,approved_at=? WHERE id=?",
            (d.get('alasan', ''), (user or {}).get('nama', ''), str(datetime.now()), kid))
    catat_log('koreksi_pembelian_reject', 'Pembelian_Koreksi', req['no_faktur'],
              f"Tolak koreksi pembelian faktur {req['no_faktur']}: {d.get('alasan', '')}")
    return jsonify({'ok': True})


@app.route('/api/barang-beli')
def get_barang_beli():
    """Search master Barang untuk pembelian (kode konsisten, tak harus ada stok)."""
    # kd_barang di MySQL pakai collation utf8mb4_bin (case-SENSITIVE, sengaja biar
    # exact-match/join presisi) — LOWER() di kedua sisi wajib spy pencarian tetap
    # case-insensitive spt kolom lain (lihat _kata_cocok_kolom() utk penjelasan lengkap).
    q = request.args.get('q', '')
    like = f'%{q.lower()}%'
    rows = query("""
        SELECT b.kd_barang, b.nm_barang, b.nm_generik, b.satuan, b.harga_jual,
               COALESCE((SELECT MAX(harga_beli) FROM Stok_barang s
                         WHERE s.kd_barang=b.kd_barang AND s.kd_lokasi='INDUK'),0) AS harga_beli,
               COALESCE((SELECT SUM(stok_sisa) FROM Stok_barang s
                         WHERE s.kd_barang=b.kd_barang AND s.kd_lokasi='INDUK'),0) AS total_stok
        FROM Barang b
        WHERE b.aktif=1 AND (LOWER(b.kd_barang) LIKE ? OR LOWER(b.nm_barang) LIKE ? OR LOWER(b.nm_generik) LIKE ?)
        ORDER BY b.nm_barang LIMIT 20
    """, (like, like, like))
    return jsonify(rows_to_list(rows))


# ─── MASTER BARANG ───────────────────────────────────────────
@app.route('/api/barang')
def get_barang_master():
    """List/search master Barang. ?q= filter, ?all=1 termasuk nonaktif."""
    # kd_barang collation utf8mb4_bin (case-sensitive) — LOWER() wajib, lihat
    # _kata_cocok_kolom() utk penjelasan lengkap kenapa.
    q    = request.args.get('q', '')
    incl = request.args.get('all', '') == '1'
    like = f'%{q.lower()}%'
    where = "(LOWER(b.kd_barang) LIKE ? OR LOWER(b.nm_barang) LIKE ? OR LOWER(b.nm_generik) LIKE ? OR LOWER(b.merk) LIKE ?)"
    params = [like, like, like, like]
    if not incl:
        where += " AND b.aktif=1"
    rows = query(f"""
        SELECT b.*,
               COALESCE((SELECT SUM(stok_sisa) FROM Stok_barang s
                         WHERE s.kd_barang=b.kd_barang AND s.kd_lokasi='INDUK'),0) AS stok_induk
        FROM Barang b
        WHERE {where}
        ORDER BY b.nm_barang LIMIT 500
    """, params)
    return jsonify(rows_to_list(rows))


@app.route('/api/barang', methods=['POST'])
def create_barang_master():
    d  = request.json or {}
    kd = (d.get('kd_barang') or '').strip()
    nm = (d.get('nm_barang') or '').strip() or kd
    if not kd:
        return jsonify({'ok': False, 'msg': 'Kode barang wajib diisi'}), 400
    dup = query("SELECT kd_barang FROM Barang WHERE kd_barang=?", (kd,), one=True)
    if dup:
        return jsonify({'ok': False, 'msg': f'Kode "{kd}" sudah ada'}), 400
    execute("""INSERT INTO Barang
               (kd_barang,nm_barang,nm_generik,jenis,merk,isi,satuan,harga_jual,aktif,non_obat)
               VALUES (?,?,?,?,?,?,?,?,1,?)""",
            (kd, nm, d.get('nm_generik', ''), d.get('jenis', ''), d.get('merk', ''),
             float(d.get('isi', 1) or 1), d.get('satuan', 'Pcs') or 'Pcs',
             float(d.get('harga_jual', 0) or 0), int(d.get('non_obat', 0) or 0)))
    catat_log('barang_master_buat', 'Barang', kd, f"Buat master barang baru '{kd}' - {nm}",
              data_baru={'kd_barang': kd, 'nm_barang': nm, 'harga_jual': d.get('harga_jual', 0)})
    return jsonify({'ok': True, 'kd_barang': kd})


@app.route('/api/barang', methods=['PUT'])
def update_barang_master():
    d  = request.json or {}
    kd = (d.get('kd_barang') or '').strip()
    if not kd:
        return jsonify({'ok': False, 'msg': 'Kode barang tidak ada'}), 400
    row = query("SELECT * FROM Barang WHERE kd_barang=?", (kd,), one=True)
    if not row:
        return jsonify({'ok': False, 'msg': 'Barang tidak ditemukan'}), 404
    execute("""UPDATE Barang SET nm_barang=?, nm_generik=?, jenis=?, merk=?,
               isi=?, satuan=?, harga_jual=?, aktif=?, non_obat=? WHERE kd_barang=?""",
            ((d.get('nm_barang') or kd).strip(), d.get('nm_generik', ''),
             d.get('jenis', ''), d.get('merk', ''), float(d.get('isi', 1) or 1),
             d.get('satuan', 'Pcs') or 'Pcs', float(d.get('harga_jual', 0) or 0),
             int(d.get('aktif', 1)), int(d.get('non_obat', 0) or 0), kd))
    # Sinkron nama tampil & harga jual ke stok INDUK yang masih default
    execute("UPDATE Stok_barang SET nm_barang=? WHERE kd_barang=? AND kd_lokasi='INDUK'",
            ((d.get('nm_barang') or kd).strip(), kd))
    catat_log('barang_master_update', 'Barang', kd, f"Update master barang '{kd}'",
              data_lama={'nm_barang': row['nm_barang'], 'harga_jual': row['harga_jual'], 'aktif': row['aktif']},
              data_baru={'nm_barang': (d.get('nm_barang') or kd).strip(), 'harga_jual': d.get('harga_jual', 0),
                         'aktif': int(d.get('aktif', 1))})
    return jsonify({'ok': True})


@app.route('/api/barang', methods=['DELETE'])
def delete_barang_master():
    kd = (request.args.get('kd') or '').strip()
    if not kd:
        return jsonify({'ok': False, 'msg': 'Kode barang tidak ada'}), 400
    # Soft delete: nonaktifkan agar tidak muncul di pencarian pembelian
    execute("UPDATE Barang SET aktif=0 WHERE kd_barang=?", (kd,))
    catat_log('barang_master_nonaktifkan', 'Barang', kd, f"Nonaktifkan master barang '{kd}'")
    return jsonify({'ok': True})


# ─── SUPLIER ─────────────────────────────────────────────────
@app.route('/api/suplier')
def get_suplier():
    q = request.args.get('q', '')
    rows = query("SELECT * FROM Suplier WHERE nm_suplier LIKE ? OR kd_suplier LIKE ? ORDER BY nm_suplier",
                 (f'%{q}%', f'%{q}%'))
    return jsonify(rows_to_list(rows))


@app.route('/api/suplier', methods=['POST'])
def create_suplier():
    d  = request.json
    kd = d.get('kd_suplier', '').strip()
    if not kd:
        last = query("SELECT kd_suplier FROM Suplier ORDER BY kd_suplier DESC LIMIT 1", one=True)
        num  = (int(last['kd_suplier'].replace('SUP', '')) + 1) if last else 1
        kd   = f"SUP{num:03d}"
    execute("INSERT OR IGNORE INTO Suplier (kd_suplier,nm_suplier,alamat,telp,kontak) VALUES (?,?,?,?,?)",
            (kd, d.get('nm_suplier', ''), d.get('alamat', ''), d.get('telp', ''), d.get('kontak', '')))
    return jsonify({'ok': True, 'kd_suplier': kd})


@app.route('/api/suplier/<kd>', methods=['PUT'])
def update_suplier(kd):
    d = request.json
    execute("UPDATE Suplier SET nm_suplier=?,alamat=?,telp=?,kontak=? WHERE kd_suplier=?",
            (d.get('nm_suplier', ''), d.get('alamat', ''), d.get('telp', ''), d.get('kontak', ''), kd))
    return jsonify({'ok': True})


@app.route('/api/suplier/<kd>', methods=['DELETE'])
def delete_suplier(kd):
    execute("DELETE FROM Suplier WHERE kd_suplier=?", (kd,))
    return jsonify({'ok': True})

# ─── LAPORAN ─────────────────────────────────────────────────
def _hitung_produk_terjual(tgl1, tgl2, jenis, q):
    """Inti query Laporan Produk Terjual — dipakai bareng oleh endpoint JSON &
    endpoint export Excel, supaya filter & hasilnya selalu identik."""
    jenis_filter = ""
    if jenis in ('barang', 'jasa', 'racikan'):
        jenis_filter = f"AND u.jenis = '{jenis}'"

    # Pencarian kombinasi kode & nama (spt POS) — lihat _kata_cocok_kolom()
    q_where, q_params = _kata_cocok_kolom('u.prod_key', 'u.nama', q.split())
    q_filter = f"AND {q_where}"

    # Sumber: barang & jasa langsung + BAHAN yang dipakai dalam racikan (DetailResepRacikan).
    # Nilai bahan = sub_total bahan; total seluruh bahan = harga racikan → omzet tetap konsisten.
    # Racikan yang sudah diretur (ResepRacikan.status='returned') otomatis terkecuali.
    rows = query(f"""
        SELECT u.prod_key, u.nama, u.jenis,
               SUM(u.qty)                  AS total_qty,
               SUM(u.omzet)                AS total_omzet,
               COUNT(DISTINCT u.no_faktur) AS total_trx,
               ROUND(AVG(u.harga), 0)      AS avg_harga,
               MAX(u.tanggal)              AS tgl_terakhir
        FROM (
            SELECT rp.kd_barang AS prod_key, rp.nm_barang AS nama, rp.jenis AS jenis,
                   rp.qty AS qty, rp.sub_total AS omzet, rp.harga AS harga,
                   p.tanggal AS tanggal, rp.no_faktur AS no_faktur
            FROM R_Penjualan rp
            JOIN Penjualan p ON rp.no_faktur = p.no_faktur
            WHERE p.status='done' AND rp.jenis IN ('barang','jasa')
              AND p.tanggal BETWEEN ? AND ?

            UNION ALL

            SELECT drr.kd_barang AS prod_key, drr.nm_bahan AS nama, 'racikan' AS jenis,
                   drr.qty AS qty, drr.sub_total AS omzet, drr.harga_satuan AS harga,
                   p.tanggal AS tanggal, rr.no_faktur AS no_faktur
            FROM DetailResepRacikan drr
            JOIN ResepRacikan rr ON drr.id_resep = rr.id
            JOIN Penjualan p ON rr.no_faktur = p.no_faktur
            WHERE p.status='done' AND rr.status='confirmed'
              AND p.tanggal BETWEEN ? AND ?

            UNION ALL

            -- Net retur barang & jasa langsung (qty negatif, sub_total retur sudah negatif).
            -- Racikan tak perlu di sini: racikan diretur → ResepRacikan.status='returned' (sudah keluar di atas).
            SELECT rp.kd_barang AS prod_key, rp.nm_barang AS nama,
                   CASE WHEN rp.jenis='retur_jasa' THEN 'jasa' ELSE 'barang' END AS jenis,
                   -rp.qty AS qty, rp.sub_total AS omzet, rp.harga AS harga,
                   p.tanggal AS tanggal, rp.no_faktur AS no_faktur
            FROM R_Penjualan rp
            JOIN Penjualan p ON rp.no_faktur = p.no_faktur
            WHERE p.status='done' AND rp.jenis IN ('retur_barang','retur_jasa')
              AND p.tanggal BETWEEN ? AND ?
        ) u
        WHERE 1=1 {jenis_filter} {q_filter}
        GROUP BY u.prod_key, u.nama, u.jenis
        HAVING total_qty <> 0 OR total_omzet <> 0
        ORDER BY total_omzet DESC
    """, [tgl1, tgl2, tgl1, tgl2, tgl1, tgl2] + q_params)
    return rows_to_list(rows)

@app.route('/api/laporan/produk-terjual')
def laporan_produk_terjual():
    """Laporan produk terjual — dikelompokkan per produk/perawatan/racikan."""
    tgl1  = request.args.get('dari',   str(date.today()))
    tgl2  = request.args.get('sampai', str(date.today()))
    jenis = request.args.get('jenis',  'semua')   # semua | barang | jasa | racikan
    q     = request.args.get('q',      '').strip()

    result = _hitung_produk_terjual(tgl1, tgl2, jenis, q)
    total_omzet = sum(r['total_omzet'] for r in result)
    total_qty   = sum(r['total_qty']   for r in result)
    max_omzet   = max((r['total_omzet'] for r in result), default=1)

    # Tambah pct bar (0–100) untuk visualisasi relatif
    for r in result:
        r['pct'] = round(r['total_omzet'] / max_omzet * 100) if max_omzet else 0

    return jsonify({
        'rows':        result,
        'total_omzet': total_omzet,
        'total_qty':   total_qty,
        'total_sku':   len(result),
    })

@app.route('/api/laporan/produk-terjual/export')
def laporan_produk_terjual_export():
    """Export Laporan Produk Terjual ke .xlsx — filter sama persis dgn endpoint JSON."""
    tgl1  = request.args.get('dari',   str(date.today()))
    tgl2  = request.args.get('sampai', str(date.today()))
    jenis = request.args.get('jenis',  'semua')
    q     = request.args.get('q',      '').strip()

    result = _hitung_produk_terjual(tgl1, tgl2, jenis, q)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Produk Terjual'
    kolom = ['Kode', 'Nama', 'Jenis', 'Qty Terjual', 'Total Omzet', 'Jumlah Transaksi',
             'Harga Rata-rata', 'Tanggal Terakhir']
    ws.append(kolom)
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='1A1A2E')
    for c in range(1, len(kolom) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    ws.freeze_panes = 'A2'

    _JENIS_LABEL_EXPORT = {'barang': 'Barang', 'jasa': 'Jasa', 'racikan': 'Racikan (bahan)'}
    for r in result:
        ws.append([
            r['prod_key'], r['nama'], _JENIS_LABEL_EXPORT.get(r['jenis'], r['jenis']),
            r['total_qty'], r['total_omzet'], r['total_trx'], r['avg_harga'], r['tgl_terakhir'],
        ])
        row_i = ws.max_row
        for c in (5, 7):
            ws.cell(row=row_i, column=c).number_format = '#,##0'

    lebar = [22, 30, 14, 12, 16, 16, 16, 15]
    for idx, w_ in enumerate(lebar, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w_

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    fname = f"Laporan_Produk_Terjual_{tgl1}_sd_{tgl2}.xlsx"
    from flask import Response
    return Response(out.getvalue(),
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     headers={'Content-Disposition': f'attachment; filename="{fname}"'})

@app.route('/api/laporan/produk-terjual/detail')
def laporan_produk_detail():
    """Detail transaksi per produk — untuk accordion drill-down."""
    tgl1     = request.args.get('dari',   str(date.today()))
    tgl2     = request.args.get('sampai', str(date.today()))
    prod_key = request.args.get('prod_key', '')
    jenis    = request.args.get('jenis', 'barang')

    if jenis == 'racikan':
        # prod_key = kd_barang bahan; tampilkan tiap pemakaian bahan ini di racikan
        rows = query("""
            SELECT rr.no_faktur, p.tanggal, ps.kd_pasien, ps.nm_pasien, drr.qty,
                   drr.harga_satuan AS harga, drr.sub_total, '0' AS disc,
                   rr.nama_produk
            FROM DetailResepRacikan drr
            JOIN ResepRacikan rr ON drr.id_resep = rr.id
            JOIN Penjualan p    ON rr.no_faktur  = p.no_faktur
            LEFT JOIN Pasien ps ON p.kd_pasien   = ps.kd_pasien
            WHERE p.status='done' AND rr.status='confirmed'
              AND p.tanggal BETWEEN ? AND ? AND drr.kd_barang=?
            ORDER BY p.tanggal DESC, drr.id DESC
        """, (tgl1, tgl2, prod_key))
    else:
        # Penjualan langsung + baris retur (qty & sub_total negatif) agar detail ikut net
        rows = query("""
            SELECT rp.no_faktur, p.tanggal, ps.kd_pasien, ps.nm_pasien, rp.qty,
                   rp.harga, rp.sub_total, rp.disc
            FROM R_Penjualan rp
            JOIN Penjualan p  ON rp.no_faktur  = p.no_faktur
            LEFT JOIN Pasien ps ON p.kd_pasien = ps.kd_pasien
            WHERE p.status='done' AND p.tanggal BETWEEN ? AND ?
              AND rp.kd_barang=? AND rp.jenis=?

            UNION ALL

            SELECT rp.no_faktur, p.tanggal, ps.kd_pasien, ps.nm_pasien, -rp.qty AS qty,
                   rp.harga, rp.sub_total, rp.disc
            FROM R_Penjualan rp
            JOIN Penjualan p  ON rp.no_faktur  = p.no_faktur
            LEFT JOIN Pasien ps ON p.kd_pasien = ps.kd_pasien
            WHERE p.status='done' AND p.tanggal BETWEEN ? AND ?
              AND rp.kd_barang=? AND rp.jenis = ?
            ORDER BY tanggal DESC
        """, (tgl1, tgl2, prod_key, jenis, tgl1, tgl2, prod_key, 'retur_' + jenis))

    return jsonify(rows_to_list(rows))

@app.route('/api/laporan/penjualan')
def laporan_penjualan():
    tgl1 = request.args.get('dari', str(date.today()))
    tgl2 = request.args.get('sampai', str(date.today()))
    faktur = request.args.get('faktur', '').strip()
    dokter = request.args.get('dokter', '').strip()
    _sel = """
        SELECT p.no_faktur, p.tanggal, p.kd_pasien, ps.nm_pasien,
               p.dokter, p.total_harga AS total_harga_stored, p.metode_bayar, p.status,
               (SELECT SUM(rp.sub_total) FROM R_Penjualan rp
                WHERE rp.no_faktur=p.no_faktur AND rp.jenis!='bahan_racikan') AS total_item
        FROM Penjualan p LEFT JOIN Pasien ps ON p.kd_pasien=ps.kd_pasien
    """
    # Pencarian dokter BERKAITAN dgn filter yg sedang aktif (bukan pengganti spt
    # pencarian no faktur yg mengabaikan rentang tanggal) — jadi selalu ditambahkan
    # sbg klausa AND tambahan, baik saat mode faktur maupun mode rentang tanggal.
    dokter_filter = " AND p.dokter LIKE ?" if dokter else ""
    dokter_param = [f'%{dokter}%'] if dokter else []
    if faktur:
        # Cari lintas semua periode berdasarkan no faktur — rentang tanggal diabaikan.
        raw_rows = query(_sel + f" WHERE p.no_faktur LIKE ? {dokter_filter} ORDER BY p.tanggal DESC LIMIT 200",
                         [f'%{faktur}%'] + dokter_param)
    else:
        raw_rows = query(_sel + f" WHERE p.tanggal BETWEEN ? AND ? {dokter_filter} ORDER BY p.tanggal DESC",
                         [tgl1, tgl2] + dokter_param)
    # Sebagian faktur lama (migrasi) punya Penjualan.total_harga TIDAK SINKRON dgn
    # baris R_Penjualan aslinya (header 0 padahal item riil ada) — selalu percaya
    # SUM baris asli kalau tersedia, cuma fallback ke kolom lama kalau memang tidak
    # ada baris sama sekali (mis. transaksi yang benar-benar kosong).
    rows = rows_to_list(raw_rows)
    for r in rows:
        r['total_harga'] = r['total_item'] if r['total_item'] is not None else r['total_harga_stored']
        del r['total_item']; del r['total_harga_stored']
    # metode_bayar='Internal' = transaksi gratis (internal/testing) — item & nilainya
    # tetap ada di baris tabel (transparan, tetap bisa dicari/diaudit), tapi SENGAJA
    # dikecualikan dari total omzet & breakdown per metode/bank (bukan transaksi riil).
    total = sum(r['total_harga'] for r in rows if r['status']=='done' and (r['metode_bayar'] or '').strip() != 'Internal')

    # Breakdown per metode & per bank — split-aware (2026-10-25): pecah per baris
    # Penjualan_Pembayaran (bukan per faktur), supaya faktur split (sebagian Tunai +
    # sebagian Transfer) kebagi benar ke kedua metode sesuai nominal masing-masing,
    # bukan ditempel seluruhnya ke 1 metode. Faktur lama tanpa baris di tabel baru
    # fallback ke 1 baris sintetis dari metode_bayar+total_harga (persis perilaku lama).
    per_metode = {}
    per_bank = {}
    done_fakturs = [r['no_faktur'] for r in rows if r['status'] == 'done']
    lines_map = _bayar_lines_map(done_fakturs)
    # Jumlah POINT yg benar² terpotong (bukan cuma nilai Rupiah-nya) — utk monitoring/audit
    # supaya breakdown 'Point Member' bisa disilang-cek, lihat _point_dipakai_map().
    pm_fakturs = [r['no_faktur'] for r in rows if r['status'] == 'done'
                  and (r['metode_bayar'] or '').strip() == 'Point Member']
    point_map = _point_dipakai_map(pm_fakturs)
    for r in rows:
        r['point_dipakai'] = point_map.get(r['no_faktur'], 0) if r['no_faktur'] in pm_fakturs else 0
        if r['status'] != 'done':
            continue
        for ln in _bayar_lines_for_row(r['no_faktur'], r['metode_bayar'], r['total_harga'], lines_map):
            metode = (ln['metode'] or 'Tunai').strip() or 'Tunai'
            bank = (ln['bank'] or '').strip()
            # 'Internal' tetap DIKELOMPOKKAN sbg metode sendiri (biar kelihatan berapa kali
            # dipakai) — tapi kontribusinya ke 'total' SENGAJA 0, bukan nominal asli,
            # supaya kartu breakdown-nya jujur menunjukkan "gratis" (bukan omzet riil).
            nilai = 0.0 if metode == 'Internal' else ln['nominal']
            m = per_metode.setdefault(metode, {'metode': metode, 'total': 0.0, 'jumlah': 0, 'point': 0})
            m['total'] += nilai; m['jumlah'] += 1
            if metode == 'Point Member':
                m['point'] += point_map.get(r['no_faktur'], 0)
            if bank:
                key = (metode, bank)
                b = per_bank.setdefault(key, {'metode': metode, 'bank': bank, 'total': 0.0, 'jumlah': 0})
                b['total'] += nilai; b['jumlah'] += 1

    per_metode_list = sorted(per_metode.values(), key=lambda x: -x['total'])
    per_bank_list = sorted(per_bank.values(), key=lambda x: -x['total'])

    return jsonify({
        'rows': rows, 'total': total,
        'per_metode': per_metode_list, 'per_bank': per_bank_list,
    })

@app.route('/api/laporan/penjualan/export')
def laporan_penjualan_export():
    """Export Laporan Penjualan ke file .xlsx asli (openpyxl) — filter SAMA PERSIS dgn
    /api/laporan/penjualan (rentang tanggal ATAU pencarian no faktur). 1 baris per
    item (barang/jasa/racikan/retur), termasuk rincian bahan racikan (kalau ada,
    ditandai kolom 'Racikan Induk'), dan kolom Metode Bayar per transaksi."""
    tgl1 = request.args.get('dari', str(date.today()))
    tgl2 = request.args.get('sampai', str(date.today()))
    faktur = request.args.get('faktur', '').strip()
    dokter = request.args.get('dokter', '').strip()

    # 1 query JOIN (bukan 1 query per faktur / N+1) — rentang bisa mencakup ratusan
    # ribu transaksi (data hasil migrasi), jadi wajib hindari round-trip per baris.
    _sel = """
        SELECT p.no_faktur, p.tanggal, ps.kd_pasien, ps.nm_pasien, p.dokter, p.metode_bayar,
               p.status, p.total_harga,
               rp.kd_barang, rp.nm_barang, rp.harga, rp.qty, rp.disc, rp.sub_total,
               rp.jenis, rp.racikan_ref
        FROM Penjualan p
        LEFT JOIN Pasien ps ON p.kd_pasien=ps.kd_pasien
        LEFT JOIN R_Penjualan rp ON rp.no_faktur=p.no_faktur
    """
    # Sama spt endpoint JSON — pencarian dokter BERKAITAN dgn filter aktif (AND
    # tambahan), bukan pengganti spt pencarian no faktur yg mengabaikan tanggal.
    dokter_filter = " AND p.dokter LIKE ?" if dokter else ""
    dokter_param = [f'%{dokter}%'] if dokter else []
    if faktur:
        # MySQL tidak izinkan LIMIT langsung di dalam subquery IN — bungkus 1 lapis
        # derived table lagi supaya berlaku di kedua backend.
        rows = query(_sel + f""" WHERE p.no_faktur IN (
                        SELECT no_faktur FROM (
                          SELECT no_faktur FROM Penjualan WHERE no_faktur LIKE ?
                          ORDER BY tanggal DESC LIMIT 200) t)
                      {dokter_filter}
                      ORDER BY p.tanggal DESC, p.no_faktur, rp.id""",
                     [f'%{faktur}%'] + dokter_param)
    else:
        rows = query(_sel + f" WHERE p.tanggal BETWEEN ? AND ? {dokter_filter} ORDER BY p.tanggal DESC, p.no_faktur, rp.id",
                     [tgl1, tgl2] + dokter_param)
    rows = rows_to_list(rows)

    _JENIS_LABEL = {'barang': 'Barang', 'jasa': 'Jasa', 'racikan': 'Racikan',
                    'bahan_racikan': 'Bahan Racikan'}

    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    # write_only=True — rentang bisa mencakup ratusan ribu baris (data hasil migrasi);
    # Workbook normal (in-memory sel per-sel) jauh lebih lambat & lebih berat memori
    # utk ukuran segitu, jadi wajib pakai mode tulis-langsung openpyxl.
    # Point dipakai (2026-10-25) — utk monitoring/audit transaksi yg dibayar via Point
    # Member, lihat _point_dipakai_map(). Dihitung sekali di awal utk SEMUA faktur unik
    # bermetode Point Member di hasil export ini (bukan per-item/N+1).
    pm_fakturs_export = list({r['no_faktur'] for r in rows if (r['metode_bayar'] or '').strip() == 'Point Member'})
    point_map_export = _point_dipakai_map(pm_fakturs_export)

    wb = Workbook(write_only=True)
    ws = wb.create_sheet('Laporan Penjualan')
    kolom = ['No Faktur', 'Tanggal', 'Kode Pasien', 'Pasien', 'Dokter', 'Metode Bayar', 'Point Dipakai', 'Status',
             'Jenis Item', 'Kode Barang', 'Nama Item', 'Racikan Induk', 'Harga', 'Qty', 'Disc (%)',
             'Subtotal', 'Total Transaksi']
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='1A1A2E')
    header_align = Alignment(horizontal='center')
    header_cells = []
    for name in kolom:
        c = WriteOnlyCell(ws, value=name)
        c.font = header_font; c.fill = header_fill; c.alignment = header_align
        header_cells.append(c)
    ws.append(header_cells)
    ws.freeze_panes = 'A2'

    _NUM_IDX = (12, 15, 16)   # index 0-based: Harga, Subtotal, Total Transaksi -> format ribuan
    _POINT_IDX = 6            # index 0-based: Point Dipakai -> merah bold (monitoring)
    _POINT_FONT = Font(bold=True, color='FFDC2626')
    def _row_cells(vals):
        out_cells = []
        for idx, v in enumerate(vals):
            if idx in _NUM_IDX:
                c = WriteOnlyCell(ws, value=v)
                c.number_format = '#,##0'
                out_cells.append(c)
            elif idx == _POINT_IDX and v:
                c = WriteOnlyCell(ws, value=v)
                c.font = _POINT_FONT
                out_cells.append(c)
            else:
                out_cells.append(v)
        return out_cells

    # Kelompokkan baris flat hasil JOIN per no_faktur (row sudah terurut per transaksi).
    i = 0
    n = len(rows)
    while i < n:
        faktur_cur = rows[i]['no_faktur']
        items = []
        while i < n and rows[i]['no_faktur'] == faktur_cur:
            items.append(rows[i]); i += 1
        h = items[0]
        status_label = 'Selesai' if h['status'] == 'done' else 'Proses'
        is_point_member = (h['metode_bayar'] or '').strip() == 'Point Member'
        base = [h['no_faktur'], h['tanggal'], h['kd_pasien'] or '', h['nm_pasien'] or '', h['dokter'] or '',
                h['metode_bayar'] or 'Tunai',
                point_map_export.get(h['no_faktur'], 0) if is_point_member else '',
                status_label]
        if items[0]['kd_barang'] is None:
            # Transaksi tanpa item sama sekali (LEFT JOIN tidak match)
            ws.append(_row_cells(base + ['', '', '', '', '', '', '', '', h['total_harga']]))
            continue
        # Nama racikan induk per id_resep (racikan_ref) — dipakai utk label baris bahan_racikan.
        racikan_nama = {it['racikan_ref']: it['nm_barang'] for it in items
                         if it['jenis'] == 'racikan' and it.get('racikan_ref')}
        for it in items:
            jenis = it['jenis'] or 'barang'
            if jenis.startswith('retur'):
                jenis_label = f"Retur ({jenis.replace('retur_', '')})"
            else:
                jenis_label = _JENIS_LABEL.get(jenis, jenis)
            racikan_induk = racikan_nama.get(it['racikan_ref'], '') if jenis == 'bahan_racikan' else ''
            ws.append(_row_cells(base + [
                jenis_label, it['kd_barang'] or '', it['nm_barang'], racikan_induk,
                it['harga'], it['qty'], it['disc'], it['sub_total'], h['total_harga'],
            ]))

    lebar = [16, 11, 14, 20, 20, 16, 13, 10, 14, 20, 24, 20, 12, 7, 9, 13, 15]
    for idx, w_ in enumerate(lebar, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w_

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    fname = (f"Laporan_Penjualan_cari_{faktur}.xlsx" if faktur
             else f"Laporan_Penjualan_{tgl1}_sd_{tgl2}.xlsx")
    from flask import Response
    return Response(out.getvalue(),
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     headers={'Content-Disposition': f'attachment; filename="{fname}"'})

_KAT_LABEL_KOMPREHENSIF = {'barang': 'Barang', 'jasa': 'Jasa / Treatment', 'racikan': 'Racikan (bahan)'}

def _laporan_komprehensif_data(tgl1, tgl2):
    """Inti hitung Laporan Audit Lengkap — dipakai BERSAMA oleh endpoint JSON (tampil
    gamblang di layar form, tanpa perlu unduh Excel) & endpoint export Excel, supaya
    angka yg dilihat pemeriksa di layar PERSIS sama dgn file unduhan (tidak ada 2 versi
    kebenaran). Tidak termasuk 'Detail Transaksi' (ratusan/ribuan baris per-item) —
    itu Excel-only, lihat laporan_komprehensif_export()."""
    raw = query("""
        SELECT p.no_faktur, p.status, p.metode_bayar, p.kd_pasien, p.total_harga AS total_harga_stored,
               (SELECT SUM(rp.sub_total) FROM R_Penjualan rp
                WHERE rp.no_faktur=p.no_faktur AND rp.jenis!='bahan_racikan') AS total_item
        FROM Penjualan p WHERE p.tanggal BETWEEN ? AND ?
    """, (tgl1, tgl2))
    rows = rows_to_list(raw)
    for r in rows:
        r['total_harga'] = r['total_item'] if r['total_item'] is not None else r['total_harga_stored']
    done_rows = [r for r in rows if r['status'] == 'done']
    riil_rows = [r for r in done_rows if (r['metode_bayar'] or '').strip() != 'Internal']
    total_omzet  = sum(r['total_harga'] for r in riil_rows)
    total_trx    = len(done_rows)
    total_pasien = len(set(r['kd_pasien'] for r in done_rows if r['kd_pasien']))
    avg_trx      = round(total_omzet / len(riil_rows)) if riil_rows else 0

    total_komisi_row = query("""SELECT COALESCE(SUM(komisi),0) t FROM Point_karyawan
                                WHERE tanggal BETWEEN ? AND ? AND COALESCE(status,'released')='released'""",
                             (tgl1, tgl2), one=True)
    total_komisi = total_komisi_row['t'] or 0

    # Split-aware (2026-10-25) — sama pola dgn laporan_penjualan(): pecah per baris
    # Penjualan_Pembayaran, fallback 1 baris sintetis utk faktur lama pra-fitur.
    per_metode, per_bank = {}, {}
    lines_map = _bayar_lines_map([r['no_faktur'] for r in done_rows])
    pm_fakturs = [r['no_faktur'] for r in done_rows if (r['metode_bayar'] or '').strip() == 'Point Member']
    point_map = _point_dipakai_map(pm_fakturs)
    for r in done_rows:
        for ln in _bayar_lines_for_row(r['no_faktur'], r['metode_bayar'], r['total_harga'], lines_map):
            metode = (ln['metode'] or 'Tunai').strip() or 'Tunai'
            bank = (ln['bank'] or '').strip()
            nilai = 0.0 if metode == 'Internal' else ln['nominal']
            m = per_metode.setdefault(metode, {'metode': metode, 'total': 0.0, 'jumlah': 0, 'point': 0})
            m['total'] += nilai; m['jumlah'] += 1
            if metode == 'Point Member':
                m['point'] += point_map.get(r['no_faktur'], 0)
            if bank:
                b = per_bank.setdefault((metode, bank), {'metode': metode, 'bank': bank, 'total': 0.0, 'jumlah': 0})
                b['total'] += nilai; b['jumlah'] += 1
    per_metode_list = sorted(per_metode.values(), key=lambda x: -x['total'])
    per_bank_list   = sorted(per_bank.values(), key=lambda x: -x['total'])

    produk_rows = sorted(_hitung_produk_terjual(tgl1, tgl2, 'semua', ''), key=lambda x: -x['total_omzet'])
    for p in produk_rows:
        p['label_jenis'] = _KAT_LABEL_KOMPREHENSIF.get(p['jenis'], p['jenis'])
    kat_agg = {}
    for p in produk_rows:
        k = kat_agg.setdefault(p['jenis'], {'jenis': p['jenis'], 'sku': 0, 'qty': 0, 'omzet': 0})
        k['sku'] += 1; k['qty'] += p['total_qty']; k['omzet'] += p['total_omzet']
    for k in kat_agg.values():
        k['label'] = _KAT_LABEL_KOMPREHENSIF.get(k['jenis'], k['jenis'])
    kategori_list = sorted(kat_agg.values(), key=lambda x: -x['omzet'])

    dokter_rows = _monitor_grup_hitung(_monitor_grup_mentah(tgl1, tgl2, 'dokter'), 'dokter', exclude_internal=True)

    komisi_rows = rows_to_list(query("""
        SELECT pk.kd_karyawan, k.nm_karyawan, k.nickname, k.pekerjaan,
               SUM(pk.komisi) as total_komisi, COUNT(*) as jumlah_tindakan
        FROM Point_karyawan pk LEFT JOIN Karyawan k ON pk.kd_karyawan=k.kd_karyawan
        WHERE pk.tanggal BETWEEN ? AND ? AND COALESCE(pk.status,'released')='released'
        GROUP BY pk.kd_karyawan ORDER BY total_komisi DESC
    """, (tgl1, tgl2)))

    return {
        'ringkasan': {
            'total_omzet': total_omzet, 'total_trx': total_trx, 'total_pasien': total_pasien,
            'avg_trx': avg_trx, 'total_komisi': total_komisi,
        },
        'per_metode': per_metode_list, 'per_bank': per_bank_list, 'per_kategori': kategori_list,
        'dokter': dokter_rows, 'komisi': komisi_rows, 'produk': produk_rows,
    }

@app.route('/api/laporan/komprehensif')
def laporan_komprehensif():
    """Versi JSON Laporan Audit Lengkap — utk ditampilkan GAMBLANG di layar form-nya
    sendiri, bukan cuma lewat unduh Excel. Pakai _laporan_komprehensif_data() yg sama
    persis dgn yg dipakai export Excel, jadi angkanya selalu identik di 2 tempat."""
    tgl1 = request.args.get('dari', str(date.today()))
    tgl2 = request.args.get('sampai', str(date.today()))
    return jsonify(_laporan_komprehensif_data(tgl1, tgl2))

@app.route('/api/laporan/komprehensif/export')
def laporan_komprehensif_export():
    """Laporan AUDIT lengkap — 1 file Excel multi-sheet gabungan SEMUA aspek yg biasa
    dicek pemeriksa/akuntan sekaligus (bukan harus buka banyak laporan terpisah):
    Ringkasan, Per Dokter, Komisi Karyawan, Produk & Jasa Terjual, Detail Transaksi.
    Sheet 1-4 pakai data dari _laporan_komprehensif_data() (SAMA PERSIS dgn versi JSON
    yg tampil di layar) — Detail Transaksi (per-item) dihitung khusus di sini."""
    tgl1 = request.args.get('dari', str(date.today()))
    tgl2 = request.args.get('sampai', str(date.today()))
    data = _laporan_komprehensif_data(tgl1, tgl2)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from flask import Response

    HEADER_FONT  = Font(bold=True, color='FFFFFF')
    HEADER_FILL  = PatternFill('solid', fgColor='1A1A2E')
    HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
    TITLE_FONT   = Font(bold=True, size=14, color='1A1A2E')
    SUB_FONT     = Font(italic=True, size=10, color='666666')
    SECTION_FONT = Font(bold=True, size=12, color='946C0E')
    TOTAL_FONT   = Font(bold=True)
    POINT_FONT   = Font(bold=True, color='FFDC2626')   # merah bold — Point Dipakai (monitoring)

    def _header_row(ws, row_idx, labels, widths=None):
        for ci, label in enumerate(labels, start=1):
            c = ws.cell(row=row_idx, column=ci, value=label)
            c.font = HEADER_FONT; c.fill = HEADER_FILL; c.alignment = HEADER_ALIGN
        if widths:
            for ci, w in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(ci)].width = w

    wb = Workbook()

    # ── SHEET 1 — RINGKASAN ─────────────────────────────────────────
    ws0 = wb.active
    ws0.title = 'Ringkasan'
    ws0['A1'] = 'LAPORAN AUDIT LENGKAP — RACHEL CLINIC'
    ws0['A1'].font = TITLE_FONT
    ws0['A2'] = f'Periode: {tgl1} s/d {tgl2}  ·  Dicetak: {datetime.now().strftime("%d/%m/%Y %H:%M")}'
    ws0['A2'].font = SUB_FONT

    ring = data['ringkasan']
    rr = 4
    for label, val in [
        ('Total Omzet (riil, tidak termasuk transaksi Internal/gratis)', ring['total_omzet']),
        ('Total Transaksi (status selesai)', ring['total_trx']),
        ('Total Pasien Unik', ring['total_pasien']),
        ('Rata-rata Omzet per Transaksi', ring['avg_trx']),
        ('Total Komisi Karyawan Dibayarkan (semua jabatan)', ring['total_komisi']),
    ]:
        ws0.cell(row=rr, column=1, value=label).font = Font(bold=True)
        c = ws0.cell(row=rr, column=2, value=val)
        if isinstance(val, (int, float)):
            c.number_format = '#,##0'
        rr += 1

    rr += 1
    ws0.cell(row=rr, column=1, value='BREAKDOWN PER METODE PEMBAYARAN').font = SECTION_FONT
    rr += 1
    _header_row(ws0, rr, ['Metode', 'Bank', 'Jumlah Transaksi', 'Total Nilai', 'Point Dipakai'])
    rr += 1
    for v in data['per_metode']:
        ws0.cell(row=rr, column=1, value=v['metode'])
        ws0.cell(row=rr, column=3, value=v['jumlah'])
        c = ws0.cell(row=rr, column=4, value=v['total']); c.number_format = '#,##0'
        if v['metode'] == 'Point Member':
            pc = ws0.cell(row=rr, column=5, value=v.get('point', 0)); pc.font = POINT_FONT
        rr += 1
        for bv in data['per_bank']:
            if bv['metode'] != v['metode']:
                continue
            ws0.cell(row=rr, column=2, value='   ' + bv['bank'])
            ws0.cell(row=rr, column=3, value=bv['jumlah'])
            c = ws0.cell(row=rr, column=4, value=bv['total']); c.number_format = '#,##0'
            rr += 1

    rr += 1
    ws0.cell(row=rr, column=1, value='BREAKDOWN PER KATEGORI').font = SECTION_FONT
    rr += 1
    _header_row(ws0, rr, ['Kategori', 'Jumlah SKU', 'Total Qty', 'Total Omzet'])
    rr += 1
    for v in data['per_kategori']:
        ws0.cell(row=rr, column=1, value=v['label'])
        ws0.cell(row=rr, column=2, value=v['sku'])
        ws0.cell(row=rr, column=3, value=v['qty'])
        c = ws0.cell(row=rr, column=4, value=v['omzet']); c.number_format = '#,##0'
        rr += 1
    for ci, w in enumerate([48, 16, 18, 18, 16], start=1):
        ws0.column_dimensions[get_column_letter(ci)].width = w

    # ── SHEET 2 — PER DOKTER ─────────────────────────────────────────
    ws1 = wb.create_sheet('Per Dokter')
    dokter_rows = data['dokter']
    _header_row(ws1, 1, ['Dokter', 'Jumlah Transaksi', 'Jumlah Pasien', 'Total Omzet', 'Tunai', 'Non-Tunai'],
                [30, 16, 14, 16, 16, 16])
    rr = 2
    for d in dokter_rows:
        ws1.cell(row=rr, column=1, value=d['dokter'])
        ws1.cell(row=rr, column=2, value=d['total_trx'])
        ws1.cell(row=rr, column=3, value=d['total_pasien'])
        for ci, key in [(4, 'total_nilai'), (5, 'tunai'), (6, 'non_tunai')]:
            c = ws1.cell(row=rr, column=ci, value=d[key]); c.number_format = '#,##0'
        rr += 1
    ws1.cell(row=rr, column=1, value='TOTAL').font = TOTAL_FONT
    ws1.cell(row=rr, column=2, value=sum(d['total_trx'] for d in dokter_rows)).font = TOTAL_FONT
    for ci, key in [(4, 'total_nilai'), (5, 'tunai'), (6, 'non_tunai')]:
        c = ws1.cell(row=rr, column=ci, value=sum(d[key] for d in dokter_rows))
        c.number_format = '#,##0'; c.font = TOTAL_FONT
    ws1.freeze_panes = 'A2'

    # ── SHEET 3 — KOMISI KARYAWAN ─────────────────────────────────────
    ws2 = wb.create_sheet('Komisi Karyawan')
    komisi_rows = data['komisi']
    _header_row(ws2, 1, ['Kode', 'Nama', 'Nickname', 'Jabatan', 'Jumlah Tindakan', 'Total Komisi'],
                [10, 28, 14, 14, 16, 16])
    rr = 2
    for k in komisi_rows:
        ws2.cell(row=rr, column=1, value=k['kd_karyawan'])
        ws2.cell(row=rr, column=2, value=k['nm_karyawan'] or k['kd_karyawan'])
        ws2.cell(row=rr, column=3, value=k['nickname'] or '')
        ws2.cell(row=rr, column=4, value=k['pekerjaan'] or '')
        ws2.cell(row=rr, column=5, value=k['jumlah_tindakan'])
        c = ws2.cell(row=rr, column=6, value=k['total_komisi']); c.number_format = '#,##0'
        rr += 1
    ws2.cell(row=rr, column=1, value='TOTAL').font = TOTAL_FONT
    c = ws2.cell(row=rr, column=6, value=sum(k['total_komisi'] for k in komisi_rows))
    c.number_format = '#,##0'; c.font = TOTAL_FONT
    ws2.freeze_panes = 'A2'

    # ── SHEET 4 — PRODUK & JASA TERJUAL ──────────────────────────────
    ws3 = wb.create_sheet('Produk & Jasa Terjual')
    _header_row(ws3, 1, ['Kode', 'Nama', 'Jenis', 'Qty Terjual', 'Total Omzet', 'Jumlah Transaksi',
                         'Harga Rata-rata', 'Tanggal Terakhir'], [22, 30, 16, 12, 16, 14, 14, 14])
    rr = 2
    for p in data['produk']:
        ws3.cell(row=rr, column=1, value=p['prod_key'])
        ws3.cell(row=rr, column=2, value=p['nama'])
        ws3.cell(row=rr, column=3, value=p['label_jenis'])
        ws3.cell(row=rr, column=4, value=p['total_qty'])
        c = ws3.cell(row=rr, column=5, value=p['total_omzet']); c.number_format = '#,##0'
        ws3.cell(row=rr, column=6, value=p['total_trx'])
        c2 = ws3.cell(row=rr, column=7, value=p['avg_harga']); c2.number_format = '#,##0'
        ws3.cell(row=rr, column=8, value=p['tgl_terakhir'])
        rr += 1
    ws3.freeze_panes = 'A2'

    # ── SHEET 5 — DETAIL TRANSAKSI (logic PERSIS laporan_penjualan_export) ───
    ws4 = wb.create_sheet('Detail Transaksi')
    det_rows = rows_to_list(query("""
        SELECT p.no_faktur, p.tanggal, ps.kd_pasien, ps.nm_pasien, p.dokter, p.metode_bayar,
               p.status, p.total_harga,
               rp.kd_barang, rp.nm_barang, rp.harga, rp.qty, rp.disc, rp.sub_total,
               rp.jenis, rp.racikan_ref
        FROM Penjualan p
        LEFT JOIN Pasien ps ON p.kd_pasien=ps.kd_pasien
        LEFT JOIN R_Penjualan rp ON rp.no_faktur=p.no_faktur
        WHERE p.tanggal BETWEEN ? AND ?
        ORDER BY p.tanggal DESC, p.no_faktur, rp.id
    """, (tgl1, tgl2)))
    _JENIS_LABEL2 = {'barang': 'Barang', 'jasa': 'Jasa', 'racikan': 'Racikan', 'bahan_racikan': 'Bahan Racikan'}
    # Point dipakai (2026-10-25) — sama pola dgn laporan_penjualan_export(), lihat
    # _point_dipakai_map(). Dihitung sekali di awal utk semua faktur unik bermetode
    # Point Member dlm rentang tanggal sheet ini.
    pm_fakturs_det = list({r['no_faktur'] for r in det_rows if (r['metode_bayar'] or '').strip() == 'Point Member'})
    point_map_det = _point_dipakai_map(pm_fakturs_det)
    _header_row(ws4, 1, ['No Faktur', 'Tanggal', 'Kode Pasien', 'Pasien', 'Dokter', 'Metode Bayar', 'Point Dipakai', 'Status',
                         'Jenis Item', 'Kode Barang', 'Nama Item', 'Racikan Induk', 'Harga', 'Qty', 'Disc (%)',
                         'Subtotal', 'Total Transaksi'],
                [16, 11, 14, 20, 20, 16, 13, 10, 13, 20, 24, 20, 12, 7, 9, 13, 15])
    _NUM_COLS = (13, 16, 17)
    rr = 2
    i, n = 0, len(det_rows)
    while i < n:
        faktur_cur = det_rows[i]['no_faktur']
        items = []
        while i < n and det_rows[i]['no_faktur'] == faktur_cur:
            items.append(det_rows[i]); i += 1
        h = items[0]
        status_label = 'Selesai' if h['status'] == 'done' else 'Proses'
        is_point_member_det = (h['metode_bayar'] or '').strip() == 'Point Member'
        base = [h['no_faktur'], h['tanggal'], h['kd_pasien'] or '', h['nm_pasien'] or '',
                h['dokter'] or '', h['metode_bayar'] or 'Tunai',
                point_map_det.get(h['no_faktur'], 0) if is_point_member_det else '',
                status_label]
        if items[0]['kd_barang'] is None:
            for ci, v in enumerate(base + ['', '', '', '', '', '', '', '', h['total_harga']], start=1):
                c = ws4.cell(row=rr, column=ci, value=v)
                if ci in _NUM_COLS: c.number_format = '#,##0'
                if ci == 7 and v: c.font = POINT_FONT
            rr += 1
            continue
        racikan_nama = {it['racikan_ref']: it['nm_barang'] for it in items
                         if it['jenis'] == 'racikan' and it.get('racikan_ref')}
        for it in items:
            jenis = it['jenis'] or 'barang'
            jenis_label = f"Retur ({jenis.replace('retur_', '')})" if jenis.startswith('retur') \
                else _JENIS_LABEL2.get(jenis, jenis)
            racikan_induk = racikan_nama.get(it['racikan_ref'], '') if jenis == 'bahan_racikan' else ''
            vals = base + [jenis_label, it['kd_barang'] or '', it['nm_barang'], racikan_induk, it['harga'], it['qty'],
                           it['disc'], it['sub_total'], h['total_harga']]
            for ci, v in enumerate(vals, start=1):
                c = ws4.cell(row=rr, column=ci, value=v)
                if ci in _NUM_COLS: c.number_format = '#,##0'
                if ci == 7 and v: c.font = POINT_FONT
            rr += 1
    ws4.freeze_panes = 'A2'

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    fname = f"Laporan_Audit_Lengkap_{tgl1}_sd_{tgl2}.xlsx"
    return Response(out.getvalue(),
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     headers={'Content-Disposition': f'attachment; filename="{fname}"'})

@app.route('/api/riwayat-barang/<kd>')
def riwayat_barang(kd):
    rows = query("""
        SELECT r.*, s.nm_barang FROM RiwayatBarang r
        LEFT JOIN (SELECT kd_barang, nm_barang FROM Stok_barang GROUP BY kd_barang) s
            ON r.kd_barang = s.kd_barang
        WHERE r.kd_barang = ?
        ORDER BY r.exp ASC, r.tanggal DESC, r.id DESC
        LIMIT 200
    """, (kd,))
    return jsonify(rows_to_list(rows))

@app.route('/api/laporan/stok')
def laporan_stok():
    """Laporan Stok Cabang. SELECT dipersempit ke kolom yg BENERAN dipakai halamannya
    (dulu 'SELECT *' — ikut narik DiscMem/DiscLain/pemilik/stok_awal/stok_gudang/id/dll
    yg sama sekali tak dipakai frontend) — payload ~2400 baris turun signifikan, biar
    device kasir yg spek-nya lbh lemah drpd dev machine tidak berat wkt buka halaman ini
    (keluhan 2026-09-28)."""
    kd_lokasi = request.args.get('kd_lokasi', '')
    _sel = """SELECT kd_barang, nm_barang, kd_lokasi, stok_sisa, stok_min, harga_beli, harga, exp,
                     CASE WHEN stok_sisa<=stok_min THEN 'warning' ELSE 'ok' END as kondisi
              FROM Stok_barang"""
    if kd_lokasi:
        rows = query(f"{_sel} WHERE kd_lokasi=? ORDER BY stok_sisa DESC", (kd_lokasi,))
    else:
        # Tanpa filter (admin/supervisor) = SEMUA CABANG, JANGAN ikutkan INDUK
        # (ini Laporan Stok Cabang; stok induk punya laporannya sendiri)
        rows = query(f"{_sel} WHERE kd_lokasi!='INDUK' ORDER BY stok_sisa DESC")
    return jsonify(rows_to_list(rows))

@app.route('/api/laporan/komisi')
def laporan_komisi():
    tgl1 = request.args.get('dari', str(date.today()))
    tgl2 = request.args.get('sampai', str(date.today()))
    rows = query("""
        SELECT pk.kd_karyawan, k.nm_karyawan, k.nickname, k.pekerjaan,
               SUM(pk.komisi) as total_komisi, COUNT(*) as jumlah_tindakan
        FROM Point_karyawan pk LEFT JOIN Karyawan k ON pk.kd_karyawan=k.kd_karyawan
        WHERE pk.tanggal BETWEEN ? AND ? AND COALESCE(pk.status,'released')='released'
        GROUP BY pk.kd_karyawan ORDER BY total_komisi DESC
    """, (tgl1, tgl2))
    return jsonify(rows_to_list(rows))

@app.route('/api/laporan/komisi-detail/<kd_karyawan>')
def laporan_komisi_detail(kd_karyawan):
    tgl1 = request.args.get('dari', str(date.today()))
    tgl2 = request.args.get('sampai', str(date.today()))
    # Semua status ikut tampil (termasuk 'diretur'/'pending') supaya perawat lihat riwayat
    # lengkap & transparan — yg TIDAK dihitung di total cuma yg 'released' (lihat laporan_komisi()).
    # Ikut sertakan siapa lawan serah-terima (kalau treatment ini pakai bahan kunci) supaya
    # masing2 perawat (pelaksana maupun pemegang kunci) ingat pernah handshake dgn siapa.
    rows = query("""
        SELECT pk.id, pk.no_faktur, pk.tanggal, pk.kd_pasien,
               pk.nm_pelayanan, pk.komisi, COALESCE(pk.status,'released') as status,
               p.nm_pasien,
               (SELECT stk.kd_pelaksana FROM Serah_Terima_Kunci stk
                WHERE stk.no_faktur=pk.no_faktur AND stk.kd_pelayanan=pk.kd_pelayanan LIMIT 1) as st_pelaksana,
               (SELECT stk.nm_pelaksana FROM Serah_Terima_Kunci stk
                WHERE stk.no_faktur=pk.no_faktur AND stk.kd_pelayanan=pk.kd_pelayanan LIMIT 1) as st_nm_pelaksana,
               (SELECT COALESCE(NULLIF(stk.kd_k_final,''),stk.kd_k_dipilih) FROM Serah_Terima_Kunci stk
                WHERE stk.no_faktur=pk.no_faktur AND stk.kd_pelayanan=pk.kd_pelayanan LIMIT 1) as st_kd_k,
               (SELECT COALESCE(k.nickname,k.nm_karyawan) FROM Serah_Terima_Kunci stk
                LEFT JOIN Karyawan k ON k.kd_karyawan=COALESCE(NULLIF(stk.kd_k_final,''),stk.kd_k_dipilih)
                WHERE stk.no_faktur=pk.no_faktur AND stk.kd_pelayanan=pk.kd_pelayanan LIMIT 1) as st_nm_k
        FROM Point_karyawan pk
        LEFT JOIN Pasien p ON pk.kd_pasien = p.kd_pasien
        WHERE pk.kd_karyawan=? AND pk.tanggal BETWEEN ? AND ?
        ORDER BY pk.tanggal DESC, pk.id DESC
    """, (kd_karyawan, tgl1, tgl2))
    result = rows_to_list(rows)
    for r in result:
        if r.get('st_pelaksana') and kd_karyawan == r['st_pelaksana']:
            r['handshake_dengan'] = r.get('st_nm_k') or r.get('st_kd_k') or ''
            r['handshake_peran'] = 'pelaksana'
        elif r.get('st_kd_k') and kd_karyawan == r['st_kd_k']:
            r['handshake_dengan'] = r.get('st_nm_pelaksana') or r.get('st_pelaksana') or ''
            r['handshake_peran'] = 'kunci'
        else:
            r['handshake_dengan'] = ''
            r['handshake_peran'] = ''
        for k in ('st_pelaksana', 'st_nm_pelaksana', 'st_kd_k', 'st_nm_k'):
            r.pop(k, None)
    return jsonify(result)


# ─── PENGATURAN UMUM (identitas klinik, faktur, logo) ─
_SETTING_KEYS = ['faktur_prefix','klinik_nama','klinik_sub','klinik_alamat','klinik_telp','klinik_kota']

def _logo_info():
    path = os.path.join(_base_dir(), 'static', 'logo.png')
    if os.path.exists(path):
        return True, '/static/logo.png?v=' + str(int(os.path.getmtime(path)))
    return False, ''

@app.route('/api/pengaturan', methods=['GET'])
def get_pengaturan():
    out = {k: get_setting(k, '') for k in _SETTING_KEYS}
    out['point_rate'] = get_point_rate()
    out['point_redeem_rate'] = get_point_redeem_rate()
    has_logo, logo_url = _logo_info()
    out['has_logo'] = has_logo
    out['logo_url'] = logo_url
    # Prefix faktur per cabang (semua lokasi non-induk)
    cabang = rows_to_list(query(
        "SELECT kd_lokasi, nm_lokasi FROM Lokasi WHERE aktif=1 AND kd_lokasi!='INDUK' AND tipe!='induk' ORDER BY kd_lokasi"))
    for c in cabang:
        c['prefix'] = get_setting(f"faktur_prefix_{c['kd_lokasi']}", '')
    out['cabang_prefix'] = cabang
    return jsonify(out)

@app.route('/api/pengaturan', methods=['PUT'])
def save_pengaturan():
    d = request.json or {}
    lama = {k: get_setting(k, '') for k in _SETTING_KEYS if k in d}
    for k in _SETTING_KEYS:
        if k in d:
            set_setting(k, str(d[k]).strip())
    # Prefix faktur per cabang: {kd_lokasi: prefix}
    cp = d.get('cabang_prefix')
    if isinstance(cp, dict):
        for kd, val in cp.items():
            set_setting(f"faktur_prefix_{kd}", str(val).strip().upper())
    catat_log('pengaturan_update', 'Pengaturan', '', "Update pengaturan identitas klinik/faktur",
              data_lama=lama, data_baru={k: str(d[k]).strip() for k in _SETTING_KEYS if k in d})
    return jsonify({'ok': True})

# ─── HAK AKSES MENU PER ROLE ─────────────────────────────────
@app.route('/api/menu-access')
def get_menu_access():
    """Config hak akses menu per role: {role: [menu_key,...]}. Kosong = pakai default frontend."""
    raw = get_setting('menu_access', '')
    try:
        data = json.loads(raw) if raw else {}
    except Exception:
        data = {}
    return jsonify(data if isinstance(data, dict) else {})

@app.route('/api/menu-access', methods=['PUT'])
def save_menu_access():
    user = get_current_user()
    if not user or user.get('role') != 'admin':
        return jsonify({'ok': False, 'msg': 'Hanya admin yang boleh mengubah hak akses menu'}), 403
    d = request.json or {}
    cfg = d.get('config', d)   # terima {config:{...}} atau langsung {...}
    if not isinstance(cfg, dict):
        return jsonify({'ok': False, 'msg': 'Format config tidak valid'}), 400
    lama_raw = get_setting('menu_access', '')
    set_setting('menu_access', json.dumps(cfg))
    catat_log('hak_akses_update', 'Pengaturan', 'menu_access', "Update konfigurasi hak akses menu per role",
              data_lama={'raw': lama_raw}, data_baru=cfg)
    return jsonify({'ok': True})

# ─── LOG AKTIVITAS (audit trail — admin only) ─────────────────
@app.route('/api/log-aktivitas')
def get_log_aktivitas():
    user = get_current_user()
    if not user or user.get('role') != 'admin':
        return jsonify({'ok': False, 'msg': 'Hanya admin yang boleh melihat log aktivitas'}), 403
    tgl1     = request.args.get('dari', str(date.today()))
    tgl2     = request.args.get('sampai', str(date.today()))
    aksi     = request.args.get('aksi', '').strip()
    username = request.args.get('username', '').strip()
    q        = request.args.get('q', '').strip()

    where  = ["waktu >= ? AND waktu < ?"]
    params = [tgl1 + ' 00:00:00', tgl2 + ' 23:59:59']
    if aksi:
        where.append("aksi=?"); params.append(aksi)
    if username:
        where.append("username=?"); params.append(username)
    if q:
        where.append("(keterangan LIKE ? OR ref LIKE ? OR nama_user LIKE ?)")
        params.extend([f'%{q}%', f'%{q}%', f'%{q}%'])

    rows = query(f"""SELECT * FROM Aktivitas_Log WHERE {' AND '.join(where)}
                     ORDER BY id DESC LIMIT 500""", params)
    aksi_list = query("SELECT DISTINCT aksi FROM Aktivitas_Log ORDER BY aksi")
    user_list = query("SELECT DISTINCT username, nama_user FROM Aktivitas_Log WHERE username!='' ORDER BY nama_user")
    return jsonify({
        'rows': rows_to_list(rows),
        'aksi_list': [r['aksi'] for r in aksi_list],
        'user_list': rows_to_list(user_list),
    })

@app.route('/api/pengaturan/logo', methods=['POST'])
def upload_logo():
    d = request.json or {}
    img = d.get('image_data', '')
    if not img:
        return jsonify({'ok': False, 'msg': 'Tidak ada gambar'}), 400
    try:
        raw = re.sub(r'^data:image/\w+;base64,', '', img)
        img_bytes = base64.b64decode(raw)
        if HAS_PIL:
            im = Image.open(_io.BytesIO(img_bytes))
            im = ImageOps.exif_transpose(im)
            if im.width > 600 or im.height > 600:
                im.thumbnail((600, 600), Image.LANCZOS)
            buf = _io.BytesIO()
            if im.mode in ('RGBA', 'LA', 'P'):
                im.convert('RGBA').save(buf, format='PNG')          # pertahankan transparansi
            else:
                im.convert('RGB').save(buf, format='PNG')
            img_bytes = buf.getvalue()
        path = os.path.join(_base_dir(), 'static', 'logo.png')
        with open(path, 'wb') as f:
            f.write(img_bytes)
        return jsonify({'ok': True, 'logo_url': '/static/logo.png?v=' + str(int(datetime.now().timestamp()))})
    except Exception as e:
        return jsonify({'ok': False, 'msg': str(e)}), 500


# ─── PENGATURAN POINT ────────────────────────────────
@app.route('/api/pengaturan/point', methods=['GET'])
def get_point_setting():
    return jsonify({'point_rate': get_point_rate(), 'point_redeem_rate': get_point_redeem_rate()})

@app.route('/api/pengaturan/point', methods=['PUT'])
def set_point_setting():
    d = request.json or {}
    try:
        rate = float(d.get('point_rate'))
    except (TypeError, ValueError):
        return jsonify({'ok': False, 'msg': 'Nilai rate tidak valid'}), 400
    if rate <= 0:
        return jsonify({'ok': False, 'msg': 'Rate harus lebih besar dari 0'}), 400
    set_setting('point_rate', rate)
    # Kurs redeem (opsional, field terpisah) — kalau tidak dikirim, tidak diubah.
    redeem_rate = get_point_redeem_rate()
    if 'point_redeem_rate' in d:
        try:
            redeem_rate = float(d.get('point_redeem_rate'))
        except (TypeError, ValueError):
            return jsonify({'ok': False, 'msg': 'Nilai kurs redeem tidak valid'}), 400
        if redeem_rate <= 0:
            return jsonify({'ok': False, 'msg': 'Kurs redeem harus lebih besar dari 0'}), 400
        set_setting('point_redeem_rate', redeem_rate)
    return jsonify({'ok': True, 'point_rate': rate, 'point_redeem_rate': redeem_rate})


# ─── LAPORAN HISTORY POINT PASIEN ────────────────────
@app.route('/api/laporan/point')
def laporan_point():
    """Rekap capaian point per pasien — SELALU akumulatif semua waktu, TIDAK difilter
    tanggal (point itu saldo berjalan dari seluruh transaksi pasien, bukan sesuatu yg
    wajar "dipotong" per periode — keputusan user 2026-10-08). Filter tanggal tetap
    relevan & dipakai di drill-down transaksi per pasien (laporan_point_detail()), cuma
    tidak di list ringkasan ini. ?q= cari nama/kode pasien — hasilnya SELALU menyeluruh
    (seluruh riwayat), tidak terpengaruh tanggal apapun."""
    q = request.args.get('q', '').strip()
    where, params = '', []
    if q:
        where = "WHERE (p.nm_pasien LIKE ? OR dp.kd_pasien LIKE ?)"
        params = [f'%{q}%', f'%{q}%']
    rows = query(f"""
        SELECT dp.kd_pasien, p.nm_pasien, p.telp,
               COALESCE(SUM(dp.debit),0)  AS point_masuk,
               COALESCE(SUM(dp.kredit),0) AS point_keluar,
               COALESCE(SUM(dp.debit),0) - COALESCE(SUM(dp.kredit),0) AS saldo_total,
               COUNT(*) AS jml_transaksi,
               MAX(dp.tanggal) AS terakhir
        FROM DetailPoint dp
        LEFT JOIN Pasien p ON dp.kd_pasien = p.kd_pasien
        {where}
        GROUP BY dp.kd_pasien
        ORDER BY saldo_total DESC, terakhir DESC
    """, params)
    result = rows_to_list(rows)
    summary = {
        'total_masuk':  sum(r['point_masuk']  for r in result),
        'total_keluar': sum(r['point_keluar'] for r in result),
        'jml_pasien':   len(result),
        'point_rate':   get_point_rate(),
        'point_redeem_rate': get_point_redeem_rate(),
    }
    return jsonify({'rows': result, 'summary': summary})


@app.route('/api/laporan/point/detail/<path:kd_pasien>')
def laporan_point_detail(kd_pasien):
    """Rincian mutasi point satu pasien."""
    dari   = request.args.get('dari', '')
    sampai = request.args.get('sampai', '')
    filters, params = ["dp.kd_pasien=?"], [kd_pasien]
    if dari:   filters.append("dp.tanggal >= ?"); params.append(dari)
    if sampai: filters.append("dp.tanggal <= ?"); params.append(sampai)
    where = "WHERE " + " AND ".join(filters)
    rows = query(f"""
        SELECT dp.id, dp.tanggal, dp.keterangan, dp.debit, dp.kredit,
               pj.total_harga, pj.metode_bayar
        FROM DetailPoint dp
        LEFT JOIN Penjualan pj ON dp.keterangan = pj.no_faktur
        {where}
        ORDER BY dp.tanggal ASC, dp.id ASC
    """, params)
    pasien = query("SELECT kd_pasien, nm_pasien, telp FROM Pasien WHERE kd_pasien=?", (kd_pasien,), one=True)
    saldo  = query("SELECT COALESCE(SUM(debit),0)-COALESCE(SUM(kredit),0) AS s FROM DetailPoint WHERE kd_pasien=?",
                   (kd_pasien,), one=True)
    return jsonify({
        'pasien': dict(pasien) if pasien else {'kd_pasien': kd_pasien},
        'saldo_total': saldo['s'] if saldo else 0,
        'rows': rows_to_list(rows),
    })


# ─── LAPORAN STOK PENJUALAN GLOBAL ───────────────────
def _hitung_stok_penjualan(tgl1, tgl2, tipe, q):
    """Inti query Laporan Stok Penjualan — dipakai bareng oleh endpoint JSON &
    endpoint export Excel, supaya filter & hasilnya selalu identik."""
    # Pencarian kombinasi kode & nama (spt POS) — lihat _kata_cocok_kolom(). Dibuat 2
    # varian kolom krn di WHERE MySQL tidak boleh pakai alias SELECT (mis. "nm_barang"
    # hasil "drr.nm_bahan AS nm_barang" di cabang racikan) — harus kolom asli per cabang.
    # q_params SAMA utk keduanya (murni dari kata kunci, bukan nama kolom).
    words = q.split()
    q_where_std,     q_params = _kata_cocok_kolom('rp.kd_barang', 'rp.nm_barang', words)
    q_where_racikan, _        = _kata_cocok_kolom('drr.kd_barang', 'drr.nm_bahan', words)
    q_filter_std     = f"AND {q_where_std}"
    q_filter_racikan = f"AND {q_where_racikan}"

    tipe_filter = ""
    if tipe == 'langsung':
        tipe_filter = "WHERE tipe_keluar='langsung'"
    elif tipe == 'racikan':
        tipe_filter = "WHERE tipe_keluar='racikan'"

    rows = query(f"""
        SELECT
            kd_barang, nm_barang,
            SUM(CASE WHEN tipe_keluar='langsung' THEN qty ELSE 0 END) AS qty_langsung,
            SUM(CASE WHEN tipe_keluar='racikan'  THEN qty ELSE 0 END) AS qty_racikan,
            SUM(qty) AS qty_total,
            SUM(CASE WHEN tipe_keluar='langsung' THEN nilai ELSE 0 END) AS omzet_langsung,
            SUM(CASE WHEN tipe_keluar='racikan'  THEN nilai ELSE 0 END) AS hpp_racikan,
            SUM(nilai) AS total_nilai,
            COUNT(DISTINCT no_faktur) AS total_trx,
            MAX(tanggal) AS tgl_terakhir
        FROM (
            SELECT rp.kd_barang, rp.nm_barang, rp.qty,
                   rp.sub_total AS nilai, 'langsung' AS tipe_keluar,
                   p.tanggal, rp.no_faktur
            FROM R_Penjualan rp
            JOIN Penjualan p ON rp.no_faktur = p.no_faktur
            WHERE p.status='done' AND rp.jenis='barang'
              AND p.tanggal BETWEEN ? AND ?
              {q_filter_std}

            UNION ALL

            SELECT drr.kd_barang, drr.nm_bahan AS nm_barang, drr.qty,
                   drr.sub_total AS nilai, 'racikan' AS tipe_keluar,
                   p.tanggal, rr.no_faktur
            FROM DetailResepRacikan drr
            JOIN ResepRacikan rr ON drr.id_resep = rr.id
            JOIN Penjualan p ON rr.no_faktur = p.no_faktur
            WHERE rr.status='confirmed' AND p.status='done'
              AND p.tanggal BETWEEN ? AND ?
              {q_filter_racikan}

            UNION ALL

            -- Retur barang langsung → kurangi qty & nilai (qty negatif, sub_total retur sudah negatif)
            SELECT rp.kd_barang, rp.nm_barang, -rp.qty AS qty,
                   rp.sub_total AS nilai, 'langsung' AS tipe_keluar,
                   p.tanggal, rp.no_faktur
            FROM R_Penjualan rp
            JOIN Penjualan p ON rp.no_faktur = p.no_faktur
            WHERE p.status='done' AND rp.jenis='retur_barang'
              AND p.tanggal BETWEEN ? AND ?
              {q_filter_std}
            -- Catatan: retur RACIKAN tidak perlu cabang netting di sini — saat racikan
            -- diretur penuh, ResepRacikan.status diset 'returned' (di proses_retur),
            -- sehingga otomatis tereksklusi dari source racikan di atas (filter status='confirmed').
        ) combined
        {tipe_filter}
        GROUP BY kd_barang, nm_barang
        HAVING qty_total <> 0 OR total_nilai <> 0
        ORDER BY qty_total DESC
    """, [tgl1, tgl2] + q_params + [tgl1, tgl2] + q_params
       + [tgl1, tgl2] + q_params)
    return rows_to_list(rows)

@app.route('/api/laporan/stok-penjualan')
def laporan_stok_penjualan():
    """Semua barang keluar: langsung terjual + masuk ke dalam racikan."""
    tgl1 = request.args.get('dari',   str(date.today()))
    tgl2 = request.args.get('sampai', str(date.today()))
    q    = request.args.get('q',      '').strip()
    tipe = request.args.get('tipe',   'semua')   # semua | langsung | racikan

    result = _hitung_stok_penjualan(tgl1, tgl2, tipe, q)
    total_qty   = sum(r['qty_total']   for r in result)
    total_nilai = sum(r['total_nilai'] for r in result)
    max_qty     = max((r['qty_total']  for r in result), default=1)
    for r in result:
        r['pct_langsung'] = round(r['qty_langsung'] / r['qty_total'] * 100) if r['qty_total'] else 0
        r['pct_bar']      = round(r['qty_total'] / max_qty * 100) if max_qty else 0

    return jsonify({
        'rows':        result,
        'total_qty':   total_qty,
        'total_nilai': total_nilai,
        'total_sku':   len(result),
    })

@app.route('/api/laporan/stok-penjualan/export')
def laporan_stok_penjualan_export():
    """Export Laporan Stok Penjualan ke .xlsx — filter sama persis dgn endpoint JSON."""
    tgl1 = request.args.get('dari',   str(date.today()))
    tgl2 = request.args.get('sampai', str(date.today()))
    q    = request.args.get('q',      '').strip()
    tipe = request.args.get('tipe',   'semua')

    result = _hitung_stok_penjualan(tgl1, tgl2, tipe, q)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Stok Penjualan'
    kolom = ['Kode', 'Nama', 'Qty Langsung', 'Qty Racikan', 'Qty Total',
             'Omzet Langsung', 'HPP Racikan', 'Total Nilai', 'Jumlah Transaksi', 'Tanggal Terakhir']
    ws.append(kolom)
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='1A1A2E')
    for c in range(1, len(kolom) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    ws.freeze_panes = 'A2'

    for r in result:
        ws.append([
            r['kd_barang'], r['nm_barang'], r['qty_langsung'], r['qty_racikan'], r['qty_total'],
            r['omzet_langsung'], r['hpp_racikan'], r['total_nilai'], r['total_trx'], r['tgl_terakhir'],
        ])
        row_i = ws.max_row
        for c in (6, 7, 8):
            ws.cell(row=row_i, column=c).number_format = '#,##0'

    lebar = [22, 30, 12, 12, 11, 15, 14, 14, 16, 15]
    for idx, w_ in enumerate(lebar, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w_

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    fname = f"Laporan_Stok_Penjualan_{tgl1}_sd_{tgl2}.xlsx"
    from flask import Response
    return Response(out.getvalue(),
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     headers={'Content-Disposition': f'attachment; filename="{fname}"'})


@app.route('/api/laporan/stok-penjualan/transaksi')
def laporan_stok_penjualan_transaksi():
    """Detail per-transaksi untuk satu produk (langsung + dalam racikan)."""
    tgl1      = request.args.get('dari',      str(date.today()))
    tgl2      = request.args.get('sampai',    str(date.today()))
    kd_barang = request.args.get('kd_barang', '')

    rows = query("""
        SELECT p.tanggal, rp.no_faktur, ps.kd_pasien, ps.nm_pasien, p.dokter,
               rp.qty, rp.harga, rp.sub_total,
               rp.disc, 'langsung' AS tipe, NULL AS nama_racikan
        FROM R_Penjualan rp
        JOIN Penjualan p   ON rp.no_faktur  = p.no_faktur
        LEFT JOIN Pasien ps ON p.kd_pasien  = ps.kd_pasien
        WHERE p.status='done' AND rp.jenis='barang'
          AND rp.kd_barang=? AND p.tanggal BETWEEN ? AND ?

        UNION ALL

        SELECT p.tanggal, rr.no_faktur, ps.kd_pasien, ps.nm_pasien, p.dokter,
               drr.qty, drr.harga_satuan AS harga, drr.sub_total,
               '0' AS disc, 'racikan' AS tipe, rr.nama_produk AS nama_racikan
        FROM DetailResepRacikan drr
        JOIN ResepRacikan rr ON drr.id_resep = rr.id
        JOIN Penjualan p     ON rr.no_faktur = p.no_faktur
        LEFT JOIN Pasien ps  ON p.kd_pasien  = ps.kd_pasien
        WHERE rr.status='confirmed' AND p.status='done'
          AND drr.kd_barang=? AND p.tanggal BETWEEN ? AND ?

        UNION ALL

        -- Retur barang langsung (qty & subtotal negatif)
        SELECT p.tanggal, rp.no_faktur, ps.kd_pasien, ps.nm_pasien, p.dokter,
               -rp.qty AS qty, rp.harga, rp.sub_total,
               rp.disc, 'retur' AS tipe, NULL AS nama_racikan
        FROM R_Penjualan rp
        JOIN Penjualan p   ON rp.no_faktur  = p.no_faktur
        LEFT JOIN Pasien ps ON p.kd_pasien  = ps.kd_pasien
        WHERE p.status='done' AND rp.jenis='retur_barang'
          AND rp.kd_barang=? AND p.tanggal BETWEEN ? AND ?
        -- Retur racikan tidak perlu cabang sendiri: ResepRacikan.status jadi 'returned'
        -- saat diretur → otomatis hilang dari source racikan di atas.

        ORDER BY tanggal DESC
    """, (kd_barang, tgl1, tgl2, kd_barang, tgl1, tgl2,
          kd_barang, tgl1, tgl2))

    return jsonify(rows_to_list(rows))


# ─── LAPORAN TREATMENT / PERAWATAN ───────────────────
# Klausa "baris ini treatment beneran" — dipakai laporan_treatment() & laporan_treatment_transaksi().
# JANGAN cuma andalkan rp.jenis='jasa': data lama hasil migrasi (sebelum app ini bisa
# bedakan jenis 'jasa') tersimpan pakai jenis='barang', padahal isinya treatment beneran
# (mis. faktur 0102230001W: "SPARKLING SR FACIAL" tersimpan jenis='barang'). Cuma 21 dari
# ribuan baris treatment yg literally berlabel 'jasa' — sisanya HARUS dikenali lewat
# Point_karyawan (petugas pelaksana treatment, sumber otoritatif yg TIDAK bergantung ke
# label jenis di R_Penjualan): kalau ada baris Point_karyawan dgn faktur & kd_pelayanan
# yg sama, baris R_Penjualan itu pasti treatment meski jenis-nya kebetulan 'barang'.
_TREATMENT_ROW_WHERE = """rp.jenis IN ('barang','jasa') AND (
    rp.jenis='jasa' OR EXISTS (
        SELECT 1 FROM Point_karyawan pk2
        WHERE pk2.no_faktur=rp.no_faktur AND pk2.kd_pelayanan=rp.kd_barang
    )
)"""

@app.route('/api/laporan/treatment')
def laporan_treatment():
    """Laporan penjualan perawatan — ringkasan per layanan & per perawat."""
    tgl1 = request.args.get('dari',   str(date.today()))
    tgl2 = request.args.get('sampai', str(date.today()))

    by_layanan = query(f"""
        SELECT rp.kd_barang AS kd_pelayanan,
               rp.nm_barang AS nm_pelayanan,
               SUM(rp.qty)             AS total_sesi,
               SUM(rp.sub_total)       AS total_omzet,
               COUNT(DISTINCT p.kd_pasien)  AS total_pasien,
               COUNT(DISTINCT rp.no_faktur) AS total_trx,
               MAX(p.tanggal)          AS tgl_terakhir,
               ROUND(AVG(rp.harga),0)  AS avg_harga
        FROM R_Penjualan rp
        JOIN Penjualan p ON rp.no_faktur = p.no_faktur
        WHERE p.status='done' AND {_TREATMENT_ROW_WHERE}
          AND p.tanggal BETWEEN ? AND ?
        GROUP BY rp.kd_barang, rp.nm_barang
        ORDER BY total_omzet DESC
    """, (tgl1, tgl2))

    by_karyawan = query("""
        SELECT pk.kd_karyawan, k.nm_karyawan, k.nickname, k.pekerjaan,
               COUNT(DISTINCT pk.no_faktur)   AS total_trx,
               COUNT(*)                        AS total_sesi,
               COUNT(DISTINCT pk.kd_pasien)   AS total_pasien,
               COUNT(DISTINCT pk.kd_pelayanan) AS jenis_layanan,
               SUM(pk.komisi)                  AS total_komisi
        FROM Point_karyawan pk
        LEFT JOIN Karyawan k ON pk.kd_karyawan = k.kd_karyawan
        LEFT JOIN Penjualan p ON pk.no_faktur  = p.no_faktur
        WHERE p.status='done' AND pk.tanggal BETWEEN ? AND ?
        GROUP BY pk.kd_karyawan
        ORDER BY total_sesi DESC
    """, (tgl1, tgl2))

    # Per DOKTER — beda sumber dari by_karyawan di atas: ini dokter yg tercatat di
    # Penjualan.dokter (dokter yg menangani KUNJUNGAN/rekam medik pasien), BUKAN
    # petugas pelaksana treatment yg dapat komisi (itu Point_karyawan/by_karyawan).
    # Satu dokter bisa menangani banyak layanan tanpa harus jadi "pelaksana" tercatat.
    by_dokter = query(f"""
        SELECT p.dokter AS nm_dokter,
               COUNT(DISTINCT rp.no_faktur) AS total_trx,
               SUM(rp.qty)                  AS total_sesi,
               COUNT(DISTINCT p.kd_pasien)  AS total_pasien,
               SUM(rp.sub_total)            AS total_omzet,
               MAX(p.tanggal)               AS tgl_terakhir
        FROM R_Penjualan rp
        JOIN Penjualan p ON rp.no_faktur = p.no_faktur
        WHERE p.status='done' AND {_TREATMENT_ROW_WHERE}
          AND p.dokter IS NOT NULL AND p.dokter != ''
          AND p.tanggal BETWEEN ? AND ?
        GROUP BY p.dokter
        ORDER BY total_omzet DESC
    """, (tgl1, tgl2))

    total_pasien_q = query(f"""
        SELECT COUNT(DISTINCT p.kd_pasien) AS c
        FROM R_Penjualan rp
        JOIN Penjualan p ON rp.no_faktur = p.no_faktur
        WHERE p.status='done' AND {_TREATMENT_ROW_WHERE}
          AND p.tanggal BETWEEN ? AND ?
    """, (tgl1, tgl2), one=True)

    result_layanan = rows_to_list(by_layanan)
    total_omzet = sum(r['total_omzet'] for r in result_layanan)
    total_sesi  = sum(r['total_sesi']  for r in result_layanan)
    max_omzet   = max((r['total_omzet'] for r in result_layanan), default=1)
    for r in result_layanan:
        r['pct'] = round(r['total_omzet'] / max_omzet * 100) if max_omzet else 0

    return jsonify({
        'by_layanan':   result_layanan,
        'by_karyawan':  rows_to_list(by_karyawan),
        'by_dokter':    rows_to_list(by_dokter),
        'total_omzet':  total_omzet,
        'total_sesi':   total_sesi,
        'total_pasien': total_pasien_q['c'] if total_pasien_q else 0,
    })


@app.route('/api/laporan/treatment/transaksi')
def laporan_treatment_transaksi():
    """Detail transaksi treatment — bisa filter per layanan atau per perawat."""
    tgl1         = request.args.get('dari',         str(date.today()))
    tgl2         = request.args.get('sampai',       str(date.today()))
    kd_karyawan  = request.args.get('kd_karyawan',  '')
    kd_pelayanan = request.args.get('kd_pelayanan', '')
    dokter       = request.args.get('dokter', '')

    filters = ["p.status='done'", _TREATMENT_ROW_WHERE, "p.tanggal BETWEEN ? AND ?"]
    params  = [tgl1, tgl2]

    if kd_karyawan:
        filters.append("rp.kd_karyawan=?")
        params.append(kd_karyawan)
    if kd_pelayanan:
        filters.append("rp.kd_barang=?")
        params.append(kd_pelayanan)
    if dokter:
        filters.append("p.dokter=?")
        params.append(dokter)

    rows = query(f"""
        SELECT p.tanggal, rp.no_faktur,
               ps.nm_pasien, ps.kd_pasien,
               p.dokter,
               rp.nm_barang AS nm_pelayanan,
               rp.kd_barang AS kd_pelayanan,
               rp.qty, rp.harga, rp.sub_total,
               rp.kd_karyawan,
               k.nm_karyawan, k.pekerjaan,
               COALESCE(pk.komisi, 0) AS komisi
        FROM R_Penjualan rp
        JOIN Penjualan p      ON rp.no_faktur   = p.no_faktur
        LEFT JOIN Pasien ps   ON p.kd_pasien    = ps.kd_pasien
        LEFT JOIN Karyawan k  ON rp.kd_karyawan = k.kd_karyawan
        LEFT JOIN Point_karyawan pk
               ON pk.no_faktur   = rp.no_faktur
              AND pk.kd_karyawan = rp.kd_karyawan
              AND pk.kd_pelayanan = rp.kd_barang
        WHERE {" AND ".join(filters)}
        ORDER BY p.tanggal DESC, rp.id DESC
    """, params)

    return jsonify(rows_to_list(rows))


# ─── POINT KARYAWAN ──────────────────────────────────
@app.route('/api/point-karyawan/batch', methods=['POST'])
def save_point_karyawan_batch():
    d = request.json
    records = d.get('records', [])
    today = str(date.today())
    for r in records:
        status = r.get('status') or 'released'   # 'pending' kalau treatment ini ada bahan kunci blm serah-terima
        existing = query("""SELECT id FROM Point_karyawan
            WHERE no_faktur=? AND kd_karyawan=? AND kd_pelayanan=?""",
            (r['no_faktur'], r['kd_karyawan'], r['kd_pelayanan']), one=True)
        if existing:
            execute("UPDATE Point_karyawan SET komisi=?, status=? WHERE id=?",
                    (r['komisi'], status, existing['id']))
        else:
            execute("""INSERT INTO Point_karyawan
                (no_faktur,tanggal,kd_pasien,kd_pelayanan,nm_pelayanan,komisi,kd_karyawan,status)
                VALUES (?,?,?,?,?,?,?,?)""",
                (r['no_faktur'], today, r.get('kd_pasien',''),
                 r['kd_pelayanan'], r['nm_pelayanan'], r['komisi'], r['kd_karyawan'], status))
    if records:
        total_komisi = sum(float(r.get('komisi', 0) or 0) for r in records)
        fkt = records[0].get('no_faktur', '')
        catat_log('simpan_komisi_petugas', 'Point_karyawan', fkt,
                  f"Simpan komisi {len(records)} petugas utk transaksi {fkt} — total {total_komisi:g}",
                  data_baru={'records': [{'kd_karyawan': r.get('kd_karyawan',''),
                                          'kd_pelayanan': r.get('kd_pelayanan',''),
                                          'komisi': r.get('komisi',0),
                                          'status': r.get('status') or 'released'} for r in records]})
    return jsonify({'ok': True, 'saved': len(records)})

# ─── RACIKAN ─────────────────────────────────────────
@app.route('/api/bahan-racikan')
def get_bahan_racikan():
    """Cari bahan untuk racikan — batch exp paling awal yg masih ada stok (FEFO) jadi
    baris utama (harga/exp acuan), sekaligus lampirkan rincian stok SEMUA batch (all_stok,
    urut FEFO) + totalnya — supaya kasir kelihatan di pencarian kalau qty besar bisa
    digabung dari beberapa batch (bukan cuma dibatasi batch pertama). Auto-split saat
    nambah bahan (tambahBahan() di frontend) pakai /api/barang-stok-batches/<kd>.

    PENTING: urutan FEFO dihitung di PYTHON pakai _exp_key() (parsing tanggal asli),
    BUKAN 'ORDER BY exp'/'MIN(exp)' di SQL — exp disimpan 'M/D/YYYY' tanpa leading zero,
    jadi sort SQL itu SORT STRING MENTAH & salah tiap kali panjang digit beda (mis.
    '1/1/2027' ke-anggap lbh awal drpd '12/1/2026' scr string, padahal sebaliknya scr
    tanggal asli) — bisa bikin batch yg SEBENARNYA lbh baru malah dipilih/ditawarkan
    duluan drpd yg beneran lbh dulu expired.
    Pencarian mengkombinasi kode & nama — lihat _kata_cocok_kolom()."""
    q = request.args.get('q', '').strip()
    kd_lokasi = request.args.get('kd_lokasi', 'CBG1')
    where, params = _kata_cocok_kolom('s.kd_barang', 's.nm_barang', q.split())
    prio, prio_params = _prioritas_kode('s.kd_barang', q)

    kd_rows = query(f"""
        SELECT DISTINCT s.kd_barang FROM Stok_barang s
        WHERE s.stok_sisa > 0 AND s.kd_lokasi=? AND {where}
        ORDER BY {prio}, s.kd_barang LIMIT 20
    """, (kd_lokasi, *params, *prio_params))
    kds = [r['kd_barang'] for r in kd_rows]
    if not kds:
        return jsonify([])

    ph = ','.join('?' * len(kds))
    all_batches = rows_to_list(query(f"""
        SELECT * FROM Stok_barang WHERE kd_barang IN ({ph}) AND stok_sisa > 0 AND kd_lokasi=?
    """, (*kds, kd_lokasi)))
    per_kd = {}
    for b in all_batches:
        per_kd.setdefault(b['kd_barang'], []).append(b)
    for lst in per_kd.values():
        lst.sort(key=lambda b: _exp_key(b['exp']))   # FEFO yg benar

    result = []
    for kd in kds:   # pertahankan urutan prioritas pencarian dari kd_rows
        batches = per_kd.get(kd)
        if not batches:
            continue
        main = dict(batches[0])   # batch FEFO paling awal → baris utama
        main['all_stok']   = [b['stok_sisa'] for b in batches]
        main['total_stok'] = sum(main['all_stok'])
        result.append(main)
    return jsonify(result)

@app.route('/api/racikan', methods=['POST'])
def simpan_racikan():
    d = request.json
    no_faktur  = d['no_faktur']
    id_rm      = d['id_rm']
    kd_pasien  = d['kd_pasien']
    nama_produk = d['nama_produk']
    cara_pakai  = d.get('cara_pakai', '')
    bahans      = d.get('bahans', [])  # [{kd_barang, nm_bahan, qty, satuan, harga_satuan, disc, exp}]
    today       = str(date.today())

    if not bahans:
        return jsonify({'ok': False, 'msg': 'Tidak ada bahan'}), 400
    if not _kd_lokasi_boleh_jual(d.get('kd_lokasi', 'CBG1')):
        return jsonify({'ok': False, 'msg': 'Penjualan ke pasien hanya boleh dari Cabang — Apotek Induk tidak bisa transaksi langsung ke pasien.'}), 403

    # Hitung total dari bahan — disc (%) per bahan opsional, bahan bisa punya diskon sendiri.
    harga_total = sum(float(b['harga_satuan']) * float(b['qty']) * (1 - float(b.get('disc', 0) or 0) / 100)
                       for b in bahans)

    con = get_db()
    try:
        cur = con.cursor()

        # 0. Pastikan header Penjualan ada (racikan bisa jadi item pertama/satu-satunya
        #    dalam transaksi → tanpa ini header tak pernah terbuat & transaksi hilang dari laporan)
        if not cur.execute("SELECT 1 FROM Penjualan WHERE no_faktur=?", (no_faktur,)).fetchone():
            cur.execute("""INSERT INTO Penjualan (no_faktur,tanggal,total_harga,hpp,dokter,kd_pasien,kasir,status)
                VALUES (?,?,0,0,?,?,?,'open')""",
                (no_faktur, today, d.get('dokter',''), kd_pasien, d.get('kasir','')))

        # 1. Simpan ResepRacikan
        cur.execute("""INSERT INTO ResepRacikan
            (id_rm, no_faktur, kd_pasien, nama_produk, cara_pakai, harga_total)
            VALUES (?,?,?,?,?,?)""",
            (id_rm, no_faktur, kd_pasien, nama_produk, cara_pakai, harga_total))
        id_resep = cur.lastrowid

        # 2. Simpan header R_Penjualan (jenis='racikan') — tampil di faktur
        cur.execute("""INSERT INTO R_Penjualan
            (no_faktur,tanggal,kd_barang,nm_barang,harga,harga_asli,qty,disc,sub_total,
             total_harbeli,cashier,exp,jenis,racikan_ref)
            VALUES (?,?,?,?,?,?,1,'0',?,0,?,?,?,?)""",
            (no_faktur, today, f'RCK-{id_resep}', nama_produk, harga_total, harga_total,
             harga_total, d.get('kasir',''), today, 'racikan', id_resep))
        r_penjualan_id = cur.lastrowid

        # 3. Simpan detail bahan + R_Penjualan (jenis='bahan_racikan') + potong stok
        for b in bahans:
            qty        = float(b['qty'])
            harga_sat  = float(b['harga_satuan'])
            disc_pct   = float(b.get('disc', 0) or 0)
            sub        = harga_sat * qty * (1 - disc_pct / 100)
            exp_bahan  = b.get('exp', today)

            cur.execute("""INSERT INTO DetailResepRacikan
                (id_resep, kd_barang, nm_bahan, qty, satuan, harga_satuan, disc, sub_total, exp)
                VALUES (?,?,?,?,?,?,?,?,?)""",
                (id_resep, b['kd_barang'], b['nm_bahan'], qty,
                 b.get('satuan',''), harga_sat, str(disc_pct), sub, exp_bahan))

            cur.execute("""INSERT INTO R_Penjualan
                (no_faktur,tanggal,kd_barang,nm_barang,harga,harga_asli,qty,disc,sub_total,
                 total_harbeli,cashier,exp,jenis,racikan_ref)
                VALUES (?,?,?,?,?,?,?,?,?,0,?,?,?,?)""",
                (no_faktur, today, b['kd_barang'], b['nm_bahan'], harga_sat, harga_sat,
                 qty, str(disc_pct), sub, d.get('kasir',''), exp_bahan, 'bahan_racikan', id_resep))

            # Potong stok cabang + catat RiwayatBarang
            kd_lok = d.get('kd_lokasi', 'CBG1')
            cur.execute("""UPDATE Stok_barang SET stok_sisa = stok_sisa - ?
                WHERE kd_barang=? AND exp=? AND kd_lokasi=?""", (qty, b['kd_barang'], exp_bahan, kd_lok))
            stok_r = cur.execute("SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                                 (b['kd_barang'], exp_bahan, kd_lok)).fetchone()
            if stok_r:
                cur.execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                            (today, b['kd_barang'], exp_bahan,
                             f"Racikan {nama_produk} [{no_faktur}]",
                             -qty, stok_r[0], stok_r[1] or 0))

        # 4. Update total header Penjualan
        cur.execute("""UPDATE Penjualan SET
            total_harga = (SELECT COALESCE(SUM(sub_total),0) FROM R_Penjualan
                           WHERE no_faktur=? AND jenis != 'bahan_racikan')
            WHERE no_faktur=?""", (no_faktur, no_faktur))

        con.commit()
    finally:
        con.close()

    catat_log('racikan_buat', 'ResepRacikan', no_faktur,
              f"Buat racikan '{nama_produk}' di transaksi {no_faktur} — {len(bahans)} bahan, total {harga_total:g}",
              data_baru={'id_resep': id_resep, 'nama_produk': nama_produk, 'no_faktur': no_faktur,
                         'harga_total': harga_total,
                         'bahans': [{'kd_barang': b['kd_barang'], 'nm_bahan': b['nm_bahan'],
                                     'qty': b['qty'], 'satuan': b.get('satuan','')} for b in bahans]})
    return jsonify({'ok': True, 'id_resep': id_resep, 'r_penjualan_id': r_penjualan_id,
                    'harga_total': harga_total})

@app.route('/api/racikan/<int:id_resep>/detail')
def get_racikan_detail(id_resep):
    resep   = query("SELECT * FROM ResepRacikan WHERE id=?", (id_resep,), one=True)
    details = query("SELECT * FROM DetailResepRacikan WHERE id_resep=? ORDER BY id", (id_resep,))
    return jsonify({
        'resep':   dict(resep) if resep else None,
        'details': rows_to_list(details)
    })

# ─── ANTRIAN RACIKAN (apoteker) ────────────────────────────────
@app.route('/api/racikan/antrian')
def racikan_antrian():
    """List racikan yg perlu/(sudah) diracik — utk halaman Antrian Racikan (apoteker).
    Bahan detail-nya TIDAK ikut di sini (dipakai endpoint lama /detail saat expand).
    HANYA racikan dari transaksi yg SUDAH DIBAYAR (Penjualan.status='done') — racikan
    yg baru ditambah kasir ke keranjang (transaksi masih 'open'/belum dibayar, bisa
    saja batal/berubah) sengaja belum masuk antrian apoteker."""
    status = request.args.get('status', 'pending')
    dari   = request.args.get('dari', '')
    sampai = request.args.get('sampai', '')
    where, params = ["COALESCE(rr.status_racik,'pending')=?", "pj.status='done'"], [status]
    if dari:
        where.append("pj.tanggal >= ?"); params.append(dari)
    if sampai:
        where.append("pj.tanggal <= ?"); params.append(sampai)
    rows = query(f"""
        SELECT rr.id, rr.no_faktur, rr.nama_produk, rr.cara_pakai, rr.kd_pasien,
               COALESCE(rr.status_racik,'pending') as status_racik,
               COALESCE(p.nm_pasien, rr.kd_pasien) as nm_pasien,
               pj.tanggal, pj.kd_lokasi,
               (SELECT COUNT(*) FROM DetailResepRacikan WHERE id_resep=rr.id) as jml_bahan
        FROM ResepRacikan rr
        LEFT JOIN Pasien p ON rr.kd_pasien=p.kd_pasien
        LEFT JOIN Penjualan pj ON rr.no_faktur=pj.no_faktur
        WHERE {' AND '.join(where)}
        ORDER BY rr.id DESC LIMIT 300
    """, params)
    return jsonify(rows_to_list(rows))

@app.route('/api/racikan/<int:id_resep>/tandai-selesai', methods=['POST'])
def racikan_tandai_selesai(id_resep):
    r = query("SELECT no_faktur, nama_produk FROM ResepRacikan WHERE id=?", (id_resep,), one=True)
    if not r:
        return jsonify({'ok': False, 'msg': 'Racikan tidak ditemukan'}), 404
    execute("UPDATE ResepRacikan SET status_racik='selesai' WHERE id=?", (id_resep,))
    catat_log('racikan_selesai', 'ResepRacikan', r['no_faktur'],
              f"Racikan '{r['nama_produk']}' ({r['no_faktur']}) ditandai selesai diracik")
    return jsonify({'ok': True})

@app.route('/api/racikan/<int:id_resep>', methods=['PUT'])
def update_racikan(id_resep):
    d = request.json
    no_faktur   = d['no_faktur']
    nama_produk = d['nama_produk']
    cara_pakai  = d.get('cara_pakai', '')
    bahans      = d.get('bahans', [])
    today       = str(date.today())

    if not bahans:
        return jsonify({'ok': False, 'msg': 'Tidak ada bahan'}), 400

    harga_total = sum(float(b['harga_satuan']) * float(b['qty']) * (1 - float(b.get('disc', 0) or 0) / 100)
                       for b in bahans)

    con = get_db()
    try:
        cur = con.cursor()

        # 1. Kembalikan stok semua bahan lama + catat RiwayatBarang
        kd_lok = d.get('kd_lokasi', 'CBG1')
        old = cur.execute(
            "SELECT kd_barang, qty, exp FROM DetailResepRacikan WHERE id_resep=?",
            (id_resep,)).fetchall()
        for row in old:
            cur.execute("UPDATE Stok_barang SET stok_sisa=stok_sisa+? WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                        (row[1], row[0], row[2], kd_lok))
            sr = cur.execute("SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                             (row[0], row[2], kd_lok)).fetchone()
            if sr:
                cur.execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                            (today, row[0], row[2], f"Edit Racikan (kembali) [{no_faktur}]",
                             row[1], sr[0], sr[1] or 0))

        # 2. Hapus detail lama + bahan_racikan lama di R_Penjualan
        cur.execute("DELETE FROM DetailResepRacikan WHERE id_resep=?", (id_resep,))
        cur.execute("DELETE FROM R_Penjualan WHERE racikan_ref=? AND jenis='bahan_racikan'", (id_resep,))

        # 3. Update header racikan di R_Penjualan + ResepRacikan — pertahankan disc
        # LEVEL TRANSAKSI yg mungkin sudah diterapkan kasir di keranjang (beda dari
        # disc per-bahan di dalam racikan), supaya edit susunan bahan tidak diam-diam
        # menghapus diskon baris itu.
        existing_disc_row = cur.execute(
            "SELECT disc FROM R_Penjualan WHERE racikan_ref=? AND jenis='racikan'", (id_resep,)).fetchone()
        existing_disc = float((existing_disc_row[0] if existing_disc_row else 0) or 0)
        sub_total_racikan = harga_total * (1 - existing_disc / 100)
        cur.execute("""UPDATE R_Penjualan SET nm_barang=?,harga=?,harga_asli=?,sub_total=?
            WHERE racikan_ref=? AND jenis='racikan'""",
            (nama_produk, harga_total, harga_total, sub_total_racikan, id_resep))
        cur.execute("UPDATE ResepRacikan SET nama_produk=?,cara_pakai=?,harga_total=? WHERE id=?",
                    (nama_produk, cara_pakai, harga_total, id_resep))

        # 4. Insert detail baru + deduct stok baru
        for b in bahans:
            qty       = float(b['qty'])
            harga_sat = float(b['harga_satuan'])
            disc_pct  = float(b.get('disc', 0) or 0)
            sub       = harga_sat * qty * (1 - disc_pct / 100)
            exp_b     = b.get('exp', today)
            cur.execute("""INSERT INTO DetailResepRacikan
                (id_resep,kd_barang,nm_bahan,qty,satuan,harga_satuan,disc,sub_total,exp)
                VALUES (?,?,?,?,?,?,?,?,?)""",
                (id_resep, b['kd_barang'], b['nm_bahan'], qty,
                 b.get('satuan',''), harga_sat, str(disc_pct), sub, exp_b))
            cur.execute("""INSERT INTO R_Penjualan
                (no_faktur,tanggal,kd_barang,nm_barang,harga,harga_asli,qty,disc,sub_total,
                 total_harbeli,cashier,exp,jenis,racikan_ref)
                VALUES (?,?,?,?,?,?,?,?,?,0,?,?,?,?)""",
                (no_faktur, today, b['kd_barang'], b['nm_bahan'], harga_sat, harga_sat,
                 qty, str(disc_pct), sub, d.get('kasir',''), exp_b, 'bahan_racikan', id_resep))
            cur.execute("UPDATE Stok_barang SET stok_sisa=stok_sisa-? WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                        (qty, b['kd_barang'], exp_b, kd_lok))
            sr2 = cur.execute("SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                              (b['kd_barang'], exp_b, kd_lok)).fetchone()
            if sr2:
                cur.execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                            (today, b['kd_barang'], exp_b,
                             f"Edit Racikan {nama_produk} [{no_faktur}]",
                             -qty, sr2[0], sr2[1] or 0))

        # 5. Update total Penjualan
        cur.execute("""UPDATE Penjualan SET
            total_harga=(SELECT COALESCE(SUM(sub_total),0) FROM R_Penjualan
                         WHERE no_faktur=? AND jenis!='bahan_racikan')
            WHERE no_faktur=?""", (no_faktur, no_faktur))

        con.commit()
    finally:
        con.close()

    catat_log('racikan_edit', 'ResepRacikan', no_faktur,
              f"Edit racikan '{nama_produk}' di transaksi {no_faktur} — {len(bahans)} bahan, total {harga_total:g}",
              data_baru={'id_resep': id_resep, 'nama_produk': nama_produk, 'no_faktur': no_faktur,
                         'harga_total': harga_total,
                         'bahans': [{'kd_barang': b['kd_barang'], 'nm_bahan': b['nm_bahan'],
                                     'qty': b['qty'], 'satuan': b.get('satuan','')} for b in bahans]})
    return jsonify({'ok': True, 'harga_total': harga_total})

@app.route('/api/racikan/<int:id_resep>', methods=['DELETE'])
def hapus_racikan(id_resep):
    no_faktur = request.args.get('no_faktur', '')
    kd_lok    = request.args.get('kd_lokasi', 'CBG1')
    today = str(date.today())
    _rck_lama = query("SELECT nama_produk, harga_total FROM ResepRacikan WHERE id=?", (id_resep,), one=True)
    if not _rck_lama:
        # Sudah dihapus sebelumnya (klik dobel yg lolos guard frontend, dll) — jangan proses
        # ulang (stok bahan sudah dikembalikan di panggilan pertama, tidak boleh dobel).
        return jsonify({'ok': True, 'sudah_dihapus': True})
    con = get_db()
    try:
        cur = con.cursor()
        # Kembalikan stok bahan cabang + catat RiwayatBarang
        details = cur.execute(
            "SELECT kd_barang, qty, exp FROM DetailResepRacikan WHERE id_resep=?",
            (id_resep,)).fetchall()
        for row in details:
            cur.execute("""UPDATE Stok_barang SET stok_sisa = stok_sisa + ?
                WHERE kd_barang=? AND exp=? AND kd_lokasi=?""", (row[1], row[0], row[2], kd_lok))
            sr = cur.execute("SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                             (row[0], row[2], kd_lok)).fetchone()
            if sr:
                cur.execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                            (today, row[0], row[2], f"Hapus Racikan [{no_faktur}]",
                             row[1], sr[0], sr[1] or 0))

        # Hapus semua R_Penjualan terkait racikan ini
        cur.execute("DELETE FROM R_Penjualan WHERE racikan_ref=? AND no_faktur=?",
                    (id_resep, no_faktur))

        # Hapus detail & header resep
        cur.execute("DELETE FROM DetailResepRacikan WHERE id_resep=?", (id_resep,))
        cur.execute("DELETE FROM ResepRacikan WHERE id=?", (id_resep,))

        # Update total
        cur.execute("""UPDATE Penjualan SET
            total_harga = (SELECT COALESCE(SUM(sub_total),0) FROM R_Penjualan
                           WHERE no_faktur=? AND jenis != 'bahan_racikan')
            WHERE no_faktur=?""", (no_faktur, no_faktur))
        con.commit()
    finally:
        con.close()
    catat_log('racikan_hapus', 'ResepRacikan', no_faktur,
              f"Hapus racikan '{_rck_lama['nama_produk'] if _rck_lama else id_resep}' dari transaksi {no_faktur}"
              + (f" (total {_rck_lama['harga_total']:g}, stok bahan dikembalikan)" if _rck_lama else ''),
              data_lama=(dict(_rck_lama) if _rck_lama else {'id_resep': id_resep}))
    _bersihkan_penjualan_kosong(no_faktur)
    return jsonify({'ok': True})

# ─── RESEP DIGITAL (dokumen teks, tidak terkait stok) ───────────────────────

@app.route('/api/resep-digital', methods=['POST'])
def buat_resep_digital():
    d = request.json
    kd_pasien  = d['kd_pasien']
    id_rm      = d.get('id_rm')
    nm_dokter  = d.get('nm_dokter', '')
    created_by = d.get('created_by', '')
    catatan    = d.get('catatan', '')
    no_faktur  = d.get('no_faktur', '')
    obat_list  = d.get('obat_list', [])
    today      = str(date.today())

    if not obat_list:
        return jsonify({'ok': False, 'msg': 'Minimal 1 obat/bahan'}), 400

    con = get_db()
    try:
        cur = con.cursor()
        count = cur.execute("SELECT COUNT(*) FROM ResepDigital WHERE tanggal=?", (today,)).fetchone()[0]
        no_resep = f"R/{datetime.now().strftime('%d%m%y')}/{(count+1):03d}"
        cur.execute("""INSERT INTO ResepDigital
            (no_resep, kd_pasien, id_rm, nm_dokter, tanggal, created_by, catatan, no_faktur)
            VALUES (?,?,?,?,?,?,?,?)""",
            (no_resep, kd_pasien, id_rm, nm_dokter, today, created_by, catatan, no_faktur))
        id_resep = cur.lastrowid
        for i, o in enumerate(obat_list, 1):
            cur.execute("""INSERT INTO DetailResepDigital
                (id_resep, nama_obat, dosis, satuan, cara_pakai, urutan)
                VALUES (?,?,?,?,?,?)""",
                (id_resep, o.get('nama_obat',''), o.get('dosis',''),
                 o.get('satuan',''), o.get('cara_pakai',''), i))
        con.commit()
    finally:
        con.close()
    return jsonify({'ok': True, 'id': id_resep, 'no_resep': no_resep})

@app.route('/api/resep-digital/<int:id>', methods=['PUT'])
def update_resep_digital(id):
    d = request.json
    nm_dokter  = d.get('nm_dokter', '')
    catatan    = d.get('catatan', '')
    obat_list  = d.get('obat_list', [])
    if not obat_list:
        return jsonify({'ok': False, 'msg': 'Minimal 1 obat/bahan'}), 400
    con = get_db()
    try:
        cur = con.cursor()
        cur.execute("UPDATE ResepDigital SET nm_dokter=?, catatan=? WHERE id=?",
                    (nm_dokter, catatan, id))
        cur.execute("DELETE FROM DetailResepDigital WHERE id_resep=?", (id,))
        for i, o in enumerate(obat_list, 1):
            cur.execute("""INSERT INTO DetailResepDigital
                (id_resep, nama_obat, dosis, satuan, cara_pakai, urutan)
                VALUES (?,?,?,?,?,?)""",
                (id, o.get('nama_obat',''), o.get('dosis',''),
                 o.get('satuan',''), o.get('cara_pakai',''), i))
        con.commit()
    finally:
        con.close()
    return jsonify({'ok': True})

@app.route('/api/resep-digital/<int:id>', methods=['DELETE'])
def hapus_resep_digital(id):
    con = get_db()
    try:
        cur = con.cursor()
        cur.execute("DELETE FROM DetailResepDigital WHERE id_resep=?", (id,))
        cur.execute("DELETE FROM ResepDigital WHERE id=?", (id,))
        con.commit()
    finally:
        con.close()
    return jsonify({'ok': True})

@app.route('/api/resep-digital/pasien/<path:kd>')
def get_resep_digital_pasien(kd):
    rows = query("""
        SELECT rd.*, p.nm_pasien FROM ResepDigital rd
        JOIN Pasien p ON rd.kd_pasien = p.kd_pasien
        WHERE rd.kd_pasien=? ORDER BY rd.tanggal DESC, rd.id DESC LIMIT 30
    """, (kd,))
    result = []
    for r in rows:
        row = dict(r)
        row['obat_list'] = rows_to_list(query(
            "SELECT * FROM DetailResepDigital WHERE id_resep=? ORDER BY urutan", (r['id'],)))
        result.append(row)
    return jsonify(result)

@app.route('/api/resep-digital/list')
def list_resep_digital():
    """Daftar & cari resep digital (untuk halaman lihat/cetak resep)."""
    q      = request.args.get('q', '')
    dari   = request.args.get('dari', '')
    sampai = request.args.get('sampai', '')
    where, params = ["1=1"], []
    if dari:   where.append("rd.tanggal>=?"); params.append(dari)
    if sampai: where.append("rd.tanggal<=?"); params.append(sampai)
    if q:
        where.append("(rd.no_resep LIKE ? OR rd.nm_dokter LIKE ? OR p.nm_pasien LIKE ? OR rd.kd_pasien LIKE ? OR rd.no_faktur LIKE ?)")
        params += [f'%{q}%'] * 5
    rows = query(f"""
        SELECT rd.id, rd.no_resep, rd.tanggal, rd.kd_pasien, p.nm_pasien,
               rd.nm_dokter, rd.no_faktur, rd.catatan, rd.created_by,
               (SELECT COUNT(*) FROM DetailResepDigital d WHERE d.id_resep=rd.id) AS jml_item
        FROM ResepDigital rd LEFT JOIN Pasien p ON rd.kd_pasien=p.kd_pasien
        WHERE {' AND '.join(where)}
        ORDER BY rd.tanggal DESC, rd.id DESC LIMIT 300
    """, params)
    return jsonify(rows_to_list(rows))

@app.route('/api/resep-digital/faktur/<faktur>')
def get_resep_digital_faktur(faktur):
    """Cek apakah faktur ini sudah punya resep digital (untuk gate di POS)."""
    r = query("SELECT id, no_resep FROM ResepDigital WHERE no_faktur=? ORDER BY id DESC LIMIT 1",
              (faktur,), one=True)
    return jsonify(dict(r) if r else None)

@app.route('/api/resep-digital/<int:id>')
def get_resep_digital(id):
    r = query("""
        SELECT rd.*, p.nm_pasien, p.tgl_lahir, p.jenis_kelamin
        FROM ResepDigital rd JOIN Pasien p ON rd.kd_pasien=p.kd_pasien
        WHERE rd.id=?
    """, (id,), one=True)
    if not r:
        return jsonify(None)
    row = dict(r)
    row['obat_list'] = rows_to_list(query(
        "SELECT * FROM DetailResepDigital WHERE id_resep=? ORDER BY urutan", (id,)))
    return jsonify(row)

# ─── RESEP DRAFT (dokter digital) ────────────────────

@app.route('/api/resep-draft/<path:kd_pasien>')
def get_resep_draft(kd_pasien):
    """Ambil semua resep draft untuk pasien hari ini (belum dikonfirmasi kasir)."""
    today = str(date.today())
    rows = query("""
        SELECT rr.*, rm.nm_dokter, rm.tanggal as tgl_rm
        FROM ResepRacikan rr
        LEFT JOIN RekamMedik rm ON rr.id_rm = rm.id
        WHERE rr.kd_pasien=? AND rr.status='draft' AND DATE(rm.tanggal)=?
        ORDER BY rr.id
    """, (kd_pasien, today))
    result = []
    for r in rows:
        row = dict(r)
        details = query("SELECT * FROM DetailResepRacikan WHERE id_resep=? ORDER BY id", (r['id'],))
        row['bahans'] = rows_to_list(details)
        result.append(row)
    return jsonify(result)

@app.route('/api/racikan-draft', methods=['POST'])
def simpan_racikan_draft():
    """Dokter simpan resep sebagai draft — tidak potong stok, tidak butuh no_faktur."""
    d = request.json
    id_rm          = d['id_rm']
    kd_pasien      = d['kd_pasien']
    cara_pakai     = d.get('cara_pakai', '')
    catatan_kasir  = d.get('catatan_kasir', '')
    bahans         = d.get('bahans', [])
    created_by     = d.get('created_by', '')

    if not bahans:
        return jsonify({'ok': False, 'msg': 'Tidak ada bahan'}), 400

    harga_total = sum(float(b['harga_satuan']) * float(b['qty']) for b in bahans)

    con = get_db()
    try:
        cur = con.cursor()
        cur.execute("""INSERT INTO ResepRacikan
            (id_rm, no_faktur, kd_pasien, nama_produk, cara_pakai, catatan_kasir, harga_total, status, created_by)
            VALUES (?,?,?,?,?,?,?,'draft',?)""",
            (id_rm, '', kd_pasien, '', cara_pakai, catatan_kasir, harga_total, created_by))
        id_resep = cur.lastrowid

        for b in bahans:
            qty = float(b['qty'])
            harga_sat = float(b['harga_satuan'])
            cur.execute("""INSERT INTO DetailResepRacikan
                (id_resep, kd_barang, nm_bahan, qty, satuan, harga_satuan, sub_total, exp)
                VALUES (?,?,?,?,?,?,?,?)""",
                (id_resep, b['kd_barang'], b['nm_bahan'], qty,
                 b.get('satuan', ''), harga_sat, harga_sat * qty, b.get('exp', '')))
        con.commit()
    finally:
        con.close()

    return jsonify({'ok': True, 'id_resep': id_resep, 'harga_total': harga_total})

@app.route('/api/racikan-draft/<int:id_resep>', methods=['PUT'])
def update_racikan_draft(id_resep):
    """Edit resep draft — ganti bahan, tidak ada perubahan stok."""
    d = request.json
    nama_produk = d['nama_produk']
    bahans = d.get('bahans', [])
    if not bahans:
        return jsonify({'ok': False, 'msg': 'Tidak ada bahan'}), 400

    harga_total = sum(float(b['harga_satuan']) * float(b['qty']) for b in bahans)

    con = get_db()
    try:
        cur = con.cursor()
        existing = cur.execute("SELECT status FROM ResepRacikan WHERE id=?", (id_resep,)).fetchone()
        if not existing or existing[0] != 'draft':
            return jsonify({'ok': False, 'msg': 'Resep sudah dikonfirmasi atau tidak ditemukan'}), 400

        cur.execute("UPDATE ResepRacikan SET nama_produk=?, cara_pakai=?, harga_total=? WHERE id=?",
                    (nama_produk, d.get('cara_pakai', ''), harga_total, id_resep))
        cur.execute("DELETE FROM DetailResepRacikan WHERE id_resep=?", (id_resep,))
        for b in bahans:
            qty = float(b['qty'])
            harga_sat = float(b['harga_satuan'])
            cur.execute("""INSERT INTO DetailResepRacikan
                (id_resep, kd_barang, nm_bahan, qty, satuan, harga_satuan, sub_total, exp)
                VALUES (?,?,?,?,?,?,?,?)""",
                (id_resep, b['kd_barang'], b['nm_bahan'], qty,
                 b.get('satuan', ''), harga_sat, harga_sat * qty, b.get('exp', '')))
        con.commit()
    finally:
        con.close()

    return jsonify({'ok': True, 'harga_total': harga_total})

@app.route('/api/racikan-draft/<int:id_resep>', methods=['DELETE'])
def hapus_racikan_draft(id_resep):
    """Hapus resep draft — tidak ada stok yang perlu dikembalikan."""
    con = get_db()
    try:
        cur = con.cursor()
        existing = cur.execute("SELECT status FROM ResepRacikan WHERE id=?", (id_resep,)).fetchone()
        if not existing or existing[0] != 'draft':
            return jsonify({'ok': False, 'msg': 'Bukan draft atau sudah dikonfirmasi'}), 400
        cur.execute("DELETE FROM DetailResepRacikan WHERE id_resep=?", (id_resep,))
        cur.execute("DELETE FROM ResepRacikan WHERE id=?", (id_resep,))
        con.commit()
    finally:
        con.close()
    return jsonify({'ok': True})

@app.route('/api/racikan-draft/<int:id_resep>/confirm', methods=['POST'])
def confirm_racikan_draft(id_resep):
    """Kasir konfirmasi resep draft: buat R_Penjualan, potong stok, ikat ke no_faktur."""
    d = request.json
    no_faktur   = d['no_faktur']
    kasir       = d.get('kasir', '')
    # Kasir yang menentukan nama produk/racikan saat konfirmasi
    nama_produk = d.get('nama_produk', '').strip()
    today = str(date.today())

    con = get_db()
    try:
        cur = con.cursor()
        resep = cur.execute("SELECT * FROM ResepRacikan WHERE id=? AND status='draft'",
                            (id_resep,)).fetchone()
        if not resep:
            return jsonify({'ok': False, 'msg': 'Draft tidak ditemukan'}), 404
        if not nama_produk:
            return jsonify({'ok': False, 'msg': 'Nama produk racikan wajib diisi oleh kasir'}), 400

        bahans = cur.execute("SELECT * FROM DetailResepRacikan WHERE id_resep=? ORDER BY id",
                             (id_resep,)).fetchall()
        harga_total = resep['harga_total']

        # Pastikan header Penjualan ada (racikan bisa jadi item satu-satunya → tanpa ini
        # transaksi hilang dari laporan penjualan walau stok sudah terpotong)
        if not cur.execute("SELECT 1 FROM Penjualan WHERE no_faktur=?", (no_faktur,)).fetchone():
            cur.execute("""INSERT INTO Penjualan (no_faktur,tanggal,total_harga,hpp,dokter,kd_pasien,kasir,status)
                VALUES (?,?,0,0,?,?,?,'open')""",
                (no_faktur, today, d.get('dokter',''), resep['kd_pasien'], kasir))

        # Buat header R_Penjualan (jenis='racikan') — tampil di faktur
        cur.execute("""INSERT INTO R_Penjualan
            (no_faktur,tanggal,kd_barang,nm_barang,harga,harga_asli,qty,disc,sub_total,
             total_harbeli,cashier,exp,jenis,racikan_ref)
            VALUES (?,?,?,?,?,?,1,'0',?,0,?,?,?,?)""",
            (no_faktur, today, f'RCK-{id_resep}', nama_produk, harga_total, harga_total,
             harga_total, kasir, today, 'racikan', id_resep))
        r_penjualan_id = cur.lastrowid

        kd_lok = d.get('kd_lokasi', 'CBG1')
        for b in bahans:
            qty = float(b['qty'])
            harga_sat = float(b['harga_satuan'])
            disc_pct = float(b['disc'] or 0) if 'disc' in b.keys() else 0.0
            sub = harga_sat * qty * (1 - disc_pct / 100)
            exp_bahan = b['exp'] or today

            cur.execute("""INSERT INTO R_Penjualan
                (no_faktur,tanggal,kd_barang,nm_barang,harga,harga_asli,qty,disc,sub_total,
                 total_harbeli,cashier,exp,jenis,racikan_ref)
                VALUES (?,?,?,?,?,?,?,?,?,0,?,?,?,?)""",
                (no_faktur, today, b['kd_barang'], b['nm_bahan'], harga_sat, harga_sat,
                 qty, str(disc_pct), sub, kasir, exp_bahan, 'bahan_racikan', id_resep))

            cur.execute("UPDATE Stok_barang SET stok_sisa=stok_sisa-? WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                        (qty, b['kd_barang'], exp_bahan, kd_lok))
            stok_r = cur.execute("SELECT stok_sisa, stok_gudang FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                                 (b['kd_barang'], exp_bahan, kd_lok)).fetchone()
            if stok_r:
                cur.execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                            (today, b['kd_barang'], exp_bahan,
                             f"Racikan {nama_produk} [{no_faktur}]",
                             -qty, stok_r[0], stok_r[1] or 0))

        cur.execute("UPDATE ResepRacikan SET no_faktur=?, nama_produk=?, status='confirmed' WHERE id=?",
                    (no_faktur, nama_produk, id_resep))

        cur.execute("""UPDATE Penjualan SET
            total_harga=(SELECT COALESCE(SUM(sub_total),0) FROM R_Penjualan
                         WHERE no_faktur=? AND jenis!='bahan_racikan')
            WHERE no_faktur=?""", (no_faktur, no_faktur))

        con.commit()
    finally:
        con.close()

    catat_log('racikan_buat', 'ResepRacikan', no_faktur,
              f"Konfirmasi racikan draft '{nama_produk}' ke transaksi {no_faktur} — {len(bahans)} bahan, total {harga_total:g}",
              data_baru={'id_resep': id_resep, 'nama_produk': nama_produk, 'no_faktur': no_faktur,
                         'harga_total': harga_total, 'via': 'draft'})
    return jsonify({'ok': True, 'id_resep': id_resep, 'harga_total': harga_total,
                    'nama_produk': nama_produk, 'r_penjualan_id': r_penjualan_id})

# ─── REKAM MEDIK ─────────────────────────────────────
@app.route('/api/rekam-medik/today/<path:kd>')
def get_rm_today(kd):
    today = str(date.today())
    # Hanya RM yang belum terhubung ke transaksi (open) yang boleh di-reuse
    open_rm = query(
        "SELECT * FROM RekamMedik WHERE kd_pasien=? AND tanggal=? AND (no_faktur='' OR no_faktur IS NULL) ORDER BY id DESC",
        (kd, today), one=True)
    # Hitung kunjungan yang sudah selesai hari ini (sudah punya faktur)
    done = query(
        "SELECT COUNT(*) as c FROM RekamMedik WHERE kd_pasien=? AND tanggal=? AND no_faktur!='' AND no_faktur IS NOT NULL",
        (kd, today), one=True)
    return jsonify({
        'open_rm': dict(open_rm) if open_rm else None,
        'completed_today': done['c'] if done else 0
    })

@app.route('/api/rekam-medik/pasien/<path:kd>')
def get_rm_pasien(kd):
    rows = query("SELECT * FROM RekamMedik WHERE kd_pasien=? ORDER BY tanggal DESC, id DESC LIMIT 20", (kd,))
    return jsonify(rows_to_list(rows))

@app.route('/api/rekam-medik', methods=['POST'])
def create_rm():
    d = request.json
    today = str(date.today())
    count = query("SELECT COUNT(*) as c FROM RekamMedik WHERE tanggal=?", (today,), one=True)
    no_rm = f"RM{datetime.now().strftime('%d%m%y')}{(count['c']+1):04d}"
    id_ = execute("""INSERT INTO RekamMedik
        (no_rm,kd_pasien,tanggal,nm_dokter,S_keluhan,O_jenis_kulit,O_kondisi_kulit,A_diagnosa,P_tindakan,P_catatan,status)
        VALUES (?,?,?,?,?,?,?,?,?,?,'open')""",
        (no_rm, d['kd_pasien'], today, d.get('nm_dokter',''),
         d.get('S_keluhan',''), d.get('O_jenis_kulit',''), d.get('O_kondisi_kulit',''),
         d.get('A_diagnosa',''), d.get('P_tindakan',''), d.get('P_catatan','')))
    # Backfill kasus "foto dulu, baru POS/RM": foto hari-yang-sama milik pasien ini
    # yang belum tertaut RM (id_rm null) diikat ke RM PERTAMA hari itu, biar
    # tautannya permanen (lewat id_rm), bukan sekadar cocok-tanggal.
    try:
        first_rm = query("""SELECT id, no_rm FROM RekamMedik
            WHERE kd_pasien=? AND tanggal=? ORDER BY id ASC LIMIT 1""",
            (d['kd_pasien'], today), one=True)
        if first_rm:
            execute("""UPDATE Foto_Pasien SET id_rm=?, no_rm=?
                WHERE kd_pasien=? AND tanggal=? AND (id_rm IS NULL OR id_rm='')""",
                (first_rm['id'], first_rm['no_rm'] or '', d['kd_pasien'], today))
    except Exception:
        pass

    rm = query("SELECT * FROM RekamMedik WHERE id=?", (id_,), one=True)
    return jsonify(dict(rm))

@app.route('/api/rekam-medik/<int:id>', methods=['PUT'])
def update_rm(id):
    d = request.json
    execute("""UPDATE RekamMedik SET nm_dokter=?,S_keluhan=?,O_jenis_kulit=?,
        O_kondisi_kulit=?,A_diagnosa=?,P_tindakan=?,P_catatan=?,no_faktur=? WHERE id=?""",
        (d.get('nm_dokter',''), d.get('S_keluhan',''), d.get('O_jenis_kulit',''),
         d.get('O_kondisi_kulit',''), d.get('A_diagnosa',''),
         d.get('P_tindakan',''), d.get('P_catatan',''), d.get('no_faktur',''), id))
    return jsonify({'ok': True})

@app.route('/api/rekam-medik/<int:id>/detail')
def get_rm_detail(id):
    """RM lengkap + item transaksi + resep racikan untuk satu kunjungan."""
    rm = query("""
        SELECT rm.*, p.nm_pasien, p.tgl_lahir, p.jenis_kelamin, p.telp
        FROM RekamMedik rm
        LEFT JOIN Pasien p ON rm.kd_pasien = p.kd_pasien
        WHERE rm.id = ?
    """, (id,), one=True)
    if not rm:
        return jsonify(None)

    result = dict(rm)
    result['items'] = []
    result['reseps'] = []
    result['penjualan'] = None

    if rm['no_faktur']:
        penjualan = query("SELECT * FROM Penjualan WHERE no_faktur=?", (rm['no_faktur'],), one=True)
        result['penjualan'] = dict(penjualan) if penjualan else None

        # Baris RETUR ITU SENDIRI (yg cuma artefak akuntansi, sub_total negatif) tetap
        # dikeluarkan dari daftar — tapi item ASLI-nya TETAP TAMPIL, ditandai `diretur`
        # (bukan disembunyikan) supaya riwayat kunjungan tetap utuh/transparan, konsisten
        # dgn pola DIRETUR yg dipakai di Laporan Komisi & Antrian Racikan.
        items_raw = query("""
            SELECT id, kd_barang, nm_barang, jenis, qty, harga, disc, sub_total, exp, racikan_ref
            FROM R_Penjualan
            WHERE no_faktur=? AND jenis NOT IN ('bahan_racikan') AND jenis NOT LIKE 'retur%'
            ORDER BY id
        """, (rm['no_faktur'],))
        items_list = []
        for it in rows_to_list(items_raw):
            if it['jenis'] == 'racikan' and it.get('racikan_ref'):
                rr = query("SELECT status FROM ResepRacikan WHERE id=?", (it['racikan_ref'],), one=True)
                it['diretur'] = bool(rr and rr['status'] == 'returned')
            else:
                ret = query("SELECT 1 FROM R_Penjualan WHERE kd_karyawan=? LIMIT 1",
                           (f"RETUR:{it['id']}",), one=True)
                it['diretur'] = bool(ret)
            items_list.append(it)
        result['items'] = items_list

        reseps = query("""
            SELECT rr.id, rr.nama_produk, rr.cara_pakai, rr.harga_total, rr.status
            FROM ResepRacikan rr
            WHERE rr.no_faktur=?
            ORDER BY rr.id
        """, (rm['no_faktur'],))
        resep_list = []
        for r in reseps:
            rd = dict(r)
            rd['diretur'] = (r['status'] == 'returned')
            bahans = query(
                "SELECT nm_bahan, qty, satuan, harga_satuan, sub_total FROM DetailResepRacikan WHERE id_resep=? ORDER BY id",
                (r['id'],)
            )
            rd['bahans'] = rows_to_list(bahans)
            resep_list.append(rd)
        result['reseps'] = resep_list

    # Resep digital yang terkait RM ini
    rd_rows = query("SELECT * FROM ResepDigital WHERE id_rm=? ORDER BY id", (id,))
    rd_list = []
    for rd in rd_rows:
        rdrow = dict(rd)
        rdrow['obat_list'] = rows_to_list(
            query("SELECT * FROM DetailResepDigital WHERE id_resep=? ORDER BY urutan", (rd['id'],)))
        # obat_list itu TEKS bebas (hasil salinan dari cart saat resep dibuat) — tidak
        # tertaut per-baris ke R_Penjualan/ResepRacikan, jadi tak bisa ditandai per-item.
        # Cukup 1 catatan di level dokumen: transaksi ini pernah ada retur atau tidak.
        rdrow['ada_retur'] = bool(rd['no_faktur'] and query(
            "SELECT 1 FROM R_Penjualan WHERE no_faktur=? AND jenis LIKE 'retur%' LIMIT 1",
            (rd['no_faktur'],), one=True))
        rd_list.append(rdrow)
    result['resep_digital'] = rd_list

    return jsonify(result)

@app.route('/api/pasien/<path:kd>/riwayat-lengkap')
def get_riwayat_lengkap(kd):
    """Timeline lengkap pasien: RM + transaksi tanpa RM, urut terbaru dulu."""
    # Kunjungan dengan Rekam Medik
    rm_rows = query("""
        SELECT 'rm'           AS tipe,
               rm.id          AS id,
               rm.no_rm       AS no_rm,
               rm.tanggal     AS tanggal,
               rm.nm_dokter   AS nm_dokter,
               rm.status      AS status,
               rm.S_keluhan, rm.O_jenis_kulit, rm.O_kondisi_kulit,
               rm.A_diagnosa, rm.P_tindakan, rm.P_catatan,
               rm.no_faktur   AS no_faktur,
               pj.total_harga AS total_harga_stored,
               (SELECT SUM(rp.sub_total) FROM R_Penjualan rp
                WHERE rp.no_faktur=rm.no_faktur AND rp.jenis!='bahan_racikan') AS total_item,
               pj.metode_bayar AS metode_bayar,
               pj.kasir       AS kasir
        FROM RekamMedik rm
        LEFT JOIN Penjualan pj ON rm.no_faktur = pj.no_faktur
        WHERE rm.kd_pasien = ?
    """, (kd,))

    # Transaksi tanpa RM (orphan)
    orphan_rows = query("""
        SELECT 'trx'          AS tipe,
               NULL           AS id,
               NULL           AS no_rm,
               pj.tanggal     AS tanggal,
               pj.kasir       AS nm_dokter,
               pj.status      AS status,
               NULL AS S_keluhan, NULL AS O_jenis_kulit, NULL AS O_kondisi_kulit,
               NULL AS A_diagnosa, NULL AS P_tindakan, NULL AS P_catatan,
               pj.no_faktur   AS no_faktur,
               pj.total_harga AS total_harga_stored,
               (SELECT SUM(rp.sub_total) FROM R_Penjualan rp
                WHERE rp.no_faktur=pj.no_faktur AND rp.jenis!='bahan_racikan') AS total_item,
               pj.metode_bayar AS metode_bayar,
               pj.kasir       AS kasir
        FROM Penjualan pj
        WHERE pj.kd_pasien = ?
          AND pj.status = 'done'
          AND NOT EXISTS (
              SELECT 1 FROM RekamMedik rm WHERE rm.no_faktur = pj.no_faktur
          )
    """, (kd,))

    # Resep digital yang TIDAK terkait RM (berdiri sendiri di timeline)
    rd_rows = query("""
        SELECT 'resep'        AS tipe,
               rd.id          AS id,
               NULL           AS no_rm,
               rd.tanggal     AS tanggal,
               rd.nm_dokter   AS nm_dokter,
               'done'         AS status,
               NULL AS S_keluhan, NULL AS O_jenis_kulit, NULL AS O_kondisi_kulit,
               NULL AS A_diagnosa, NULL AS P_tindakan, NULL AS P_catatan,
               NULL           AS no_faktur,
               NULL           AS total_harga,
               NULL           AS metode_bayar,
               rd.created_by  AS kasir,
               rd.no_resep    AS no_resep,
               rd.catatan     AS catatan_resep
        FROM ResepDigital rd
        WHERE rd.kd_pasien = ?
          AND (rd.id_rm IS NULL OR rd.id_rm = 0)
    """, (kd,))

    rm_list = rows_to_list(rm_rows)
    orphan_list = rows_to_list(orphan_rows)
    # total_harga stored di Penjualan bisa tidak sinkron dgn baris R_Penjualan
    # asli (mis. faktur lama migrasi) — selalu percaya SUM baris asli kalau ada.
    for r in rm_list + orphan_list:
        total_item = r.pop('total_item', None)
        total_stored = r.pop('total_harga_stored', None)
        r['total_harga'] = total_item if total_item is not None else total_stored

    all_rows = rm_list + orphan_list + rows_to_list(rd_rows)
    all_rows.sort(key=lambda r: (r['tanggal'] or '', r['id'] or 0), reverse=True)
    return jsonify(all_rows)

@app.route('/api/rekam-medik/pasien/<path:kd>/ringkasan')
def get_rm_pasien_ringkasan(kd):
    """Ringkasan kunjungan pasien — untuk quick-view di POS kasir."""
    limit = int(request.args.get('limit', 5))
    rows = query("""
        SELECT rm.id, rm.no_rm, rm.tanggal, rm.nm_dokter, rm.status,
               rm.S_keluhan, rm.A_diagnosa, rm.P_tindakan, rm.no_faktur,
               p.total_harga AS total_harga_stored,
               (SELECT SUM(rp.sub_total) FROM R_Penjualan rp
                WHERE rp.no_faktur=rm.no_faktur AND rp.jenis!='bahan_racikan') AS total_item,
               p.metode_bayar
        FROM RekamMedik rm
        LEFT JOIN Penjualan p ON rm.no_faktur = p.no_faktur
        WHERE rm.kd_pasien = ?
        ORDER BY rm.tanggal DESC, rm.id DESC
        LIMIT ?
    """, (kd, limit))
    out = rows_to_list(rows)
    for r in out:
        total_item = r.pop('total_item', None)
        total_stored = r.pop('total_harga_stored', None)
        r['total_harga'] = total_item if total_item is not None else total_stored
    return jsonify(out)

# ─── LOKASI ──────────────────────────────────────────────────
@app.route('/api/lokasi')
def get_lokasi():
    rows = query("SELECT * FROM Lokasi WHERE aktif=1 ORDER BY tipe DESC, nm_lokasi")
    return jsonify(rows_to_list(rows))

@app.route('/api/lokasi', methods=['POST'])
def create_lokasi():
    d = request.json
    kd = d.get('kd_lokasi','').strip().upper()
    nm = d.get('nm_lokasi','').strip()
    if not kd or not nm:
        return jsonify({'ok': False, 'msg': 'kd_lokasi dan nm_lokasi wajib'}), 400
    try:
        execute("INSERT INTO Lokasi (kd_lokasi,nm_lokasi,tipe,alamat) VALUES (?,?,?,?)",
                (kd, nm, d.get('tipe','cabang'), d.get('alamat','')))
        return jsonify({'ok': True, 'kd_lokasi': kd})
    except Exception as e:
        return jsonify({'ok': False, 'msg': str(e)}), 400

@app.route('/api/lokasi/<kd>', methods=['PUT'])
def update_lokasi(kd):
    d = request.json
    execute("UPDATE Lokasi SET nm_lokasi=?,tipe=?,alamat=?,aktif=? WHERE kd_lokasi=?",
            (d.get('nm_lokasi',''), d.get('tipe','cabang'), d.get('alamat',''), d.get('aktif',1), kd))
    return jsonify({'ok': True})

# ─── MASTER PEMILIK STOK INDUK ───────────────────────────────
@app.route('/api/induk-pemilik')
def get_induk_pemilik():
    """List pemilik stok induk. ?all=1 termasuk nonaktif. ?stats=1 sertakan jumlah SKU & nilai stok."""
    incl  = request.args.get('all', '') == '1'
    stats = request.args.get('stats', '') == '1'
    where = "" if incl else "WHERE aktif=1"
    rows = rows_to_list(query(f"SELECT * FROM Induk_Pemilik {where} ORDER BY aktif DESC, nm_pemilik"))
    if stats:
        for r in rows:
            s = query("""SELECT COUNT(DISTINCT kd_barang) AS sku,
                                COALESCE(SUM(stok_sisa),0) AS unit,
                                COALESCE(ROUND(SUM(stok_sisa*harga_beli),0),0) AS nilai
                         FROM Stok_barang WHERE kd_lokasi='INDUK' AND pemilik=?""",
                      (r['kd_pemilik'],), one=True)
            r['sku']   = s['sku']   if s else 0
            r['unit']  = s['unit']  if s else 0
            r['nilai'] = s['nilai'] if s else 0
    return jsonify(rows)

@app.route('/api/induk-pemilik', methods=['POST'])
def create_induk_pemilik():
    d  = request.json or {}
    nm = (d.get('nm_pemilik') or '').strip()
    if not nm:
        return jsonify({'ok': False, 'msg': 'Nama pemilik wajib diisi'}), 400
    # Auto kode P{n} berikutnya
    ada = [r['kd_pemilik'] for r in query("SELECT kd_pemilik FROM Induk_Pemilik")]
    n = 1
    while f'P{n}' in ada:
        n += 1
    kd = f'P{n}'
    execute("INSERT INTO Induk_Pemilik (kd_pemilik,nm_pemilik,keterangan,aktif) VALUES (?,?,?,1)",
            (kd, nm, (d.get('keterangan') or '').strip()))
    return jsonify({'ok': True, 'kd_pemilik': kd})

@app.route('/api/induk-pemilik/<kd>', methods=['PUT'])
def update_induk_pemilik(kd):
    d = request.json or {}
    nm = (d.get('nm_pemilik') or '').strip()
    if not nm:
        return jsonify({'ok': False, 'msg': 'Nama pemilik wajib diisi'}), 400
    execute("UPDATE Induk_Pemilik SET nm_pemilik=?, keterangan=?, aktif=? WHERE kd_pemilik=?",
            (nm, (d.get('keterangan') or '').strip(), int(d.get('aktif', 1)), kd))
    return jsonify({'ok': True})

@app.route('/api/induk-pemilik/<kd>', methods=['DELETE'])
def delete_induk_pemilik(kd):
    if kd == 'P1':
        return jsonify({'ok': False, 'msg': 'Pemilik utama (P1) tidak boleh dihapus'}), 400
    cnt = query("SELECT COUNT(*) AS c FROM Stok_barang WHERE kd_lokasi='INDUK' AND pemilik=? AND stok_sisa>0",
                (kd,), one=True)
    if cnt and cnt['c'] > 0:
        return jsonify({'ok': False, 'msg': f'Masih ada {cnt["c"]} batch berstok milik pemilik ini. Pindahkan/habiskan dulu.'}), 400
    execute("UPDATE Induk_Pemilik SET aktif=0 WHERE kd_pemilik=?", (kd,))
    return jsonify({'ok': True})

@app.route('/api/stok-induk')
def get_stok_induk():
    """Stok di Apotek Induk, grouped by kd_barang dengan ringkasan batch."""
    q = request.args.get('q', '')
    filter_type = request.args.get('filter', 'semua')
    pemilik = request.args.get('pemilik', '').strip()   # '' = semua pemilik
    today = str(date.today())
    thirty_days = str(date.today().replace(day=1))  # placeholder; computed per row
    from datetime import timedelta
    exp_warn_date = str(date.today() + timedelta(days=30))
    pemilik_sql = " AND pemilik=?" if pemilik else ""
    # kd_barang collation utf8mb4_bin (case-sensitive) — LOWER() wajib, lihat
    # _kata_cocok_kolom() utk penjelasan lengkap kenapa.
    q_like = f'%{q.lower()}%'
    sql_base = f"""
        SELECT kd_barang, MAX(nm_barang) as nm_barang, MAX(satuan) as satuan,
               SUM(stok_sisa)                                      as stok_sisa,
               SUM(stok_awal)                                      as stok_awal,
               MAX(stok_min)                                       as stok_min,
               ROUND(SUM(stok_sisa * harga_beli),0)               as nilai_stok,
               MIN(CASE WHEN stok_sisa>0 THEN harga_beli END)     as harga_beli_min,
               MAX(harga_beli)                                     as harga_beli_max,
               MAX(harga)                                          as harga_jual,
               MAX(disc)                                           as disc,
               MIN(CASE WHEN stok_sisa>0 THEN exp END)            as exp_fefo,
               MAX(exp)                                            as exp_max,
               COUNT(*)                                            as jml_batch,
               SUM(CASE WHEN stok_sisa>0 THEN 1 ELSE 0 END)       as batch_aktif,
               GROUP_CONCAT(DISTINCT pemilik)                      as pemilik_list,
               SUM(CASE WHEN exp < ? AND stok_sisa>0 THEN stok_sisa ELSE 0 END) as stok_exp_soon
        FROM Stok_barang
        WHERE kd_lokasi='INDUK'
          AND (LOWER(kd_barang) LIKE ? OR LOWER(nm_barang) LIKE ?){pemilik_sql}
        GROUP BY kd_barang
    """
    params = [exp_warn_date, q_like, q_like]
    if pemilik:
        params.append(pemilik)
    if filter_type == 'habis':
        sql_base += " HAVING SUM(stok_sisa) = 0"
    elif filter_type == 'ada':
        sql_base += " HAVING SUM(stok_sisa) > 0"
    elif filter_type == 'mau_habis':
        sql_base += " HAVING SUM(stok_sisa) > 0 AND SUM(stok_sisa) <= MAX(stok_min)"
    elif filter_type == 'exp_soon':
        sql_base += f" HAVING SUM(CASE WHEN exp < '{exp_warn_date}' AND stok_sisa>0 THEN stok_sisa ELSE 0 END) > 0"
    sql_base += " ORDER BY nm_barang"
    rows = query(sql_base, params)
    return jsonify(rows_to_list(rows))

@app.route('/api/stok-induk/cari')
def cari_stok_induk():
    """Cari barang terdaftar di stok INDUK — autocomplete formula produksi & konversi.
    Opsional ?pemilik= (batasi 1 pemilik induk) & ?instok=1 (hanya yang stok > 0)."""
    q = request.args.get('q', '').strip()
    pemilik = request.args.get('pemilik', '').strip()
    instok  = request.args.get('instok', '') == '1'
    # kd_barang collation utf8mb4_bin (case-sensitive) — LOWER() wajib, lihat
    # _kata_cocok_kolom() utk penjelasan lengkap kenapa.
    where = "kd_lokasi='INDUK' AND (LOWER(kd_barang) LIKE ? OR LOWER(nm_barang) LIKE ?)"
    params = [f'%{q.lower()}%', f'%{q.lower()}%']
    if pemilik:
        where += " AND pemilik=?"; params.append(pemilik)
    having = " HAVING SUM(stok_sisa) > 0" if instok else ""
    rows = query(f"""
        SELECT kd_barang, MAX(nm_barang) as nm_barang,
               SUM(stok_sisa) as stok_sisa, MAX(satuan) as satuan
        FROM Stok_barang
        WHERE {where}
        GROUP BY kd_barang
        {having}
        ORDER BY nm_barang LIMIT 15
    """, params)
    return jsonify(rows_to_list(rows))

@app.route('/api/stok-induk/batches/<kd_barang>')
def get_stok_induk_batches(kd_barang):
    pemilik = request.args.get('pemilik', '').strip()
    pem_sql = " AND s.pemilik=?" if pemilik else ""
    params = [kd_barang]
    if pemilik:
        params.append(pemilik)
    rows = query(f"""
        SELECT s.id, s.kd_barang, s.nm_barang, s.exp, s.stok_sisa, s.stok_awal, s.stok_gudang,
               s.harga_beli, s.harga_beli_awal, s.satuan, s.pemilik,
               COALESCE(ip.nm_pemilik, s.pemilik) AS nm_pemilik,
               ROUND(s.stok_sisa * s.harga_beli, 0) as nilai_batch
        FROM Stok_barang s
        LEFT JOIN Induk_Pemilik ip ON s.pemilik = ip.kd_pemilik
        WHERE s.kd_barang=? AND s.kd_lokasi='INDUK'{pem_sql}
        ORDER BY s.pemilik, s.exp ASC
    """, params)
    return jsonify(rows_to_list(rows))

# ─── KONVERSI SATUAN STOK INDUK (box→pcs dll, pindah antar pemilik) ───
def _next_no_konversi():
    today = str(date.today()).replace('-', '')[2:]   # yymmdd
    pfx = f'KV{today}'
    last = query("SELECT no_konversi FROM Konversi_Stok WHERE no_konversi LIKE ? ORDER BY no_konversi DESC LIMIT 1",
                 (f'{pfx}%',), one=True)
    seq = int(last['no_konversi'][-3:]) + 1 if last else 1
    return f'{pfx}{seq:03d}'

@app.route('/api/konversi/eksekusi', methods=['POST'])
def konversi_stok_eksekusi():
    """Konversi satuan stok induk (mis. 1 box → 10 pcs) sekaligus pindah pemilik.
    Stok sumber (kd+exp+pemilik) berkurang, stok tujuan (kd+exp+pemilik) bertambah.
    Nilai/HPP total dipertahankan: harga_beli tujuan = (harga_beli sumber × qty_sumber) / qty_target."""
    d = request.json or {}
    sumber_kd      = (d.get('sumber_kd') or '').strip()
    sumber_pemilik = (d.get('sumber_pemilik') or '').strip()
    sumber_exp     = (d.get('sumber_exp') or '').strip()
    target_kd      = (d.get('target_kd') or '').strip()
    target_pemilik = (d.get('target_pemilik') or '').strip()
    by             = (d.get('by') or '').strip()
    try:
        qty_sumber = float(d.get('qty_sumber') or 0)
        qty_target = float(d.get('qty_target') or 0)
    except (TypeError, ValueError):
        return jsonify({'ok': False, 'msg': 'Jumlah tidak valid'}), 400

    if not sumber_kd or not target_kd:
        return jsonify({'ok': False, 'msg': 'Pilih produk sumber & tujuan'}), 400
    if not sumber_pemilik or not target_pemilik:
        return jsonify({'ok': False, 'msg': 'Pilih pemilik sumber & tujuan'}), 400
    if qty_sumber <= 0 or qty_target <= 0:
        return jsonify({'ok': False, 'msg': 'Jumlah sumber & hasil harus lebih dari 0'}), 400
    if sumber_kd == target_kd and sumber_pemilik == target_pemilik:
        return jsonify({'ok': False, 'msg': 'Sumber & tujuan tidak boleh identik'}), 400

    # Produk tujuan wajib berasal dari master Barang — cegah kode/nama "ngarang" hasil
    # ketik bebas (kalau frontend kelewatan/dibypass) yang bikin Stok_barang berantakan.
    if not query("SELECT 1 FROM Barang WHERE kd_barang=?", (target_kd,), one=True):
        return jsonify({'ok': False, 'msg': f"Produk tujuan '{target_kd}' tidak ditemukan di master barang. Pilih dari daftar saran, jangan ketik manual."}), 400

    no_konv = _next_no_konversi()
    today   = str(date.today())
    con = get_db()
    try:
        cur = con.cursor()
        # Ambil batch sumber (exact exp bila diberikan; FEFO bila tidak)
        if sumber_exp:
            src = cur.execute("""SELECT * FROM Stok_barang
                WHERE kd_barang=? AND kd_lokasi='INDUK' AND pemilik=? AND exp=?
                ORDER BY stok_sisa DESC LIMIT 1""",
                (sumber_kd, sumber_pemilik, sumber_exp)).fetchone()
        else:
            src = cur.execute("""SELECT * FROM Stok_barang
                WHERE kd_barang=? AND kd_lokasi='INDUK' AND pemilik=? AND stok_sisa>0
                ORDER BY exp ASC LIMIT 1""",
                (sumber_kd, sumber_pemilik)).fetchone()
        if not src:
            return jsonify({'ok': False, 'msg': 'Batch sumber tidak ditemukan'}), 400
        if (src['stok_sisa'] or 0) < qty_sumber:
            return jsonify({'ok': False, 'msg': f"Stok sumber tidak cukup (ada {src['stok_sisa']}, butuh {qty_sumber})"}), 400

        exp_val    = src['exp']            # exp ikut pindah ke produk tujuan
        nm_sumber  = src['nm_barang']
        sat_sumber = src['satuan'] or ''
        hb_sumber  = src['harga_beli'] or 0
        hb_target  = round((hb_sumber * qty_sumber) / qty_target, 2) if qty_target else 0

        # 1) Kurangi stok sumber
        saldo_src = (src['stok_sisa'] or 0) - qty_sumber
        cur.execute("UPDATE Stok_barang SET stok_sisa=stok_sisa-? WHERE id=?", (qty_sumber, src['id']))
        cur.execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                    (today, sumber_kd, exp_val, f"Konversi {no_konv} → {target_kd}", -qty_sumber, saldo_src, 0))
        cur.execute("INSERT INTO Transfer_Stok_Log (no_referensi,tipe_referensi,kd_barang,nm_barang,kd_lokasi_dari,kd_lokasi_ke,qty,exp,created_by) VALUES (?,?,?,?,?,?,?,?,?)",
                    (no_konv, 'konversi_keluar', sumber_kd, nm_sumber or '', 'INDUK', 'KONVERSI', qty_sumber, exp_val, by))

        # 2) Info master tujuan
        mt = cur.execute("SELECT nm_barang, satuan, harga_jual FROM Barang WHERE kd_barang=?", (target_kd,)).fetchone()
        nm_target  = ((mt['nm_barang'] if mt else None) or '').strip() or target_kd
        sat_target = ((mt['satuan'] if mt else None) or '').strip()
        hj_target  = (mt['harga_jual'] if mt else 0) or 0

        # 3) Tambah stok tujuan (gabung bila batch kd+exp+pemilik sudah ada)
        tgt = cur.execute("""SELECT id, stok_sisa FROM Stok_barang
            WHERE kd_barang=? AND exp=? AND kd_lokasi='INDUK' AND pemilik=?""",
            (target_kd, exp_val, target_pemilik)).fetchone()
        if tgt:
            saldo_tgt = (tgt['stok_sisa'] or 0) + qty_target
            cur.execute("""UPDATE Stok_barang SET stok_sisa=stok_sisa+?, stok_awal=stok_awal+?, stok_gudang=stok_gudang+?
                           WHERE id=?""", (qty_target, qty_target, qty_target, tgt['id']))
        else:
            if not hj_target:
                hjr = cur.execute("SELECT MAX(harga) h FROM Stok_barang WHERE kd_barang=? AND kd_lokasi='INDUK'", (target_kd,)).fetchone()
                hj_target = (hjr['h'] if hjr else 0) or 0
            if not sat_target:
                sr = cur.execute("SELECT satuan FROM Stok_barang WHERE kd_barang=? AND satuan IS NOT NULL AND satuan!='' LIMIT 1", (target_kd,)).fetchone()
                sat_target = (sr['satuan'] if sr else '') or ''
            saldo_tgt = qty_target
            cur.execute("""INSERT INTO Stok_barang
                (kd_barang,nm_barang,harga_beli,harga_beli_awal,stok_awal,stok_sisa,stok_gudang,stok_min,harga,disc,satuan,exp,DiscMem,DiscLain,kd_lokasi,pemilik)
                VALUES (?,?,?,?,?,?,?,0,?,'0',?,?,0,0,'INDUK',?)""",
                (target_kd, nm_target, hb_target, hb_target, qty_target, qty_target, qty_target, hj_target, sat_target, exp_val, target_pemilik))
        cur.execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                    (today, target_kd, exp_val, f"Konversi {no_konv} dari {sumber_kd}", qty_target, saldo_tgt, 0))
        cur.execute("INSERT INTO Transfer_Stok_Log (no_referensi,tipe_referensi,kd_barang,nm_barang,kd_lokasi_dari,kd_lokasi_ke,qty,exp,created_by) VALUES (?,?,?,?,?,?,?,?,?)",
                    (no_konv, 'konversi_masuk', target_kd, nm_target, 'KONVERSI', 'INDUK', qty_target, exp_val, by))

        # 4) Catat riwayat konversi
        cur.execute("""INSERT INTO Konversi_Stok
            (no_konversi,tgl,sumber_kd,sumber_nm,sumber_pemilik,qty_sumber,satuan_sumber,exp,
             target_kd,target_nm,target_pemilik,qty_target,satuan_target,harga_beli_sumber,harga_beli_target,created_by,created_at)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (no_konv, today, sumber_kd, nm_sumber, sumber_pemilik, qty_sumber, sat_sumber, exp_val,
             target_kd, nm_target, target_pemilik, qty_target, sat_target, hb_sumber, hb_target, by, str(datetime.now())))
        con.commit()
    except Exception as e:
        con.rollback()
        con.close()
        return jsonify({'ok': False, 'msg': str(e)}), 500
    finally:
        try: con.close()
        except Exception: pass
    catat_log('konversi_stok', 'Stok_barang', no_konv,
              f"Konversi {no_konv}: {qty_sumber} {sumber_kd} ({sumber_pemilik}) → {qty_target} {target_kd} ({target_pemilik})",
              data_baru={'sumber_kd': sumber_kd, 'qty_sumber': qty_sumber, 'target_kd': target_kd, 'qty_target': qty_target})
    return jsonify({'ok': True, 'no_konversi': no_konv})

@app.route('/api/konversi/riwayat')
def konversi_riwayat():
    try:
        limit = int(request.args.get('limit', 50))
    except (TypeError, ValueError):
        limit = 50
    rows = query("""
        SELECT k.*,
               COALESCE(ps.nm_pemilik, k.sumber_pemilik) AS sumber_nm_pemilik,
               COALESCE(pt.nm_pemilik, k.target_pemilik) AS target_nm_pemilik
        FROM Konversi_Stok k
        LEFT JOIN Induk_Pemilik ps ON k.sumber_pemilik = ps.kd_pemilik
        LEFT JOIN Induk_Pemilik pt ON k.target_pemilik = pt.kd_pemilik
        ORDER BY k.id DESC LIMIT ?
    """, (limit,))
    return jsonify(rows_to_list(rows))

@app.route('/api/stok-induk/<kd_barang>/harga-jual', methods=['PUT'])
def set_harga_jual_induk(kd_barang):
    """Set harga jual & stok_min untuk semua batch produk ini di INDUK."""
    d = request.json
    harga  = float(d.get('harga', 0))
    disc   = str(d.get('disc', '0'))
    stok_min = float(d.get('stok_min', 0))
    lama = query("SELECT harga,disc,stok_min FROM Stok_barang WHERE kd_barang=? AND kd_lokasi='INDUK' LIMIT 1",
                 (kd_barang,), one=True)
    execute("""UPDATE Stok_barang SET harga=?, disc=?, stok_min=?
               WHERE kd_barang=? AND kd_lokasi='INDUK'""",
            (harga, disc, stok_min, kd_barang))
    catat_log('ubah_harga_induk', 'Stok_barang', kd_barang, f"Ubah harga jual INDUK '{kd_barang}'",
              data_lama=dict(lama) if lama else None,
              data_baru={'harga': harga, 'disc': disc, 'stok_min': stok_min})
    return jsonify({'ok': True})

@app.route('/api/stok-induk/kartu/<kd_barang>')
def kartu_stok_induk(kd_barang):
    """Kartu stok per produk di INDUK: semua mutasi masuk & keluar + saldo berjalan."""
    dari    = request.args.get('dari', '')
    sampai  = request.args.get('sampai', str(date.today()))

    # Info produk
    info = query("""SELECT MAX(nm_barang) as nm_barang, MAX(satuan) as satuan,
                           SUM(stok_sisa) as stok_sisa_aktual
                    FROM Stok_barang WHERE kd_barang=? AND kd_lokasi='INDUK'""",
                 (kd_barang,), one=True)

    # Hitung saldo awal (semua mutasi sebelum tanggal 'dari')
    saldo_awal = 0.0
    if dari:
        r = query("""SELECT
              COALESCE(SUM(CASE WHEN kd_lokasi_ke='INDUK'   THEN qty ELSE 0 END),0) as masuk,
              COALESCE(SUM(CASE WHEN kd_lokasi_dari='INDUK' THEN qty ELSE 0 END),0) as keluar
           FROM Transfer_Stok_Log
           WHERE kd_barang=? AND (kd_lokasi_ke='INDUK' OR kd_lokasi_dari='INDUK')
             AND date(tgl) < ?""", (kd_barang, dari), one=True)
        saldo_awal = round((r['masuk'] - r['keluar']) if r else 0, 4)

    # Ambil semua mutasi pada periode
    if dari:
        entries_raw = query("""
            SELECT id, tgl, no_referensi, tipe_referensi, nm_barang, exp,
                   CASE WHEN kd_lokasi_ke='INDUK'   THEN qty ELSE 0 END as masuk,
                   CASE WHEN kd_lokasi_dari='INDUK' THEN qty ELSE 0 END as keluar,
                   kd_lokasi_dari, kd_lokasi_ke, created_by
            FROM Transfer_Stok_Log
            WHERE kd_barang=? AND (kd_lokasi_ke='INDUK' OR kd_lokasi_dari='INDUK')
              AND date(tgl) BETWEEN ? AND ?
            ORDER BY tgl ASC, id ASC
        """, (kd_barang, dari, sampai))
    else:
        entries_raw = query("""
            SELECT id, tgl, no_referensi, tipe_referensi, nm_barang, exp,
                   CASE WHEN kd_lokasi_ke='INDUK'   THEN qty ELSE 0 END as masuk,
                   CASE WHEN kd_lokasi_dari='INDUK' THEN qty ELSE 0 END as keluar,
                   kd_lokasi_dari, kd_lokasi_ke, created_by
            FROM Transfer_Stok_Log
            WHERE kd_barang=? AND (kd_lokasi_ke='INDUK' OR kd_lokasi_dari='INDUK')
              AND date(tgl) <= ?
            ORDER BY tgl ASC, id ASC
        """, (kd_barang, sampai))

    # Hitung saldo berjalan
    saldo = saldo_awal
    entries = []
    for e in entries_raw:
        row = dict(e)
        saldo = round(saldo + row['masuk'] - row['keluar'], 4)
        row['saldo'] = saldo
        # Label keterangan yang ramah
        tipe_label = {
            'pembelian_masuk': 'Pembelian dari Vendor',
            'pembelian':       'Pembelian dari Vendor',
            'pembelian_hapus': 'Batal/Hapus Pembelian',
            'produksi_bahan':  'Dipakai Produksi',
            'produksi_hasil':  'Hasil Produksi',
            'order_transfer':  f"Kirim ke {row['kd_lokasi_ke']}",
        }.get(row['tipe_referensi'], row['tipe_referensi'])
        row['keterangan'] = tipe_label
        entries.append(row)

    return jsonify({
        'kd_barang': kd_barang,
        'nm_barang': info['nm_barang'] if info else kd_barang,
        'satuan': info['satuan'] if info else '',
        'stok_sisa_aktual': info['stok_sisa_aktual'] if info else 0,
        'saldo_awal': saldo_awal,
        'saldo_akhir': saldo,
        'entries': entries,
        'dari': dari, 'sampai': sampai,
    })

@app.route('/api/laporan/stok-induk')
def laporan_stok_induk():
    """Laporan lengkap stok induk: ringkasan nilai + per-batch + riwayat transfer bulan ini."""
    from datetime import timedelta
    today = str(date.today())
    exp_warn = str(date.today() + timedelta(days=30))
    bulan_ini = str(date.today())[:7]  # YYYY-MM

    summary = query("""
        SELECT COUNT(DISTINCT kd_barang) as total_sku,
               SUM(stok_sisa) as total_unit,
               ROUND(SUM(stok_sisa * harga_beli),0) as total_nilai,
               SUM(CASE WHEN stok_sisa=0 THEN 1 ELSE 0 END) as stok_habis,
               SUM(CASE WHEN stok_sisa>0 AND stok_sisa<=stok_min THEN 1 ELSE 0 END) as stok_mau_habis,
               SUM(CASE WHEN exp < ? AND stok_sisa>0 THEN 1 ELSE 0 END) as batch_exp_soon
        FROM Stok_barang WHERE kd_lokasi='INDUK'
    """, (exp_warn,), one=True)

    produk = query("""
        SELECT kd_barang, MAX(nm_barang) as nm_barang, MAX(satuan) as satuan,
               SUM(stok_sisa) as stok_sisa, MAX(stok_min) as stok_min,
               ROUND(SUM(stok_sisa * harga_beli),0) as nilai_stok,
               MIN(CASE WHEN stok_sisa>0 THEN exp END) as exp_fefo,
               COUNT(*) as jml_batch
        FROM Stok_barang WHERE kd_lokasi='INDUK'
        GROUP BY kd_barang
        ORDER BY nilai_stok DESC
    """)

    batches = query("""
        SELECT kd_barang, nm_barang, exp, stok_sisa, harga_beli,
               ROUND(stok_sisa*harga_beli,0) as nilai_batch, satuan
        FROM Stok_barang WHERE kd_lokasi='INDUK'
        ORDER BY kd_barang, exp ASC
    """)

    # Transfer_Stok_Log TABEL BERSAMA dipakai hampir semua fitur mutasi stok (Pembelian,
    # Produksi, Order Transfer, Terima Kiriman, Kit Perawat, Input Stok Manual, dst) —
    # dibedakan lewat kd_lokasi_dari/kd_lokasi_ke, BUKAN tabel khusus induk. Dulu query
    # ini cuma filter bulan (tanpa filter lokasi) jadi ikut narik mutasi cabang yg sama
    # sekali tak menyangkut Induk (mis. Input Stok Manual ke CBG1) — dibetulkan 2026-10-06,
    # sekarang HARUS melibatkan INDUK di salah satu sisi (dari ATAU ke).
    _month_expr = "DATE_FORMAT(tgl, '%Y-%m')" if DB_BACKEND == 'mysql' else "strftime('%Y-%m', tgl)"
    transfer_log = query(f"""
        SELECT * FROM Transfer_Stok_Log
        WHERE {_month_expr} = ? AND (kd_lokasi_dari='INDUK' OR kd_lokasi_ke='INDUK')
        ORDER BY id DESC LIMIT 200
    """, (bulan_ini,))

    return jsonify({
        'summary': dict(summary) if summary else {},
        'produk': rows_to_list(produk),
        'batches': rows_to_list(batches),
        'transfer_log': rows_to_list(transfer_log)
    })

# ─── LAPORAN PENGELUARAN INDUK → CABANG ──────────────────────
@app.route('/api/laporan/pengeluaran-induk')
def laporan_pengeluaran_induk():
    dari   = request.args.get('dari',   str(date.today()))
    sampai = request.args.get('sampai', str(date.today()))
    kd_lok = request.args.get('kd_lokasi', '')

    # Ubah sampai agar inklusif sampai akhir hari
    sampai_dt = sampai + ' 23:59:59'
    dari_dt   = dari   + ' 00:00:00'

    lok_clause = " AND o.kd_lokasi_asal=?" if kd_lok else ""
    base_params = [dari_dt, sampai_dt] + ([kd_lok] if kd_lok else [])

    # Harga jual per produk di INDUK
    harga_sub = """
        (SELECT MAX(harga) FROM Stok_barang
         WHERE kd_barang=d.kd_barang AND kd_lokasi='INDUK')
    """

    # Summary keseluruhan
    summary = query(f"""
        SELECT
          COUNT(DISTINCT o.id)                                              as total_order,
          COUNT(DISTINCT d.kd_barang)                                       as total_sku,
          COALESCE(SUM(d.qty_fulfilled), 0)                                 as total_qty,
          COALESCE(SUM(d.qty_fulfilled * {harga_sub}), 0)                   as total_nilai,
          COUNT(DISTINCT o.kd_lokasi_asal)                                  as total_cabang
        FROM Order_Transfer o
        JOIN Order_Transfer_Detail d ON d.id_order=o.id
        WHERE o.status IN ('fulfilled','partial')
          AND COALESCE(o.fulfilled_at, o.created_at) BETWEEN ? AND ?
          {lok_clause}
    """, base_params, one=True)

    # Per cabang
    per_cabang = query(f"""
        SELECT o.kd_lokasi_asal, l.nm_lokasi,
               COUNT(DISTINCT o.id)                                          as jml_order,
               COALESCE(SUM(d.qty_fulfilled), 0)                             as total_qty,
               COALESCE(SUM(d.qty_fulfilled * {harga_sub}), 0)               as total_nilai
        FROM Order_Transfer o
        JOIN Order_Transfer_Detail d ON d.id_order=o.id
        LEFT JOIN Lokasi l ON l.kd_lokasi=o.kd_lokasi_asal
        WHERE o.status IN ('fulfilled','partial')
          AND COALESCE(o.fulfilled_at, o.created_at) BETWEEN ? AND ?
          {lok_clause}
        GROUP BY o.kd_lokasi_asal
        ORDER BY total_nilai DESC
    """, base_params)

    # Per produk — total keluar
    per_produk = query(f"""
        SELECT d.kd_barang, MAX(d.nm_barang) as nm_barang,
               COALESCE(SUM(d.qty_fulfilled), 0)                             as total_qty,
               COALESCE(SUM(d.qty_fulfilled * {harga_sub}), 0)               as total_nilai,
               MAX({harga_sub})                                               as harga_jual,
               COUNT(DISTINCT o.id)                                          as jml_order
        FROM Order_Transfer o
        JOIN Order_Transfer_Detail d ON d.id_order=o.id
        WHERE o.status IN ('fulfilled','partial')
          AND COALESCE(o.fulfilled_at, o.created_at) BETWEEN ? AND ?
          {lok_clause}
          AND d.qty_fulfilled > 0
        GROUP BY d.kd_barang
        ORDER BY total_nilai DESC
    """, base_params)

    # Per order — dengan item detail di dalamnya
    orders_raw = query(f"""
        SELECT o.id, o.no_order, o.kd_lokasi_asal, l.nm_lokasi, o.status,
               o.fulfilled_at, o.fulfilled_by,
               o.approved_supervisor_by, o.catatan_cabang, o.catatan_supervisor,
               d.id as d_id, d.kd_barang, d.nm_barang,
               d.qty_request, d.qty_approved, d.qty_fulfilled,
               d.status_item,
               {harga_sub}                                                   as harga_jual,
               ROUND(d.qty_fulfilled * COALESCE({harga_sub}, 0), 0)         as nilai_item
        FROM Order_Transfer o
        JOIN Order_Transfer_Detail d ON d.id_order=o.id
        LEFT JOIN Lokasi l ON l.kd_lokasi=o.kd_lokasi_asal
        WHERE o.status IN ('fulfilled','partial')
          AND COALESCE(o.fulfilled_at, o.created_at) BETWEEN ? AND ?
          {lok_clause}
        ORDER BY COALESCE(o.fulfilled_at, o.created_at) DESC, o.id, d.id
    """, base_params)

    # Susun orders dict
    orders_map = {}
    for row in rows_to_list(orders_raw):
        oid = row['id']
        if oid not in orders_map:
            orders_map[oid] = {
                'id': oid, 'no_order': row['no_order'],
                'kd_lokasi_asal': row['kd_lokasi_asal'],
                'nm_lokasi': row['nm_lokasi'] or row['kd_lokasi_asal'],
                'status': row['status'],
                'fulfilled_at': row['fulfilled_at'],
                'fulfilled_by': row['fulfilled_by'],
                'approved_supervisor_by': row['approved_supervisor_by'],
                'catatan_cabang': row['catatan_cabang'],
                'catatan_supervisor': row['catatan_supervisor'],
                'items': [],
                'total_qty': 0,
                'total_nilai': 0,
            }
        if row['qty_fulfilled'] and row['qty_fulfilled'] > 0:
            orders_map[oid]['items'].append({
                'kd_barang':    row['kd_barang'],
                'nm_barang':    row['nm_barang'],
                'qty_request':  row['qty_request'],
                'qty_approved': row['qty_approved'],
                'qty_fulfilled':row['qty_fulfilled'],
                'status_item':  row['status_item'],
                'harga_jual':   row['harga_jual'] or 0,
                'nilai_item':   row['nilai_item'] or 0,
            })
            orders_map[oid]['total_qty']   += row['qty_fulfilled'] or 0
            orders_map[oid]['total_nilai'] += row['nilai_item'] or 0

    return jsonify({
        'summary':    dict(summary) if summary else {},
        'per_cabang': rows_to_list(per_cabang),
        'per_produk': rows_to_list(per_produk),
        'orders':     list(orders_map.values()),
    })

# ─── FORMULA PRODUKSI (BOM) ───────────────────────────────────
@app.route('/api/formula-produksi')
def list_formula():
    rows = query("""
        SELECT f.*,
               COALESCE(f.nm_produk_hasil,
                   (SELECT nm_barang FROM Stok_barang WHERE kd_barang=f.kd_produk_hasil LIMIT 1),
                   f.kd_produk_hasil) as nm_produk_hasil,
               (SELECT COUNT(*) FROM Formula_Produksi_Detail WHERE id_formula=f.id) as jml_bahan,
               (SELECT COUNT(*) FROM Formula_Produksi_Detail fd
                WHERE fd.id_formula=f.id
                  AND fd.qty > COALESCE((SELECT SUM(stok_sisa) FROM Stok_barang
                                         WHERE kd_barang=fd.kd_barang AND kd_lokasi='INDUK'), 0)
               ) as jml_bahan_kurang
        FROM Formula_Produksi f
        WHERE f.aktif=1
        ORDER BY f.nm_formula
    """)
    return jsonify(rows_to_list(rows))

@app.route('/api/formula-produksi', methods=['POST'])
def create_formula():
    d = request.json
    # Validasi: semua bahan harus terdaftar di stok INDUK
    for b in d.get('bahans', []):
        ada = query("SELECT 1 FROM Stok_barang WHERE kd_barang=? AND kd_lokasi='INDUK' LIMIT 1",
                    (b.get('kd_barang',''),), one=True)
        if not ada:
            return jsonify({'ok': False,
                'msg': f"Bahan '{b.get('nm_barang') or b.get('kd_barang')}' tidak ditemukan di Stok Induk. Pastikan sudah dibeli terlebih dahulu."}), 400
    fid = execute("""INSERT INTO Formula_Produksi
                     (kd_produk_hasil,nm_produk_hasil,nm_formula,qty_hasil,satuan_hasil,keterangan)
                     VALUES (?,?,?,?,?,?)""",
                  (d['kd_produk_hasil'], d.get('nm_produk_hasil',''), d['nm_formula'],
                   d.get('qty_hasil',1), d.get('satuan_hasil','pcs'), d.get('keterangan','')))
    for b in d.get('bahans', []):
        execute("INSERT INTO Formula_Produksi_Detail (id_formula,kd_barang,nm_barang,qty,satuan) VALUES (?,?,?,?,?)",
                (fid, b['kd_barang'], b.get('nm_barang',''), b['qty'], b.get('satuan','')))
    return jsonify({'ok': True, 'id': fid})

@app.route('/api/formula-produksi/<int:fid>')
def get_formula(fid):
    f = query("SELECT * FROM Formula_Produksi WHERE id=?", (fid,), one=True)
    d = query("""
        SELECT fd.*,
               COALESCE((SELECT SUM(stok_sisa) FROM Stok_barang
                         WHERE kd_barang=fd.kd_barang AND kd_lokasi='INDUK'), 0) as stok_sisa,
               COALESCE((SELECT satuan FROM Stok_barang
                         WHERE kd_barang=fd.kd_barang AND kd_lokasi='INDUK' LIMIT 1),
                        fd.satuan) as satuan_stok,
               COALESCE((SELECT MIN(exp) FROM Stok_barang
                         WHERE kd_barang=fd.kd_barang AND kd_lokasi='INDUK'
                         AND exp > ''), '') as exp_fefo
        FROM Formula_Produksi_Detail fd
        WHERE fd.id_formula=? ORDER BY fd.id
    """, (fid,))
    return jsonify({'formula': dict(f) if f else None, 'detail': rows_to_list(d)})

@app.route('/api/formula-produksi/<int:fid>', methods=['PUT'])
def update_formula(fid):
    d = request.json
    # Validasi: semua bahan harus terdaftar di stok INDUK
    for b in d.get('bahans', []):
        ada = query("SELECT 1 FROM Stok_barang WHERE kd_barang=? AND kd_lokasi='INDUK' LIMIT 1",
                    (b.get('kd_barang',''),), one=True)
        if not ada:
            return jsonify({'ok': False,
                'msg': f"Bahan '{b.get('nm_barang') or b.get('kd_barang')}' tidak ditemukan di Stok Induk. Pastikan sudah dibeli terlebih dahulu."}), 400
    execute("""UPDATE Formula_Produksi
               SET kd_produk_hasil=?,nm_produk_hasil=?,nm_formula=?,qty_hasil=?,satuan_hasil=?,keterangan=?
               WHERE id=?""",
            (d['kd_produk_hasil'], d.get('nm_produk_hasil',''), d['nm_formula'],
             d.get('qty_hasil',1), d.get('satuan_hasil','pcs'), d.get('keterangan',''), fid))
    execute("DELETE FROM Formula_Produksi_Detail WHERE id_formula=?", (fid,))
    for b in d.get('bahans', []):
        execute("INSERT INTO Formula_Produksi_Detail (id_formula,kd_barang,nm_barang,qty,satuan) VALUES (?,?,?,?,?)",
                (fid, b['kd_barang'], b.get('nm_barang',''), b['qty'], b.get('satuan','')))
    return jsonify({'ok': True})

@app.route('/api/formula-produksi/<int:fid>', methods=['DELETE'])
def delete_formula(fid):
    execute("UPDATE Formula_Produksi SET aktif=0 WHERE id=?", (fid,))
    return jsonify({'ok': True})

# ─── PRODUKSI ─────────────────────────────────────────────────
def _next_no_produksi():
    today = str(date.today()).replace('-','')[2:]
    last = query("SELECT no_produksi FROM Produksi WHERE no_produksi LIKE ? ORDER BY no_produksi DESC LIMIT 1",
                 (f'PRD{today}%',), one=True)
    seq = int(last['no_produksi'][-3:]) + 1 if last else 1
    return f'PRD{today}{seq:03d}'

@app.route('/api/produksi')
def list_produksi():
    dari   = request.args.get('dari', '')
    sampai = request.args.get('sampai', '')
    status = request.args.get('status', '')
    params, clauses = [], []
    if dari:   clauses.append("p.tgl_produksi >= ?"); params.append(dari)
    if sampai: clauses.append("p.tgl_produksi <= ?"); params.append(sampai)
    if status: clauses.append("p.status = ?");        params.append(status)
    where = ('WHERE ' + ' AND '.join(clauses)) if clauses else ''
    rows = query(f"""
        SELECT p.*, f.nm_formula, f.kd_produk_hasil, f.qty_hasil, f.satuan_hasil,
               (SELECT nm_barang FROM Stok_barang
                WHERE kd_barang=f.kd_produk_hasil AND kd_lokasi='INDUK' LIMIT 1) as nm_produk_hasil
        FROM Produksi p
        JOIN Formula_Produksi f ON p.id_formula=f.id
        {where}
        ORDER BY p.created_at DESC LIMIT 200
    """, params)
    result = rows_to_list(rows)
    for r in result:
        if not r.get('nm_produk_hasil'):
            nb = query("SELECT nm_barang FROM Stok_barang WHERE kd_barang=? LIMIT 1", (r['kd_produk_hasil'],), one=True)
            r['nm_produk_hasil'] = nb['nm_barang'] if nb else r['kd_produk_hasil']
    return jsonify(result)

@app.route('/api/produksi', methods=['POST'])
def create_produksi():
    d = request.json
    id_formula = d['id_formula']
    qty_batch  = float(d.get('qty_batch', 1) or 1)

    # Cek ketersediaan SEMUA bahan untuk jumlah batch ini. Bila ada yang kurang,
    # produksi TIDAK dibuat sama sekali — kembalikan daftar bahan yang kurang.
    bahans = query("SELECT * FROM Formula_Produksi_Detail WHERE id_formula=?", (id_formula,))
    kurang = []
    for b in bahans:
        needed = (b['qty'] or 0) * qty_batch
        av = query("SELECT COALESCE(SUM(stok_sisa),0) s FROM Stok_barang WHERE kd_barang=? AND kd_lokasi='INDUK'",
                   (b['kd_barang'],), one=True)
        avail = (av['s'] if av else 0) or 0
        if avail < needed:
            kurang.append({
                'kd_barang': b['kd_barang'],
                'nm_barang': b['nm_barang'] or b['kd_barang'],
                'perlu':  round(needed, 4),
                'ada':    round(avail, 4),
                'kurang': round(needed - avail, 4),
                'satuan': b['satuan'] or '',
            })
    if kurang:
        nm = ', '.join(f"{k['nm_barang']} (perlu {k['perlu']}, ada {k['ada']} {k['satuan']})" for k in kurang)
        return jsonify({'ok': False, 'kurang': kurang,
                        'msg': f"Bahan tidak cukup untuk {qty_batch} batch: {nm}"}), 400

    no = _next_no_produksi()
    pid = execute("""INSERT INTO Produksi (no_produksi,id_formula,qty_batch,tgl_produksi,tgl_expired,keterangan,created_by)
                     VALUES (?,?,?,?,?,?,?)""",
                  (no, id_formula, qty_batch, d.get('tgl_produksi', str(date.today())),
                   d.get('tgl_expired',''), d.get('keterangan',''), d.get('created_by','')))
    return jsonify({'ok': True, 'id': pid, 'no_produksi': no})

@app.route('/api/produksi/<int:pid>/approve', methods=['POST'])
def approve_produksi(pid):
    d = request.json
    p = query("SELECT * FROM Produksi WHERE id=?", (pid,), one=True)
    if not p or p['status'] != 'draft':
        return jsonify({'ok': False, 'msg': 'Hanya status draft yang bisa di-approve'}), 400
    execute("UPDATE Produksi SET status='approved',approved_by=?,approved_at=? WHERE id=?",
            (d.get('approved_by',''), str(datetime.now()), pid))
    return jsonify({'ok': True})

@app.route('/api/produksi/<int:pid>/reject', methods=['POST'])
def reject_produksi(pid):
    d = request.json
    execute("UPDATE Produksi SET status='batal',approved_by=?,approved_at=? WHERE id=?",
            (d.get('by',''), str(datetime.now()), pid))
    return jsonify({'ok': True})

@app.route('/api/produksi/<int:pid>/selesai', methods=['POST'])
def selesai_produksi(pid):
    """Eksekusi produksi: kurangi bahan baku INDUK, tambah produk hasil ke INDUK."""
    d   = request.json
    p   = query("SELECT * FROM Produksi WHERE id=?", (pid,), one=True)
    if not p or p['status'] != 'approved':
        return jsonify({'ok': False, 'msg': 'Hanya status approved yang bisa diselesaikan'}), 400
    formula = query("SELECT * FROM Formula_Produksi WHERE id=?", (p['id_formula'],), one=True)
    bahans  = query("SELECT * FROM Formula_Produksi_Detail WHERE id_formula=?", (p['id_formula'],))
    qty_batch = float(p['qty_batch'])
    today = str(date.today())

    # Cek ketersediaan bahan baku di INDUK
    for b in bahans:
        needed = b['qty'] * qty_batch
        avail  = query("SELECT COALESCE(SUM(stok_sisa),0) as s FROM Stok_barang WHERE kd_barang=? AND kd_lokasi='INDUK'",
                       (b['kd_barang'],), one=True)
        if (avail['s'] if avail else 0) < needed:
            return jsonify({'ok': False, 'msg': f"Stok bahan {b['nm_barang']} tidak cukup (perlu {needed}, ada {avail['s'] if avail else 0})"}), 400

    con = get_db()
    try:
        cur = con.cursor()
        # Kurangi bahan baku dari INDUK (FEFO)
        for b in bahans:
            sisa = b['qty'] * qty_batch
            batches = cur.execute("""SELECT id, stok_sisa, exp FROM Stok_barang
                WHERE kd_barang=? AND kd_lokasi='INDUK' AND stok_sisa>0 ORDER BY exp ASC""",
                (b['kd_barang'],)).fetchall()
            for batch in batches:
                if sisa <= 0: break
                take = min(sisa, batch['stok_sisa'])
                cur.execute("UPDATE Stok_barang SET stok_sisa=stok_sisa-? WHERE id=?", (take, batch['id']))
                cur.execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                            (today, b['kd_barang'], batch['exp'], f"Produksi {p['no_produksi']}", -take, batch['stok_sisa']-take, 0))
                cur.execute("INSERT INTO Transfer_Stok_Log (no_referensi,tipe_referensi,kd_barang,nm_barang,kd_lokasi_dari,kd_lokasi_ke,qty,exp,created_by) VALUES (?,?,?,?,?,?,?,?,?)",
                            (p['no_produksi'], 'produksi_bahan', b['kd_barang'], b['nm_barang'] or '', 'INDUK', 'PRODUKSI', take, batch['exp'], d.get('by','')))
                sisa -= take

        # Tambah produk hasil ke INDUK
        qty_hasil = float(formula['qty_hasil']) * qty_batch
        kd_hasil  = formula['kd_produk_hasil']
        exp_hasil = p['tgl_expired'] or today
        # Utamakan nama yang SUDAH terpakai di stok INDUK utk kd_barang ini (biar konsisten,
        # tak pecah jadi baris terpisah di Laporan Stok Induk krn beda ejaan) — baru fallback
        # ke nm_produk_hasil formula, lalu ke kode, kalau produk ini benar2 baru di INDUK.
        nb = query("SELECT nm_barang FROM Stok_barang WHERE kd_barang=? AND kd_lokasi='INDUK' LIMIT 1",
                   (kd_hasil,), one=True)
        nm_hasil = (nb['nm_barang'] if nb else None) \
            or (formula['nm_produk_hasil'] or '').strip() or kd_hasil

        existing = cur.execute("SELECT id,stok_sisa FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi='INDUK' AND pemilik='P1'",
                               (kd_hasil, exp_hasil)).fetchone()
        if existing:
            cur.execute("UPDATE Stok_barang SET stok_sisa=stok_sisa+?,stok_awal=stok_awal+?,stok_gudang=stok_gudang+? WHERE id=?",
                        (qty_hasil, qty_hasil, qty_hasil, existing['id']))
        else:
            cur.execute("""INSERT INTO Stok_barang (kd_barang,nm_barang,harga_beli,harga_beli_awal,stok_awal,stok_sisa,
                           stok_gudang,stok_min,harga,disc,satuan,exp,DiscMem,DiscLain,kd_lokasi,pemilik)
                           VALUES (?,?,0,0,?,?,?,0,0,'0',?,?,0,0,'INDUK','P1')""",
                        (kd_hasil, nm_hasil, qty_hasil, qty_hasil, qty_hasil, formula['satuan_hasil'] or 'pcs', exp_hasil))
        cur.execute("INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g) VALUES (?,?,?,?,?,?,?)",
                    (today, kd_hasil, exp_hasil, f"Hasil Produksi {p['no_produksi']}", qty_hasil, qty_hasil, qty_hasil))
        cur.execute("INSERT INTO Transfer_Stok_Log (no_referensi,tipe_referensi,kd_barang,nm_barang,kd_lokasi_dari,kd_lokasi_ke,qty,exp,created_by) VALUES (?,?,?,?,?,?,?,?,?)",
                    (p['no_produksi'], 'produksi_hasil', kd_hasil, nm_hasil, 'PRODUKSI', 'INDUK', qty_hasil, exp_hasil, d.get('by','')))

        cur.execute("UPDATE Produksi SET status='selesai',selesai_by=?,selesai_at=? WHERE id=?",
                    (d.get('by',''), str(datetime.now()), pid))
        con.commit()
    finally:
        con.close()
    return jsonify({'ok': True})

# ─── ORDER TRANSFER ───────────────────────────────────────────
def _next_no_order(kd_lokasi):
    today = str(date.today()).replace('-','')[2:]
    pfx = f'OT{kd_lokasi}{today}'
    last = query("SELECT no_order FROM Order_Transfer WHERE no_order LIKE ? ORDER BY no_order DESC LIMIT 1",
                 (f'{pfx}%',), one=True)
    seq = int(last['no_order'][-3:]) + 1 if last else 1
    return f'{pfx}{seq:03d}'

@app.route('/api/order-transfer')
def list_order_transfer():
    kd_lokasi = request.args.get('kd_lokasi', '')
    status    = request.args.get('status', '')
    params, clauses = [], []
    if kd_lokasi: clauses.append("o.kd_lokasi_asal=?"); params.append(kd_lokasi)
    if status:    clauses.append("o.status=?");         params.append(status)
    where = ('WHERE ' + ' AND '.join(clauses)) if clauses else ''
    rows = query(f"""
        SELECT o.*, l.nm_lokasi,
               (SELECT COUNT(*) FROM Order_Transfer_Detail WHERE id_order=o.id) as jml_item
        FROM Order_Transfer o
        LEFT JOIN Lokasi l ON o.kd_lokasi_asal=l.kd_lokasi
        {where}
        ORDER BY o.created_at DESC LIMIT 200
    """, params)
    return jsonify(rows_to_list(rows))

@app.route('/api/order-transfer', methods=['POST'])
def create_order_transfer():
    d  = request.json
    kd = d.get('kd_lokasi_asal','CBG1')
    no = _next_no_order(kd)
    oid = execute("""INSERT INTO Order_Transfer (no_order,kd_lokasi_asal,tgl_order,catatan_cabang,created_by)
                     VALUES (?,?,?,?,?)""",
                  (no, kd, d.get('tgl_order', str(date.today())),
                   d.get('catatan_cabang',''), d.get('created_by','')))
    for item in d.get('items', []):
        execute("INSERT INTO Order_Transfer_Detail (id_order,kd_barang,nm_barang,qty_request) VALUES (?,?,?,?)",
                (oid, item['kd_barang'], item.get('nm_barang',''), item['qty_request']))
    return jsonify({'ok': True, 'id': oid, 'no_order': no})

def _stok_induk_efektif(kd_barang, exclude_order_id):
    """Stok INDUK yang BENAR-BENAR bisa dijanjikan ke order transfer ini = stok
    fisik dikurangi apa yang SUDAH dijanjikan (approved/partial, belum sepenuhnya
    dikirim) ke order transfer LAIN utk barang yang sama. Mencegah 2+ order
    sama-sama dianggap "stok cukup" & di-approve dobel padahal fisiknya cuma
    cukup utk salah satu (mis. stok 2, order A & B sama-sama minta 2)."""
    s = query("SELECT COALESCE(SUM(stok_sisa),0) as s FROM Stok_barang WHERE kd_barang=? AND kd_lokasi='INDUK'",
              (kd_barang,), one=True)
    total_stok = s['s'] if s else 0
    reserved = query("""
        SELECT COALESCE(SUM(COALESCE(otd.qty_approved,0) - COALESCE(otd.qty_fulfilled,0)),0) as r
        FROM Order_Transfer_Detail otd
        JOIN Order_Transfer ot ON otd.id_order=ot.id
        WHERE otd.kd_barang=? AND ot.id!=? AND ot.status IN ('approved','partial')
          AND otd.status_item!='cancelled'
    """, (kd_barang, exclude_order_id), one=True)
    reserved_qty = reserved['r'] if reserved else 0
    return max(0, total_stok - reserved_qty)

@app.route('/api/order-transfer/<int:oid>')
def get_order_transfer(oid):
    o = query("SELECT o.*,l.nm_lokasi FROM Order_Transfer o LEFT JOIN Lokasi l ON o.kd_lokasi_asal=l.kd_lokasi WHERE o.id=?", (oid,), one=True)
    d = query("SELECT * FROM Order_Transfer_Detail WHERE id_order=? ORDER BY id", (oid,))
    stok_induk = {row['kd_barang']: _stok_induk_efektif(row['kd_barang'], oid) for row in d}
    return jsonify({'order': dict(o) if o else None, 'detail': rows_to_list(d), 'stok_induk': stok_induk})

@app.route('/api/order-transfer/<int:oid>/submit', methods=['POST'])
def submit_order_transfer(oid):
    o = query("SELECT status FROM Order_Transfer WHERE id=?", (oid,), one=True)
    if not o or o['status'] != 'draft':
        return jsonify({'ok': False, 'msg': 'Hanya status draft yang bisa disubmit'}), 400
    execute("UPDATE Order_Transfer SET status='submitted' WHERE id=?", (oid,))
    return jsonify({'ok': True})

@app.route('/api/order-transfer/<int:oid>/approve', methods=['POST'])
def approve_order_transfer(oid):
    d = request.json
    o = query("SELECT status,no_order FROM Order_Transfer WHERE id=?", (oid,), one=True)
    if not o or o['status'] != 'submitted':
        return jsonify({'ok': False, 'msg': 'Hanya status submitted yang bisa di-approve'}), 400
    # Update qty_approved per item (bisa dikurangi supervisor, TAPI TIDAK BOLEH
    # melebihi qty_request MAUPUN stok yang benar-benar masih tersedia setelah
    # dikurangi reservasi order lain — dicek ulang di server, bukan cuma percaya
    # input dari form, supaya tidak bisa "dipaksa" via API langsung.
    for item in d.get('items', []):
        detail = query("SELECT kd_barang, qty_request FROM Order_Transfer_Detail WHERE id=?", (item['id'],), one=True)
        qty_req = detail['qty_request'] if detail else 0
        stok_efektif = _stok_induk_efektif(detail['kd_barang'], oid) if detail else 0
        try:
            qty_approved = float(item.get('qty_approved') or 0)
        except (TypeError, ValueError):
            qty_approved = 0
        qty_approved = max(0.0, min(qty_approved, qty_req, stok_efektif))
        execute("UPDATE Order_Transfer_Detail SET qty_approved=?,catatan=? WHERE id=?",
                (qty_approved, item.get('catatan',''), item['id']))
    execute("""UPDATE Order_Transfer SET status='approved',
               catatan_supervisor=?,approved_supervisor_by=?,approved_supervisor_at=?
               WHERE id=?""",
            (d.get('catatan_supervisor',''), d.get('approved_by',''), str(datetime.now()), oid))
    catat_log('ot_approve', 'Order_Transfer', o['no_order'], f"Approve order transfer {o['no_order']}",
              data_baru={'items': d.get('items', [])})
    return jsonify({'ok': True})

@app.route('/api/order-transfer/<int:oid>/reject', methods=['POST'])
def reject_order_transfer(oid):
    d = request.json
    o = query("SELECT no_order FROM Order_Transfer WHERE id=?", (oid,), one=True)
    execute("""UPDATE Order_Transfer SET status='batal',catatan_supervisor=?,
               approved_supervisor_by=?,approved_supervisor_at=? WHERE id=?""",
            (d.get('catatan',''), d.get('by',''), str(datetime.now()), oid))
    catat_log('ot_reject', 'Order_Transfer', o['no_order'] if o else oid,
              f"Tolak order transfer {o['no_order'] if o else oid}: {d.get('catatan','')}")
    return jsonify({'ok': True})

@app.route('/api/order-transfer/<int:oid>/fulfill', methods=['POST'])
def fulfill_order_transfer(oid):
    """Apotek Induk eksekusi pengiriman: potong stok INDUK, tambah stok cabang."""
    d = request.json
    o = query("SELECT * FROM Order_Transfer WHERE id=?", (oid,), one=True)
    if not o or o['status'] != 'approved':
        return jsonify({'ok': False, 'msg': 'Hanya status approved yang bisa di-fulfill'}), 400

    items   = query("SELECT * FROM Order_Transfer_Detail WHERE id_order=?", (oid,))
    kd_tuju = o['kd_lokasi_asal']
    today   = str(date.today())
    fulfillments = d.get('items', {})  # {id: qty_kirim}

    fulfill_log_items = []
    con = get_db()
    try:
        cur = con.cursor()
        all_done = True
        for item in items:
            if item['status_item'] == 'cancelled': continue
            qty_approved = item['qty_approved'] if item['qty_approved'] is not None else item['qty_request']
            try:
                qty_kirim = float(fulfillments.get(str(item['id']), qty_approved))
            except (TypeError, ValueError):
                qty_kirim = 0
            # Tidak boleh kirim melebihi yang sudah disetujui — dicek ulang di
            # server (bukan cuma percaya input form) supaya tidak bisa "dipaksa".
            qty_kirim = max(0.0, min(qty_kirim, qty_approved))
            if qty_kirim <= 0:
                all_done = False
                continue

            # Ambil dari INDUK (FEFO)
            sisa = qty_kirim
            batches = cur.execute("""SELECT id, stok_sisa, exp, harga_beli, nm_barang, satuan
                FROM Stok_barang WHERE kd_barang=? AND kd_lokasi='INDUK' AND stok_sisa>0
                ORDER BY exp ASC""", (item['kd_barang'],)).fetchall()
            # Harga jual: ambil dari INDUK (sudah diset admin), fallback ke cabang yg ada
            harga_ref = cur.execute("SELECT harga,disc,DiscMem,DiscLain,stok_min FROM Stok_barang WHERE kd_barang=? AND kd_lokasi='INDUK' AND harga>0 LIMIT 1",
                                    (item['kd_barang'],)).fetchone()
            if not harga_ref:
                harga_ref = cur.execute("SELECT harga,disc,DiscMem,DiscLain,stok_min FROM Stok_barang WHERE kd_barang=? AND kd_lokasi=? AND harga>0 LIMIT 1",
                                        (item['kd_barang'], kd_tuju)).fetchone()
            actual = 0
            for batch in batches:
                if sisa <= 0: break
                take = min(sisa, batch['stok_sisa'])
                # Kurangi dari INDUK
                cur.execute("UPDATE Stok_barang SET stok_sisa=stok_sisa-? WHERE id=?", (take, batch['id']))
                # Tambah ke stok cabang
                existing_cab = cur.execute("SELECT id FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                                           (item['kd_barang'], batch['exp'], kd_tuju)).fetchone()
                if existing_cab:
                    cur.execute("UPDATE Stok_barang SET stok_sisa=stok_sisa+?,stok_awal=stok_awal+?,stok_gudang=stok_gudang+? WHERE id=?",
                                (take, take, take, existing_cab['id']))
                else:
                    harga_jual = harga_ref['harga'] if harga_ref else 0
                    disc      = harga_ref['disc']    if harga_ref else '0'
                    disc_mem  = harga_ref['DiscMem'] if harga_ref else 0
                    disc_lain = harga_ref['DiscLain'] if harga_ref else 0
                    stok_min  = harga_ref['stok_min'] if harga_ref else 0
                    cur.execute("""INSERT INTO Stok_barang (kd_barang,nm_barang,harga_beli,harga_beli_awal,
                                   stok_awal,stok_sisa,stok_gudang,stok_min,harga,disc,satuan,exp,DiscMem,DiscLain,kd_lokasi)
                                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                                (item['kd_barang'], batch['nm_barang'], batch['harga_beli'], batch['harga_beli'],
                                 take, take, take, stok_min, harga_jual, disc, batch['satuan'], batch['exp'],
                                 disc_mem, disc_lain, kd_tuju))
                cur.execute("INSERT INTO Transfer_Stok_Log (no_referensi,tipe_referensi,kd_barang,nm_barang,kd_lokasi_dari,kd_lokasi_ke,qty,exp,created_by) VALUES (?,?,?,?,?,?,?,?,?)",
                            (o['no_order'], 'order_transfer', item['kd_barang'], batch['nm_barang'], 'INDUK', kd_tuju, take, batch['exp'], d.get('by','')))
                sisa -= take
                actual += take

            status_item = 'fulfilled' if actual >= qty_approved else ('partial' if actual > 0 else 'pending')
            if status_item != 'fulfilled': all_done = False
            cur.execute("UPDATE Order_Transfer_Detail SET qty_fulfilled=?,status_item=? WHERE id=?",
                        (item['qty_fulfilled'] + actual, status_item, item['id']))
            fulfill_log_items.append({'kd_barang': item['kd_barang'], 'diminta': qty_kirim,
                                       'terkirim': actual, 'status_item': status_item})

        new_status = 'fulfilled' if all_done else 'partial'
        cur.execute("UPDATE Order_Transfer SET status=?,fulfilled_by=?,fulfilled_at=? WHERE id=?",
                    (new_status, d.get('by',''), str(datetime.now()), oid))
        con.commit()
    finally:
        con.close()
    catat_log('ot_fulfill', 'Order_Transfer', o['no_order'],
              f"Kirim barang order transfer {o['no_order']} — status akhir: {new_status}",
              data_baru={'items': fulfill_log_items, 'status': new_status})
    return jsonify({'ok': True, 'status': new_status})

@app.route('/api/order-transfer/<int:oid>/export')
def export_order_transfer(oid):
    """Export order sebagai JSON file untuk cabang beda gedung."""
    o = query("SELECT o.*,l.nm_lokasi FROM Order_Transfer o LEFT JOIN Lokasi l ON o.kd_lokasi_asal=l.kd_lokasi WHERE o.id=?", (oid,), one=True)
    d = query("SELECT * FROM Order_Transfer_Detail WHERE id_order=?", (oid,))
    if not o:
        return jsonify({'ok': False, 'msg': 'Order tidak ditemukan'}), 404
    payload = {
        'schema_version': '1.0',
        'tipe': 'order_transfer',
        'order': dict(o),
        'detail': rows_to_list(d),
        'exported_at': str(datetime.now())
    }
    from flask import Response
    return Response(
        json.dumps(payload, ensure_ascii=False, indent=2),
        mimetype='application/json',
        headers={'Content-Disposition': f'attachment; filename=OT_{o["no_order"]}.json'}
    )

@app.route('/api/order-transfer/import', methods=['POST'])
def import_order_transfer():
    """Import JSON file order dari cabang beda gedung."""
    try:
        payload = request.json
        if payload.get('schema_version') != '1.0' or payload.get('tipe') != 'order_transfer':
            return jsonify({'ok': False, 'msg': 'Format file tidak valid'}), 400
        o = payload['order']
        # Cek duplikat
        existing = query("SELECT id FROM Order_Transfer WHERE no_order=?", (o['no_order'],), one=True)
        if existing:
            return jsonify({'ok': False, 'msg': f"Order {o['no_order']} sudah ada di sistem"}), 400
        oid = execute("""INSERT INTO Order_Transfer
            (no_order,kd_lokasi_asal,tgl_order,status,catatan_cabang,created_by,is_from_import)
            VALUES (?,?,?,'submitted',?,?,1)""",
            (o['no_order'], o['kd_lokasi_asal'], o.get('tgl_order', str(date.today())),
             o.get('catatan_cabang',''), o.get('created_by','')))
        for item in payload.get('detail', []):
            execute("INSERT INTO Order_Transfer_Detail (id_order,kd_barang,nm_barang,qty_request) VALUES (?,?,?,?)",
                    (oid, item['kd_barang'], item.get('nm_barang',''), item['qty_request']))
        return jsonify({'ok': True, 'id': oid, 'no_order': o['no_order']})
    except Exception as e:
        return jsonify({'ok': False, 'msg': str(e)}), 400

@app.route('/api/monitoring/retur-mencurigakan')
def monitoring_retur_mencurigakan():
    """Deteksi pola kecurangan 'retur bayangan' (sweethearting): kasir meretur
    barang/jasa/racikan, LALU kasir yg SAMA menjual barang yg SAMA lagi di HARI
    YANG SAMA dgn diskon. Retur dipakai utk menutup jejak stok/nominal padahal
    uang dari transaksi asli tidak benar-benar dikembalikan ke pasien — selisihnya
    (nilai retur vs diskon di penjualan ulang) berpotensi dikantongi kasir.
    TIDAK menyimpulkan kecurangan pasti — cuma menandai pola yg PANTAS DIPERIKSA."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor'}), 403

    tgl1 = request.args.get('dari', str(date.today() - timedelta(days=30)))
    tgl2 = request.args.get('sampai', str(date.today()))

    retur_rows = rows_to_list(query("""
        SELECT r.id, r.no_faktur, r.kd_barang, r.nm_barang, r.jenis, r.qty, r.sub_total,
               r.tanggal, r.cashier, p.metode_bayar, p.kd_pasien, ps.nm_pasien
        FROM R_Penjualan r
        JOIN Penjualan p ON r.no_faktur=p.no_faktur
        LEFT JOIN Pasien ps ON p.kd_pasien=ps.kd_pasien
        WHERE r.jenis LIKE 'retur%' AND r.tanggal BETWEEN ? AND ?
        ORDER BY r.tanggal DESC
    """, (tgl1, tgl2)))

    hasil = []
    for r in retur_rows:
        if not r['cashier']:
            continue
        jual_ulang_raw = query("""
            SELECT rp.no_faktur, rp.qty, rp.disc, rp.harga, rp.sub_total, rp.tanggal, p2.metode_bayar
            FROM R_Penjualan rp
            JOIN Penjualan p2 ON rp.no_faktur=p2.no_faktur
            WHERE rp.kd_barang=? AND rp.cashier=? AND rp.tanggal=?
              AND rp.jenis NOT LIKE 'retur%' AND rp.no_faktur != ?
        """, (r['kd_barang'], r['cashier'], r['tanggal'], r['no_faktur']))
        cocok = []
        for j in rows_to_list(jual_ulang_raw):
            try:
                disc_pct = float(j['disc'] or 0)
            except (TypeError, ValueError):
                disc_pct = 0.0
            if disc_pct > 0:
                j['disc'] = disc_pct
                j['potensi_selisih'] = round(float(j['harga']) * float(j['qty']) * disc_pct / 100, 2)
                cocok.append(j)
        if cocok:
            # Sandingkan dgn Log Aktivitas — semua kejadian yg menyentuh faktur retur
            # MAUPUN faktur jual-ulangnya, diurutkan waktu, supaya kelihatan jeda waktu
            # asli antar kejadian (mis. retur cuma berselang detik dari pembayaran —
            # itu sendiri sudah jadi indikasi kuat, terlepas dari nominal selisihnya).
            fakturs = [r['no_faktur']] + [j['no_faktur'] for j in cocok]
            ph = ','.join('?' * len(fakturs))
            log_rows = rows_to_list(query(f"""
                SELECT waktu, nama_user, username, role, aksi, keterangan, ref
                FROM Aktivitas_Log WHERE ref IN ({ph})
                ORDER BY waktu ASC
            """, fakturs))
            hasil.append({
                'retur': r,
                'jual_ulang': cocok,
                'potensi_selisih_total': round(sum(j['potensi_selisih'] for j in cocok), 2),
                'log': log_rows,
            })
    hasil.sort(key=lambda x: -x['potensi_selisih_total'])
    return jsonify({'ok': True, 'items': hasil,
                    'total_kasus': len(hasil),
                    'total_potensi_selisih': round(sum(h['potensi_selisih_total'] for h in hasil), 2)})

@app.route('/api/approval-queue')
def approval_queue():
    """Semua item pending approval untuk supervisor."""
    produksi = query("""
        SELECT p.*, f.nm_formula, f.kd_produk_hasil, f.qty_hasil, f.satuan_hasil,
               (SELECT nm_barang FROM Stok_barang
                WHERE kd_barang=f.kd_produk_hasil AND kd_lokasi='INDUK' LIMIT 1) as nm_produk_hasil
        FROM Produksi p
        JOIN Formula_Produksi f ON p.id_formula=f.id
        WHERE p.status='draft'
        ORDER BY p.created_at ASC
    """)
    orders = query("""
        SELECT o.*, l.nm_lokasi,
               (SELECT COUNT(*) FROM Order_Transfer_Detail WHERE id_order=o.id) as jml_item,
               (SELECT COALESCE(SUM(qty_request),0) FROM Order_Transfer_Detail WHERE id_order=o.id) as total_qty
        FROM Order_Transfer o
        LEFT JOIN Lokasi l ON o.kd_lokasi_asal=l.kd_lokasi
        WHERE o.status='submitted'
        ORDER BY o.created_at ASC
    """)
    koreksi_pembelian = query("""
        SELECT k.*, p.nm_suplier,
               r.nm_barang as item_nm_barang, r.qty as item_qty, r.exp as item_exp
        FROM Pembelian_Koreksi k
        JOIN Pembelian p ON k.no_faktur=p.no_faktur
        LEFT JOIN R_Pembelian r ON k.id_item=r.id
        WHERE k.status='pending'
        ORDER BY k.requested_at ASC
    """)
    pembelian_pending = query("""
        SELECT p.*, (SELECT COUNT(*) FROM R_Pembelian WHERE no_faktur=p.no_faktur) as jml_item
        FROM Pembelian p
        WHERE p.status='submitted'
        ORDER BY p.tanggal ASC
    """)
    retur_pending_raw = query("""
        SELECT rr.*, p.tanggal as tgl_faktur, ps.nm_pasien
        FROM Retur_Request rr
        LEFT JOIN Penjualan p ON rr.no_faktur=p.no_faktur
        LEFT JOIN Pasien ps ON p.kd_pasien=ps.kd_pasien
        WHERE rr.status='pending'
        ORDER BY rr.diajukan_at ASC
    """)
    retur_pending = rows_to_list(retur_pending_raw)
    for r in retur_pending:
        items = json.loads(r['items_json']) if r.get('items_json') else []
        r['jml_item'] = len(items)
        try:
            r['umur_hari'] = (date.today() - datetime.strptime(r['tgl_faktur'], '%Y-%m-%d').date()).days
        except (ValueError, TypeError):
            r['umur_hari'] = None
        del r['items_json']
    return jsonify({'produksi': rows_to_list(produksi), 'orders': rows_to_list(orders),
                     'koreksi_pembelian': rows_to_list(koreksi_pembelian),
                     'pembelian_pending': rows_to_list(pembelian_pending),
                     'retur_pending': retur_pending})

# ─── TERIMA KIRIMAN STOK DARI GEDUNG LAIN (import file _P/_R per faktur) ──────
# Cabang di gedung lain tidak realtime → kirim stok lewat sepasang file teks:
#   <faktur>_P.txt : 1 baris header  -> no_faktur, tgl(M/D/YYYY), col3, total, metode
#   <faktur>_R.txt : n baris rincian -> nm_display(=kd_barang kita), nm_generik, harga, qty, exp
# Kolom-1 file = kd_barang di sistem ini (kd_barang memang teks nama-brand, bukan angka).

def _ki_num(s):
    try:
        return float(str(s).strip().replace(',', '') or 0)
    except (TypeError, ValueError):
        return 0.0

def _ki_tgl(s):
    """Normalkan tgl exp dari file ke format yang DIPAKAI aplikasi ini di Stok_barang.exp
    — yaitu 'M/D/YYYY' TANPA leading zero (mis. '6/30/2028', '1/1/2029').
    PENTING: JANGAN ubah ke 'YYYY-MM-DD' — akan bikin batch dgn tgl exp sama terpecah &
    urutan FEFO (ORDER BY exp ASC = string sort) kacau karena beda format.

    Urutan format yg dicoba: D/M/Y (tanggal/bulan/tahun) DULUAN — sesuai arahan user
    (2026-09-18) bahwa gedung sumber biasanya menulis tanggal duluan. Baru fallback ke
    M/D/Y kalau D/M/Y tidak valid (mis. field ke-2 > 12 sehingga cuma masuk akal sbg
    bulan, bukan tanggal) — supaya file lama yg terbukti M/D/Y (faktur 0909260001G,
    sudah diterima & datanya sudah dipakai di stok) tetap terbaca benar, tidak berubah
    mundur. Field yg AMBIGU (kedua angka <=12) sekarang dianggap D/M/Y sesuai arahan ini."""
    s = (s or '').strip()
    if not s:
        return ''
    for fmt in ('%d/%m/%Y', '%d/%m/%y', '%m/%d/%Y', '%m/%d/%y', '%Y-%m-%d', '%d-%m-%Y'):
        try:
            dt = datetime.strptime(s, fmt)
            return f"{dt.month}/{dt.day}/{dt.year}"
        except ValueError:
            pass
    return s

def _parse_kiriman_files(text_p, text_r):
    """Murni parse sepasang file kiriman — TIDAK menyentuh DB. Raise ValueError kalau
    format tidak valid. Dipakai endpoint upload & bisa dipanggil langsung utk dry-run."""
    p_rows = [r for r in csv.reader((text_p or '').splitlines()) if r and any(c.strip() for c in r)]
    if not p_rows:
        raise ValueError('File _P (header) kosong / tidak terbaca')
    p = (list(p_rows[0]) + [''] * 5)[:5]
    no_faktur  = p[0].strip()
    tgl_faktur = _ki_tgl(p[1])
    disc_pct   = _ki_num(p[2])   # % diskon header (mis. "10" = 10%) — dulu diabaikan
    total_header = _ki_num(p[3])
    metode = p[4].strip()
    if not no_faktur:
        raise ValueError('No faktur di file _P kosong')

    items, total_rincian = [], 0.0
    for r in csv.reader((text_r or '').splitlines()):
        if not r or not any(c.strip() for c in r):
            continue
        r = (list(r) + [''] * 5)[:5]
        kd = r[0].strip()
        if not kd:
            continue
        harga = _ki_num(r[2]); qty = _ki_num(r[3])
        total_rincian += harga * qty
        items.append({'kd_barang': kd, 'nm_file': kd, 'nm_generik': r[1].strip(),
                      'harga': harga, 'qty': qty, 'exp': _ki_tgl(r[4])})
    if not items:
        raise ValueError('File _R (rincian) tidak berisi baris barang')

    # Disc header (%) diterapkan ke total rincian mentah SEBELUM dibandingkan ke total
    # header file — supaya diskon yg memang ada di file tidak dianggap selisih/mismatch.
    total_setelah_disc = total_rincian * (1 - disc_pct / 100)
    return {
        'no_faktur': no_faktur, 'tgl_faktur': tgl_faktur, 'metode': metode,
        'disc_pct':      round(disc_pct, 2),
        'total_header':  round(total_header, 2),
        'total_rincian': round(total_rincian, 2),
        'total_setelah_disc': round(total_setelah_disc, 2),
        'selisih':       round(total_setelah_disc - total_header, 2),
        'items': items,
    }

def _ki_expired(exp_str, today_date=None):
    """Cek apakah tgl exp (format M/D/YYYY hasil _ki_tgl) sudah lewat — pakai PARSING
    tanggal beneran, bukan string compare. exp_str 'M/D/YYYY' vs today ISO 'YYYY-MM-DD'
    beda format, jadi '<' string bisa salah total (mis. '2/28/2029' < '2026-09-18' krn
    '/' < '0' scr ASCII, padahal 2029 masih jauh di depan)."""
    if not exp_str:
        return False
    try:
        exp_dt = datetime.strptime(exp_str, '%m/%d/%Y').date()
    except ValueError:
        return False
    return exp_dt < (today_date or date.today())

def _correlate_kiriman_item(kd_barang):
    """'ok' = ada di master Barang · 'unregistered' = tak di master tapi ada di Stok_barang
    · 'notfound' = tak dikenal sama sekali. LOWER() di kedua sisi — kd_barang di MySQL
    collation utf8mb4_bin (case-sensitive), tapi kode barang dari file kiriman gedung lain
    tidak boleh beda hasil cuma krn beda huruf besar/kecil dari master (user 2026-10-06:
    "tidak ada perbedaan disini antara besar dan kecil")."""
    if query("SELECT 1 FROM Barang WHERE LOWER(kd_barang)=LOWER(?) LIMIT 1", (kd_barang,), one=True):
        return 'ok'
    if query("SELECT 1 FROM Stok_barang WHERE LOWER(kd_barang)=LOWER(?) LIMIT 1", (kd_barang,), one=True):
        return 'unregistered'
    return 'notfound'

def _tk_item_sukses(it):
    """True kalau baris Terima_Kiriman_Detail ini BENAR-BENAR sudah diposting ke stok saat
    konfirmasi (mis logic persis di terima_kiriman_konfirmasi()) — dipakai utk mendeteksi
    barang yg TERLEWAT (gagal masuk stok walau faktur sudah 'diterima'), supaya bisa
    diproses ulang tanpa menyentuh/dobel yg sudah berhasil."""
    if not it.get('diterima'):
        return False
    if it.get('match_status') != 'notfound':
        return True
    return bool((it.get('kd_barang_final') or '').strip())

def _tk_root_id(hdr):
    """id leluhur paling atas (kiriman ASLI) — kalau hdr sendiri sudah berupa susulan,
    induk_id-nya SELALU sudah merujuk leluhur plg atas (lihat terima_kiriman_proses_ulang),
    jadi tinggal pakai itu, bukan hdr['id'] sendiri."""
    return hdr['induk_id'] or hdr['id']

def _tk_family_ids(root_id):
    """id leluhur + SEMUA susulannya (anak langsung — app ini sengaja cuma 1 level nesting,
    lihat _tk_root_id)."""
    anak = query("SELECT id FROM Terima_Kiriman WHERE induk_id=?", (root_id,))
    return [root_id] + [r['id'] for r in anak]

def _tk_belum_sukses_items(root_id):
    """Daftar baris Terima_Kiriman_Detail (dari record ASLI/leluhur) yg nm_file-nya BELUM
    berhasil diproses di MANA PUN dalam 1 keluarga (leluhur + semua susulannya) — bukan
    cuma dicek di 1 record doang. Ini yg bikin tombol "Proses Barang Terlewat" TIDAK
    muncul lagi kalau susulan sebelumnya sudah berhasil menyelesaikannya (ditemukan
    2026-10-01: tanpa pengecekan lintas-keluarga ini, user bisa bikin susulan baru yg
    isinya PERSIS sama dgn susulan yg sudah sukses sebelumnya — kelihatan spt "dobel")."""
    family_ids = _tk_family_ids(root_id)
    ph = ','.join('?' * len(family_ids))
    all_items = rows_to_list(query(
        f"SELECT nm_file, diterima, match_status, kd_barang_final FROM Terima_Kiriman_Detail WHERE id_kiriman IN ({ph})",
        family_ids))
    sukses_nmfile = {it['nm_file'] for it in all_items if _tk_item_sukses(it)}
    root_items = rows_to_list(query(
        "SELECT * FROM Terima_Kiriman_Detail WHERE id_kiriman=? ORDER BY id", (root_id,)))
    return [it for it in root_items if it['nm_file'] not in sukses_nmfile]

def _stok_cabang_masuk(cur, kd_barang, nm_barang, exp, qty, harga_beli, kd_lokasi,
                       no_ref, tipe_ref, by, kd_lokasi_dari='LUAR'):
    """Stok masuk ke `kd_lokasi` (upsert batch kd_barang+exp) + RiwayatBarang + Transfer_Stok_Log.
    Meniru pola stock-in di fulfill_order_transfer(). Dipakai di dalam transaksi (cur aktif).

    CATATAN HARGA (dikonfirmasi user 2026-10-06): kolom `harga` di file kiriman (param
    `harga_beli` di sini) itu SUDAH harga jual cabang — bukan harga pokok/beli terpisah,
    krn harga dari induk/kiriman = harga jual juga di cabang. Jadi kalau barang ini BENAR²
    baru (blm pernah ada referensi harga di lokasi manapun), harga jual (`Stok_barang.harga`)
    ikut diisi dari situ jg — SEBELUMNYA dibiarkan 0 (nunggu referensi yg tak akan pernah
    ketemu utk barang baru), bikin barang bisa "gratis" tak sengaja kalau langsung dijual
    tanpa dibetulkan manual dulu (ketemu dari kasus nyata faktur 2909260003G)."""
    today = str(date.today())
    row = cur.execute("SELECT id, stok_sisa, harga FROM Stok_barang WHERE kd_barang=? AND exp=? AND kd_lokasi=?",
                      (kd_barang, exp, kd_lokasi)).fetchone()
    if row:
        # Kalau batch ini sendiri kebetulan masih harga=0 dari bug lama di atas, ikut
        # dibetulkan skrg saat topup qty — drpd dibiarkan 0 selamanya sampai ada yg sadar.
        harga_jual = harga_beli if not (row['harga'] or 0) else row['harga']
        cur.execute("""UPDATE Stok_barang SET stok_sisa=stok_sisa+?, stok_awal=stok_awal+?,
                       stok_gudang=stok_gudang+?, harga_beli=?, harga=? WHERE id=?""",
                    (qty, qty, qty, harga_beli, harga_jual, row['id']))
        saldo = (row['stok_sisa'] or 0) + qty
    else:
        hr = cur.execute("""SELECT harga,disc,DiscMem,DiscLain,stok_min,satuan FROM Stok_barang
                            WHERE kd_barang=? AND harga>0
                            ORDER BY CASE WHEN kd_lokasi='INDUK' THEN 0 ELSE 1 END LIMIT 1""",
                         (kd_barang,)).fetchone()
        cur.execute("""INSERT INTO Stok_barang (kd_barang,nm_barang,harga_beli,harga_beli_awal,
                       stok_awal,stok_sisa,stok_gudang,stok_min,harga,disc,satuan,exp,DiscMem,DiscLain,kd_lokasi)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (kd_barang, nm_barang, harga_beli, harga_beli, qty, qty, qty,
                     (hr['stok_min'] if hr else 0), ((hr['harga'] if hr else 0) or harga_beli),
                     (hr['disc'] if hr else '0'), (hr['satuan'] if hr else ''), exp,
                     (hr['DiscMem'] if hr else 0), (hr['DiscLain'] if hr else 0), kd_lokasi))
        saldo = qty
    cur.execute("""INSERT INTO RiwayatBarang (tanggal,kd_barang,exp,keterangan,nilai,saldo_e,saldo_g)
                   VALUES (?,?,?,?,?,?,?)""",
                (today, kd_barang, exp, f"Terima Kiriman {no_ref}", qty, saldo, saldo))
    cur.execute("""INSERT INTO Transfer_Stok_Log (no_referensi,tipe_referensi,kd_barang,nm_barang,
                   kd_lokasi_dari,kd_lokasi_ke,qty,exp,created_by) VALUES (?,?,?,?,?,?,?,?,?)""",
                (no_ref, tipe_ref, kd_barang, nm_barang, kd_lokasi_dari, kd_lokasi, qty, exp, by))


@app.route('/api/terima-kiriman/upload', methods=['POST'])
def terima_kiriman_upload():
    """Terima file_p + file_r, parse & korelasi, buat staging Terima_Kiriman (pending_review).
    BELUM menyentuh stok."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor', 'kasir'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor/kasir'}), 403
    fp = request.files.get('file_p')
    fr = request.files.get('file_r')
    if not fp or not fr:
        return jsonify({'ok': False, 'msg': 'Wajib unggah kedua file: _P (header) dan _R (rincian)'}), 400
    try:
        text_p = fp.read().decode('utf-8-sig', errors='replace')
        text_r = fr.read().decode('utf-8-sig', errors='replace')
        parsed = _parse_kiriman_files(text_p, text_r)
    except ValueError as e:
        return jsonify({'ok': False, 'msg': str(e)}), 400
    except Exception as e:
        return jsonify({'ok': False, 'msg': f'Gagal membaca file: {e}'}), 400

    no_faktur = parsed['no_faktur']
    prev = query("SELECT id, status FROM Terima_Kiriman WHERE no_faktur=?", (no_faktur,), one=True)
    if prev and prev['status'] == 'diterima':
        return jsonify({'ok': False, 'msg': f'Faktur {no_faktur} sudah pernah diterima — tidak bisa import ulang'}), 409
    if prev and prev['status'] == 'pending_review':
        # buang staging lama, ganti dgn upload terbaru
        execute("DELETE FROM Terima_Kiriman_Detail WHERE id_kiriman=?", (prev['id'],))
        execute("DELETE FROM Terima_Kiriman WHERE id=?", (prev['id'],))

    today = str(date.today())
    sumber = 'Induk Utama'
    kid = execute("""INSERT INTO Terima_Kiriman
        (no_faktur,tgl_faktur,kd_lokasi_tujuan,total_header,total_rincian,disc_pct,
         total_setelah_disc,selisih,metode,sumber,status,dibuat_oleh,created_at)
        VALUES (?,?,?,?,?,?,?,?,?,?,'pending_review',?,?)""",
        (no_faktur, parsed['tgl_faktur'] or today, 'CBG1', parsed['total_header'],
         parsed['total_rincian'], parsed['disc_pct'], parsed['total_setelah_disc'],
         parsed['selisih'], parsed['metode'], sumber,
         (user.get('nama') if user else ''), str(datetime.now())))

    n_ok = n_unreg = n_notfound = n_exp = 0
    for it in parsed['items']:
        ms = _correlate_kiriman_item(it['kd_barang'])
        expired = 1 if _ki_expired(it['exp']) else 0
        if ms == 'ok': n_ok += 1
        elif ms == 'unregistered': n_unreg += 1
        else: n_notfound += 1
        if expired: n_exp += 1
        execute("""INSERT INTO Terima_Kiriman_Detail
            (id_kiriman,kd_barang,nm_file,nm_generik,harga,qty,exp,match_status,expired,diterima)
            VALUES (?,?,?,?,?,?,?,?,?,1)""",
            (kid, it['kd_barang'], it['nm_file'], it['nm_generik'], it['harga'],
             it['qty'], it['exp'], ms, expired))

    return jsonify({'ok': True, 'id': kid, 'no_faktur': no_faktur,
                    'ringkasan': {'total': len(parsed['items']), 'ok': n_ok,
                                  'unregistered': n_unreg, 'notfound': n_notfound,
                                  'expired': n_exp, 'disc_pct': parsed['disc_pct'],
                                  'total_setelah_disc': parsed['total_setelah_disc'],
                                  'selisih': parsed['selisih']}})


@app.route('/api/terima-kiriman')
def terima_kiriman_list():
    """List kiriman — record 'susulan' (induk_id terisi, lihat terima_kiriman_proses_ulang)
    DIKELUARKAN dari baris utama & dilampirkan sbg sub-list (`susulan`) di bawah induknya,
    supaya tidak kelihatan spt "dobel" entri utk 1 kiriman yg sama (keluhan user
    2026-10-01) — tetap bisa diakses penuh via GET /api/terima-kiriman/<id>."""
    status = request.args.get('status', '')
    where, params = [], []
    if status:
        where.append("status=?"); params.append(status)
    w = ('WHERE ' + ' AND '.join(where)) if where else ''
    rows = rows_to_list(query(f"""SELECT tk.*,
                        (SELECT COUNT(*) FROM Terima_Kiriman_Detail d WHERE d.id_kiriman=tk.id) as jml_item
                     FROM Terima_Kiriman tk {w} ORDER BY tk.id DESC LIMIT 200""", params))

    induk_rows  = [r for r in rows if not r.get('induk_id')]
    induk_ids   = {r['id'] for r in induk_rows}
    susulan_map = {}
    for r in rows:
        if r.get('induk_id'):
            susulan_map.setdefault(r['induk_id'], []).append(r)
    for r in induk_rows:
        r['susulan'] = susulan_map.get(r['id'], [])
    # Jaga-jaga: kalau filter ?status= bikin induknya kefilter tapi susulannya lolos
    # (mis. induk sudah 'diterima' tapi susulan masih 'pending_review') — jangan
    # sampai susulan itu hilang senyap, tampilkan sbg baris berdiri sendiri.
    for induk_id, subs in susulan_map.items():
        if induk_id not in induk_ids:
            induk_rows.extend(subs)
    induk_rows.sort(key=lambda r: r['id'], reverse=True)
    return jsonify(induk_rows)


@app.route('/api/terima-kiriman/<int:kid>')
def terima_kiriman_detail(kid):
    hdr = query("SELECT * FROM Terima_Kiriman WHERE id=?", (kid,), one=True)
    if not hdr:
        return jsonify(None)
    res = dict(hdr)
    res['items'] = rows_to_list(query(
        "SELECT * FROM Terima_Kiriman_Detail WHERE id_kiriman=? ORDER BY id", (kid,)))
    # Dihitung LINTAS KELUARGA (leluhur + semua susulannya, lihat _tk_belum_sukses_items)
    # — bukan cuma dari baris record ini sendiri — supaya kalau susulan SEBELUMNYA sudah
    # berhasil menuntaskan barang yg terlewat, tombol "Proses Ulang" tidak muncul lagi.
    res['jml_belum_sukses'] = len(_tk_belum_sukses_items(_tk_root_id(hdr))) \
        if hdr['status'] == 'diterima' else 0
    return jsonify(res)


@app.route('/api/terima-kiriman/<int:kid>/review', methods=['POST'])
def terima_kiriman_review(kid):
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor', 'kasir'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor/kasir'}), 403
    hdr = query("SELECT status FROM Terima_Kiriman WHERE id=?", (kid,), one=True)
    if not hdr:
        return jsonify({'ok': False, 'msg': 'Data tidak ditemukan'}), 404
    if hdr['status'] != 'pending_review':
        return jsonify({'ok': False, 'msg': f"Sudah {hdr['status']} — tidak bisa diubah"}), 400
    d = request.json or {}
    for it in d.get('items', []):
        execute("""UPDATE Terima_Kiriman_Detail
                   SET diterima=?, kd_barang_final=?, catatan_review=?
                   WHERE id=? AND id_kiriman=?""",
                (1 if it.get('diterima') else 0, (it.get('kd_barang_final') or '').strip(),
                 (it.get('catatan_review') or '').strip(), it.get('id'), kid))
    if 'catatan' in d:
        execute("UPDATE Terima_Kiriman SET catatan=? WHERE id=?", (d.get('catatan', ''), kid))
    return jsonify({'ok': True})


@app.route('/api/terima-kiriman/<int:kid>/batal', methods=['POST'])
def terima_kiriman_batal(kid):
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor', 'kasir'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor/kasir'}), 403
    hdr = query("SELECT no_faktur, status FROM Terima_Kiriman WHERE id=?", (kid,), one=True)
    if not hdr:
        return jsonify({'ok': False, 'msg': 'Tidak ditemukan'}), 404
    if hdr['status'] != 'pending_review':
        return jsonify({'ok': False, 'msg': 'Hanya import berstatus pending review yang bisa dibatalkan'}), 400
    execute("DELETE FROM Terima_Kiriman_Detail WHERE id_kiriman=?", (kid,))
    execute("DELETE FROM Terima_Kiriman WHERE id=?", (kid,))
    catat_log('terima_kiriman_batal', 'Terima_Kiriman', hdr['no_faktur'],
              f"Batalkan import kiriman {hdr['no_faktur']} (belum diterima)")
    return jsonify({'ok': True})


@app.route('/api/terima-kiriman/<int:kid>/konfirmasi', methods=['POST'])
def terima_kiriman_konfirmasi(kid):
    """Eksekusi: stok masuk ke CBG1 utk tiap detail diterima=1. Item 'unregistered' yg
    disetujui → didaftarkan ke master Barang. Item 'notfound' wajib punya kd_barang_final."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor', 'kasir'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor/kasir'}), 403
    hdr = query("SELECT * FROM Terima_Kiriman WHERE id=?", (kid,), one=True)
    if not hdr:
        return jsonify({'ok': False, 'msg': 'Tidak ditemukan'}), 404
    if hdr['status'] != 'pending_review':
        return jsonify({'ok': False, 'msg': f"Sudah {hdr['status']}"}), 400
    dets = query("SELECT * FROM Terima_Kiriman_Detail WHERE id_kiriman=?", (kid,))
    kd_lokasi = hdr['kd_lokasi_tujuan'] or 'CBG1'
    by = user.get('nama') if user else ''
    diproses, dilewati = [], []

    con = get_db()
    try:
        cur = con.cursor()
        for it in dets:
            if not it['diterima']:
                dilewati.append({'kd_barang': it['kd_barang'], 'alasan': 'tidak dicentang'})
                continue
            kd = it['kd_barang']
            ms = it['match_status']
            # Semua pencocokan kd_barang di sini case-INSENSITIVE (LOWER() di query) —
            # tapi hasil AKHIR `kd` yg dipakai nulis ke Stok_barang/Barang SELALU disamakan
            # ke casing yg SUDAH tersimpan di master, bukan casing mentah dari file/input
            # kiriman — kalau tidak, _stok_cabang_masuk() (yg exact-match krn kd_barang
            # collation utf8mb4_bin) bisa bikin BATCH BARU terpisah hanya krn beda huruf
            # besar/kecil, padahal barangnya sama (user 2026-10-06: "tidak ada perbedaan
            # disini antara besar dan kecil").
            if ms == 'notfound':
                kd_fin = (it['kd_barang_final'] or '').strip()
                if not kd_fin:
                    dilewati.append({'kd_barang': kd, 'alasan': 'tak dikenal & belum dipetakan'})
                    continue
                kd = kd_fin
                # kd_barang_final BISA = kd aslinya sendiri (berarti admin memutuskan ini
                # genuinely barang BARU, bukan typo/salah nama — lihat tombol "✓ Daftarkan
                # sbg Barang Baru" di frontend) — kalau itu kasusnya & belum ada di master,
                # daftarkan jg spt jalur 'unregistered' di bawah (dulu HANYA 'unregistered'
                # yg otomatis terdaftar, 'notfound' yg dipetakan ke kode sendiri tidak —
                # ketemu 2026-10-01 pas kasus nyata 'Lotion Uriage 200ml'/'Cream Uriage PSO').
                ada = cur.execute("SELECT kd_barang FROM Barang WHERE LOWER(kd_barang)=LOWER(?)", (kd,)).fetchone()
                if ada:
                    kd = ada['kd_barang']
                else:
                    nm_master = (it['catatan_review'] or '').strip() or it['nm_generik'] or it['nm_file'] or kd
                    cur.execute("""INSERT INTO Barang (kd_barang,nm_barang,nm_generik,harga_jual,aktif)
                                   VALUES (?,?,?,?,1)""",
                                (kd, nm_master, it['nm_generik'] or '', it['harga'] or 0))
            elif ms == 'unregistered':
                # 'unregistered' artinya PASTI sudah ada di Stok_barang (itu definisinya) —
                # samakan dulu ke casing yg sudah ada di sana sblm cek/daftarkan ke Barang.
                stok_row = cur.execute("SELECT kd_barang FROM Stok_barang WHERE LOWER(kd_barang)=LOWER(?) LIMIT 1",
                                       (kd,)).fetchone()
                if stok_row:
                    kd = stok_row['kd_barang']
                ada = cur.execute("SELECT 1 FROM Barang WHERE LOWER(kd_barang)=LOWER(?)", (kd,)).fetchone()
                if not ada:
                    nm_master = (it['catatan_review'] or '').strip() or it['nm_generik'] or it['nm_file'] or kd
                    cur.execute("""INSERT INTO Barang (kd_barang,nm_barang,nm_generik,harga_jual,aktif)
                                   VALUES (?,?,?,?,1)""",
                                (kd, nm_master, it['nm_generik'] or '', it['harga'] or 0))
            else:
                # 'ok' — sudah ada di master Barang, samakan casing-nya jg.
                ok_row = cur.execute("SELECT kd_barang FROM Barang WHERE LOWER(kd_barang)=LOWER(?) LIMIT 1",
                                    (kd,)).fetchone()
                if ok_row:
                    kd = ok_row['kd_barang']
            nm_barang = it['nm_generik'] or it['nm_file'] or kd
            _stok_cabang_masuk(cur, kd, nm_barang, it['exp'] or '', it['qty'] or 0,
                               it['harga'] or 0, kd_lokasi, hdr['no_faktur'],
                               'terima_kiriman', by)
            diproses.append({'kd_barang': kd, 'qty': it['qty'], 'exp': it['exp']})
        cur.execute("UPDATE Terima_Kiriman SET status='diterima', diterima_at=? WHERE id=?",
                    (str(datetime.now()), kid))
        con.commit()
    except Exception as e:
        con.rollback()
        return jsonify({'ok': False, 'msg': f'Gagal: {e}'}), 500
    finally:
        con.close()

    catat_log('terima_kiriman', 'Terima_Kiriman', hdr['no_faktur'],
              f"Terima kiriman stok {hdr['no_faktur']} ke {kd_lokasi} — {len(diproses)} barang masuk"
              + (f", {len(dilewati)} dilewati" if dilewati else ''),
              data_baru={'no_faktur': hdr['no_faktur'], 'kd_lokasi': kd_lokasi,
                         'diproses': diproses, 'dilewati': dilewati,
                         'selisih_total': hdr['selisih']})
    return jsonify({'ok': True, 'diproses': len(diproses), 'dilewati': dilewati})

@app.route('/api/terima-kiriman/<int:kid>/proses-ulang', methods=['POST'])
def terima_kiriman_proses_ulang(kid):
    """Barang yg TERLEWAT (diceklis tapi gagal/tak dipetakan) saat konfirmasi pertama —
    biar bisa diproses ulang TANPA perlu upload file lagi & TANPA menyentuh/menduplikasi
    barang yg sudah berhasil. Dipakai utk kasus spt faktur 2909260003G (2 barang baru
    'tak dikenal' terlewat krn belum dipetakan sebelum konfirmasi).

    Caranya: bikin record Terima_Kiriman BARU (child, no_faktur diberi akhiran '-R<jam>'
    krn kolom no_faktur UNIQUE) isinya HANYA baris2 yg belum sukses dari record asli —
    lewat alur review+konfirmasi yg SAMA PERSIS spt import biasa (reuse 100%, bukan
    logic baru). no_referensi yg ditulis ke Transfer_Stok_Log/RiwayatBarang nanti tetap
    no_faktur ASLI (bukan yg ber-akhiran), krn itu field tersendiri milik tiap baris hasil
    _stok_cabang_masuk — lihat terima_kiriman_konfirmasi()."""
    user = get_current_user()
    role = user.get('role', '') if user else ''
    if role not in ('admin', 'supervisor', 'kasir'):
        return jsonify({'ok': False, 'msg': 'Hanya admin/supervisor/kasir'}), 403
    hdr = query("SELECT * FROM Terima_Kiriman WHERE id=?", (kid,), one=True)
    if not hdr:
        return jsonify({'ok': False, 'msg': 'Tidak ditemukan'}), 404
    if hdr['status'] != 'diterima':
        return jsonify({'ok': False, 'msg': 'Cuma kiriman yg sudah diterima yg bisa diproses ulang'}), 400

    # Dihitung LINTAS KELUARGA (bukan cuma baris record kid ini sendiri) — supaya kalau
    # susulan SEBELUMNYA (dibuat dari record lain dlm keluarga yg sama) sudah berhasil
    # menuntaskan sebagian/semua barang terlewat, tombol ini tidak bikin susulan BARU
    # yg isinya persis sama & bikin "dobel" percobaan utk barang yg sama (ditemukan
    # 2026-10-01 — lihat _tk_belum_sukses_items).
    root_id = _tk_root_id(hdr)
    root_hdr = query("SELECT * FROM Terima_Kiriman WHERE id=?", (root_id,), one=True)
    belum_sukses = _tk_belum_sukses_items(root_id)
    if not belum_sukses:
        return jsonify({'ok': False, 'msg': 'Semua barang di faktur ini sudah berhasil diproses — tidak ada yang perlu diproses ulang'}), 400

    total_partial = sum((it['harga'] or 0) * (it['qty'] or 0) for it in belum_sukses)
    no_faktur_child = f"{root_hdr['no_faktur']}-R{datetime.now().strftime('%H%M%S')}"
    new_kid = execute("""INSERT INTO Terima_Kiriman
        (no_faktur,tgl_faktur,kd_lokasi_tujuan,total_header,total_rincian,disc_pct,
         total_setelah_disc,selisih,metode,sumber,status,catatan,dibuat_oleh,created_at,induk_id)
        VALUES (?,?,?,?,?,0,?,0,?,?,'pending_review',?,?,?,?)""",
        (no_faktur_child, root_hdr['tgl_faktur'], root_hdr['kd_lokasi_tujuan'], total_partial, total_partial,
         total_partial, root_hdr['metode'], root_hdr['sumber'],
         f"Susulan dari {root_hdr['no_faktur']} — proses ulang {len(belum_sukses)} barang yg terlewat sebelumnya",
         (user.get('nama') if user else ''), str(datetime.now()), root_id))
    for it in belum_sukses:
        execute("""INSERT INTO Terima_Kiriman_Detail
            (id_kiriman,kd_barang,nm_file,nm_generik,harga,qty,exp,match_status,expired,diterima,kd_barang_final,catatan_review)
            VALUES (?,?,?,?,?,?,?,?,?,1,?,?)""",
            (new_kid, it['kd_barang'], it['nm_file'], it['nm_generik'], it['harga'], it['qty'],
             it['exp'], it['match_status'], it['expired'], it['kd_barang_final'] or '', it['catatan_review'] or ''))

    catat_log('terima_kiriman_susulan', 'Terima_Kiriman', root_hdr['no_faktur'],
              f"Buat susulan {no_faktur_child} utk {len(belum_sukses)} barang terlewat dari {root_hdr['no_faktur']}",
              data_baru={'no_faktur_asli': root_hdr['no_faktur'], 'no_faktur_susulan': no_faktur_child,
                         'jml_item': len(belum_sukses)})
    return jsonify({'ok': True, 'id': new_kid, 'no_faktur': no_faktur_child, 'jml_item': len(belum_sukses)})


# ─── MIGRASI DATA (folder sementara) ─────────────────────────
SEMENTARA_DIR = os.path.join(_base_dir(), 'sementara')
os.makedirs(SEMENTARA_DIR, exist_ok=True)

def _migr_pd(s):
    """Parse tanggal ke YYYY-MM-DD, toleran format M/D/YYYY atau D/M/YYYY."""
    s = str(s).strip()
    if not s or s in ('None', ''):
        return ''
    if len(s) >= 10 and s[4] == '-':
        return s[:10]
    for fmt in ('%m/%d/%Y %H:%M:%S', '%d/%m/%Y %H:%M:%S',
                '%m/%d/%Y', '%d/%m/%Y', '%Y/%m/%d', '%d-%m-%Y'):
        try:
            return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except Exception:
            pass
    return s[:10]

def _migr_fv(s):
    try:
        return float(str(s).replace(',', '.').strip())
    except Exception:
        return 0.0

def _migr_sv(s):
    return str(s).strip() if s is not None else ''

def _migr_satuan_aman(s, default=''):
    """Validasi kolom 'satuan' saat migrasi CSV — HARUS teks satuan (mis. 'gr',
    'pcs', 'ml'), BUKAN angka. Ditemukan 2026-10-01: 15 kd_barang di data lama
    (hasil migrasi Access via migrate_access.py, skrip SATU KALI dulu yg sudah
    tidak dipakai lagi) ternyata kolom satuan-nya kebetulan berisi angka harga
    (CSV sumber-nya sudah salah sejak di sistem lama, tanpa validasi ikut
    terbawa). Supaya migrasi CSV BERIKUTNYA (tool /api/migrasi_*/eksekusi yg
    MASIH AKTIF dipakai) tidak mengulang pola yg sama: kalau nilainya cuma
    angka murni (tak ada huruf satuan sama sekali), jangan dipercaya — pakai
    default kosong drpd ikut menyimpan angka sbg 'satuan'."""
    v = _migr_sv(s)
    if not v:
        return default
    try:
        float(v)
        return default   # angka murni = bukan teks satuan yg valid, buang
    except ValueError:
        return v

def _migr_disc_aman(s, default='0'):
    """Validasi kolom 'disc' (persentase diskon, dipakai sbg (1-disc/100) di
    seluruh rumus harga POS/racikan/retur) saat migrasi CSV — HARUS 0-100.
    Ditemukan 2026-10-01 bareng bug satuan di atas: data lama ada yg disc-nya
    >100% (mis. '5400', '7500') — kalau barang itu SEMPAT laku tanpa dikoreksi
    dulu, rumus harga*(1-disc/100) bisa menghasilkan harga NEGATIF. Klem ke
    rentang 0-100 drpd percaya mentah2 dari CSV."""
    v = _migr_fv(s)
    if v < 0 or v > 100:
        return default
    # Pertahankan representasi teks aslinya kalau valid (kolom disc TEXT, bukan
    # REAL) — supaya format spt '12.5' tidak berubah jadi '12.5' lewat float lagi
    # secara tak perlu; cukup validasi rentangnya, bukan reformat nilainya.
    sv = _migr_sv(s)
    return sv if sv else default

_MIGRASI_KNOWN_COLS = {
    'no_faktur','tanggal','total_harga','hpp','dokter','kd_pasien','kasir',
    'metode_bayar','status','disc_trx','disc_trx_nilai','kd_lokasi',
    'kd_barang','nm_barang','harga','harga_asli','qty','disc','sub_total',
    'total_harbeli','kd_karyawan','cashier','exp','jenis','racikan_ref',
}
# Sumber lawas (VB6/Access) tanpa baris header, urutan kolom TETAP per file — file
# ASLINYA TIDAK PERNAH diubah; aplikasi yg beradaptasi baca posisinya di sini saja.
# 'dup_col' opsional: index kolom yg mau disalin jadi kolom virtual tambahan
# (kd_barang di R_penjualan.csv jadi rangkap nm_barang, krn di app ini kd_barang=nm_barang).
_LEGACY_LAYOUTS = {
    'Penjualan.csv': {
        'ncol': 7,
        'header': ["no_faktur","tanggal","total_harga","hpp","metode_bayar","kolom6_abaikan","kd_pasien"],
        'dup_col': None,
    },
    'R_penjualan.csv': {
        'ncol': 14,
        'header': ["no_faktur","petugas_abaikan","tanggal","kd_barang","nm_generik_abaikan",
                   "qty","harga","disc","sub_total","kd_karyawan","cashier","harga_asli",
                   "total_harbeli","exp","nm_barang"],
        'dup_col': 3,   # kd_barang (posisi ke-4, index 3) -> disalin jadi kolom virtual nm_barang
    },
}

def _migrasi_cek_header(hdrs, mapped):
    """Kalau file BERISI kolom (row-1 terbaca) tapi TIDAK SATU PUN cocok ke skema tujuan,
    kemungkinan besar file ini tidak punya baris header sama sekali (row pertama malah
    data asli, ke-anggap header) — pola yg sama persis dgn kasus Penjualan.csv/
    R_penjualan.csv. Dipakai di SEMUA tool migrasi (bukan cuma Penjualan) supaya kalau
    suatu saat file lawas serupa diupload utk Pembelian/Pasien/Stok/dst, yg muncul
    peringatan JELAS — bukan '0 kolom cocok' tanpa penjelasan spt sebelumnya."""
    if hdrs and not mapped:
        return ("File ini kemungkinan TIDAK PUNYA baris header (baris pertama berisi "
                "data asli, bukan nama kolom) — makanya 0 kolom cocok. JANGAN eksekusi "
                "dulu. Kalau ini memang format lawas tanpa header, sampaikan ke "
                "admin/AI supaya urutan kolomnya dipetakan manual dulu (spt yg sudah "
                "dilakukan utk Penjualan.csv/R_penjualan.csv).")
    return None


def _load_sementara_csv(filename):
    path = os.path.join(SEMENTARA_DIR, filename)
    if not os.path.exists(path):
        return None, None, f"File {filename} tidak ditemukan di folder sementara"
    with open(path, encoding='utf-8-sig', errors='replace') as f:
        reader = csv.reader(f)
        rows = list(reader)
    if not rows:
        return [], [], None

    first_row_vals = {c.strip().lower() for c in rows[0]}
    if not (first_row_vals & _MIGRASI_KNOWN_COLS):
        # Baris-1 TIDAK mengandung nama kolom yg dikenal -> file export lawas tanpa
        # header (VB6/Access). Kalau jumlah kolomnya cocok dgn layout yg sudah
        # dikonfirmasi, pakai header sintetis di memori — SEMUA baris (termasuk
        # baris-1) dianggap data, file di disk tidak disentuh sama sekali.
        legacy = _LEGACY_LAYOUTS.get(filename)
        if legacy and len(rows[0]) == legacy['ncol']:
            data_rows = rows
            if legacy['dup_col'] is not None:
                idx = legacy['dup_col']
                data_rows = [list(r) + [r[idx]] for r in rows]
            return legacy['header'], data_rows, None

    headers = [h.strip() for h in rows[0]]
    return headers, rows[1:], None

import csv as _csv_mod
import csv

@app.route('/api/migrasi/upload', methods=['POST'])
def migrasi_upload():
    fp = request.files.get('file_penjualan')
    fr = request.files.get('file_r_penjualan')
    if not fp and not fr:
        return jsonify({'ok': False, 'msg': 'Tidak ada file diunggah'}), 400
    if fp:
        fp.save(os.path.join(SEMENTARA_DIR, 'Penjualan.csv'))
    if fr:
        fr.save(os.path.join(SEMENTARA_DIR, 'R_penjualan.csv'))
    return jsonify({'ok': True})


@app.route('/api/migrasi/status')
def migrasi_status():
    def _st(fn):
        path = os.path.join(SEMENTARA_DIR, fn)
        if os.path.exists(path):
            return {'exists': True, 'size': os.path.getsize(path)}
        return {'exists': False}
    return jsonify({'penjualan': _st('Penjualan.csv'), 'r_penjualan': _st('R_penjualan.csv')})


@app.route('/api/migrasi/info')
def migrasi_info():
    """Hitung pemetaan kolom CSV Penjualan & R_Penjualan → kolom DB."""
    result = {}

    # ── Penjualan ────────────────────────────────────────────
    hdrs_p, rows_p, err_p = _load_sementara_csv('Penjualan.csv')
    if err_p:
        result['penjualan'] = {'error': err_p}
    else:
        row_count_p = len(rows_p or [])
        db_cols_p = ['no_faktur','tanggal','total_harga','hpp','dokter',
                     'kd_pasien','kasir','metode_bayar','status',
                     'disc_trx','disc_trx_nilai','kd_lokasi']
        csv_lower_p = [h.lower() for h in (hdrs_p or [])]
        db_lower_p  = [c.lower() for c in db_cols_p]
        mapped_p    = [(h, c) for h in hdrs_p
                       for c in db_cols_p if h.lower() == c.lower()]
        skipped_p   = [h for h in hdrs_p
                       if h.lower() not in db_lower_p]
        missing_p   = [c for c in db_cols_p
                       if c.lower() not in csv_lower_p]
        result['penjualan'] = {
            'file': 'Penjualan.csv',
            'row_count': row_count_p,
            'csv_columns': hdrs_p or [],
            'db_columns': db_cols_p,
            'mapped': [{'csv': a, 'db': b} for a, b in mapped_p],
            'skipped_csv': skipped_p,
            'missing_db': missing_p,
        }

    # ── R_Penjualan ──────────────────────────────────────────
    hdrs_r, rows_r, err_r = _load_sementara_csv('R_penjualan.csv')
    if err_r:
        result['r_penjualan'] = {'error': err_r}
        lokasi_list = rows_to_list(query("SELECT kd_lokasi, nm_lokasi FROM Lokasi WHERE aktif=1 ORDER BY nm_lokasi"))
        result['lokasi_options'] = lokasi_list
        return jsonify(result)
    row_count_r = len(rows_r or [])
    db_cols_r = ['no_faktur','tanggal','kd_barang','nm_barang','harga',
                 'harga_asli','qty','disc','sub_total','total_harbeli',
                 'kd_karyawan','cashier','exp','jenis','racikan_ref']
    mapped_r = []
    for h in (hdrs_r or []):
        for c in db_cols_r:
            if h.lower() == c.lower():
                mapped_r.append((h, c))
                break
    skipped_r  = [h for h in (hdrs_r or []) if not any(h == a for a, _ in mapped_r)]
    missing_r  = [c for c in db_cols_r if not any(c == b for _, b in mapped_r)]
    result['r_penjualan'] = {
        'file': 'R_penjualan.csv',
        'row_count': row_count_r,
        'csv_columns': hdrs_r or [],
        'db_columns': db_cols_r,
        'mapped': [{'csv': a, 'db': b} for a, b in mapped_r],
        'skipped_csv': skipped_r,
        'missing_db': missing_r,
    }

    # Lokasi tersedia untuk dropdown
    lokasi_list = rows_to_list(query("SELECT kd_lokasi, nm_lokasi FROM Lokasi WHERE aktif=1 ORDER BY nm_lokasi"))
    result['lokasi_options'] = lokasi_list

    return jsonify(result)


@app.route('/api/migrasi/eksekusi', methods=['POST'])
def migrasi_eksekusi():
    try:
        body     = request.json or {}
        mode     = body.get('mode', 'append')   # 'append' | 'replace'
        defaults = body.get('defaults') or {}
        def_p    = defaults.get('penjualan') or {}
        def_r    = defaults.get('r_penjualan') or {}

        hdrs_p, rows_p, err_p = _load_sementara_csv('Penjualan.csv')
        if err_p:
            return jsonify({'ok': False, 'msg': err_p}), 400
        hdrs_r, rows_r, err_r = _load_sementara_csv('R_penjualan.csv')
        if err_r:
            return jsonify({'ok': False, 'msg': err_r}), 400

        con = get_raw_db(timeout=30)
        cur = con.cursor()
        cur.execute("PRAGMA foreign_keys = OFF")
        cur.execute("PRAGMA journal_mode = WAL")
        cur.execute("PRAGMA synchronous = NORMAL")

        if mode == 'replace':
            cur.execute("DELETE FROM R_Penjualan WHERE no_faktur IN (SELECT no_faktur FROM Penjualan)")
            cur.execute("DELETE FROM Penjualan")
            con.commit()

        sql_p = """INSERT OR IGNORE INTO Penjualan
            (no_faktur,tanggal,total_harga,hpp,dokter,kd_pasien,
             kasir,metode_bayar,status,disc_trx,disc_trx_nilai,kd_lokasi)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?)"""
        cnt_p = 0; skip_p = 0
        buf = []
        for r in rows_p:
            d = {h.lower(): v for h, v in zip(hdrs_p, r)}
            no_f = _migr_sv(d.get('no_faktur',''))
            if not no_f:
                skip_p += 1
                continue
            buf.append((
                no_f,
                _migr_pd(d.get('tanggal','')),
                _migr_fv(d.get('total_harga', 0)),
                _migr_fv(d.get('hpp', 0)),
                _migr_sv(d.get('dokter','')),
                _migr_sv(d.get('kd_pasien','')),
                _migr_sv(def_p.get('kasir', d.get('dokter',''))),
                # 'Tunai' (kapital) — konsisten dgn konvensi seluruh sistem (DEFAULT skema
                # tabel, tombol POS, validasi bayar_penjualan()). Dulu default di sini
                # 'tunai' huruf kecil — ketahuan 2026-09-27 krn bikin Laporan Penjualan
                # pecah jadi 2 kartu metode terpisah ('tunai' vs 'Tunai') padahal maksudnya
                # sama (MySQL collation kolom ini case-insensitive, tapi breakdown laporan
                # dihitung di Python yg case-sensitive).
                _migr_sv(def_p.get('metode_bayar','Tunai')),
                _migr_sv(def_p.get('status','done')),
                _migr_fv(def_p.get('disc_trx', 0)),
                _migr_fv(def_p.get('disc_trx_nilai', 0)),
                _migr_sv(def_p.get('kd_lokasi','INDUK')),
            ))
            if len(buf) >= 2000:
                _bulk_executemany(con, cur, sql_p, buf)
                cnt_p += len(buf)
                buf = []
        if buf:
            _bulk_executemany(con, cur, sql_p, buf)
            cnt_p += len(buf)
        con.commit()

        kd_pelayanan_set = {
            row[0] for row in cur.execute(
                "SELECT kd_pelayanan FROM produk_pelayanan"
            ).fetchall()
        }

        sql_r = """INSERT INTO R_Penjualan
            (no_faktur,tanggal,kd_barang,nm_barang,harga,harga_asli,qty,
             disc,sub_total,total_harbeli,kd_karyawan,cashier,exp,jenis,racikan_ref)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)"""
        cnt_r = 0; skip_r = 0; cnt_r_jasa = 0
        buf = []
        for r in rows_r:
            d = {h.lower(): v for h, v in zip(hdrs_r, r)}
            no_f = _migr_sv(d.get('no_faktur',''))
            kd_b = _migr_sv(d.get('kd_barang',''))
            if not no_f or not kd_b:
                skip_r += 1
                continue
            # Auto-deteksi: kd_barang ada di produk_pelayanan → jasa, selain itu → barang
            jenis = 'jasa' if kd_b in kd_pelayanan_set else 'barang'
            if jenis == 'jasa':
                cnt_r_jasa += 1
            buf.append((
                no_f,
                _migr_pd(d.get('tanggal','')),
                kd_b,
                _migr_sv(d.get('nm_barang','')),
                _migr_fv(d.get('harga', 0)),
                _migr_fv(d.get('harga_asli', 0)),
                _migr_fv(d.get('qty', 0)),
                # disc dipakai ULANG kalau faktur ini diretur nanti (_eksekusi_retur baca
                # R_Penjualan.disc utk hitung harga_eff) — kalau >100%/negatif dari CSV
                # rusak, retur thd faktur migrasi ini bisa salah hitung (bahkan negatif).
                # Ditemukan 2026-10-02 saat audit semua tool migrasi pasca kasus satuan/disc.
                _migr_disc_aman(d.get('disc', 0)),
                _migr_fv(d.get('sub_total', 0)),
                _migr_fv(d.get('total_harbeli', 0)),
                _migr_sv(d.get('kd_karyawan','')),
                _migr_sv(d.get('cashier','')),
                _migr_sv(d.get('exp','')),
                jenis,
                _migr_sv(def_r.get('racikan_ref','')),
            ))
            if len(buf) >= 2000:
                _bulk_executemany(con, cur, sql_r, buf)
                cnt_r += len(buf)
                buf = []
        if buf:
            _bulk_executemany(con, cur, sql_r, buf)
            cnt_r += len(buf)
        con.commit()
        cur.execute("PRAGMA foreign_keys = ON")
        con.commit()
        con.close()
        return jsonify({
            'ok': True,
            'penjualan_count':    cnt_p,
            'penjualan_skip':     skip_p,
            'r_penjualan_count':  cnt_r,
            'r_penjualan_skip':   skip_r,
            'r_penjualan_jasa':   cnt_r_jasa,
            'r_penjualan_barang': cnt_r - cnt_r_jasa,
        })
    except Exception as e:
        try:
            con.close()
        except Exception:
            pass
        return jsonify({'ok': False, 'msg': str(e)}), 500


# ─── MIGRASI PASIEN ──────────────────────────────────────────

def _load_sementara_csv_auto(filename):
    """Seperti _load_sementara_csv tapi auto-detect delimiter (koma atau pipa)."""
    path = os.path.join(SEMENTARA_DIR, filename)
    if not os.path.exists(path):
        return None, None, f"File {filename} tidak ditemukan di folder sementara"
    with open(path, encoding='utf-8-sig', errors='replace') as f:
        first_line = f.readline()
        f.seek(0)
        delim = '|' if first_line.count('|') > first_line.count(',') else ','
        reader = csv.reader(f, delimiter=delim)
        rows = list(reader)
    if not rows:
        return [], [], None
    headers = [h.strip() for h in rows[0]]
    return headers, rows[1:], None


def _pasien_alias_map():
    """Alias mapping: csv_col_lower -> db_col (None = skip)."""
    return {
        'telepon':    'telp',
        'kelamin':    'jenis_kelamin',
        't_lahir':    'tgl_lahir',
    }


@app.route('/migrasi_pasien')
def migrasi_pasien_page():
    return send_from_directory('static', 'migrasi_pasien.html')


@app.route('/api/migrasi_pasien/upload', methods=['POST'])
def migrasi_pasien_upload():
    f = request.files.get('file')
    if not f:
        return jsonify({'ok': False, 'msg': 'File tidak ditemukan'}), 400
    f.save(os.path.join(SEMENTARA_DIR, 'Pasien.csv'))
    return jsonify({'ok': True})


@app.route('/api/migrasi_pasien/status')
def migrasi_pasien_status():
    path = os.path.join(SEMENTARA_DIR, 'Pasien.csv')
    if os.path.exists(path):
        return jsonify({'exists': True, 'size': os.path.getsize(path)})
    return jsonify({'exists': False})


@app.route('/api/migrasi_pasien/info')
def migrasi_pasien_info():
    """Hitung pemetaan kolom CSV Pasien → kolom DB."""
    hdrs, rows, err = _load_sementara_csv_auto('Pasien.csv')
    if err:
        return jsonify({'ok': False, 'msg': err}), 400
    row_count = len(rows or [])

    db_cols   = ['kd_pasien', 'nm_pasien', 'alamat', 'telp', 'hp',
                 'tgl_lahir', 'jenis_kelamin', 'created_at']
    alias_map = _pasien_alias_map()

    mapped  = []
    skipped = []
    for h in (hdrs or []):
        h_low = h.lower()
        # exact match (case-insensitive)
        exact = next((c for c in db_cols if h_low == c.lower()), None)
        if exact:
            mapped.append((h, exact))
            continue
        # alias match
        if h_low in alias_map:
            db_target = alias_map[h_low]
            if db_target:
                mapped.append((h, db_target))
            else:
                skipped.append(h)
        else:
            skipped.append(h)

    mapped_targets = {b for _, b in mapped}
    missing        = [c for c in db_cols if c not in mapped_targets]
    has_alias      = any(a.lower() != b.lower() for a, b in mapped)

    return jsonify({
        'pasien': {
            'file':        'Pasien.csv',
            'row_count':   row_count,
            'csv_columns': hdrs or [],
            'db_columns':  db_cols,
            'mapped':      [{'csv': a, 'db': b} for a, b in mapped],
            'skipped_csv': skipped,
            'missing_db':  missing,
            'has_alias':   has_alias,
            'notes': [
                'Kolom <b>Telepon</b> di-rename ke <b>telp</b>.',
                'Kolom <b>HP</b> dipetakan langsung ke kolom <b>hp</b> di database.',
                'Kolom <b>Kelamin</b> di-rename ke <b>jenis_kelamin</b>.',
                'Kolom <b>T_lahir</b> di-rename ke <b>tgl_lahir</b>.',
                'Format tanggal didukung: M/D/YYYY, D/M/YYYY, dan YYYY-MM-DD.',
            ],
        }
    })


@app.route('/api/migrasi_pasien/eksekusi', methods=['POST'])
def migrasi_pasien_eksekusi():
    try:
        body     = request.json or {}
        defaults = body.get('defaults', {})
        mode     = body.get('mode', 'append')
        hdrs, rows, err = _load_sementara_csv_auto('Pasien.csv')
        if err:
            return jsonify({'ok': False, 'msg': err}), 400
        default_created = _migr_sv(defaults.get('created_at', str(date.today())))

        con = get_raw_db(timeout=30)
        cur = con.cursor()
        cur.execute("PRAGMA foreign_keys = OFF")
        cur.execute("PRAGMA journal_mode = WAL")
        cur.execute("PRAGMA synchronous = NORMAL")

        if mode == 'replace':
            cur.execute("DELETE FROM Pasien")
            con.commit()

        cnt  = 0
        skip = 0

        if rows:
            sql = """INSERT OR IGNORE INTO Pasien
                (kd_pasien, nm_pasien, alamat, telp, hp, tgl_lahir, jenis_kelamin, created_at)
                VALUES (?,?,?,?,?,?,?,?)"""
            buf = []
            for r in rows:
                d    = {h.strip().lower(): v.strip() for h, v in zip(hdrs, r)}
                kd   = _migr_sv(d.get('kd_pasien', ''))
                nm   = _migr_sv(d.get('nm_pasien', ''))
                if not kd or not nm:
                    skip += 1
                    continue

                def _clean_phone(s):
                    v = _migr_sv(s)
                    return '' if (not v or v == '---' or v.startswith('---')) else v

                buf.append((
                    kd,
                    nm,
                    _migr_sv(d.get('alamat', '')),
                    _clean_phone(d.get('telepon', '')),
                    _clean_phone(d.get('hp', '')),
                    _migr_pd(d.get('t_lahir', '')),
                    _migr_sv(d.get('kelamin', '')),
                    default_created,
                ))
                if len(buf) >= 2000:
                    _bulk_executemany(con, cur, sql, buf)
                    cnt += len(buf)
                    buf = []
            if buf:
                _bulk_executemany(con, cur, sql, buf)
                cnt += len(buf)
            con.commit()

        cur.execute("PRAGMA foreign_keys = ON")
        con.commit()
        con.close()

        return jsonify({'ok': True, 'pasien_count': cnt, 'pasien_skip': skip})
    except Exception as e:
        try:
            con.close()
        except Exception:
            pass
        return jsonify({'ok': False, 'msg': str(e)}), 500


# ─── MIGRASI PEMBELIAN ────────────────────────────────────────

@app.route('/migrasi_pembelian')
def migrasi_pembelian_page():
    return send_from_directory('static', 'migrasi_pembelian.html')


@app.route('/api/migrasi_pembelian/upload', methods=['POST'])
def migrasi_pembelian_upload():
    fp = request.files.get('file_pembelian')
    fr = request.files.get('file_r_pembelian')
    if not fp and not fr:
        return jsonify({'ok': False, 'msg': 'Tidak ada file diunggah'}), 400
    if fp:
        fp.save(os.path.join(SEMENTARA_DIR, 'Pembelian.csv'))
    if fr:
        fr.save(os.path.join(SEMENTARA_DIR, 'R_pembelian.csv'))
    return jsonify({'ok': True})


@app.route('/api/migrasi_pembelian/status')
def migrasi_pembelian_status():
    def _st(fn):
        path = os.path.join(SEMENTARA_DIR, fn)
        if os.path.exists(path):
            return {'exists': True, 'size': os.path.getsize(path)}
        return {'exists': False}
    return jsonify({'pembelian': _st('Pembelian.csv'), 'r_pembelian': _st('R_pembelian.csv')})


@app.route('/api/migrasi_pembelian/info')
def migrasi_pembelian_info():
    """Hitung pemetaan kolom CSV Pembelian & R_Pembelian → kolom DB."""
    result = {}

    # ── Pembelian ─────────────────────────────────────────────
    hdrs_p, rows_p, err_p = _load_sementara_csv('Pembelian.csv')
    if err_p:
        result['pembelian'] = {'error': err_p}
    else:
        row_count_p = len(rows_p or [])
        db_cols_p = ['no_faktur', 'tanggal', 'total_harga', 'sistem_bayar',
                     'lunas', 'tempo', 'nm_suplier', 'kd_lokasi']
        mapped_p = []
        for h in (hdrs_p or []):
            for c in db_cols_p:
                if h.lower() == c.lower():
                    mapped_p.append((h, c))
                    break
        skipped_p = [h for h in (hdrs_p or []) if not any(h == a for a, _ in mapped_p)]
        missing_p = [c for c in db_cols_p if not any(c == b for _, b in mapped_p)]
        result['pembelian'] = {
            'file': 'Pembelian.csv',
            'row_count': row_count_p,
            'csv_columns': hdrs_p or [],
            'db_columns': db_cols_p,
            'mapped': [{'csv': a, 'db': b} for a, b in mapped_p],
            'skipped_csv': skipped_p,
            'missing_db': missing_p,
        }

    # ── R_Pembelian ───────────────────────────────────────────
    hdrs_r, rows_r, err_r = _load_sementara_csv('R_pembelian.csv')
    if err_r:
        result['r_pembelian'] = {'error': err_r}
        lokasi_list = rows_to_list(query("SELECT kd_lokasi, nm_lokasi FROM Lokasi WHERE aktif=1 ORDER BY nm_lokasi"))
        result['lokasi_options'] = lokasi_list
        result['pemilik_options'] = rows_to_list(query(
            "SELECT kd_pemilik, nm_pemilik FROM Induk_Pemilik WHERE aktif=1 ORDER BY kd_pemilik"))
        return jsonify(result)
    row_count_r = len(rows_r or [])
    db_cols_r = ['no_faktur', 'tanggal', 'kd_suplier', 'nm_suplier', 'sistem_bayar',
                 'kd_barang', 'nm_barang', 'disc', 'harga_beli_kotor', 'harga_beli_bersih',
                 'sub_total', 'ppn', 'tempo', 'qty', 'harga_asli', 'satuan', 'exp']
    mapped_r = []
    for h in (hdrs_r or []):
        for c in db_cols_r:
            if h.lower() == c.lower():
                mapped_r.append((h, c))
                break
    skipped_r = [h for h in (hdrs_r or []) if not any(h == a for a, _ in mapped_r)]
    missing_r = [c for c in db_cols_r if not any(c == b for _, b in mapped_r)]
    result['r_pembelian'] = {
        'file': 'R_pembelian.csv',
        'row_count': row_count_r,
        'csv_columns': hdrs_r or [],
        'db_columns': db_cols_r,
        'mapped': [{'csv': a, 'db': b} for a, b in mapped_r],
        'skipped_csv': skipped_r,
        'missing_db': missing_r,
    }

    lokasi_list = rows_to_list(query(
        "SELECT kd_lokasi, nm_lokasi FROM Lokasi WHERE aktif=1 ORDER BY nm_lokasi"))
    result['lokasi_options'] = lokasi_list
    result['pemilik_options'] = rows_to_list(query(
        "SELECT kd_pemilik, nm_pemilik FROM Induk_Pemilik WHERE aktif=1 ORDER BY kd_pemilik"))
    return jsonify(result)


@app.route('/api/migrasi_pembelian/eksekusi', methods=['POST'])
def migrasi_pembelian_eksekusi():
    try:
        body     = request.json or {}
        mode     = body.get('mode', 'append')
        defaults = body.get('defaults') or {}

        hdrs_p, rows_p, err_p = _load_sementara_csv('Pembelian.csv')
        if err_p:
            return jsonify({'ok': False, 'msg': err_p}), 400
        hdrs_r, rows_r, err_r = _load_sementara_csv('R_pembelian.csv')
        if err_r:
            return jsonify({'ok': False, 'msg': err_r}), 400

        con = get_raw_db(timeout=30)
        cur = con.cursor()
        cur.execute("PRAGMA foreign_keys = OFF")
        cur.execute("PRAGMA journal_mode = WAL")
        cur.execute("PRAGMA synchronous = NORMAL")

        if mode == 'replace':
            cur.execute("DELETE FROM R_Pembelian WHERE no_faktur IN (SELECT no_faktur FROM Pembelian)")
            cur.execute("DELETE FROM Pembelian")
            con.commit()

        sql_p = """INSERT OR IGNORE INTO Pembelian
            (no_faktur, tanggal, total_harga, sistem_bayar, lunas, tempo, nm_suplier, kd_lokasi, pemilik)
            VALUES (?,?,?,?,?,?,?,?,?)"""
        cnt_p = 0; skip_p = 0
        buf = []
        for r in rows_p:
            d    = {h.lower(): v for h, v in zip(hdrs_p, r)}
            no_f = _migr_sv(d.get('no_faktur', ''))
            if not no_f:
                skip_p += 1
                continue
            buf.append((
                no_f,
                _migr_pd(d.get('tanggal', '')),
                _migr_fv(d.get('total_harga', 0)),
                _migr_sv(d.get('sistem_bayar', '')),
                _migr_sv(d.get('lunas', 'Belum')),
                _migr_pd(d.get('tempo', '')),
                _migr_sv(d.get('nm_suplier', '')),
                _migr_sv(defaults.get('kd_lokasi', 'INDUK')),
                _migr_sv(defaults.get('pemilik', 'P1')),
            ))
            if len(buf) >= 2000:
                _bulk_executemany(con, cur, sql_p, buf)
                cnt_p += len(buf)
                buf = []
        if buf:
            _bulk_executemany(con, cur, sql_p, buf)
            cnt_p += len(buf)
        con.commit()

        # Tidak ada DELETE di sini — sudah dibersihkan di atas (mode=replace), persis
        # spt alur lama (R_Pembelian pakai plain INSERT tanpa OR IGNORE, tidak ada
        # dedup selain validasi no_faktur+kd_barang wajib ada).
        sql_r = """INSERT INTO R_Pembelian
            (no_faktur, tanggal, kd_suplier, nm_suplier, sistem_bayar,
             kd_barang, nm_barang, disc, harga_beli_kotor, harga_beli_bersih,
             sub_total, ppn, tempo, qty, harga_asli, satuan, exp)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)"""
        cnt_r = 0; skip_r = 0
        buf = []
        for r in rows_r:
            d    = {h.lower(): v for h, v in zip(hdrs_r, r)}
            no_f = _migr_sv(d.get('no_faktur', ''))
            kd_b = _migr_sv(d.get('kd_barang', ''))
            if not no_f or not kd_b:
                skip_r += 1
                continue
            buf.append((
                no_f,
                _migr_pd(d.get('tanggal', '')),
                _migr_sv(d.get('kd_suplier', '')),
                _migr_sv(d.get('nm_suplier', '')),
                _migr_sv(d.get('sistem_bayar', '')),
                kd_b,
                _migr_sv(d.get('nm_barang', '')),
                _migr_fv(d.get('disc', 0)),
                _migr_fv(d.get('harga_beli_kotor', 0)),
                _migr_fv(d.get('harga_beli_bersih', 0)),
                _migr_fv(d.get('sub_total', 0)),
                _migr_fv(d.get('ppn', 0)),
                _migr_pd(d.get('tempo', '')),
                _migr_fv(d.get('qty', 0)),
                _migr_fv(d.get('harga_asli', 0)),
                _migr_satuan_aman(d.get('satuan', '')),
                _migr_sv(d.get('exp', '')),
            ))
            if len(buf) >= 2000:
                _bulk_executemany(con, cur, sql_r, buf)
                cnt_r += len(buf)
                buf = []
        if buf:
            _bulk_executemany(con, cur, sql_r, buf)
            cnt_r += len(buf)
        con.commit()
        cur.execute("PRAGMA foreign_keys = ON")
        con.commit()
        con.close()
        return jsonify({
            'ok': True,
            'pembelian_count':   cnt_p,
            'pembelian_skip':    skip_p,
            'r_pembelian_count': cnt_r,
            'r_pembelian_skip':  skip_r,
        })
    except Exception as e:
        try:
            con.close()
        except Exception:
            pass
        return jsonify({'ok': False, 'msg': str(e)}), 500


@app.route('/migrasi_treatmen')
def migrasi_treatmen_page():
    return send_from_directory('static', 'migrasi_treatmen.html')


# Alias: csv_col_lower → db_col
_TREATMEN_ALIAS = {
    'komisi':  'komisi_dokter',
    'komisi2': 'komisi_supervisor',
    'komisi3': 'komisi_perawat',
}

@app.route('/api/migrasi_treatmen/upload', methods=['POST'])
def migrasi_treatmen_upload():
    f = request.files.get('file')
    if not f:
        return jsonify({'ok': False, 'msg': 'File tidak ditemukan'}), 400
    f.save(os.path.join(SEMENTARA_DIR, 'Produk_pelayanan.csv'))
    return jsonify({'ok': True})


@app.route('/api/migrasi_treatmen/status')
def migrasi_treatmen_status():
    path = os.path.join(SEMENTARA_DIR, 'Produk_pelayanan.csv')
    if os.path.exists(path):
        return jsonify({'exists': True, 'size': os.path.getsize(path)})
    return jsonify({'exists': False})


@app.route('/api/migrasi_treatmen/info')
def migrasi_treatmen_info():
    """Hitung pemetaan kolom CSV Produk_pelayanan → kolom DB."""
    hdrs, rows, err = _load_sementara_csv('Produk_pelayanan.csv')
    if err:
        return jsonify({'ok': False, 'msg': err}), 400
    row_count = len(rows or [])
    db_cols = ['kd_pelayanan', 'nm_pelayanan', 'harga',
               'komisi_dokter', 'komisi_supervisor', 'komisi_perawat']

    mapped  = []
    aliased = []
    for h in (hdrs or []):
        hl = h.lower()
        for c in db_cols:
            if hl == c.lower():
                mapped.append({'csv': h, 'db': c})
                break
        else:
            if hl in _TREATMEN_ALIAS:
                aliased.append({'csv': h, 'db': _TREATMEN_ALIAS[hl]})

    mapped_csv  = {m['csv'] for m in mapped}
    aliased_csv = {a['csv'] for a in aliased}
    skipped     = [h for h in (hdrs or []) if h not in mapped_csv and h not in aliased_csv]
    all_covered = {m['db'] for m in mapped} | {a['db'] for a in aliased}
    missing     = [c for c in db_cols if c not in all_covered]

    return jsonify({
        'produk_pelayanan': {
            'file':        'Produk_pelayanan.csv',
            'row_count':   row_count,
            'csv_columns': hdrs or [],
            'db_columns':  db_cols,
            'mapped':      mapped,
            'aliased':     aliased,
            'skipped_csv': skipped,
            'missing_db':  missing,
        }
    })


@app.route('/api/migrasi_treatmen/eksekusi', methods=['POST'])
def migrasi_treatmen_eksekusi():
    try:
        body = request.json or {}
        mode = body.get('mode', 'append')
        hdrs, rows, err = _load_sementara_csv('Produk_pelayanan.csv')
        if err:
            return jsonify({'ok': False, 'msg': err}), 400

        con = get_raw_db(timeout=30)
        cur = con.cursor()
        cur.execute("PRAGMA foreign_keys = OFF")
        cur.execute("PRAGMA journal_mode = WAL")
        cur.execute("PRAGMA synchronous = NORMAL")

        if mode == 'replace':
            cur.execute("DELETE FROM produk_pelayanan")
            con.commit()

        sql = """INSERT OR IGNORE INTO produk_pelayanan
            (kd_pelayanan, nm_pelayanan, harga,
             komisi_dokter, komisi_supervisor, komisi_perawat)
            VALUES (?,?,?,?,?,?)"""

        # bangun lookup: db_col → csv_col (gabung exact + alias)
        col_map = {}
        for h in (hdrs or []):
            hl = h.lower()
            for c in ['kd_pelayanan', 'nm_pelayanan', 'harga',
                      'komisi_dokter', 'komisi_supervisor', 'komisi_perawat']:
                if hl == c.lower():
                    col_map[c] = h
                    break
            if hl in _TREATMEN_ALIAS:
                col_map[_TREATMEN_ALIAS[hl]] = h

        buf      = []
        inserted = 0
        skipped  = 0
        total    = 0

        for r in rows:
            total += 1
            d  = {h: v for h, v in zip(hdrs, r)}
            kd = _migr_sv(d.get(col_map.get('kd_pelayanan', ''), ''))
            if not kd:
                skipped += 1
                continue

            def _cv(db_col):
                csv_col = col_map.get(db_col, '')
                return d.get(csv_col, '') if csv_col else ''

            buf.append((
                kd,
                _migr_sv(_cv('nm_pelayanan')),
                _migr_fv(_cv('harga')),
                _migr_fv(_cv('komisi_dokter')),
                _migr_fv(_cv('komisi_supervisor')),
                _migr_fv(_cv('komisi_perawat')),
            ))
            if len(buf) >= 2000:
                _bulk_executemany(con, cur, sql, buf)
                inserted += len(buf)
                buf = []

        if buf:
            _bulk_executemany(con, cur, sql, buf)
            inserted += len(buf)

        cur.execute("PRAGMA foreign_keys = ON")
        con.commit()
        con.close()

        return jsonify({'ok': True, 'inserted': inserted, 'skipped': skipped, 'total': total})
    except Exception as e:
        try: con.close()
        except Exception: pass
        return jsonify({'ok': False, 'msg': str(e)}), 500


@app.route('/migrasi_stok')
def migrasi_stok_page():
    return send_from_directory('static', 'migrasi_stok.html')


@app.route('/api/migrasi_stok/upload', methods=['POST'])
def migrasi_stok_upload():
    f = request.files.get('file')
    if not f:
        return jsonify({'ok': False, 'msg': 'File tidak ditemukan'}), 400
    f.save(os.path.join(SEMENTARA_DIR, 'Stok_barang.csv'))
    return jsonify({'ok': True})


@app.route('/api/migrasi_stok/status')
def migrasi_stok_status():
    path = os.path.join(SEMENTARA_DIR, 'Stok_barang.csv')
    if os.path.exists(path):
        return jsonify({'exists': True, 'size': os.path.getsize(path)})
    return jsonify({'exists': False})


@app.route('/api/migrasi_stok/info')
def migrasi_stok_info():
    """Hitung pemetaan kolom CSV Stok_barang → kolom DB.
    lokasi_options/pemilik_options TETAP query DB langsung (bukan dari CSV)."""
    hdrs, rows, err = _load_sementara_csv('Stok_barang.csv')
    if err:
        return jsonify({'ok': False, 'msg': err}), 400
    row_count = len(rows or [])
    db_cols = ['kd_barang', 'nm_barang', 'harga_beli', 'harga_beli_awal',
               'stok_awal', 'stok_sisa', 'stok_gudang', 'stok_min',
               'harga', 'disc', 'satuan', 'exp', 'DiscMem', 'DiscLain', 'kd_lokasi']
    pemilik_list = rows_to_list(query(
        "SELECT kd_pemilik, nm_pemilik FROM Induk_Pemilik WHERE aktif=1 ORDER BY kd_pemilik"))

    db_lower  = [c.lower() for c in db_cols]
    mapped    = []
    for h in (hdrs or []):
        for c in db_cols:
            if h.lower() == c.lower():
                mapped.append({'csv': h, 'db': c})
                break
    mapped_csv = {m['csv'] for m in mapped}
    skipped    = [h for h in (hdrs or []) if h not in mapped_csv]
    missing    = [c for c in db_cols if not any(m['db'] == c for m in mapped)]

    lokasi_list = rows_to_list(query(
        "SELECT kd_lokasi, nm_lokasi, tipe FROM Lokasi WHERE aktif=1 ORDER BY nm_lokasi"))
    return jsonify({
        'stok_barang': {
            'file':        'Stok_barang.csv',
            'row_count':   row_count,
            'csv_columns': hdrs or [],
            'db_columns':  db_cols,
            'mapped':      mapped,
            'skipped_csv': skipped,
            'missing_db':  missing,
        },
        'lokasi_options':  lokasi_list,
        'pemilik_options': pemilik_list,
    })


@app.route('/api/migrasi_stok/eksekusi', methods=['POST'])
def migrasi_stok_eksekusi():
    try:
        body     = request.json or {}
        mode     = body.get('mode', 'append')
        defaults = body.get('defaults', {})
        hdrs, rows, err = _load_sementara_csv('Stok_barang.csv')
        if err:
            return jsonify({'ok': False, 'msg': err}), 400

        con = get_raw_db(timeout=30)
        cur = con.cursor()
        cur.execute("PRAGMA foreign_keys = OFF")
        cur.execute("PRAGMA journal_mode = WAL")
        cur.execute("PRAGMA synchronous = NORMAL")

        if mode == 'replace':
            cur.execute("DELETE FROM Stok_barang")
            con.commit()

        sql = """INSERT OR IGNORE INTO Stok_barang
            (kd_barang, nm_barang, harga_beli, harga_beli_awal,
             stok_awal, stok_sisa, stok_gudang, stok_min,
             harga, disc, satuan, exp, DiscMem, DiscLain, kd_lokasi, pemilik)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)"""

        kd_lokasi_default = _migr_sv(defaults.get('kd_lokasi', 'CBG1'))
        # Pemilik hanya relevan utk stok INDUK; ke cabang dikosongkan
        pemilik_default = _migr_sv(defaults.get('pemilik', 'P1')) if kd_lokasi_default == 'INDUK' else ''

        buf      = []
        inserted = 0
        skipped  = 0
        total    = 0

        for r in rows:
            total += 1
            d = {h.lower(): v for h, v in zip(hdrs, r)}
            kd = _migr_sv(d.get('kd_barang', ''))
            if not kd:
                skipped += 1
                continue
            buf.append((
                kd,
                _migr_sv(d.get('nm_barang', '')),
                _migr_fv(d.get('harga_beli', 0)),
                _migr_fv(d.get('harga_beli_awal', 0)),
                _migr_fv(d.get('stok_awal', 0)),
                _migr_fv(d.get('stok_sisa', 0)),
                _migr_fv(d.get('stok_gudang', 0)),
                _migr_fv(d.get('stok_min', 0)),
                _migr_fv(d.get('harga', 0)),
                _migr_disc_aman(d.get('disc', '0')),
                _migr_satuan_aman(d.get('satuan', '')),
                _migr_sv(d.get('exp', '')),
                _migr_fv(d.get('discmem', 0)),
                _migr_fv(d.get('disclain', 0)),
                kd_lokasi_default,
                pemilik_default,
            ))
            if len(buf) >= 2000:
                _bulk_executemany(con, cur, sql, buf)
                inserted += len(buf)
                buf = []

        if buf:
            _bulk_executemany(con, cur, sql, buf)
            inserted += len(buf)

        cur.execute("PRAGMA foreign_keys = ON")
        con.commit()
        con.close()

        return jsonify({
            'ok':       True,
            'inserted': inserted,
            'skipped':  skipped,
            'total':    total,
        })
    except Exception as e:
        try:
            con.close()
        except Exception:
            pass
        return jsonify({'ok': False, 'msg': str(e)}), 500


# ─── MIGRASI FORMULA PRODUKSI (BOM racikan) ────────────────────
_FORMULA_CSV = 'FormRacik.csv'

@app.route('/migrasi_formula')
def migrasi_formula_page():
    return send_from_directory('static', 'migrasi_formula.html')

def _formula_idx_val(hdrs):
    idx = {h.lower(): i for i, h in enumerate(hdrs or [])}
    def val(r, key):
        i = idx.get(key)
        return (r[i].strip() if (i is not None and i < len(r) and r[i] is not None) else '')
    return val

@app.route('/api/migrasi_formula/upload', methods=['POST'])
def migrasi_formula_upload():
    f = request.files.get('file')
    if not f:
        return jsonify({'ok': False, 'msg': 'File tidak ditemukan'}), 400
    f.save(os.path.join(SEMENTARA_DIR, _FORMULA_CSV))
    return jsonify({'ok': True})


@app.route('/api/migrasi_formula/status')
def migrasi_formula_status():
    path = os.path.join(SEMENTARA_DIR, _FORMULA_CSV)
    if os.path.exists(path):
        return jsonify({'exists': True, 'size': os.path.getsize(path)})
    return jsonify({'exists': False})


@app.route('/api/migrasi_formula/info')
def migrasi_formula_info():
    """Hitung pemetaan kolom + sinkronisasi (formula/bahan/produk vs master Barang)."""
    hdrs, rows, err = _load_sementara_csv(_FORMULA_CSV)
    if err:
        return jsonify({'ok': False, 'msg': err}), 400
    row_count = len(rows or [])
    db_cols = ['kd_racik', 'kd_barang', 'nm_barang', 'komposisi', 'Satuan', 'nm_racik']
    mapped = []
    for h in (hdrs or []):
        for c in db_cols:
            if h.lower() == c.lower():
                mapped.append({'csv': h, 'db': c}); break
    mapped_csv = {m['csv'] for m in mapped}
    skipped = [h for h in (hdrs or []) if h not in mapped_csv]
    missing = [c for c in db_cols if not any(m['db'] == c for m in mapped)]

    val = _formula_idx_val(hdrs)
    bahan = set(); produk = set(); formula_count = 0
    seen_racik = set()
    for r in (rows or []):
        kr = val(r, 'kd_racik')
        if not kr:
            continue
        if kr not in seen_racik:
            seen_racik.add(kr); produk.add(kr); formula_count += 1
        kb = val(r, 'kd_barang')
        if kb:
            bahan.add(kb)

    ada = {x['kd_barang'] for x in query("SELECT kd_barang FROM Barang")}
    return jsonify({'formula': {
        'file': _FORMULA_CSV, 'row_count': row_count,
        'csv_columns': hdrs or [], 'db_columns': db_cols,
        'mapped': mapped, 'skipped_csv': skipped, 'missing_db': missing,
        'formula_count': formula_count,
        'bahan_total':  len(bahan),  'bahan_belum':  sorted([b for b in bahan if b not in ada]),
        'produk_total': len(produk), 'produk_belum': sorted([p for p in produk if p not in ada]),
    }})

@app.route('/api/migrasi_formula/eksekusi', methods=['POST'])
def migrasi_formula_eksekusi():
    from collections import Counter, OrderedDict
    con = None
    try:
        body = request.json or {}
        mode = body.get('mode', 'append')
        hdrs, rows, err = _load_sementara_csv(_FORMULA_CSV)
        if err:
            return jsonify({'ok': False, 'msg': err}), 400
        val = _formula_idx_val(hdrs)

        formulas = OrderedDict()
        for r in (rows or []):
            kr = val(r, 'kd_racik')
            if not kr:
                continue
            formulas.setdefault(kr, []).append({
                'kd_barang': val(r, 'kd_barang'), 'nm_barang': val(r, 'nm_barang'),
                'qty': _migr_fv(val(r, 'komposisi')),
                # Validasi sejak titik ini (bukan nanti pas dipakai) — Counter most_common()
                # di bawah bisa kepilih nilai GARBAGE sbg "satuan terbanyak" kalau byk baris
                # kebetulan share angka sama persis (pola yg sama spt kasus satuan=harga
                # 2026-10-01, lihat _migr_satuan_aman).
                'satuan': _migr_satuan_aman(val(r, 'satuan')),
                'nm_racik': val(r, 'nm_racik'),
            })

        con = get_raw_db(timeout=30)
        cur = con.cursor()
        cur.execute("PRAGMA foreign_keys = OFF")
        cur.execute("PRAGMA journal_mode = WAL")

        if mode == 'replace':
            cur.execute("DELETE FROM Formula_Produksi_Detail")
            cur.execute("DELETE FROM Formula_Produksi")
            con.commit()

        ada = {x[0] for x in cur.execute("SELECT kd_barang FROM Barang").fetchall()}
        cnt_f = cnt_d = skip_f = master_created = 0

        def ensure_master(kd, nm, satuan):
            nonlocal master_created
            if kd and kd not in ada:
                cur.execute("""INSERT OR IGNORE INTO Barang (kd_barang,nm_barang,nm_generik,satuan,harga_jual,aktif)
                               VALUES (?,?,?,?,0,1)""", (kd, nm or kd, nm or '', satuan or 'pcs'))
                ada.add(kd); master_created += 1

        for kr, items in formulas.items():
            if mode != 'replace':
                if cur.execute("SELECT 1 FROM Formula_Produksi WHERE kd_produk_hasil=? AND aktif=1", (kr,)).fetchone():
                    skip_f += 1
                    continue
            sat_count = Counter([it['satuan'] for it in items if it['satuan']])
            satuan_hasil = sat_count.most_common(1)[0][0] if sat_count else 'pcs'
            nm_racik = next((it['nm_racik'] for it in items if it['nm_racik']), '')
            ket = ('Alias: ' + nm_racik) if nm_racik else ''
            ensure_master(kr, kr, satuan_hasil)   # produk hasil
            cur.execute("""INSERT INTO Formula_Produksi
                           (kd_produk_hasil,nm_produk_hasil,nm_formula,qty_hasil,satuan_hasil,keterangan,aktif)
                           VALUES (?,?,?,1,?,?,1)""", (kr, kr, kr, satuan_hasil, ket))
            fid = cur.lastrowid; cnt_f += 1
            for it in items:
                if not it['kd_barang']:
                    continue
                ensure_master(it['kd_barang'], it['nm_barang'], it['satuan'])
                cur.execute("""INSERT INTO Formula_Produksi_Detail (id_formula,kd_barang,nm_barang,qty,satuan)
                               VALUES (?,?,?,?,?)""",
                            (fid, it['kd_barang'], it['nm_barang'], it['qty'], it['satuan']))
                cnt_d += 1

        cur.execute("PRAGMA foreign_keys = ON")
        con.commit(); con.close()
        return jsonify({'ok': True, 'formula_count': cnt_f, 'detail_count': cnt_d,
                        'formula_skip': skip_f, 'master_created': master_created})
    except Exception as e:
        try:
            if con: con.close()
        except Exception:
            pass
        return jsonify({'ok': False, 'msg': str(e)}), 500


# ─── MIGRASI KARYAWAN ──────────────────────────────────────────

_KARYAWAN_ALIAS = {'hp': 'telp', 'no_hp': 'telp'}
_KARYAWAN_DASH  = {'---', '--', '-', '----', 'none', 'null', 'n/a'}

@app.route('/migrasi_karyawan')
def migrasi_karyawan_page():
    return send_from_directory('static', 'migrasi_karyawan.html')


@app.route('/api/migrasi_karyawan/upload', methods=['POST'])
def migrasi_karyawan_upload():
    f = request.files.get('file')
    if not f:
        return jsonify({'ok': False, 'msg': 'File tidak ditemukan'}), 400
    f.save(os.path.join(SEMENTARA_DIR, 'Karyawan.csv'))
    return jsonify({'ok': True})


@app.route('/api/migrasi_karyawan/status')
def migrasi_karyawan_status():
    path = os.path.join(SEMENTARA_DIR, 'Karyawan.csv')
    if os.path.exists(path):
        return jsonify({'exists': True, 'size': os.path.getsize(path)})
    return jsonify({'exists': False})


@app.route('/api/migrasi_karyawan/info')
def migrasi_karyawan_info():
    """Hitung pemetaan kolom CSV Karyawan → kolom DB."""
    hdrs, rows, err = _load_sementara_csv('Karyawan.csv')
    if err:
        return jsonify({'ok': False, 'msg': err}), 400
    row_count = len(rows or [])
    pek_idx = next((i for i, h in enumerate(hdrs or []) if h.lower() == 'pekerjaan'), None)
    pekerjaan_unik = sorted({r[pek_idx].strip() for r in rows if pek_idx is not None and pek_idx < len(r) and r[pek_idx] and r[pek_idx].strip()}) if pek_idx is not None else []
    db_cols = ['kd_karyawan', 'nm_karyawan', 'nickname', 'pekerjaan', 'telp']

    mapped, aliased = [], []
    for h in (hdrs or []):
        hl = h.lower()
        for c in db_cols:
            if hl == c.lower():
                mapped.append({'csv': h, 'db': c})
                break
        else:
            if hl in _KARYAWAN_ALIAS:
                db_col = _KARYAWAN_ALIAS[hl]
                if not any(a['db'] == db_col for a in aliased):
                    aliased.append({'csv': h, 'db': db_col})

    mapped_csv  = {m['csv'] for m in mapped}
    aliased_csv = {a['csv'] for a in aliased}
    skipped     = [h for h in (hdrs or []) if h not in mapped_csv and h not in aliased_csv]
    all_covered = {m['db'] for m in mapped} | {a['db'] for a in aliased}
    missing     = [c for c in db_cols if c not in all_covered]

    return jsonify({
        'karyawan': {
            'file':           'Karyawan.csv',
            'row_count':      row_count,
            'csv_columns':    hdrs or [],
            'db_columns':     db_cols,
            'mapped':         mapped,
            'aliased':        aliased,
            'skipped_csv':    skipped,
            'missing_db':     missing,
            'pekerjaan_unik': pekerjaan_unik,
        }
    })


@app.route('/api/migrasi_karyawan/eksekusi', methods=['POST'])
def migrasi_karyawan_eksekusi():
    try:
        body = request.json or {}
        mode = body.get('mode', 'append')
        hdrs, rows, err = _load_sementara_csv('Karyawan.csv')
        if err:
            return jsonify({'ok': False, 'msg': err}), 400

        con = get_raw_db(timeout=30)
        cur = con.cursor()
        cur.execute("PRAGMA foreign_keys = OFF")
        cur.execute("PRAGMA journal_mode = WAL")
        cur.execute("PRAGMA synchronous = NORMAL")

        if mode == 'replace':
            cur.execute("DELETE FROM Karyawan")
            con.commit()

        # Bangun col_map: db_col → csv header (asli, case-sensitif)
        col_map = {}
        hdrs_low = [h.lower() for h in hdrs]
        db_cols  = ['kd_karyawan', 'nm_karyawan', 'nickname', 'pekerjaan', 'telp']
        for c in db_cols:
            for i, hl in enumerate(hdrs_low):
                if hl == c.lower():
                    col_map[c] = hdrs[i]
                    break
            if c not in col_map and c == 'telp':
                for i, hl in enumerate(hdrs_low):
                    if hl in _KARYAWAN_ALIAS and _KARYAWAN_ALIAS[hl] == 'telp':
                        col_map['telp'] = hdrs[i]
                        break

        def _clean_telp(v):
            s = str(v).strip()
            return '' if s.lower() in _KARYAWAN_DASH or not s else s

        sql = "INSERT OR IGNORE INTO Karyawan (kd_karyawan,nm_karyawan,nickname,pekerjaan,telp,aktif) VALUES (?,?,?,?,?,1)"

        count_before = cur.execute("SELECT COUNT(*) FROM Karyawan").fetchone()[0]
        buf          = []
        skip_invalid = 0
        total        = 0

        for r in rows:
            total += 1
            d  = {h: v for h, v in zip(hdrs, r)}
            kd = _migr_sv(d.get(col_map.get('kd_karyawan', ''), ''))
            nm = _migr_sv(d.get(col_map.get('nm_karyawan', ''), ''))
            if not kd or not nm:
                skip_invalid += 1
                continue
            buf.append((
                kd,
                nm,
                _migr_sv(d.get(col_map.get('nickname', ''), '')),
                _migr_sv(d.get(col_map.get('pekerjaan', ''), '')),
                _clean_telp(d.get(col_map.get('telp', ''), '')),
            ))
            if len(buf) >= 500:
                _bulk_executemany(con, cur, sql, buf)
                buf = []

        if buf:
            _bulk_executemany(con, cur, sql, buf)

        cur.execute("PRAGMA foreign_keys = ON")
        con.commit()

        count_after = cur.execute("SELECT COUNT(*) FROM Karyawan").fetchone()[0]
        con.close()

        inserted = count_after - count_before
        skipped  = (total - skip_invalid) - inserted
        return jsonify({
            'ok':       True,
            'inserted': inserted,
            'skipped':  skip_invalid + skipped,
            'total':    total,
        })
    except Exception as e:
        try:
            con.close()
        except Exception:
            pass
        return jsonify({'ok': False, 'msg': str(e)}), 500


# ─── MIGRASI KOMISI ────────────────────────────────────────────

@app.route('/migrasi_komisi')
def migrasi_komisi_page():
    return send_from_directory('static', 'migrasi_komisi.html')


@app.route('/api/migrasi_komisi/upload', methods=['POST'])
def migrasi_komisi_upload():
    f = request.files.get('file')
    if not f:
        return jsonify({'ok': False, 'msg': 'File tidak ditemukan'}), 400
    f.save(os.path.join(SEMENTARA_DIR, 'Point_karyawan.csv'))
    return jsonify({'ok': True})


@app.route('/api/migrasi_komisi/status')
def migrasi_komisi_status():
    path = os.path.join(SEMENTARA_DIR, 'Point_karyawan.csv')
    if os.path.exists(path):
        return jsonify({'exists': True, 'size': os.path.getsize(path)})
    return jsonify({'exists': False})


@app.route('/api/migrasi_komisi/info')
def migrasi_komisi_info():
    """Hitung pemetaan kolom CSV Point_karyawan → kolom DB."""
    hdrs, rows, err = _load_sementara_csv('Point_karyawan.csv')
    if err:
        return jsonify({'ok': False, 'msg': err}), 400
    row_count = len(rows or [])
    db_cols = ['no_faktur', 'tanggal', 'kd_pasien', 'kd_pelayanan', 'nm_pelayanan', 'komisi', 'kd_karyawan']

    mapped = []
    for h in (hdrs or []):
        for c in db_cols:
            if h.lower() == c.lower():
                mapped.append({'csv': h, 'db': c})
                break
    mapped_csv = {m['csv'] for m in mapped}
    skipped    = [h for h in (hdrs or []) if h not in mapped_csv]
    missing    = [c for c in db_cols if not any(m['db'] == c for m in mapped)]

    return jsonify({
        'point_karyawan': {
            'file':        'Point_karyawan.csv',
            'row_count':   row_count,
            'csv_columns': hdrs or [],
            'db_columns':  db_cols,
            'mapped':      mapped,
            'skipped_csv': skipped,
            'missing_db':  missing,
        }
    })


@app.route('/api/migrasi_komisi/eksekusi', methods=['POST'])
def migrasi_komisi_eksekusi():
    try:
        body = request.json or {}
        mode = body.get('mode', 'append')
        hdrs, rows, err = _load_sementara_csv('Point_karyawan.csv')
        if err:
            return jsonify({'ok': False, 'msg': err}), 400

        con = get_raw_db(timeout=30)
        cur = con.cursor()
        cur.execute("PRAGMA foreign_keys = OFF")
        cur.execute("PRAGMA journal_mode = WAL")
        cur.execute("PRAGMA synchronous = NORMAL")

        if mode == 'replace':
            cur.execute("DELETE FROM Point_karyawan")
            con.commit()
        existing = set(cur.execute(
            "SELECT no_faktur, kd_karyawan, kd_pelayanan FROM Point_karyawan"
        ).fetchall())

        sql = """INSERT INTO Point_karyawan
            (no_faktur, tanggal, kd_pasien, kd_pelayanan, nm_pelayanan, komisi, kd_karyawan)
            VALUES (?,?,?,?,?,?,?)"""

        buf      = []
        inserted = 0
        skipped  = 0
        total    = 0

        for r in rows:
            total += 1
            d    = {h.lower(): v for h, v in zip(hdrs, r)}
            no_f = _migr_sv(d.get('no_faktur', ''))
            kd_k = _migr_sv(d.get('kd_karyawan', ''))
            kd_p = _migr_sv(d.get('kd_pelayanan', ''))
            if not no_f or not kd_k:
                skipped += 1
                continue
            key = (no_f, kd_k, kd_p)
            if key in existing:
                skipped += 1
                continue
            existing.add(key)
            buf.append((
                no_f,
                _migr_pd(d.get('tanggal', '')),
                _migr_sv(d.get('kd_pasien', '')),
                kd_p,
                _migr_sv(d.get('nm_pelayanan', '')),
                _migr_fv(d.get('komisi', 0)),
                kd_k,
            ))
            if len(buf) >= 2000:
                _bulk_executemany(con, cur, sql, buf)
                inserted += len(buf)
                buf = []

        if buf:
            _bulk_executemany(con, cur, sql, buf)
            inserted += len(buf)

        cur.execute("PRAGMA foreign_keys = ON")
        con.commit()
        con.close()

        return jsonify({
            'ok':       True,
            'inserted': inserted,
            'skipped':  skipped,
            'total':    total,
        })
    except Exception as e:
        try:
            con.close()
        except Exception:
            pass
        return jsonify({'ok': False, 'msg': str(e)}), 500


# ─── MIGRASI MASTER BARANG (katalog produk — beda dari Stok_barang yg per-batch) ──
# Format lawas (dipakai bootstrap _seed_barang_master di Barang.csv): kolom "Nm_barang"
# sebenarnya nama GENERIK (disimpan ke Barang.nm_generik), sedangkan Barang.nm_barang
# (nama tampil di faktur/POS) memakai kd_barang itu sendiri — konvensi yg sama dipakai
# di seluruh app ini utk barang "simple" (kd_barang = teks deskriptif, bukan kode angka).
_BARANG_ALIAS = {'nm_barang': 'nm_generik', 'j_barang': 'jenis'}

@app.route('/migrasi_barang')
def migrasi_barang_page():
    return send_from_directory('static', 'migrasi_barang.html')


@app.route('/api/migrasi_barang/upload', methods=['POST'])
def migrasi_barang_upload():
    f = request.files.get('file')
    if not f:
        return jsonify({'ok': False, 'msg': 'File tidak ditemukan'}), 400
    f.save(os.path.join(SEMENTARA_DIR, 'Barang.csv'))
    return jsonify({'ok': True})


@app.route('/api/migrasi_barang/status')
def migrasi_barang_status():
    path = os.path.join(SEMENTARA_DIR, 'Barang.csv')
    if os.path.exists(path):
        return jsonify({'exists': True, 'size': os.path.getsize(path)})
    return jsonify({'exists': False})


@app.route('/api/migrasi_barang/info')
def migrasi_barang_info():
    """Hitung pemetaan kolom CSV Barang → kolom DB."""
    hdrs, rows, err = _load_sementara_csv('Barang.csv')
    if err:
        return jsonify({'ok': False, 'msg': err}), 400
    row_count = len(rows or [])
    jenis_idx = next((i for i, h in enumerate(hdrs or []) if h.lower() == 'j_barang'), None)
    if jenis_idx is None:
        jenis_idx = next((i for i, h in enumerate(hdrs or []) if h.lower() == 'jenis'), None)
    jenis_unik = sorted({r[jenis_idx].strip() for r in rows if jenis_idx is not None and jenis_idx < len(r) and r[jenis_idx] and r[jenis_idx].strip()}) if jenis_idx is not None else []
    db_cols = ['kd_barang', 'nm_generik', 'jenis', 'merk', 'isi', 'satuan', 'harga_jual']

    mapped, aliased = [], []
    for h in (hdrs or []):
        hl = h.lower()
        for c in db_cols:
            if hl == c.lower():
                mapped.append({'csv': h, 'db': c})
                break
        else:
            if hl in _BARANG_ALIAS:
                db_col = _BARANG_ALIAS[hl]
                if not any(a['db'] == db_col for a in aliased):
                    aliased.append({'csv': h, 'db': db_col})

    mapped_csv  = {m['csv'] for m in mapped}
    aliased_csv = {a['csv'] for a in aliased}
    skipped     = [h for h in (hdrs or []) if h not in mapped_csv and h not in aliased_csv]
    all_covered = {m['db'] for m in mapped} | {a['db'] for a in aliased}
    missing     = [c for c in db_cols if c not in all_covered]

    return jsonify({
        'barang': {
            'file':        'Barang.csv',
            'row_count':   row_count,
            'csv_columns': hdrs or [],
            'db_columns':  db_cols,
            'mapped':      mapped,
            'aliased':     aliased,
            'skipped_csv': skipped,
            'missing_db':  missing,
            'jenis_unik':  jenis_unik,
            'catatan':     "Kolom 'Nm_barang' di file diisi ke nm_generik. Nama tampil "
                           "(nm_barang) otomatis memakai kd_barang, sesuai konvensi aplikasi.",
        }
    })


@app.route('/api/migrasi_barang/eksekusi', methods=['POST'])
def migrasi_barang_eksekusi():
    try:
        body = request.json or {}
        mode = body.get('mode', 'append')
        hdrs, rows, err = _load_sementara_csv('Barang.csv')
        if err:
            return jsonify({'ok': False, 'msg': err}), 400

        con = get_raw_db(timeout=30)
        cur = con.cursor()
        cur.execute("PRAGMA foreign_keys = OFF")
        cur.execute("PRAGMA journal_mode = WAL")
        cur.execute("PRAGMA synchronous = NORMAL")

        if mode == 'replace':
            cur.execute("DELETE FROM Barang")
            con.commit()

        # Bangun col_map: db_col → csv header asli (case-sensitif) — langsung ATAU via alias.
        col_map  = {}
        hdrs_low = [h.lower() for h in hdrs]
        db_cols  = ['kd_barang', 'nm_generik', 'jenis', 'merk', 'isi', 'satuan', 'harga_jual']
        for c in db_cols:
            for i, hl in enumerate(hdrs_low):
                if hl == c.lower():
                    col_map[c] = hdrs[i]
                    break
            if c not in col_map:
                for i, hl in enumerate(hdrs_low):
                    if hl in _BARANG_ALIAS and _BARANG_ALIAS[hl] == c:
                        col_map[c] = hdrs[i]
                        break

        sql = """INSERT OR IGNORE INTO Barang
            (kd_barang,nm_barang,nm_generik,jenis,merk,isi,satuan,harga_jual,aktif)
            VALUES (?,?,?,?,?,?,?,?,1)"""

        count_before = cur.execute("SELECT COUNT(*) FROM Barang").fetchone()[0]
        buf          = []
        skip_invalid = 0
        total        = 0

        for r in rows:
            total += 1
            d  = {h: v for h, v in zip(hdrs, r)}
            kd = _migr_sv(d.get(col_map.get('kd_barang', ''), ''))
            if not kd:
                skip_invalid += 1
                continue
            buf.append((
                kd,
                kd,   # nm_barang (nama tampil) = kd_barang, konvensi aplikasi ini
                _migr_sv(d.get(col_map.get('nm_generik', ''), '')),
                _migr_sv(d.get(col_map.get('jenis', ''), '')),
                _migr_sv(d.get(col_map.get('merk', ''), '')),
                _migr_fv(d.get(col_map.get('isi', ''), 1)) or 1,
                _migr_satuan_aman(d.get(col_map.get('satuan', ''), ''), 'Pcs'),
                _migr_fv(d.get(col_map.get('harga_jual', ''), 0)),
            ))
            if len(buf) >= 2000:
                _bulk_executemany(con, cur, sql, buf)
                buf = []

        if buf:
            _bulk_executemany(con, cur, sql, buf)

        cur.execute("PRAGMA foreign_keys = ON")
        con.commit()

        count_after = cur.execute("SELECT COUNT(*) FROM Barang").fetchone()[0]
        con.close()

        inserted = count_after - count_before
        skipped  = (total - skip_invalid) - inserted
        return jsonify({
            'ok':       True,
            'inserted': inserted,
            'skipped':  skip_invalid + skipped,
            'total':    total,
        })
    except Exception as e:
        try:
            con.close()
        except Exception:
            pass
        return jsonify({'ok': False, 'msg': str(e)}), 500


# ─── BACKUP & RESTORE DATABASE (histori tanggal + keterangan) ──────────────
# Format snapshot: JSONL (1 objek per baris) di-gzip, PORTABEL lintas backend
# SQLite/MySQL — bukan file .db mentah / mysqldump (beda format per-backend,
# butuh binary eksternal yg belum tentu ada di semua komputer). Skema tabel
# TIDAK ikut dibackup (dipercayakan ke init_db()/init_db_mysql() versi app yg
# jalan saat restore) — yg disnapshot murni datanya, tabel demi tabel, dibaca/
# ditulis streaming baris-per-baris supaya aman utk tabel jutaan baris
# (R_Penjualan) tanpa membengkakkan memori.
import gzip

BACKUP_DIR = os.path.join(_base_dir(), 'backup')
os.makedirs(BACKUP_DIR, exist_ok=True)
_BACKUP_INDEX_PATH = os.path.join(BACKUP_DIR, '_index.json')


def _backup_load_index():
    """Histori backup (tanggal, keterangan, dst) disimpan di file INDEX TERPISAH
    dari database — supaya tetap akurat & tidak ikut berubah/hilang seandainya
    suatu saat sebuah backup lama di-restore (restore hanya mengganti isi tabel,
    file index ini di luar jangkauannya)."""
    if not os.path.exists(_BACKUP_INDEX_PATH):
        return []
    try:
        with open(_BACKUP_INDEX_PATH, encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        return []


def _backup_save_index(items):
    with open(_BACKUP_INDEX_PATH, 'w', encoding='utf-8') as f:
        json.dump(items, f, ensure_ascii=False, indent=2)


def _backup_require_admin():
    user = get_current_user()
    if not user or user.get('role') != 'admin':
        return None, (jsonify({'ok': False, 'msg': 'Khusus admin'}), 403)
    return user, None


@app.route('/api/backup/list')
def backup_list():
    _, err = _backup_require_admin()
    if err: return err
    items = sorted(_backup_load_index(), key=lambda x: x.get('created_at', ''), reverse=True)
    return jsonify({'items': items, 'backup_dir': BACKUP_DIR})


@app.route('/api/backup/create', methods=['POST'])
def backup_create():
    user, err = _backup_require_admin()
    if err: return err
    keterangan = (request.json or {}).get('keterangan', '').strip()

    ts       = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'backup_{ts}.jsonl.gz'
    path     = os.path.join(BACKUP_DIR, filename)
    table_names = [t[0] for t in _MYSQL_TABLES]   # daftar tabel kanonik (sama utk kedua backend)

    con = get_db()
    try:
        cur = con.cursor()
        meta = {'__meta__': True, 'created_at': str(datetime.now()), 'keterangan': keterangan,
                'created_by': user.get('nama', ''), 'db_backend': DB_BACKEND, 'tables': table_names}
        total_rows, tables_done = 0, 0
        with gzip.open(path, 'wt', encoding='utf-8') as gz:
            gz.write(json.dumps(meta, ensure_ascii=False) + '\n')
            for tname in table_names:
                try:
                    cur.execute(f"SELECT * FROM {tname}")
                except Exception:
                    continue   # tabel blm ada (skema versi lama) -- lewati aman
                gz.write(json.dumps({'__table__': tname}) + '\n')
                n = 0
                while True:
                    batch = cur.fetchmany(2000)
                    if not batch:
                        break
                    for row in batch:
                        gz.write(json.dumps(dict(row), ensure_ascii=False, default=str) + '\n')
                        n += 1
                gz.write(json.dumps({'__table_end__': tname, 'row_count': n}) + '\n')
                total_rows += n
                tables_done += 1
    except Exception as e:
        try: os.remove(path)
        except Exception: pass
        return jsonify({'ok': False, 'msg': f'Gagal membuat backup: {e}'}), 500
    finally:
        con.close()

    record = {
        'filename': filename, 'keterangan': keterangan, 'created_by': user.get('nama', ''),
        'created_at': str(datetime.now()), 'size_bytes': os.path.getsize(path),
        'tables_count': tables_done, 'total_rows': total_rows, 'db_backend': DB_BACKEND,
    }
    items = _backup_load_index()
    items.append(record)
    _backup_save_index(items)

    catat_log('backup_buat', 'Backup', filename,
              f"Buat backup database — {tables_done} tabel, {total_rows} baris" +
              (f' — {keterangan}' if keterangan else ''))
    return jsonify({'ok': True, **record})


@app.route('/api/backup/<filename>/download')
def backup_download(filename):
    _, err = _backup_require_admin()
    if err: return err
    filename = os.path.basename(filename)   # cegah path traversal
    if not os.path.exists(os.path.join(BACKUP_DIR, filename)):
        return jsonify({'ok': False, 'msg': 'File backup tidak ditemukan'}), 404
    return send_from_directory(BACKUP_DIR, filename, as_attachment=True)


@app.route('/api/backup/<filename>', methods=['DELETE'])
def backup_delete(filename):
    _, err = _backup_require_admin()
    if err: return err
    filename = os.path.basename(filename)
    path = os.path.join(BACKUP_DIR, filename)
    if os.path.exists(path):
        os.remove(path)
    items = [it for it in _backup_load_index() if it.get('filename') != filename]
    _backup_save_index(items)
    catat_log('backup_hapus', 'Backup', filename, f"Hapus file backup '{filename}'")
    return jsonify({'ok': True})


def _backup_restore_from_path(path):
    """Inti proses restore (dipakai baik dari file yg ada di folder backup/ MAUPUN
    dari file .jsonl.gz yg diupload ulang — mis. hasil download sebelumnya, sbg
    jaga-jaga kalau file asli di server rusak/hilang). Hapus & timpa ULANG seluruh
    isi tabel dari snapshot. Skema tabel tidak disentuh (dipercayakan ke init_db()
    versi app yg sedang jalan) — yg direstore murni datanya, tabel demi tabel.

    PENTING (bukan route lagi — lihat catatan di bawah): tabel yg ADA di skema
    SEKARANG tapi TIDAK disebut sama sekali di file backup (mis. tabel itu BARU
    dibuat setelah backup itu diambil, jadi belum eksis saat itu) ikut DIKOSONGKAN
    juga — supaya hasil restore benar2 "replace total" sesuai kondisi sumber
    (tabel yg kosong/tak ada di sumber -> ikut kosong di target), BUKAN dibiarkan
    berisi data lama yg tak relevan dgn snapshot yg sedang direstore (ditemukan
    2026-10-01: restore backup lama bikin tabel 'Retur_Request' -yg blm ada saat
    backup itu dibuat- tetap berisi data terbaru, kelihatan spt "tercampur")."""
    con = get_raw_db(timeout=60)
    try:
        cur = con.cursor()
        cur.execute("PRAGMA foreign_keys = OFF")

        tables_restored = 0
        total_rows      = 0
        tables_in_backup = set()
        cur_table       = None
        skip_current    = False
        cols            = None
        sql             = None
        buf             = []

        def _flush():
            nonlocal buf
            if buf and sql:
                _bulk_executemany(con, cur, sql, buf)
            buf = []

        with gzip.open(path, 'rt', encoding='utf-8') as gz:
            for line in gz:
                line = line.strip()
                if not line:
                    continue
                obj = json.loads(line)
                if '__meta__' in obj:
                    continue
                if '__table__' in obj:
                    cur_table = obj['__table__']
                    tables_in_backup.add(cur_table)
                    cols = None; sql = None; buf = []
                    skip_current = False
                    try:
                        cur.execute(f"DELETE FROM `{cur_table}`")
                    except Exception:
                        skip_current = True   # tabel tak ada di skema saat ini -- lewati aman
                    continue
                if '__table_end__' in obj:
                    if not skip_current:
                        _flush()
                        tables_restored += 1
                        total_rows += obj.get('row_count', 0)
                    cur_table = None
                    continue
                # baris data
                if skip_current:
                    continue
                if cols is None:
                    cols = list(obj.keys())
                    placeholders = ','.join(['?'] * len(cols))
                    # Backtick tiap nama kolom & tabel — beberapa kolom (mis. Pengaturan.key)
                    # kebetulan sama dgn reserved word SQL ("KEY"), bikin INSERT tanpa
                    # backtick gagal syntax error di MySQL. SQLite jg menerima backtick
                    # (fitur kompatibilitas MySQL bawaan SQLite) jadi aman utk 2 backend.
                    col_list = ','.join(f'`{c}`' for c in cols)
                    sql = f"INSERT INTO `{cur_table}` ({col_list}) VALUES ({placeholders})"
                buf.append(tuple(obj.get(c) for c in cols))
                if len(buf) >= 2000:
                    _flush()

        # Tabel yg eksis di skema SEKARANG tapi tak pernah disebut di file backup
        # sama sekali (bukan yg diproses-lalu-skip di atas — itu kasus kebalikannya:
        # ada di backup, tak ada di skema sekarang) -> kosongkan jg, lihat docstring.
        tables_wiped_empty = 0
        for tname in [t[0] for t in _MYSQL_TABLES]:
            if tname not in tables_in_backup:
                try:
                    cur.execute(f"DELETE FROM `{tname}`")
                    tables_wiped_empty += 1
                except Exception:
                    pass

        cur.execute("PRAGMA foreign_keys = ON")
        con.commit()
    except Exception:
        try: con.rollback()
        except Exception: pass
        raise
    finally:
        con.close()
    return tables_restored, total_rows, tables_wiped_empty


@app.route('/api/backup/<filename>/restore', methods=['POST'])
def backup_restore(filename):
    """DESTRUKTIF: hapus & timpa ULANG seluruh isi tabel dari snapshot backup yg
    SUDAH ADA di folder backup/ server. Wajib ketik 'PULIHKAN' persis sbg
    konfirmasi eksplisit krn ini mengganti SELURUH data produksi.

    BUG DITEMUKAN & DIPERBAIKI 2026-10-01: dekorator route ini SEBELUMNYA kepasang
    di atas `_backup_restore_from_path` (helper internal tanpa cek admin/konfirmasi
    sama sekali), bukan di sini — jadi endpoint ini TIDAK PERNAH benar2 jalan
    (selalu 500 krn mismatch nama parameter `filename` vs `path`). Fungsi yg benar
    (dgn pengaman lengkap) ini jadi dead code selama itu. Sekarang dekoratornya
    sudah dipindah ke tempat yg benar."""
    user, err = _backup_require_admin()
    if err: return err
    filename = os.path.basename(filename)
    path = os.path.join(BACKUP_DIR, filename)
    if not os.path.exists(path):
        return jsonify({'ok': False, 'msg': 'File backup tidak ditemukan'}), 404

    konfirmasi = (request.json or {}).get('konfirmasi', '').strip().upper()
    if konfirmasi != 'PULIHKAN':
        return jsonify({'ok': False, 'msg': "Ketik PULIHKAN persis utk konfirmasi restore"}), 400

    try:
        tables_restored, total_rows, tables_wiped_empty = _backup_restore_from_path(path)
    except Exception as e:
        return jsonify({'ok': False, 'msg': f'Gagal restore: {e}'}), 500

    catat_log('backup_restore', 'Backup', filename,
              f"⚠ RESTORE database dari backup '{filename}' — {tables_restored} tabel direstore, "
              f"{total_rows} baris, {tables_wiped_empty} tabel ikut dikosongkan (tak ada di backup)")
    return jsonify({'ok': True, 'tables_restored': tables_restored, 'total_rows': total_rows,
                    'tables_wiped_empty': tables_wiped_empty})


@app.route('/api/backup/restore-upload', methods=['POST'])
def backup_restore_upload():
    """Restore dari file .jsonl.gz yg DIUPLOAD ulang (mis. hasil download sebelumnya
    yg disimpan di flashdisk/cloud) — jaga-jaga kalau file di folder backup/ server
    rusak atau hilang. Isi & aturan konfirmasi SAMA PERSIS dgn restore dari histori."""
    user, err = _backup_require_admin()
    if err: return err

    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'msg': 'Tidak ada file yang diunggah'}), 400
    if not f.filename.lower().endswith('.gz'):
        return jsonify({'ok': False, 'msg': 'File harus berformat .jsonl.gz (hasil download backup)'}), 400

    konfirmasi = (request.form.get('konfirmasi', '') or '').strip().upper()
    if konfirmasi != 'PULIHKAN':
        return jsonify({'ok': False, 'msg': "Ketik PULIHKAN persis utk konfirmasi restore"}), 400

    orig_name = os.path.basename(f.filename)
    tmp_path  = os.path.join(BACKUP_DIR, f'.upload_tmp_{datetime.now().strftime("%Y%m%d_%H%M%S_%f")}.jsonl.gz')
    f.save(tmp_path)

    # Validasi ringan: baris pertama wajib metadata backup asli — cegah file sembarangan
    try:
        with gzip.open(tmp_path, 'rt', encoding='utf-8') as gz:
            first = json.loads(gz.readline().strip() or '{}')
        if not first.get('__meta__'):
            os.remove(tmp_path)
            return jsonify({'ok': False, 'msg': 'File ini bukan hasil backup Rachel Clinic yang valid'}), 400
    except Exception as e:
        try: os.remove(tmp_path)
        except Exception: pass
        return jsonify({'ok': False, 'msg': f'File tidak bisa dibaca (rusak / bukan .jsonl.gz): {e}'}), 400

    try:
        tables_restored, total_rows, tables_wiped_empty = _backup_restore_from_path(tmp_path)
    except Exception as e:
        return jsonify({'ok': False, 'msg': f'Gagal restore: {e}'}), 500
    finally:
        try: os.remove(tmp_path)
        except Exception: pass

    catat_log('backup_restore', 'Backup', orig_name,
              f"⚠ RESTORE database dari file upload '{orig_name}' — {tables_restored} tabel direstore, "
              f"{total_rows} baris, {tables_wiped_empty} tabel ikut dikosongkan (tak ada di backup)")
    return jsonify({'ok': True, 'tables_restored': tables_restored, 'total_rows': total_rows,
                    'tables_wiped_empty': tables_wiped_empty})


if __name__ == '__main__':
    import socket
    init_db()
    _cek_format_exp_asing()
    try:
        host_ip = socket.gethostbyname(socket.gethostname())
    except Exception:
        host_ip = '0.0.0.0'
    print(f"\n{'='*52}")
    print(f"  Rachel Clinic Server")
    print(f"  Database                : {DB_BACKEND.upper()}")
    print(f"  Akses dari komputer ini : http://localhost:5000")
    print(f"  Akses dari komputer lain: http://{host_ip}:5000")
    print(f"{'='*52}\n")
    app.run(host='0.0.0.0', port=5000, debug=False, threaded=True)
